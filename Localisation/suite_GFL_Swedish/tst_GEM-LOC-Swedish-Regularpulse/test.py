# -*- coding: utf-8 -*-

import names
import re
import unicodedata
from openpyxl import load_workbook


EXCEL_PATH = (
    "/home/ezen/Test_Automation_Gemini/Localisation/"
    "suite_GFL_Swedish/"
    "tst_GEM-LOC-Swedish-PeakPower-SoftTissueCharacters/"
    "testdata/Gemini Split strings.xlsx"
)

ENGLISH_COLUMN = 2
Swedish_COLUMN = 25
SHEET_NAME = "Treatment Modes"

workbook = None
sheet = None
Swedish_strings = []


def get_object_text(obj):
    for attribute in (
        "text",
        "plainText",
        "title",
        "windowTitle",
        "accessibleName"
    ):
        try:
            value = getattr(obj, attribute)

            if value is not None:
                value = str(value)

                if value.strip():
                    return value

        except Exception:
            pass

    try:
        value = obj.property("text")

        if value is not None:
            value = str(value)

            if value.strip():
                return value

    except Exception:
        pass

    return ""


def normalize_text(value):
    if value is None:
        return ""

    text = str(value)

    text = re.sub(
        r"<br\s*/?>",
        " ",
        text,
        flags=re.IGNORECASE
    )

    text = re.sub(
        r"<[^>]*>",
        "",
        text
    )

    text = text.replace("\u00A0", " ")
    text = text.replace("\u2007", " ")
    text = text.replace("\u202F", " ")

    text = re.sub(r"[\r\n\t]+", " ", text)

    text = text.replace("™", "")
    text = text.replace("®", "")
    text = text.replace("©", "")

    try:
        text = unicodedata.normalize("NFKC", text)
    except Exception:
        pass

    text = re.sub(r"\s+", " ", text)
    text = re.sub(r"\s+([.,;:!?])", r"\1", text)

    return text.strip().casefold()


def comparison_key(value):
    normalized = normalize_text(value)
    return re.sub(r"\s+", "", normalized)


def load_Swedish_strings():
    global workbook
    global sheet
    global Swedish_strings

    try:
        workbook = load_workbook(
            EXCEL_PATH,
            read_only=True,
            data_only=True
        )

        if SHEET_NAME not in workbook.sheetnames:
            test.fail(
                "Worksheet '"
                + SHEET_NAME
                + "' does not exist in Excel file"
            )

            workbook.close()
            workbook = None
            return False

        sheet = workbook[SHEET_NAME]
        Swedish_strings = []

        Swedish_index = Swedish_COLUMN - ENGLISH_COLUMN

        for row in sheet.iter_rows(
            min_row=2,
            min_col=ENGLISH_COLUMN,
            max_col=Swedish_COLUMN,
            values_only=True
        ):
            english_value = row[0]

            if english_value is None:
                continue

            if Swedish_index >= len(row):
                continue

            Swedish_value = row[Swedish_index]

            if Swedish_value is None:
                continue

            english_value = str(english_value).strip()
            Swedish_value = str(Swedish_value).strip()

            if not Swedish_value:
                continue

            Swedish_strings.append(
                (
                    english_value,
                    Swedish_value
                )
            )

        if not Swedish_strings:
            test.fail(
                "No Swedish localization strings found in "
                "Excel Column Y."
            )
            workbook.close()
            workbook = None
            return False

        return True

    except Exception as e:
        test.fail(
            "Failed to load Gemini Split strings Excel / "
            "Treatment Modes: "
            + str(e)
        )

        if workbook is not None:
            try:
                workbook.close()
            except Exception:
                pass

            workbook = None

        return False


def find_Swedish_text(actual_text):
    if not actual_text:
        return None

    actual_key = comparison_key(actual_text)

    if not actual_key:
        return None

    for english_value, Swedish_value in Swedish_strings:
        Swedish_key = comparison_key(Swedish_value)

        if Swedish_key == actual_key:
            return english_value, Swedish_value

    return None


def verify_text(object_name, description):
    snooze(6)

    try:
        obj = waitForObject(object_name)

    except Exception as e:
        test.fail(
            description
            + " | Object not found: "
            + str(e)
        )
        return False

    actual_text = get_object_text(obj)

    if not actual_text:
        test.fail(
            description
            + " | UI text is empty"
        )
        return False

    result = find_Swedish_text(actual_text)

    if result is None:
        test.fail(
            description
            + " | Swedish text not found in Excel: "
            + str(actual_text)
        )
        return False

    english_value, expected_text = result

    actual_key = comparison_key(actual_text)
    expected_key = comparison_key(expected_text)

    test.compare(
        actual_key,
        expected_key,
        description + " - Swedish text verification"
    )

    if actual_key == expected_key:
        return True

    test.fail(
        description
        + " | Swedish text mismatch"
        + " | Actual: "
        + str(actual_text)
        + " | Expected: "
        + str(expected_text)
        + " | English Excel: "
        + str(english_value)
    )

    return False


def verify_group(group_name):
    for object_name, description in verification_objects[group_name]:
        verify_text(
            object_name,
            description
        )


def verify_popup(group_name):
    verify_group(group_name)

    snooze(6)

    clickButton(
        waitForObject(
            names.centralWidget_OKButton_QPushButton
        )
    )


verification_objects = {
    "main": [
        (
            names.centralWidget_lbTitle_QLabel,
            "Mode Selection Title"
        ),
        (
            names.centralWidget_label_2_QLabel,
            "Standard Mode Label"
        ),
        (
            names.centralWidget_Square_QPushButton,
            "Standard Pulse Button"
        ),
        (
            names.centralWidget_CW_QPushButton,
            "Continuous Wave Button"
        ),
        (
            names.centralWidget_MRP_QPushButton,
            "MRP Button"
        ),
        (
            names.centralWidget_label_QLabel,
            "Expert Mode Label"
        ),
        (
            names.centralWidget_Fragmentation_QPushButton,
            "Fragmentation Pulse Button"
        ),
        (
            names.centralWidget_Cleaning_QPushButton,
            "Cleaning Pulse Button"
        ),
        (
            names.centralWidget_Enucleation_QPushButton,
            "Enucleation Pulse Button"
        ),
        (
            names.buttonsFrame_rejectButton_QPushButton,
            "Cancel Button"
        )
    ],

    "standard_pulse": [
        (
            names.centralWidget_TitleLabel_QLabel_3,
            "Standard Pulse Popup Title"
        ),
        (
            names.centralWidget_MessageLabel_QLabel_3,
            "Standard Pulse Popup Message"
        ),
        (
            names.centralWidget_OKButton_QPushButton,
            "Standard Pulse Popup OK Button"
        )
    ],

    "cw": [
        (
            names.centralWidget_TitleLabel_QLabel_3,
            "Continuous Wave Popup Title"
        ),
        (
            names.centralWidget_MessageLabel_QLabel_3,
            "Continuous Wave Popup Message"
        ),
        (
            names.centralWidget_OKButton_QPushButton,
            "Continuous Wave Popup OK Button"
        )
    ],

    "mrp": [
        (
            names.centralWidget_TitleLabel_QLabel_3,
            "MRP Popup Title"
        ),
        (
            names.centralWidget_MessageLabel_QLabel_3,
            "MRP Popup Message"
        ),
        (
            names.centralWidget_OKButton_QPushButton,
            "MRP Popup OK Button"
        )
    ],

    "fragmentation": [
        (
            names.centralWidget_TitleLabel_QLabel_3,
            "Fragmentation Popup Title"
        ),
        (
            names.centralWidget_MessageLabel_QLabel_3,
            "Fragmentation Popup Message"
        ),
        (
            names.centralWidget_OKButton_QPushButton,
            "Fragmentation Popup OK Button"
        )
    ],

    "cleaning": [
        (
            names.centralWidget_TitleLabel_QLabel_3,
            "Cleaning Popup Title"
        ),
        (
            names.centralWidget_MessageLabel_QLabel_3,
            "Cleaning Pulse Popup Message"
        ),
        (
            names.centralWidget_OKButton_QPushButton,
            "Cleaning Pulse Popup OK Button"
        )
    ],

    "enucleation": [
        (
            names.centralWidget_TitleLabel_QLabel_3,
            "Enucleation Popup Title"
        ),
        (
            names.centralWidget_MessageLabel_QLabel_3,
            "Enucleation Popup Message"
        ),
        (
            names.centralWidget_OKButton_QPushButton,
            "Enucleation Popup OK Button"
        )
    ]
}


def main():
    global workbook

    if not load_Swedish_strings():
        return

    try:
        doubleClick(
            waitForObject(
                names.numpad_pushButton_8_NumpadButton
            ),
            91,
            77,
            Qt.NoModifier,
            Qt.LeftButton
        )

        snooze(6)

        clickButton(
            waitForObject(
                names.numpad_pushButton_8_NumpadButton
            )
        )

        snooze(6)

        clickButton(
            waitForObject(
                names.numpad_pushButton_8_NumpadButton
            )
        )

        snooze(6)

        clickButton(
            waitForObject(
                names.mainFrame_okButton_QPushButton
            )
        )

        snooze(6)

        clickButton(
            waitForObject(
                names.homeScreen_settingsButton_QPushButton
            )
        )

        snooze(6)

        clickTab(
            waitForObject(
                names.settingsScreen_tbSystemInfo_TabWidget
            ),
            "INFO"
        )

        snooze(6)

        clickButton(
            waitForObject(
                names.gbSystemInformation_pbLanguageSettingsUpdate_QPushButton
            )
        )

        snooze(6)

        clickButton(
            waitForObject(
                names.languagesFrame_pbSwedish_QPushButton
            )
        )

        clickButton(
            waitForObject(
                names.buttonsFrame_acceptButton_QPushButton
            )
        )

        snooze(6)

        clickButton(
            waitForObject(
                names.buttonsBar_homeButton_QPushButton
            )
        )

        snooze(6)

        clickButton(
            waitForObject(
                names.gbExpert_expertButton_QPushButton
            )
        )

        snooze(6)

        mouseClick(
            waitForObject(
                names.modeButton_nameLabel_QLabel
            ),
            73,
            25,
            Qt.NoModifier,
            Qt.LeftButton
        )

        snooze(6)

        verify_group("main")

        snooze(6)

        clickButton(
            waitForObject(
                names.square_icon_QPushButton
            )
        )

        verify_popup("standard_pulse")

        snooze(6)

        clickButton(
            waitForObject(
                names.cW_icon_QPushButton
            )
        )

        verify_popup("cw")

        snooze(6)

        clickButton(
            waitForObject(
                names.mRP_icon_QPushButton
            )
        )

        verify_popup("mrp")

        snooze(6)

        clickButton(
            waitForObject(
                names.fragmentation_icon_QPushButton
            )
        )

        verify_popup("fragmentation")

        snooze(6)

        clickButton(
            waitForObject(
                names.cleaning_icon_QPushButton
            )
        )

        verify_popup("cleaning")

        snooze(6)

        clickButton(
            waitForObject(
                names.enucleation_icon_QPushButton
            )
        )

        verify_popup("enucleation")

        snooze(6)

        clickButton(
            waitForObject(
                names.buttonsFrame_rejectButton_QPushButton
            )
        )

        snooze(6)

    except Exception as e:
        test.fail(
            "Swedish Mode Selection / Popup "
            "localization test failed: "
            + str(e)
        )

    finally:
        try:
            if workbook is not None:
                workbook.close()
                workbook = None
        except Exception:
            pass