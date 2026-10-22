# -*- coding: utf-8 -*-

import names
import re
import unicodedata
from openpyxl import load_workbook
import os 

TESTCASE_DIR = os.path.dirname(os.path.abspath(__file__))
 
EXCEL_PATH = os.path.abspath(
    os.path.join( TESTCASE_DIR,
        "..",
        "shared",
        "testdata",
        "Gemini Split strings.xlsx"
    )
)




ENGLISH_COLUMN = 2
GERMAN_COLUMN = 11
SHEET_NAME = "Treatment Modes"


def normalize_text(value):
    if value is None:
        return ""

    text = str(value)
    text = re.sub(r"<[^>]+>", "", text)

    text = (
        text
        .replace("\u00A0", " ")
        .replace("\u2007", " ")
        .replace("\u202F", " ")
        .replace("\r", " ")
        .replace("\n", " ")
        .replace("\t", " ")
        .replace("™", "")
        .replace("®", "")
        .replace("©", "")
    )

    try:
        text = unicodedata.normalize("NFKC", text)
    except Exception:
        pass

    text = re.sub(r"\s+", " ", text)
    text = re.sub(r"\s+([.,;:!?])", r"\1", text)

    return text.strip().casefold()


def comparison_key(value):
    return re.sub(r"\s+", "", normalize_text(value))


def get_german_values(sheet, english_text):
    expected_english = comparison_key(english_text)
    german_values = []

    german_index = GERMAN_COLUMN - ENGLISH_COLUMN

    for row in sheet.iter_rows(
        min_row=2,
        min_col=ENGLISH_COLUMN,
        max_col=GERMAN_COLUMN,
        values_only=True
    ):
        english_value = row[0]

        if english_value is None:
            continue

        if comparison_key(english_value) != expected_english:
            continue

        if german_index >= len(row):
            continue

        german_value = row[german_index]

        if german_value is None:
            continue

        german_value = str(german_value).strip()

        if german_value and german_value not in german_values:
            german_values.append(german_value)

    return german_values


def get_object_text(obj):
    for attribute in ("text", "title", "windowTitle", "accessibleName"):
        try:
            value = getattr(obj, attribute)

            if value is not None and str(value).strip():
                return str(value)
        except Exception:
            pass

    try:
        value = obj.property("text")

        if value is not None and str(value).strip():
            return str(value)
    except Exception:
        pass

    return ""


def verify_excel_string(object_name, sheet, english_text, description):
    snooze(2)

    try:
        obj = waitForObjectExists(object_name)
    except Exception as e:
        test.fail(
            description
            + " | Object not found: "
            + str(e)
        )
        return False

    actual = get_object_text(obj).strip()

    if not actual:
        test.fail(description + " | UI text is empty")
        return False

    expected_values = get_german_values(sheet, english_text)

    if not expected_values:
        test.fail(
            description
            + " | English string not found in Column B: "
            + english_text
        )
        return False

    actual_key = comparison_key(actual)

    for expected in expected_values:
        expected_key = comparison_key(expected)

        if actual_key == expected_key:
            test.compare(
                actual_key,
                expected_key,
                description + " - German text verification"
            )
            return True

    test.fail(
        description
        + " | Expected German value from Column K: "
        + " / ".join(expected_values)
        + " | Actual: "
        + actual
    )

    return False


def verify_direct_text(actual, expected, description):
    actual_key = comparison_key(actual)
    expected_key = comparison_key(expected)

    test.compare(
        actual_key,
        expected_key,
        description
    )


def main():
    workbook = None

    try:
        workbook = load_workbook(
            EXCEL_PATH,
            read_only=True,
            data_only=True
        )

        if SHEET_NAME not in workbook.sheetnames:
            test.fail(
                "Excel sheet not found: "
                + SHEET_NAME
                + " | Available sheets: "
                + ", ".join(workbook.sheetnames)
            )
            return

        sheet = workbook[SHEET_NAME]

        clickButton(
            waitForObject(
                names.numpad_pushButton_8_NumpadButton
            )
        )

        clickButton(
            waitForObject(
                names.numpad_pushButton_8_NumpadButton
            )
        )

        clickButton(
            waitForObject(
                names.numpad_pushButton_8_NumpadButton
            )
        )

        clickButton(
            waitForObject(
                names.numpad_pushButton_8_NumpadButton
            )
        )

        clickButton(
            waitForObject(
                names.mainFrame_okButton_QPushButton
            )
        )

        clickButton(
            waitForObject(
                names.homeScreen_settingsButton_QPushButton
            )
        )

        clickTab(
            waitForObject(
                names.settingsScreen_tbSystemInfo_TabWidget
            ),
            "INFO"
        )

        clickButton(
            waitForObject(
                names.gbSystemInformation_pbLanguageSettingsUpdate_QPushButton
            )
        )

        clickButton(
            waitForObject(
                names.languagesFrame_pbGermany_QPushButton
            )
        )

        clickButton(
            waitForObject(
                names.buttonsFrame_acceptButton_QPushButton
            )
        )

        snooze(3)

        tabWidget = waitForObject(
            names.settingsScreen_tbSystemInfo_TabWidget
        )

        tabWidget.setCurrentIndex(2)
        snooze(3)

        tabWidget.setCurrentIndex(2)
        snooze(3)

        sendEvent(
            "QMoveEvent",
            waitForObject(names.o_ScreenSwitcher),
            0,
            -66,
            1285,
            233
        )

        clickButton(
            waitForObject(
                names.settingsScreen_addButton_QPushButton
            )
        )

        verify_excel_string(
            names.centralWidget_lbOperation_QLabel,
            sheet,
            "ADD PRESET",
            "Add Preset Operation"
        )

        verify_excel_string(
            names.centralWidget_pbName_QPushButton,
            sheet,
            "NAME",
            "Preset Name"
        )

        verify_excel_string(
            names.averagePowerWidget_nameLabel_QLabel,
            sheet,
            "AVERAGE POWER",
            "Average Power"
        )

        verify_excel_string(
            names.averagePowerWidget_unitsOfMeasureLabel_QLabel,
            sheet,
            "W",
            "Average Power Unit"
        )

        verify_excel_string(
            names.pulseEnergyWidget_nameLabel_QLabel,
            sheet,
            "PULSE ENERGY",
            "Pulse Energy"
        )

        mouseClick(
            waitForObject(names.centralWidget_frame_QFrame),
            271,
            13,
            Qt.NoModifier,
            Qt.LeftButton
        )

        verify_excel_string(
            names.pulseEnergyWidget_unitsOfMeasureLabel_QLabel,
            sheet,
            "J",
            "Pulse Energy Unit"
        )

        verify_excel_string(
            names.frequencyWidget_nameLabel_QLabel,
            sheet,
            "FREQUENCY",
            "Frequency"
        )

        verify_excel_string(
            names.frequencyWidget_unitsOfMeasureLabel_QLabel,
            sheet,
            "Hz",
            "Frequency Unit"
        )

        verify_excel_string(
            names.modeButton_nameLabel_QLabel_3,
            sheet,
            "REGULAR PULSE",
            "Regular Pulse"
        )

        verify_excel_string(
            names.peakPowerButton_nameLabel_QLabel_2,
            sheet,
            "PEAK POWER",
            "Peak Power"
        )

        verify_excel_string(
            names.pulseModeWidget_saveButton_QPushButton,
            sheet,
            "ADD",
            "Add Button"
        )

        mouseClick(
            waitForObject(
                names.pulseModeWidget_buttonsWidget_QWidget
            ),
            508,
            46,
            Qt.NoModifier,
            Qt.LeftButton
        )

        mouseClick(
            waitForObject(
                names.pulseModeWidget_buttonsWidget_QWidget
            ),
            411,
            1,
            Qt.NoModifier,
            Qt.LeftButton
        )

        clickButton(
            waitForObject(
                names.centralWidget_pbName_QPushButton
            )
        )

        type(
            waitForObject(
                names.keyboardPopup_wdEdit_WdLineEdit
            ),
            "preset one"
        )

        mouseClick(
            waitForObject(
                names.keyboardPopup_lbText_QLabel
            ),
            77,
            26,
            Qt.NoModifier,
            Qt.LeftButton
        )

        clickButton(
            waitForObject(
                names.pulseModeWidget_saveButton_QPushButton
            )
        )

        mouseClick(
            waitForObject(
                names.tabSettingsPresetsUserStonePresets_presetListView_PresetListView
            ),
            546,
            72,
            Qt.NoModifier,
            Qt.LeftButton
        )

        clickButton(
            waitForObject(
                names.settingsScreen_editButton_QPushButton
            )
        )

        mouseClick(
            waitForObject(
                names.tabSettingsPresetsUserStonePresets_presetListView_PresetListView
            ),
            494,
            66,
            Qt.NoModifier,
            Qt.LeftButton
        )

        verify_excel_string(
            names.centralWidget_lbOperation_QLabel,
            sheet,
            "EDIT PRESET",
            "Edit Preset Operation"
        )

        snooze(2)

        actual_edit_name = str(
            waitForObjectExists(
                names.centralWidget_pbName_QPushButton
            ).text
        )

        verify_direct_text(
            actual_edit_name,
            "preset one",
            "Edit Preset Name"
        )

        verify_excel_string(
            names.averagePowerWidget_nameLabel_QLabel,
            sheet,
            "AVERAGE POWER",
            "Edit - Average Power"
        )

        verify_excel_string(
            names.averagePowerWidget_unitsOfMeasureLabel_QLabel,
            sheet,
            "W",
            "Edit - Average Power Unit"
        )

        verify_excel_string(
            names.pulseEnergyWidget_nameLabel_QLabel,
            sheet,
            "PULSE ENERGY",
            "Edit - Pulse Energy"
        )

        verify_excel_string(
            names.pulseEnergyWidget_unitsOfMeasureLabel_QLabel,
            sheet,
            "J",
            "Edit - Pulse Energy Unit"
        )

        verify_excel_string(
            names.frequencyWidget_nameLabel_QLabel,
            sheet,
            "FREQUENCY",
            "Edit - Frequency"
        )

        verify_excel_string(
            names.frequencyWidget_unitsOfMeasureLabel_QLabel,
            sheet,
            "Hz",
            "Edit - Frequency Unit"
        )

        verify_excel_string(
            names.modeButton_nameLabel_QLabel_3,
            sheet,
            "REGULAR PULSE",
            "Edit - Regular Pulse"
        )

        verify_excel_string(
            names.peakPowerButton_nameLabel_QLabel_2,
            sheet,
            "PEAK POWER",
            "Edit - Peak Power"
        )

        verify_excel_string(
            names.pulseModeWidget_saveButton_QPushButton,
            sheet,
            "UPDATE",
            "Update Button"
        )

        mouseClick(
            waitForObject(
                names.pulseModeWidget_buttonsWidget_QWidget
            ),
            531,
            22,
            Qt.NoModifier,
            Qt.LeftButton
        )

        mouseClick(
            waitForObject(
                names.pulseModeWidget_buttonsWidget_QWidget
            ),
            509,
            30,
            Qt.NoModifier,
            Qt.LeftButton
        )

        clickButton(
            waitForObject(
                names.pulseModeWidget_saveButton_QPushButton
            )
        )

        clickButton(
            waitForObject(
                names.settingsScreen_deletePresetButton_QPushButton
            )
        )

        mouseClick(
            waitForObject(
                names.tabSettingsPresetsUserStonePresets_presetListView_PresetListView
            ),
            553,
            80,
            Qt.NoModifier,
            Qt.LeftButton
        )

        verify_excel_string(
            names.centralWidget_lbDialogTitle_QLabel,
            sheet,
            "DELETE PRESET",
            "Delete Dialog Title"
        )

        snooze(2)

        actual_delete_name = str(
            waitForObjectExists(
                names.centralWidget_questionLabel_QLabel
            ).text
        )

        verify_direct_text(
            actual_delete_name,
            "preset one",
            "Delete Preset Name"
        )

        snooze(2)

        actual_info = str(
            waitForObjectExists(
                names.infoFrame_extInfoLabel_QLabel
            ).text
        )

        expected_info = "6 J     3 Hz     18 W"

        verify_direct_text(
            actual_info,
            expected_info,
            "Preset Information"
        )

        verify_excel_string(
            names.buttonsFrame_rejectButton_QPushButton,
            sheet,
            "CANCEL",
            "Cancel Button"
        )

        verify_excel_string(
            names.centralWidget_acceptButton_QPushButton,
            sheet,
            "DELETE",
            "Delete Button"
        )
        test.vp("Delete Preset screen")
        clickButton(
            waitForObject(
                names.centralWidget_acceptButton_QPushButton
            )
        )
    except Exception as e:
        test.fail(
            "German User Presets localization test failed: "
            + str(e)
        )
        
    finally:
        if workbook is not None:
            try:
                workbook.close()
            except Exception:
                pass