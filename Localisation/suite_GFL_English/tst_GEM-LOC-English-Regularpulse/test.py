
# -*- coding: utf-8 -*-

import names
import openpyxl
import re
from openpyxl import load_workbook


EXCEL_PATH = (
    "/home/ezen/Test_Automation_Gemini/Localisation/"
    "suite_GFL_English/tst_GEM-LOC-English-PeakPower-SoftTissueCharacters/"
    "testdata/Gemini Split strings.xlsx"
)

# Gemini Split strings.xlsx
# Column B = English
ENGLISH_COLUMN = 2

SHEET_NAME = "Treatment Modes"


workbook = None
sheet = None


def get_object_text(obj):

    try:
        value = obj.text

        if value is not None:
            value = str(value)

            if value.strip():
                return value

    except:
        pass


    try:
        value = obj.title

        if value is not None:
            value = str(value)

            if value.strip():
                return value

    except:
        pass

    try:
        value = obj.windowTitle

        if value is not None:
            value = str(value)

            if value.strip():
                return value

    except:
        pass

    try:
        value = obj.accessibleName

        if value is not None:
            value = str(value)

            if value.strip():
                return value

    except:
        pass


    return ""


def normalize_text(text):

    if text is None:
        return ""

    text = str(text)

    text = re.sub(
        r"<[^>]*>",
        " ",
        text
    )


    text = text.replace(
        "\u00A0",
        " "
    )

    text = text.replace(
        "\u2007",
        " "
    )

    text = text.replace(
        "\u202F",
        " "
    )

    text = re.sub(
        r"[\r\n\t]+",
        " ",
        text
    )


    text = text.replace(
        "™",
        ""
    )

    text = text.replace(
        "®",
        ""
    )

    text = text.replace(
        "©",
        ""
    )

    try:

        import unicodedata

        text = unicodedata.normalize(
            "NFKC",
            text
        )

    except:
        pass
    text = re.sub(
        r"\s+",
        " ",
        text
    )

    text = re.sub(
        r"\s+([.,;:!?])",
        r"\1",
        text
    )

    return text.strip().casefold()


def find_english_text(actual_text):

    if sheet is None:
        return None

    actual_normalized = normalize_text(
        actual_text
    )

    if not actual_normalized:
        return None

    for row in sheet.iter_rows(
        min_row=2,
        min_col=ENGLISH_COLUMN,
        max_col=ENGLISH_COLUMN
    ):

        english_value = row[0].value

        if english_value is None:
            continue

        english_normalized = normalize_text(
            english_value
        )

        if not english_normalized:
            continue


        if actual_normalized == english_normalized:
            return english_value


    return None

def verify_text(object_name, description):

    snooze(6)


    try:

        obj = waitForObject(
            object_name
        )

    except Exception as e:

        test.fail(
            description
            + " | Object not found: "
            + str(e)
        )

        return False

    actual_text = get_object_text(
        obj
    )


    if not actual_text:

        test.fail(
            description
            + " | UI text is empty"
        )

        return False


    matched_english = find_english_text(
        actual_text
    )

    if matched_english is not None:

        actual_normalized = normalize_text(
            actual_text
        )

        expected_normalized = normalize_text(
            matched_english
        )

        test.compare(
            actual_normalized,
            expected_normalized,
            description
        )

        return True


    test.fail(
        description
        + " | English text not found in Column B: "
        + str(actual_text)
    )

    return False


def verify_group(group_name):

    for object_name, description in verification_objects[group_name]:

        verify_text(
            object_name,
            description
        )

def verify_popup(group_name):

    verify_group(
        group_name
    )

    snooze(6)

    clickButton(
        waitForObject(
            names.centralWidget_OKButton_QPushButton
        )
    )

verification_objects = {


    "main": [

        (
            names.centralWidget_lbTitle_QLabel,
            "Mode Selection Title"
        ),

        (
            names.centralWidget_label_2_QLabel,
            "Standard Mode Label"
        ),

        (
            names.centralWidget_Square_QPushButton,
            "Standard Pulse Button"
        ),

        (
            names.centralWidget_CW_QPushButton,
            "Continuous Wave Button"
        ),

        (
            names.centralWidget_MRP_QPushButton,
            "MRP Button"
        ),

        (
            names.centralWidget_label_QLabel,
            "Expert Mode Label"
        ),

        (
            names.centralWidget_Fragmentation_QPushButton,
            "Fragmentation Pulse Button"
        ),

        (
            names.centralWidget_Cleaning_QPushButton,
            "Cleaning Pulse Button"
        ),

        (
            names.centralWidget_Enucleation_QPushButton,
            "Enucleation Pulse Button"
        ),

        (
            names.buttonsFrame_rejectButton_QPushButton,
            "Cancel Button"
        )
    ],


    "standard_pulse": [

        (
            names.centralWidget_TitleLabel_QLabel_3,
            "Standard Pulse Popup Title"
        ),

        (
            names.centralWidget_MessageLabel_QLabel_3,
            "Standard Pulse Popup Message"
        ),

        (
            names.centralWidget_OKButton_QPushButton,
            "Standard Pulse Popup OK Button"
        )
    ],


    "cw": [

        (
            names.centralWidget_TitleLabel_QLabel_3,
            "Continuous Wave Popup Title"
        ),

        (
            names.centralWidget_MessageLabel_QLabel_3,
            "Continuous Wave Popup Message"
        ),

        (
            names.centralWidget_OKButton_QPushButton,
            "Continuous Wave Popup OK Button"
        )
    ],

    "mrp": [

        (
            names.centralWidget_TitleLabel_QLabel_3,
            "MRP Popup Title"
        ),

        (
            names.centralWidget_MessageLabel_QLabel_3,
            "MRP Popup Message"
        ),

        (
            names.centralWidget_OKButton_QPushButton,
            "MRP Popup OK Button"
        )
    ],

    "fragmentation": [

        (
            names.centralWidget_TitleLabel_QLabel_3,
            "Fragmentation Popup Title"
        ),

        (
            names.centralWidget_MessageLabel_QLabel_3,
            "Fragmentation Popup Message"
        ),

        (
            names.centralWidget_OKButton_QPushButton,
            "Fragmentation Popup OK Button"
        )
    ],


    "cleaning": [

        (
            names.centralWidget_TitleLabel_QLabel_3,
            "Cleaning Popup Title"
        ),

        (
            names.centralWidget_MessageLabel_QLabel_3,
            "Cleaning Popup Message"
        ),

        (
            names.centralWidget_OKButton_QPushButton,
            "Cleaning Popup OK Button"
        )
    ],


    "enucleation": [

        (
            names.centralWidget_TitleLabel_QLabel_3,
            "Enucleation Popup Title"
        ),

        (
            names.centralWidget_MessageLabel_QLabel_3,
            "Enucleation Popup Message"
        ),

        (
            names.centralWidget_OKButton_QPushButton,
            "Enucleation Popup OK Button"
        )
    ]
}


def main():

    global workbook
    global sheet


    doubleClick(
        waitForObject(
            names.numpad_pushButton_8_NumpadButton
        ),
        91,
        77,
        Qt.NoModifier,
        Qt.LeftButton
    )

    snooze(6)

    clickButton(
        waitForObject(
            names.numpad_pushButton_8_NumpadButton
        )
    )

    snooze(6)

    clickButton(
        waitForObject(
            names.numpad_pushButton_8_NumpadButton
        )
    )

    snooze(6)

    clickButton(
        waitForObject(
            names.mainFrame_okButton_QPushButton
        )
    )

    snooze(6)

    # test.log(
    #     "Opening Settings"
    # )

    clickButton(
        waitForObject(
            names.homeScreen_settingsButton_QPushButton
        )
    )


    snooze(6)

    clickTab(
        waitForObject(
            names.settingsScreen_tbSystemInfo_TabWidget
        ),
        "INFO"
    )

    snooze(6)

    clickButton(
        waitForObject(
            names.buttonsBar_homeButton_QPushButton
        )
    )


    snooze(6)

    clickButton(
        waitForObject(
            names.gbExpert_expertButton_QPushButton
        )
    )

    snooze(6)

    mouseClick(
        waitForObject(
            names.modeButton_nameLabel_QLabel
        ),
        73,
        25,
        Qt.NoModifier,
        Qt.LeftButton
    )


    snooze(2)

    # test.log(
    #     "Loading English Gemini Split strings"
    # )

    try:

        workbook = load_workbook(
            EXCEL_PATH,
            read_only=True,
            data_only=True
        )

    except Exception as e:

        test.fail(
            "Failed to load English Gemini Split strings Excel: "
            + str(e)
        )

        return

    if SHEET_NAME not in workbook.sheetnames:

        test.fail(
            "Worksheet '" +
            SHEET_NAME +
            "' does not exist. Available sheets: " +
            str(workbook.sheetnames)
        )

        workbook.close()

        return


    sheet = workbook[
        SHEET_NAME
    ]


    verify_group(
        "main"
    )

    snooze(6)

    clickButton(
        waitForObject(
            names.square_icon_QPushButton
        )
    )

    verify_popup(
        "standard_pulse"
    )

    snooze(6)

    clickButton(
        waitForObject(
            names.cW_icon_QPushButton
        )
    )

    verify_popup(
        "cw"
    )

    snooze(6)

    clickButton(
        waitForObject(
            names.mRP_icon_QPushButton
        )
    )

    verify_popup(
        "mrp"
    )

    snooze(6)

    clickButton(
        waitForObject(
            names.fragmentation_icon_QPushButton
        )
    )

    verify_popup(
        "fragmentation"
    )
    snooze(6)

    clickButton(
        waitForObject(
            names.cleaning_icon_QPushButton
        )
    )

    verify_popup(
        "cleaning"
    )
    snooze(6)

    clickButton(
        waitForObject(
            names.enucleation_icon_QPushButton
        )
    )

    verify_popup(
        "enucleation"
    )
    snooze(6)

    clickButton(
        waitForObject(
            names.buttonsFrame_rejectButton_QPushButton
        )
    )

    try:

        if workbook is not None:
            workbook.close()

    except:

        pass


    

