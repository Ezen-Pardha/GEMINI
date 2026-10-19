
# -*- coding: utf-8 -*-

import names
import openpyxl
import re
import time
import os

TESTCASE_DIR = os.path.dirname(os.path.abspath(__file__))
 
EXCEL_PATH = os.path.abspath(
    os.path.join( TESTCASE_DIR,
        "..",
        "shared",
        "testdata",
        "Gemini Split strings.xlsx"
    )
)

def normalize(text):

    if text is None:
        return ""

    text = str(text)

    # Remove HTML line break tags
    text = re.sub(
        r"<br\s*/?>",
        " ",
        text,
        flags=re.IGNORECASE
    )

    # Remove other HTML tags
    text = re.sub(
        r"<[^>]+>",
        " ",
        text
    )

    # Replace non-breaking space
    text = text.replace("\xa0", " ")

    # Remove extra spaces and new lines
    text = " ".join(text.split())

    return text.strip().casefold()

def load_english_excel(EXCEL_PATH):

    start_time = time.time()

    workbook = openpyxl.load_workbook(
        EXCEL_PATH,
        data_only=True,
        read_only=True
    )

    load_time = time.time() - start_time

    # test.log(
    #     "Excel loading completed in "
    #     + str(round(load_time, 2))
    #     + " seconds"
    # )

    # Check required worksheet
    if "Treatment Modes" not in workbook.sheetnames:

        test.fail(
            "Worksheet 'Treatment Modes' does not exist "
            "in Excel file"
        )

        workbook.close()
        return None, None, {}

    sheet = workbook["Treatment Modes"]

    # English = Column B
    english_column = 2

    english_data = {}

    start_time = time.time()

    # Read only English column
    for row in sheet.iter_rows(
        min_col=english_column,
        max_col=english_column
    ):

        english_value = row[0].value

        if english_value is None:
            continue

        english_key = normalize(english_value)

        if english_key == "":
            continue

        english_data[english_key] = str(english_value)

    dictionary_time = time.time() - start_time

    # test.log(
    #     "English Excel dictionary created in "
    #     + str(round(dictionary_time, 2))
    #     + " seconds"
    # )

    return workbook, sheet, english_data

def get_english_text(english_data, english_text):

    key = normalize(english_text)

    if key not in english_data:

        test.fail(
            "English text not found in Excel: "
            + str(english_text)
        )

        return ""

    expected = english_data[key]

    return expected

def verify_text(english_data, object_name, english_text):

    actual_text = ""

    try:

        obj = waitForObjectExists(object_name)

        actual_text = str(obj.text).strip()

    except Exception as e:

        test.fail(
            "Unable to read application object: "
            + str(object_name)
            + " | Error: "
            + str(e)
        )

        return False

    # Get expected English text from Excel
    expected_text = get_english_text(
        english_data,
        english_text
    )

    actual_normalized = normalize(actual_text)
    expected_normalized = normalize(expected_text)

    if actual_normalized == expected_normalized:

        # test.log(
        #     "PASS - English text verified: "
        #     + str(english_text)
        #     + " -> "
        #     + actual_text
        # )

        test.compare(
            actual_normalized,
            expected_normalized,
            "English text verification: "
            + str(english_text)
        )

        return True

    else:

        test.fail(
            "FAIL - English text mismatch for: "
            + str(english_text)
            + " | Actual: "
            + actual_text
            + " | Expected from Excel: "
            + str(expected_text)
        )

        return False

def main():

    # EXECL_PATH = (
    #     "/home/ezen/Test_Automation_Gemini/Localisation/"
    #     "suite_GFL_English/"
    #     "tst_GEM-LOC-English-PeakPower-SoftTissueCharacters/"
    #     "testdata/Gemini Split strings.xlsx"
    # )


   
    clickButton(
        waitForObject(
            names.numpad_pushButton_8_NumpadButton
        )
    )

    clickButton(
        waitForObject(
            names.numpad_pushButton_8_NumpadButton
        )
    )

    clickButton(
        waitForObject(
            names.numpad_pushButton_8_NumpadButton
        )
    )

    clickButton(
        waitForObject(
            names.numpad_pushButton_8_NumpadButton
        )
    )

    snooze(6)

    clickButton(
        waitForObject(
            names.mainFrame_okButton_QPushButton
        )
    )


    clickButton(
        waitForObject(
            names.homeScreen_pbSoftTissueQuickStart_QToolButton
        )
    )

    mouseClick(
        waitForObjectItem(
            names.tabLeftMainPresets_leftPresetsView_PresetListView,
            "_1"
        ),
        305,
        35,
        Qt.NoModifier,
        Qt.LeftButton
    )

    mouseClick(
        waitForObjectItem(
            names.tabRightMainPresets_rightPresetsView_PresetListView,
            "_1"
        ),
        299,
        49,
        Qt.NoModifier,
        Qt.LeftButton
    )

    clickButton(
        waitForObject(
            names.quickStartScreen_pbContinue_QPushButton
        )
    )

    clickButton(
        waitForObject(
            names.mainFrame_okButton_QPushButton_2
        )
    )


    mouseClick(
        waitForObject(
            names.peakPowerButton_nameLabel_QLabel
        ),
        34,
        34,
        Qt.NoModifier,
        Qt.LeftButton
    )

    snooze(2)


    workbook, sheet, english_data = load_english_excel(
        EXCEL_PATH
    )


    verify_text(
        english_data,
        names.centralWidget_lbTitle_QLabel,
        "SELECT PEAK POWER"
    )

    verify_text(
        english_data,
        names.centralWidget_pw500_QPushButton,
        "High"
    )

    verify_text(
        english_data,
        names.centralWidget_pw250_QPushButton,
        "Medium"
    )

    verify_text(
        english_data,
        names.centralWidget_pw125_QPushButton,
        "Low"
    )

    verify_text(
        english_data,
        names.buttonsFrame_rejectButton_QPushButton,
        "CANCEL"
    )


    mouseClick(
        waitForObject(
            names.centralWidget_mainWidget_QWidget
        ),
        460,
        549,
        Qt.NoModifier,
        Qt.LeftButton
    )

    clickButton(
        waitForObject(
            names.buttonsFrame_rejectButton_QPushButton
        )
    )

    clickButton(
        waitForObject(
            names.buttonsBar_homeButton_QPushButton_2
        )
    )

    clickButton(
        waitForObject(
            names.homeScreen_pbSoftTissueAssistant_QToolButton
        )
    )

    clickButton(
        waitForObject(
            names.gbSoftTissueProcedures_pbIncision_QPushButton
        )
    )

    clickButton(
        waitForObject(
            names.softTissueAssistantModeScreen_pbContinue_QPushButton
        )
    )

    clickButton(
        waitForObject(
            names.treatmentCharacteristicsBar_pbEdit_QPushButton
        )
    )

    snooze(2)


    verify_text(
        english_data,
        names.centralWidget_titleLabel_QLabel,
        "SOFT TISSUE CHARACTERISTICS"
    )

    verify_text(
        english_data,
        names.centralWidget_titleLabel_QLabel_2,
        "PROCEDURE"
    )

    verify_text(
        english_data,
        names.centralWidget_valueLabel_QLabel,
        "Incision"
    )

    verify_text(
        english_data,
        names.centralWidget_rejectButton_QPushButton,
        "EDIT"
    )

    verify_text(
        english_data,
        names.centralWidget_acceptButton_QPushButton,
        "OK"
    )


    if workbook is not None:
        workbook.close()

