
# -*- coding: utf-8 -*-

import names
import openpyxl
import re
from openpyxl import load_workbook


EXCEL_PATH = (
    "/home/ezen/Test_Automation_Gemini/Localisation/"
    "suite_GFL_English/tst_GEM-LOC-English-PeakPower-SoftTissueCharacters/"
    "testdata/Gemini Split strings.xlsx"
)

SHEET_NAME = "Treatment Modes"
ENGLISH_COLUMN = 2


def normalize(text):
    """
    Normalize text for comparison.
    Removes HTML, line breaks, extra spaces and non-breaking spaces.
    """
    if text is None:
        return ""

    text = str(text)

    # Convert HTML line breaks to spaces
    text = re.sub(
        r"<br\s*/?>",
        " ",
        text,
        flags=re.IGNORECASE
    )

    # Remove HTML tags
    text = re.sub(
        r"<[^>]+>",
        " ",
        text
    )

    # Replace non-breaking spaces
    text = text.replace("\xa0", " ")

    # Normalize whitespace
    text = " ".join(text.split())

    return text.strip().casefold()


def load_english_excel(excelpath):

    try:
        workbook = load_workbook(
            excelpath,
            read_only=True,
            data_only=True
        )

    except Exception as e:
        test.fail(
            "Unable to load English Excel file: " +
            str(e)
        )
        return None, {}

    if SHEET_NAME not in workbook.sheetnames:

        test.fail(
            "Worksheet '" +
            SHEET_NAME +
            "' does not exist. Available sheets: " +
            str(workbook.sheetnames)
        )

        workbook.close()

        return None, {}

    sheet = workbook[SHEET_NAME]

    english_data = {}

    # Read only English column
    for row in sheet.iter_rows(
        min_col=ENGLISH_COLUMN,
        max_col=ENGLISH_COLUMN
    ):

        english_value = row[0].value

        if english_value is None:
            continue

        english_value = str(english_value).strip()

        if english_value == "":
            continue

        key = normalize(english_value)

        if key == "":
            continue

        english_data[key] = english_value

    return workbook, english_data


def get_english_text(english_data, english_lookup):

    key = normalize(english_lookup)

    if key not in english_data:

        test.fail(
            "English string not found in Excel Column B: " +
            str(english_lookup)
        )

        return ""

    return english_data[key]


def verify_text(
    english_data,
    object_name,
    english_lookup,
    verification_name
):

    try:

        object_item = waitForObjectExists(
            object_name
        )

        actual_text = str(
            object_item.text
        ).strip()

    except Exception as e:

        test.fail(
            "Unable to read object '" +
            str(object_name) +
            "': " +
            str(e)
        )

        return False

    expected_text = get_english_text(
        english_data,
        english_lookup
    )

    actual_normalized = normalize(
        actual_text
    )

    expected_normalized = normalize(
        expected_text
    )

    test.compare(
        actual_normalized,
        expected_normalized,
        verification_name
    )

    return (
        actual_normalized ==
        expected_normalized
    )


def main():

    

    clickButton(
        waitForObject(
            names.numpad_pushButton_8_NumpadButton
        )
    )

    clickButton(
        waitForObject(
            names.numpad_pushButton_8_NumpadButton
        )
    )

    clickButton(
        waitForObject(
            names.numpad_pushButton_8_NumpadButton
        )
    )

    clickButton(
        waitForObject(
            names.numpad_pushButton_8_NumpadButton
        )
    )

    snooze(6)

    clickButton(
        waitForObject(
            names.mainFrame_okButton_QPushButton
        )
    )

    snooze(6)


    # ---------------------------------------------------------
    # OPEN SETTINGS
    # ---------------------------------------------------------

    

    clickButton(
        waitForObject(
            names.homeScreen_settingsButton_QPushButton
        )
    )

    snooze(2)


    # ---------------------------------------------------------
    # OPEN USER PRESETS
    # ---------------------------------------------------------
    # Settings contains 3 tabs.
    # User Presets is tab index 2.
    # No INFO or WIRELESS FOOTSWITCH click is required.
    # ---------------------------------------------------------

    

    tab_widget = waitForObject(
        names.settingsScreen_tbSystemInfo_TabWidget
    )

    tab_widget.setCurrentIndex(2)

    snooze(3)


    # ---------------------------------------------------------
    # OPEN EXCEL
    # ---------------------------------------------------------

    

    workbook, english_data = load_english_excel(
        EXCEL_PATH
    )

    if workbook is None:
        test.fail(
            "English Excel data could not be loaded."
        )
        return


    # ---------------------------------------------------------
    # ADD PRESET
    # ---------------------------------------------------------

   
    clickButton(
        waitForObject(
            names.settingsScreen_addButton_QPushButton
        )
    )

    snooze(2)


    # ---------------------------------------------------------
    # VERIFY ADD PRESET
    # ---------------------------------------------------------

    snooze(6)

    verify_text(
        english_data,
        names.centralWidget_lbOperation_QLabel,
        "ADD PRESET",
        "Add Preset operation text"
    )

    snooze(6)

    verify_text(
        english_data,
        names.centralWidget_pbName_QPushButton,
        "Enter preset name",
        "Preset name button text"
    )

    snooze(6)

    verify_text(
        english_data,
        names.averagePowerWidget_nameLabel_QLabel,
        "AVERAGE POWER",
        "Average Power label"
    )

    snooze(6)

    verify_text(
        english_data,
        names.averagePowerWidget_unitsOfMeasureLabel_QLabel,
        "W",
        "Average Power unit"
    )

    snooze(6)

    verify_text(
        english_data,
        names.pulseEnergyWidget_nameLabel_QLabel,
        "PULSE ENERGY",
        "Pulse Energy label"
    )

    snooze(6)

    verify_text(
        english_data,
        names.pulseEnergyWidget_unitsOfMeasureLabel_QLabel,
        "J",
        "Pulse Energy unit"
    )

    snooze(6)

    verify_text(
        english_data,
        names.frequencyWidget_nameLabel_QLabel,
        "FREQUENCY",
        "Frequency label"
    )

    snooze(6)

    verify_text(
        english_data,
        names.frequencyWidget_unitsOfMeasureLabel_QLabel,
        "Hz",
        "Frequency unit"
    )

    snooze(6)

    verify_text(
        english_data,
        names.modeButton_nameLabel_QLabel_3,
        "REGULAR PULSE",
        "Regular Pulse label"
    )

    snooze(6)

    verify_text(
        english_data,
        names.peakPowerButton_nameLabel_QLabel_2,
        "PEAK POWER",
        "Peak Power label"
    )

    snooze(6)

    verify_text(
        english_data,
        names.pulseModeWidget_saveButton_QPushButton,
        "ADD",
        "Add button text"
    )


    # ---------------------------------------------------------
    # ENTER PRESET NAME
    # ---------------------------------------------------------

    clickButton(
        waitForObject(
            names.centralWidget_pbName_QPushButton
        )
    )

    snooze(1)

    clickButton(
        waitForObject(
            names.keyboardPopup_p_KeyButton
        )
    )

    clickButton(
        waitForObject(
            names.keyboardPopup_r_KeyButton
        )
    )

    clickButton(
        waitForObject(
            names.keyboardPopup_e_KeyButton
        )
    )

    clickButton(
        waitForObject(
            names.keyboardPopup_s_KeyButton
        )
    )

    clickButton(
        waitForObject(
            names.keyboardPopup_e_KeyButton
        )
    )

    clickButton(
        waitForObject(
            names.keyboardPopup_t_KeyButton
        )
    )

    clickButton(
        waitForObject(
            names.keyboardPopup_KeyButton_5
        )
    )

    clickButton(
        waitForObject(
            names.keyboardPopup_o_KeyButton
        )
    )

    clickButton(
        waitForObject(
            names.keyboardPopup_n_KeyButton
        )
    )

    clickButton(
        waitForObject(
            names.keyboardPopup_e_KeyButton
        )
    )

    mouseClick(
        waitForObject(
            names.keyboardPopup_lbText_QLabel
        ),
        62,
        15,
        Qt.NoModifier,
        Qt.LeftButton
    )

    snooze(2)


    # ---------------------------------------------------------
    # SAVE NEW PRESET
    # ---------------------------------------------------------

    clickButton(
        waitForObject(
            names.pulseModeWidget_saveButton_QPushButton
        )
    )

    snooze(4)


    # ---------------------------------------------------------
    # SELECT PRESET
    # ---------------------------------------------------------

    mouseClick(
        waitForObject(
            names.tabSettingsPresetsUserStonePresets_presetListView_PresetListView
        ),
        416,
        71,
        Qt.NoModifier,
        Qt.LeftButton
    )

    snooze(2)


    # ---------------------------------------------------------
    # EDIT PRESET
    # ---------------------------------------------------------

    clickButton(
        waitForObject(
            names.settingsScreen_editButton_QPushButton
        )
    )

    snooze(3)


    # ---------------------------------------------------------
    # SELECT PRESET IN EDIT SCREEN
    # ---------------------------------------------------------

    mouseClick(
        waitForObject(
            names.tabSettingsPresetsUserStonePresets_presetListView_PresetListView
        ),
        416,
        71,
        Qt.NoModifier,
        Qt.LeftButton
    )

    snooze(2)


    # ---------------------------------------------------------
    # VERIFY EDIT PRESET
    # ---------------------------------------------------------

    snooze(6)

    verify_text(
        english_data,
        names.centralWidget_lbOperation_QLabel,
        "EDIT PRESET",
        "Edit Preset operation text"
    )

    snooze(6)

    verify_text(
        english_data,
        names.averagePowerWidget_nameLabel_QLabel,
        "AVERAGE POWER",
        "Edit Average Power label"
    )

    snooze(6)

    verify_text(
        english_data,
        names.averagePowerWidget_unitsOfMeasureLabel_QLabel,
        "W",
        "Edit Average Power unit"
    )

    snooze(6)

    verify_text(
        english_data,
        names.pulseEnergyWidget_nameLabel_QLabel,
        "PULSE ENERGY",
        "Edit Pulse Energy label"
    )

    snooze(6)

    verify_text(
        english_data,
        names.pulseEnergyWidget_unitsOfMeasureLabel_QLabel,
        "J",
        "Edit Pulse Energy unit"
    )

    snooze(6)

    verify_text(
        english_data,
        names.frequencyWidget_nameLabel_QLabel,
        "FREQUENCY",
        "Edit Frequency label"
    )

    snooze(6)

    verify_text(
        english_data,
        names.frequencyWidget_unitsOfMeasureLabel_QLabel,
        "Hz",
        "Edit Frequency unit"
    )

    snooze(6)

    verify_text(
        english_data,
        names.modeButton_nameLabel_QLabel_3,
        "REGULAR PULSE",
        "Edit Regular Pulse label"
    )

    snooze(6)

    verify_text(
        english_data,
        names.peakPowerButton_nameLabel_QLabel_2,
        "PEAK POWER",
        "Edit Peak Power label"
    )

    snooze(6)

    verify_text(
        english_data,
        names.pulseModeWidget_saveButton_QPushButton,
        "UPDATE",
        "Update button text"
    )


    # ---------------------------------------------------------
    # UPDATE PRESET
    # ---------------------------------------------------------

    clickButton(
        waitForObject(
            names.pulseModeWidget_saveButton_QPushButton
        )
    )

    snooze(4)


    # ---------------------------------------------------------
    # DELETE PRESET
    # ---------------------------------------------------------

    test.log("Deleting preset")

    clickButton(
        waitForObject(
            names.settingsScreen_deletePresetButton_QPushButton
        )
    )

    snooze(3)


    # ---------------------------------------------------------
    # SELECT PRESET FOR DELETE
    # ---------------------------------------------------------

    mouseClick(
        waitForObject(
            names.tabSettingsPresetsUserStonePresets_presetListView_PresetListView
        ),
        489,
        57,
        Qt.NoModifier,
        Qt.LeftButton
    )

    snooze(3)


    # ---------------------------------------------------------
    # VERIFY DELETE DIALOG
    # ---------------------------------------------------------

    snooze(6)

    verify_text(
        english_data,
        names.centralWidget_lbDialogTitle_QLabel,
        "Delete the following preset?",
        "Delete preset dialog title"
    )

    snooze(6)

    verify_text(
        english_data,
        names.infoFrame_extInfoLabel_QLabel,
        "6 J 3 Hz 18 W",
        "Delete preset information"
    )

    snooze(6)

    verify_text(
        english_data,
        names.buttonsFrame_rejectButton_QPushButton,
        "CANCEL",
        "Delete dialog Cancel button"
    )

    snooze(6)

    verify_text(
        english_data,
        names.centralWidget_acceptButton_QPushButton,
        "DELETE",
        "Delete dialog Delete button"
    )


    # ---------------------------------------------------------
    # CONFIRM DELETE
    # ---------------------------------------------------------

    clickButton(
        waitForObject(
            names.centralWidget_acceptButton_QPushButton
        )
    )

    snooze(4)


    # ---------------------------------------------------------
    # CLOSE EXCEL
    # ---------------------------------------------------------

    try:
        workbook.close()
    except Exception:
        pass

    test.log("Add, Edit, Update and Delete preset test completed")
