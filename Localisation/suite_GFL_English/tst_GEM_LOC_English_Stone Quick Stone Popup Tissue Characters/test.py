# -*- coding: utf-8 -*-

import names
import openpyxl
import re
import unicodedata
import html
import os

TESTCASE_DIR = os.path.dirname(os.path.abspath(__file__))

EXCEL_PATH = os.path.abspath(
    os.path.join(
        TESTCASE_DIR,
        "..",
        "shared",
        "testdata",
        "Gemini Split strings.xlsx"
    )
)


def normalize(text):
    if text is None:
        return ""

    text = html.unescape(str(text))
    text = re.sub(r"<br\\s*/?>", " ", text, flags=re.IGNORECASE)
    text = re.sub(r"<[^>]+>", " ", text)
    text = unicodedata.normalize("NFKC", text)

    text = text.replace("\\xa0", " ")
    text = text.replace("\\u2007", " ")
    text = text.replace("\\u202f", " ")
    text = text.replace("\\r", " ")
    text = text.replace("\\n", " ")
    text = text.replace("\\t", " ")

    text = text.replace("™", "")
    text = text.replace("®", "")
    text = text.replace("©", "")

    return " ".join(text.split()).strip().casefold()


def comparison_key(text):
    return re.sub(r"\\s+", "", normalize(text))


def load_english_excel(excel_path):
    if not os.path.isfile(excel_path):
        test.fail("Excel file not found: " + excel_path)
        return None, {}

    workbook = openpyxl.load_workbook(
        excel_path,
        data_only=True,
        read_only=True
    )

    if "Treatment Modes" not in workbook.sheetnames:
        workbook.close()
        test.fail(
            "Worksheet 'Treatment Modes' does not exist in Excel file"
        )
        return None, {}

    sheet = workbook["Treatment Modes"]
    english_column = 2
    english_data = {}

    for row in sheet.iter_rows(
        min_col=english_column,
        max_col=english_column,
        values_only=True
    ):
        english_value = row[0]

        if english_value is None:
            continue

        original_text = str(english_value).strip()
        key = comparison_key(original_text)

        if key and original_text:
            english_data.setdefault(key, original_text)

    return workbook, english_data


def get_english_text(english_data, expected_english):
    return english_data.get(comparison_key(expected_english), "")


def verify_text(english_data, object_name, expected_english):
    try:
        obj = waitForObjectExists(object_name)
        actual_text = str(obj.text).strip()
    except Exception as e:
        test.fail(
            "Unable to read application object: "
            + str(object_name)
            + " | Error: "
            + str(e)
        )
        return False

    excel_english_text = get_english_text(
        english_data,
        expected_english
    )

    if not excel_english_text:
        test.fail(
            "English text not found in Excel column B: "
            + str(expected_english)
        )
        return False

    actual_normalized = comparison_key(actual_text)
    expected_normalized = comparison_key(excel_english_text)

    test.compare(
        actual_normalized,
        expected_normalized,
        "English text verification: " + str(expected_english)
    )

    return actual_normalized == expected_normalized


def main():
    workbook = None

    try:
        clickButton(
            waitForObject(names.numpad_pushButton_8_NumpadButton)
        )
        clickButton(
            waitForObject(names.numpad_pushButton_8_NumpadButton)
        )
        clickButton(
            waitForObject(names.numpad_pushButton_8_NumpadButton)
        )
        clickButton(
            waitForObject(names.numpad_pushButton_8_NumpadButton)
        )

        clickButton(
            waitForObject(names.mainFrame_okButton_QPushButton)
        )
        clickButton(
            waitForObject(names.homeScreen_settingsButton_QPushButton)
        )

        clickTab(
            waitForObject(names.settingsScreen_tbSystemInfo_TabWidget),
            "INFO"
        )

        clickButton(
            waitForObject(
                names.gbSystemInformation_pbLanguageSettingsUpdate_QPushButton
            )
        )

        # mouseClick(waitForObject(names.centralWidget_buttonsFrame_QFrame), 418, 35, Qt.NoModifier, Qt.LeftButton)
        # clickButton(waitForObject(names.buttonsFrame_acceptButton_QPushButton))
        # clickButton(waitForObject(names.buttonsBar_homeButton_QPushButton))
        # clickButton(waitForObject(names.homeScreen_pbStoneAssistant_QToolButton))
        # clickButton(waitForObject(names.locationFrame_kidneyPcnlLocationButton_QPushButton))
        # clickButton(waitForObject(names.flowRateFrame_lowFlowRateButton_QPushButton))
        # clickButton(waitForObject(names.hardnessFrame_smallHardnessButton_QPushButton))
        # clickButton(waitForObject(names.stoneAssistantModeScreen_continueButton_QPushButton))
        # clickButton(waitForObject(names.treatmentCharacteristicsBar_pbEdit_QPushButton_2))
        # test.compare(str(waitForObjectExists(names.centralWidget_titleLabel_QLabel).text), "STONE CHARACTERISTICS")
        # test.compare(str(waitForObjectExists(names.centralWidget_titleLabel_QLabel_2).text), "LOCATION")
        # test.compare(str(waitForObjectExists(names.centralWidget_valueLabel_QLabel).text), "KIDNEY (PCNL)")
        # test.compare(str(waitForObjectExists(names.centralWidget_titleLabel_QLabel_3).text), "FLOW RATE")
        # test.compare(str(waitForObjectExists(names.centralWidget_valueLabel_QLabel_2).text), "LOW")
        # test.compare(str(waitForObjectExists(names.centralWidget_titleLabel_QLabel_4).text), "HARDNESS")
        # test.compare(str(waitForObjectExists(names.centralWidget_descriptionLabel_QLabel).text), "Hounsfield Scale <500")
        # test.compare(str(waitForObjectExists(names.centralWidget_valueLabel_QLabel_3).text), "Soft")
        # test.compare(str(waitForObjectExists(names.centralWidget_rejectButton_QPushButton).text), "EDIT")
        # test.compare(str(waitForObjectExists(names.centralWidget_acceptButton_QPushButton).text), "OK")

       
        clickButton(
            waitForObject(names.buttonsFrame_acceptButton_QPushButton)
        )
        clickButton(
            waitForObject(names.buttonsBar_homeButton_QPushButton)
        )

        clickButton(
            waitForObject(names.homeScreen_pbStoneAssistant_QToolButton)
        )
        clickButton(
            waitForObject(
                names.locationFrame_kidneyPcnlLocationButton_QPushButton
            )
        )
        clickButton(
            waitForObject(
                names.flowRateFrame_lowFlowRateButton_QPushButton
            )
        )
        clickButton(
            waitForObject(
                names.hardnessFrame_smallHardnessButton_QPushButton
            )
        )
        clickButton(
            waitForObject(
                names.stoneAssistantModeScreen_continueButton_QPushButton
            )
        )
        clickButton(
            waitForObject(
                names.treatmentCharacteristicsBar_pbEdit_QPushButton_2
            )
        )

        workbook, english_data = load_english_excel(EXCEL_PATH)

        if workbook is None:
            return

        verify_text(
            english_data,
            names.centralWidget_titleLabel_QLabel,
            "STONE CHARACTERISTICS"
        )
        verify_text(
            english_data,
            names.centralWidget_titleLabel_QLabel_2,
            "LOCATION"
        )
        verify_text(
            english_data,
            names.centralWidget_valueLabel_QLabel,
            "KIDNEY (PCNL)"
        )
        verify_text(
            english_data,
            names.centralWidget_titleLabel_QLabel_3,
            "FLOW RATE"
        )
        verify_text(
            english_data,
            names.centralWidget_valueLabel_QLabel_2,
            "LOW"
        )
        verify_text(
            english_data,
            names.centralWidget_titleLabel_QLabel_4,
            "HARDNESS"
        )
        verify_text(
            english_data,
            names.centralWidget_descriptionLabel_QLabel,
            "Hounsfield Scale < 500"
        )
        verify_text(
            english_data,
            names.centralWidget_valueLabel_QLabel_3,
            "SOFT"
        )
        verify_text(
            english_data,
            names.centralWidget_rejectButton_QPushButton,
            "EDIT"
        )
        verify_text(
            english_data,
            names.centralWidget_acceptButton_QPushButton,
            "OK"
        )

    finally:
        if workbook is not None:
            workbook.close()
