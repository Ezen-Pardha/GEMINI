
# -*- coding: utf-8 -*-

import names
import openpyxl
import re
from openpyxl import load_workbook
import os

TESTCASE_DIR = os.path.dirname(os.path.abspath(__file__))
 
EXCEL_PATH = os.path.abspath(
    os.path.join( TESTCASE_DIR,
        "..",
        "shared",
        "testdata",
        "Gemini Split strings.xlsx"
    )
)

# EXCEL_PATH = (
#     "/home/ezen/Test_Automation_Gemini/Localisation/"
#     "suite_GFL_English/"
#     "tst_GEM-LOC-English-PeakPower-SoftTissueCharacters/"
#     "testdata/Gemini Split strings.xlsx"
# )



ENGLISH_COLUMN = 2

SHEET_NAME = "Soft tissue"


def normalize(text):

    if text is None:
        return ""

    text = str(text)

    text = re.sub(
        r"<[^>]+>",
        " ",
        text
    )

    text = " ".join(
        text.split()
    )

    text = re.sub(
        r"\s*<\s*",
        " < ",
        text
    )

    text = re.sub(
        r"\s*>\s*",
        " > ",
        text
    )

    return text.strip().casefold()


def load_excel():

    try:

        
        workbook = openpyxl.load_workbook(
            EXCEL_PATH,
            data_only=True,
            read_only=True
        )

        if SHEET_NAME not in workbook.sheetnames:

            test.fail(
                "Excel sheet not found: "
                + SHEET_NAME
                + " | Available sheets: "
                + str(workbook.sheetnames)
            )

            workbook.close()

            return None, None

        sheet = workbook[SHEET_NAME]

       

        return workbook, sheet

    except Exception as e:

        test.fail(
            "Unable to load English localization Excel: "
            + str(e)
        )

        return None, None


def get_english_value(
    sheet,
    english_text
):

    expected_english = normalize(
        english_text
    )

    if not expected_english:
        return None

    for row in sheet.iter_rows(
        min_row=2,
        min_col=ENGLISH_COLUMN,
        max_col=ENGLISH_COLUMN
    ):

        english_value = row[0].value

        if english_value is None:
            continue

        english_value = str(
            english_value
        ).strip()

        if not english_value:
            continue

        if normalize(
            english_value
        ) == expected_english:

            return english_value

    return None


def verify_text(
    sheet,
    object_name,
    english_text,
    description
):

    snooze(6)

    try:

        obj = waitForObjectExists(
            object_name
        )

    except Exception as e:

        test.fail(
            description
            + " | Unable to find object: "
            + str(e)
        )

        return False

    try:

        actual_value = obj.text

        actual = ""

        if actual_value is not None:

            actual = str(
                actual_value
            ).strip()

    except Exception as e:

        test.fail(
            description
            + " | Unable to read object text: "
            + str(e)
        )

        return False

    if not actual:

        test.fail(
            description
            + " | UI text is empty"
        )

        return False

    expected_value = get_english_value(
        sheet,
        english_text
    )

    if expected_value is None:

        test.fail(
            description
            + " | English string not found in Column B: "
            + english_text
        )

        return False

    actual_normalized = normalize(
        actual
    )

    expected_normalized = normalize(
        expected_value
    )

    test.compare(
        actual_normalized,
        expected_normalized,
        description
    )

    if actual_normalized == expected_normalized:

        # test.log(
        #     description
        #     + " | PASS"
        #     + " | English: "
        #     + expected_value
        # )

        return True

    else:

        test.fail(
            description
            + " | Expected English value: "
            + expected_value
            + " | Actual: "
            + actual
        )

        return False


def get_central_widget():

    try:

        central_widget = findObject(
            {
                "name": "centralWidget",
                "type": "QFrame",
                "window": ":o_ScreenSwitcher"
            }
        )

        for i in range(20):

            try:

                if central_widget.visible:

                    test.log(
                        "centralWidget is visible"
                    )

                    return central_widget

            except:

                pass

            snooze(0.5)

        test.log(
            "centralWidget was found but is not visible"
        )

        return None

    except Exception as e:

        test.log(
            "centralWidget not found: "
            + str(e)
        )

        return None


def main():

    workbook, sheet = load_excel()

    if workbook is None or sheet is None:

        return


    

    clickButton(
        waitForObject(
            names.numpad_pushButton_8_NumpadButton
        )
    )

    snooze(2)

    doubleClick(
        waitForObject(
            names.numpad_pushButton_8_NumpadButton
        ),
        94,
        77,
        Qt.NoModifier,
        Qt.LeftButton
    )

    snooze(2)

    doubleClick(
        waitForObject(
            names.numpad_pushButton_8_NumpadButton
        ),
        94,
        77,
        Qt.NoModifier,
        Qt.LeftButton
    )

    snooze(2)

    clickButton(
        waitForObject(
            names.mainFrame_okButton_QPushButton
        )
    )


    snooze(6)

    clickButton(
        waitForObject(
            names.homeScreen_pbSoftTissueQuickStart_QToolButton
        )
    )


    snooze(6)

    mouseClick(
        waitForObjectItem(
            names.tabLeftMainPresets_leftPresetsView_PresetListView,
            "_1"
        ),
        377,
        65,
        Qt.NoModifier,
        Qt.LeftButton
    )


    snooze(6)

    mouseClick(
        waitForObjectItem(
            names.tabRightMainPresets_rightPresetsView_PresetListView,
            "_1"
        ),
        150,
        37,
        Qt.NoModifier,
        Qt.LeftButton
    )


    snooze(6)

    sendEvent(
        "QMoveEvent",
        waitForObject(
            names.o_ScreenSwitcher
        ),
        0,
        -50,
        1288,
        292
    )


    snooze(6)

    clickButton(
        waitForObject(
            names.quickStartScreen_pbContinue_QPushButton
        )
    )

    snooze(6)

    clickButton(
        waitForObject(
            names.mainFrame_okButton_QPushButton_2
        )
    )


    snooze(6)

    mouseClick(
        waitForObject(
            names.modeButton_nameLabel_QLabel
        ),
        115,
        15,
        Qt.NoModifier,
        Qt.LeftButton
    )


    verify_text(
        sheet,
        names.centralWidget_TitleLabel_QLabel_2,
        "EXIT QUICK START MODE?",
        "Quick Start Exit Title"
    )


    verify_text(
        sheet,
        names.centralWidget_MessageLabel_QLabel_2,
        "Please confirm that you want to exit quick start mode.",
        "Quick Start Exit Message"
    )


    verify_text(
        sheet,
        names.centralWidget_CancelButton_QPushButton_3,
        "CANCEL",
        "Quick Start Cancel Button"
    )

    verify_text(
        sheet,
        names.centralWidget_DisableButton_QPushButton_2,
        "CONFIRM",
        "Quick Start Confirm Button"
    )

    snooze(6)

    clickButton(
        waitForObject(
            names.centralWidget_DisableButton_QPushButton_2
        )
    )

    snooze(6)

    clickButton(
        waitForObject(
            names.buttonsBar_homeButton_QPushButton_2
        )
    )


    snooze(6)

    clickButton(
        waitForObject(
            names.homeScreen_pbSoftTissueAssistant_QToolButton
        )
    )


    snooze(6)

    clickButton(
        waitForObject(
            names.gbSoftTissueProcedures_pbIncision_QPushButton
        )
    )

    snooze(6)

    clickButton(
        waitForObject(
            names.softTissueAssistantModeScreen_pbContinue_QPushButton
        )
    )


    snooze(6)

    mouseDrag(
        waitForObject(
            names.indexPicker_valuePicker_IntervalSlider
        ),
        195,
        15,
        211,
        29,
        1,
        Qt.LeftButton
    )

    verify_text(
        sheet,
        names.centralWidget_TitleLabel_QLabel_2,
        "EXIT GUIDED MODE?",
        "Guided Mode Exit Title"
    )

    verify_text(
        sheet,
        names.centralWidget_MessageLabel_QLabel_2,
        "Please confirm that you want to exit guided mode.",
        "Guided Mode Exit Message"
    )

    verify_text(
        sheet,
        names.centralWidget_CancelButton_QPushButton_3,
        "CANCEL",
        "Guided Mode Cancel Button"
    )


    verify_text(
        sheet,
        names.centralWidget_DisableButton_QPushButton_2,
        "CONFIRM",
        "Guided Mode Confirm Button"
    )


    snooze(6)

    clickButton(
        waitForObject(
            names.centralWidget_DisableButton_QPushButton_2
        )
    )

    try:

        workbook.close()

    except:

        pass


    
