# -*- coding: utf-8 -*-

import names
import re
import unicodedata
from openpyxl import load_workbook
import os 
TESTCASE_DIR = os.path.dirname(os.path.abspath(__file__))
 
EXCEL_PATH = os.path.abspath(
    os.path.join( TESTCASE_DIR,
        "..",
        "shared",
        "testdata",
        "Gemini Split strings.xlsx"
    )
)





EXCEL_SHEET_NAME = "Treatment Modes"

ENGLISH_COLUMN = 2
JAPAN_COLUMN = 15


def normalize_text(value):

    if value is None:
        return ""

    text = str(value)

    text = re.sub(
        r"<[^>]*>",
        "",
        text
    )

    text = re.sub(
        r"<br\s*/?>",
        "",
        text,
        flags=re.IGNORECASE
    )

    try:
        text = unicodedata.normalize(
            "NFKC",
            text
        )
    except Exception:
        pass

    text = (
        text
        .replace("\n", "")
        .replace("\r", "")
        .replace("\t", "")
        .replace("\xa0", "")
        .replace("\u2007", "")
        .replace("\u202f", "")
        .replace("™", "")
        .replace("®", "")
        .replace("©", "")
    )

    text = re.sub(
        r"\s+",
        "",
        text
    )

    return text.casefold()


def get_japan_texts(sheet, english_text):

    search_text = normalize_text(
        english_text
    )

    japan_texts = []

    for row in sheet.iter_rows(
        min_row=1,
        min_col=ENGLISH_COLUMN,
        max_col=JAPAN_COLUMN,
        values_only=True
    ):

        english_value = row[0]

        if normalize_text(english_value) != search_text:
            continue

        japan_value = row[
            JAPAN_COLUMN - ENGLISH_COLUMN
        ]

        if japan_value is None:
            continue

        japan_text = str(japan_value).strip()

        if japan_text == "":
            continue

        japan_normalized = normalize_text(
            japan_text
        )

        if japan_normalized == "":
            continue

        duplicate = False

        for existing_text in japan_texts:

            if normalize_text(existing_text) == japan_normalized:
                duplicate = True
                break

        if not duplicate:
            japan_texts.append(japan_text)

    if len(japan_texts) == 0:
        test.fail(
            "English text not found or Japan translation "
            "is empty in Excel: "
            + english_text
        )

    return japan_texts


def get_actual_text(obj, description):

    actual_value = None

    try:
        actual_value = obj.text
    except Exception:
        actual_value = None

    if actual_value is None:
        try:
            actual_value = obj.property("text")
        except Exception:
            actual_value = None

    if actual_value is None:
        try:
            actual_value = obj.windowTitle
        except Exception:
            actual_value = None

    if actual_value is None:
        try:
            actual_value = obj.accessibleName
        except Exception:
            actual_value = None

    actual_text = normalize_text(
        actual_value
    )

    if actual_text == "":
        test.fail(
            "Application text is empty or None for: "
            + description
        )
        return None

    return actual_text


def verify_text(sheet, object_name, english_text, description):

    snooze(6)

    expected_texts = get_japan_texts(
        sheet,
        english_text
    )

    if len(expected_texts) == 0:
        return

    expected_normalized_values = []

    for expected_text in expected_texts:

        normalized_value = normalize_text(
            expected_text
        )

        if normalized_value != "":
            expected_normalized_values.append(
                normalized_value
            )

    if len(expected_normalized_values) == 0:
        test.fail(
            "Expected Japan text is empty for: "
            + description
        )
        return

    try:
        obj = waitForObjectExists(
            getattr(names, object_name)
        )
    except Exception as e:
        test.fail(
            description
            + " | Object not found: "
            + object_name
            + " | "
            + str(e)
        )
        return

    actual_text = get_actual_text(
        obj,
        description
    )

    if actual_text is None:
        return

    for expected_normalized in expected_normalized_values:

        if actual_text == expected_normalized:
            test.compare(
                actual_text,
                expected_normalized,
                description + " - Japan text verification"
            )
            return

    test.fail(
        description
        + " - Japan text mismatch | Expected: "
        + " / ".join(expected_texts)
        + " | Actual: "
        + actual_text
    )


def verify_left_manual_tab(sheet, english_text, description):

    snooze(6)

    expected_texts = get_japan_texts(
        sheet,
        english_text
    )

    if len(expected_texts) == 0:
        return

    expected_normalized_values = []

    for expected_text in expected_texts:

        normalized_value = normalize_text(
            expected_text
        )

        if normalized_value != "":
            expected_normalized_values.append(
                normalized_value
            )

    if len(expected_normalized_values) == 0:
        test.fail(
            "Expected Japan text is empty for: "
            + description
        )
        return

    try:
        manual_tab_object = getattr(
            names,
            "mANUEL_TabItem"
        ).copy()
    except Exception as e:
        test.fail(
            description
            + " | Manual tab object definition not found | "
            + str(e)
        )
        return

    if "text" in manual_tab_object:
        del manual_tab_object["text"]

    try:
        obj = waitForObjectExists(
            manual_tab_object
        )
    except Exception as e:
        test.fail(
            description
            + " | Manual tab object not found | "
            + str(e)
        )
        return

    actual_text = get_actual_text(
        obj,
        description
    )

    if actual_text is None:
        return

    for expected_normalized in expected_normalized_values:

        if actual_text == expected_normalized:
            test.compare(
                actual_text,
                expected_normalized,
                description + " - Japan text verification"
            )
            return

    test.fail(
        description
        + " - Japan text mismatch | Expected: "
        + " / ".join(expected_texts)
        + " | Actual: "
        + actual_text
    )


def verify_right_manual_tab(sheet, english_text, description):

    snooze(6)

    expected_texts = get_japan_texts(
        sheet,
        english_text
    )

    if len(expected_texts) == 0:
        return

    expected_normalized_values = []

    for expected_text in expected_texts:

        normalized_value = normalize_text(
            expected_text
        )

        if normalized_value != "":
            expected_normalized_values.append(
                normalized_value
            )

    if len(expected_normalized_values) == 0:
        test.fail(
            "Expected Japan text is empty for: "
            + description
        )
        return

    try:
        manual_tab_object = getattr(
            names,
            "mANUEL_TabItem_2"
        ).copy()
    except Exception as e:
        test.fail(
            description
            + " | Manual tab object definition not found | "
            + str(e)
        )
        return

    if "text" in manual_tab_object:
        del manual_tab_object["text"]

    try:
        obj = waitForObjectExists(
            manual_tab_object
        )
    except Exception as e:
        test.fail(
            description
            + " | Manual tab object not found | "
            + str(e)
        )
        return

    actual_text = get_actual_text(
        obj,
        description
    )

    if actual_text is None:
        return

    for expected_normalized in expected_normalized_values:

        if actual_text == expected_normalized:
            test.compare(
                actual_text,
                expected_normalized,
                description + " - Japan text verification"
            )
            return

    test.fail(
        description
        + " - Japan text mismatch | Expected: "
        + " / ".join(expected_texts)
        + " | Actual: "
        + actual_text
    )


def main():

    workbook = None

    try:
        workbook = load_workbook(
            EXCEL_PATH,
            read_only=True,
            data_only=True
        )

        if EXCEL_SHEET_NAME not in workbook.sheetnames:
            test.fail(
                "Excel sheet not found: "
                + EXCEL_SHEET_NAME
                + " | Available sheets: "
                + ", ".join(workbook.sheetnames)
            )
            return

        sheet = workbook[EXCEL_SHEET_NAME]

        doubleClick(
            waitForObject(
                names.numpad_pushButton_8_NumpadButton,
                100461
            ),
            96,
            56,
            Qt.NoModifier,
            Qt.LeftButton
        )

        snooze(1)

        doubleClick(
            waitForObject(
                names.numpad_pushButton_8_NumpadButton
            ),
            96,
            56,
            Qt.NoModifier,
            Qt.LeftButton
        )

        snooze(1)

        clickButton(
            waitForObject(
                names.mainFrame_okButton_QPushButton
            )
        )

        snooze(2)

        clickButton(
            waitForObject(
                names.homeScreen_settingsButton_QPushButton
            )
        )

        snooze(2)

        clickTab(
            waitForObject(
                names.settingsScreen_tbSystemInfo_TabWidget
            ),
            "INFO"
        )

        snooze(2)

        clickButton(
            waitForObject(
                names.gbSystemInformation_pbLanguageSettingsUpdate_QPushButton
            )
        )

        snooze(2)

        clickButton(
            waitForObject(
                names.languagesFrame_pbJapan_QPushButton
            )
        )

        snooze(1)

        clickButton(
            waitForObject(
                names.buttonsFrame_acceptButton_QPushButton
            )
        )

        snooze(3)

        clickButton(
            waitForObject(
                names.buttonsBar_homeButton_QPushButton
            )
        )

        snooze(2)

        clickButton(
            waitForObject(
                names.gbExpert_expertButton_QPushButton
            )
        )

        snooze(3)

        verify_left_manual_tab(
            sheet,
            "MANUAL",
            "Left Manual Tab"
        )

        verify_text(
            sheet,
            "averagePowerWidget_nameLabel_QLabel_2",
            "AVERAGE POWER",
            "Left Average Power"
        )

        verify_text(
            sheet,
            "averagePowerWidget_unitsOfMeasureLabel_QLabel_2",
            "W",
            "Left Power Unit"
        )

        verify_text(
            sheet,
            "pulseEnergyWidget_nameLabel_QLabel_2",
            "PULSE ENERGY",
            "Left Pulse Energy"
        )

        verify_text(
            sheet,
            "pulseEnergyWidget_unitsOfMeasureLabel_QLabel_2",
            "J",
            "Left Energy Unit"
        )

        verify_text(
            sheet,
            "frequencyWidget_nameLabel_QLabel_2",
            "FREQUENCY",
            "Left Frequency"
        )

        verify_text(
            sheet,
            "frequencyWidget_unitsOfMeasureLabel_QLabel_2",
            "Hz",
            "Left Frequency Unit"
        )

        verify_text(
            sheet,
            "modeButton_nameLabel_QLabel",
            "REGULAR PULSE",
            "Left Pulse Mode"
        )

        verify_text(
            sheet,
            "peakPowerButton_nameLabel_QLabel",
            "PEAK POWER",
            "Left Peak Power"
        )

        verify_text(
            sheet,
            "pulseModeWidget_saveButton_QPushButton_2",
            "SAVE",
            "Left Save Button"
        )

        verify_text(
            sheet,
            "totalEnergyFrame_nameLabel_QLabel",
            "TOTAL ENERGY",
            "Left Total Energy"
        )

        verify_text(
            sheet,
            "totalTimeFrame_nameLabel_QLabel",
            "TOTAL LASING TIME",
            "Left Total Lasing Time"
        )

        verify_right_manual_tab(
            sheet,
            "MANUAL",
            "Right Manual Tab"
        )

        verify_text(
            sheet,
            "averagePowerWidget_nameLabel_QLabel_3",
            "AVERAGE POWER",
            "Right Average Power"
        )

        verify_text(
            sheet,
            "averagePowerWidget_unitsOfMeasureLabel_QLabel_3",
            "W",
            "Right Power Unit"
        )

        verify_text(
            sheet,
            "pulseEnergyWidget_nameLabel_QLabel_3",
            "PULSE ENERGY",
            "Right Pulse Energy"
        )

        verify_text(
            sheet,
            "pulseEnergyWidget_unitsOfMeasureLabel_QLabel_3",
            "J",
            "Right Energy Unit"
        )

        verify_text(
            sheet,
            "frequencyWidget_nameLabel_QLabel_3",
            "FREQUENCY",
            "Right Frequency"
        )

        verify_text(
            sheet,
            "pulseModeWidget_saveButton_QPushButton_3",
            "SAVE",
            "Right Save Button"
        )

        verify_text(
            sheet,
            "peakPowerButton_nameLabel_QLabel_3",
            "PEAK POWER",
            "Right Peak Power"
        )

        verify_text(
            sheet,
            "modeButton_nameLabel_QLabel_2",
            "REGULAR PULSE",
            "Right Pulse Mode"
        )

        verify_text(
            sheet,
            "stateSwitch_readyButton_QPushButton",
            "READY",
            "Ready Button"
        )

        verify_text(
            sheet,
            "stateSwitch_standbyButton_QPushButton",
            "STANDBY",
            "Standby Button"
        )

        verify_text(
            sheet,
            "fiberFrame_nameLabel_QLabel",
            "FIBER",
            "Fiber"
        )

        verify_text(
            sheet,
            "fiberUsesFrame_nameLabel_QLabel",
            "FIBER USES REMAINING",
            "Fiber Uses Remaining"
        )

        test.vp("ManualStandby")

        verify_text(
            sheet,
            "bottomBar_aimingBeamButton_AimingBeamModeButton",
            "AIMING BEAM",
            "Aiming Beam"
        )

    except Exception as e:
        test.fail(
            "Japan Regular Pulse localization test failed: "
            + str(e)
        )

    finally:
        if workbook is not None:
            try:
                workbook.close()
            except Exception:
                pass
