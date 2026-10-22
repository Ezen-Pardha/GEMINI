# -*- coding: utf-8 -*-

import names
import openpyxl
import re
import unicodedata
import os 
TESTCASE_DIR = os.path.dirname(os.path.abspath(__file__))
 
EXCEL_PATH = os.path.abspath(
    os.path.join( TESTCASE_DIR,
        "..",
        "shared",
        "testdata",
        "Gemini Split strings.xlsx"
    )
)


ENGLISH_COLUMN = 2
JAPAN_COLUMN = 15
SHEET_NAME = "Soft tissue"


def normalize(text):

    if text is None:
        return ""

    text = str(text)

    text = re.sub(
        r"<[^>]+>",
        "",
        text
    )

    try:
        text = unicodedata.normalize(
            "NFKC",
            text
        )
    except Exception:
        pass

    text = (
        text
        .replace("\n", "")
        .replace("\r", "")
        .replace("\t", "")
        .replace("\xa0", "")
        .replace("\u2007", "")
        .replace("\u202f", "")
        .replace("™", "")
        .replace("®", "")
        .replace("©", "")
    )

    text = re.sub(
        r"\s+",
        "",
        text
    )

    return text.casefold()


def load_excel():

    try:

        workbook = openpyxl.load_workbook(
            EXCEL_PATH,
            data_only=True,
            read_only=True
        )

        if SHEET_NAME not in workbook.sheetnames:

            test.fail(
                "Excel sheet not found: "
                + SHEET_NAME
                + " | Available sheets: "
                + ", ".join(workbook.sheetnames)
            )

            workbook.close()
            return None, None

        sheet = workbook[SHEET_NAME]

        return workbook, sheet

    except Exception as e:

        test.fail(
            "Unable to load Japan localization Excel: "
            + str(e)
        )

        return None, None


def get_japan_values(sheet, english_text):

    expected_english = normalize(
        english_text
    )

    japan_values = []

    japan_index = JAPAN_COLUMN - ENGLISH_COLUMN

    for row in sheet.iter_rows(
        min_row=2,
        min_col=ENGLISH_COLUMN,
        max_col=JAPAN_COLUMN,
        values_only=True
    ):

        english_value = row[0]

        if english_value is None:
            continue

        if normalize(english_value) != expected_english:
            continue

        if japan_index >= len(row):
            continue

        japan_value = row[japan_index]

        if japan_value is None:
            continue

        japan_value = str(japan_value).strip()

        if not japan_value:
            continue

        normalized_japan = normalize(
            japan_value
        )

        duplicate = any(
            normalize(existing_value) == normalized_japan
            for existing_value in japan_values
        )

        if not duplicate:
            japan_values.append(japan_value)

    return japan_values


def get_object_text(obj):

    for attribute in (
        "text",
        "plainText",
        "title",
        "windowTitle",
        "accessibleName"
    ):

        try:

            value = getattr(
                obj,
                attribute
            )

            if value is not None:

                value = str(value)

                if value.strip():
                    return value

        except Exception:
            pass

    try:

        value = obj.property("text")

        if value is not None:

            value = str(value)

            if value.strip():
                return value

    except Exception:
        pass

    return None


def verify_text(
        sheet,
        object_name,
        english_text,
        description
        ):

    snooze(2)

    try:

        obj = waitForObjectExists(
            object_name
        )

    except Exception as e:

        test.fail(
            description
            + " | Unable to find object: "
            + str(e)
        )

        return False

    actual_value = get_object_text(obj)

    if actual_value is None:

        test.fail(
            description
            + " | Application text is None"
        )

        return False

    actual = str(actual_value).strip()

    if not actual:

        test.fail(
            description
            + " | Application text is empty"
        )

        return False

    expected_values = get_japan_values(
        sheet,
        english_text
    )

    if not expected_values:

        test.fail(
            description
            + " | English string not found in Column B: "
            + english_text
        )

        return False

    actual_normalized = normalize(actual)

    for expected in expected_values:

        expected_normalized = normalize(
            expected
        )

        if actual_normalized == expected_normalized:

            test.compare(
                actual_normalized,
                expected_normalized,
                description + " - Japan text verification"
            )

            return True

    test.fail(
        description
        + " | Expected Japan value from Column O: "
        + " / ".join(expected_values)
        + " | Actual: "
        + actual
    )

    return False


def main():

    workbook, sheet = load_excel()

    if workbook is None or sheet is None:
        return

    try:

        clickButton(
            waitForObject(
                names.numpad_pushButton_8_NumpadButton
            )
        )

        doubleClick(
            waitForObject(
                names.numpad_pushButton_8_NumpadButton
            ),
            94,
            77,
            Qt.NoModifier,
            Qt.LeftButton
        )

        doubleClick(
            waitForObject(
                names.numpad_pushButton_8_NumpadButton
            ),
            94,
            77,
            Qt.NoModifier,
            Qt.LeftButton
        )

        clickButton(
            waitForObject(
                names.mainFrame_okButton_QPushButton
            )
        )

        clickButton(
            waitForObject(
                names.homeScreen_settingsButton_QPushButton
            )
        )

        clickTab(
            waitForObject(
                names.settingsScreen_tbSystemInfo_TabWidget
            ),
            "INFO"
        )

        clickButton(
            waitForObject(
                names.gbSystemInformation_pbLanguageSettingsUpdate_QPushButton
            )
        )

        clickButton(
            waitForObject(
                names.languagesFrame_pbJapan_QPushButton
            )
        )

        clickButton(
            waitForObject(
                names.buttonsFrame_acceptButton_QPushButton
            )
        )

        clickButton(
            waitForObject(
                names.buttonsBar_homeButton_QPushButton
            )
        )

        clickButton(
            waitForObject(
                names.homeScreen_pbSoftTissueQuickStart_QToolButton
            )
        )

        mouseClick(
            waitForObjectItem(
                names.tabLeftMainPresets_leftPresetsView_PresetListView,
                "_1"
            ),
            377,
            65,
            Qt.NoModifier,
            Qt.LeftButton
        )

        mouseClick(
            waitForObjectItem(
                names.tabRightMainPresets_rightPresetsView_PresetListView,
                "_1"
            ),
            150,
            37,
            Qt.NoModifier,
            Qt.LeftButton
        )

        sendEvent(
            "QMoveEvent",
            waitForObject(
                names.o_ScreenSwitcher
            ),
            0,
            -50,
            1288,
            292
        )

        clickButton(
            waitForObject(
                names.quickStartScreen_pbContinue_QPushButton
            )
        )

        clickButton(
            waitForObject(
                names.mainFrame_okButton_QPushButton_2
            )
        )

        mouseClick(
            waitForObject(
                names.modeButton_nameLabel_QLabel
            ),
            115,
            15,
            Qt.NoModifier,
            Qt.LeftButton
        )

        verify_text(
            sheet,
            names.centralWidget_TitleLabel_QLabel_2,
            "EXIT QUICK START MODE?",
            "Quick Start Exit Title"
        )

        verify_text(
            sheet,
            names.centralWidget_MessageLabel_QLabel_2,
            "Please confirm that you want to exit quick start mode.",
            "Quick Start Exit Message"
        )

        verify_text(
            sheet,
            names.centralWidget_CancelButton_QPushButton_3,
            "CANCEL",
            "Quick Start Cancel Button"
        )

        verify_text(
            sheet,
            names.centralWidget_DisableButton_QPushButton_2,
            "CONFIRM",
            "Quick Start Confirm Button"
        )

        clickButton(
            waitForObject(
                names.centralWidget_DisableButton_QPushButton_2
            )
        )

        clickButton(
            waitForObject(
                names.buttonsBar_homeButton_QPushButton_2
            )
        )

        clickButton(
            waitForObject(
                names.homeScreen_pbSoftTissueAssistant_QToolButton
            )
        )

        clickButton(
            waitForObject(
                names.gbSoftTissueProcedures_pbIncision_QPushButton
            )
        )

        clickButton(
            waitForObject(
                names.softTissueAssistantModeScreen_pbContinue_QPushButton
            )
        )

        mouseDrag(
            waitForObject(
                names.indexPicker_valuePicker_IntervalSlider
            ),
            195,
            15,
            211,
            29,
            1,
            Qt.LeftButton
        )

        verify_text(
            sheet,
            names.centralWidget_TitleLabel_QLabel_2,
            "EXIT GUIDED MODE?",
            "Guided Mode Exit Title"
        )

        verify_text(
            sheet,
            names.centralWidget_MessageLabel_QLabel_2,
            "Please confirm that you want to exit guided mode.",
            "Guided Mode Exit Message"
        )

        verify_text(
            sheet,
            names.centralWidget_CancelButton_QPushButton_3,
            "CANCEL",
            "Guided Mode Cancel Button"
        )

        verify_text(
            sheet,
            names.centralWidget_DisableButton_QPushButton_2,
            "CONFIRM",
            "Guided Mode Confirm Button"
        )

        test.vp("Exit Guided Mode")

        clickButton(
            waitForObject(
                names.centralWidget_DisableButton_QPushButton_2
            )
        )

    except Exception as e:

        test.fail(
            "Japan localization test failed: "
            + str(e)
        )

    finally:

        try:
            workbook.close()
        except Exception:
            pass
