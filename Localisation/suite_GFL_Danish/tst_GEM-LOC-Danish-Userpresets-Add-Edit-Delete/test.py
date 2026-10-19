# -*- coding: utf-8 -*-

import names
import openpyxl
import re
import unicodedata
from openpyxl import load_workbook
import os

TESTCASE_DIR = os.path.dirname(os.path.abspath(__file__))
 
EXCEL_PATH = os.path.abspath(
    os.path.join( TESTCASE_DIR,
        "..",
        "shared",
        "testdata",
        "Gemini Split strings.xlsx"
    )
)

# EXCEL_PATH = (
# "/home/ezen/Test_Automation_Gemini/Localisation/"
# "suite_GFL_Danish/"
# "tst_GEM-LOC-Danish-PeakPower-SoftTissueCharacters/"
# "testdata/Gemini Split strings.xlsx"
# )

SHEET_NAME = "Treatment Modes"
LANGUAGE_NAME = "Danish"



def normalize_text(value):


    if value is None:
        return ""
    
    value = str(value)
    
    # Remove HTML/XML tags
    value = re.sub(r"<[^>]+>", "", value)
    
    # Replace special spaces
    value = (
        value.replace("\u00A0", " ")
             .replace("\u2007", " ")
             .replace("\u202F", " ")
    )
    
    # Remove symbols which may appear in localization strings
    value = (
        value.replace("™", "")
             .replace("®", "")
             .replace("©", "")
    )
    
    # Unicode normalization
    value = unicodedata.normalize("NFKC", value)
    
    # Normalize line breaks/tabs
    value = re.sub(r"[\r\n\t]+", " ", value)
    
    # Normalize repeated spaces
    value = re.sub(r"\s+", " ", value)
    
    return value.strip().casefold()
    
    
def comparison_key(value):
  
    
    return re.sub(r"\s+", "", normalize_text(value))
    
   
    
def find_language_column(sheet, language_name):
  
    target = comparison_key(language_name)
    
    # Search header area
    max_header_rows = min(sheet.max_row, 30)
    
    for row in sheet.iter_rows(
        min_row=1,
        max_row=max_header_rows
    ):
    
        for cell in row:
    
            if cell.value is None:
                continue
    
            cell_key = comparison_key(cell.value)
    
            if cell_key == target:
                return cell.column
    
      
            if target in cell_key:
                return cell.column
    
    return None
    
    
  
    
def find_english_column(sheet):
   
    
    possible_names = [
        "English",
        "EN",
        "English String",
        "English Strings",
        "English Translation",
        "English Text"
    ]
    
    normalized_names = [
        comparison_key(name)
        for name in possible_names
    ]
    
    max_header_rows = min(sheet.max_row, 30)
    
    for row in sheet.iter_rows(
        min_row=1,
        max_row=max_header_rows
    ):
    
        for cell in row:
    
            if cell.value is None:
                continue
    
            cell_key = comparison_key(cell.value)
    
            if cell_key in normalized_names:
                return cell.column
    
    return None
    
    

    
def get_translation(sheet, english_text, language_name):
   
    
    english_column = find_english_column(sheet)
    
    language_column = find_language_column(
        sheet,
        language_name
    )
    
    # --------------------------------------------------------
    # If headers are available
    # --------------------------------------------------------
    
    if english_column is not None and language_column is not None:
    
        target_key = comparison_key(english_text)
    
        for row_number in range(1, sheet.max_row + 1):
    
            english_value = sheet.cell(
                row=row_number,
                column=english_column
            ).value
    
            if comparison_key(english_value) == target_key:
    
                translation = sheet.cell(
                    row=row_number,
                    column=language_column
                ).value
    
                return "" if translation is None else str(translation)
    
  
    
    target_key = comparison_key(english_text)
    
    if language_column is not None:
    
        for row in sheet.iter_rows():
    
            for cell in row:
    
                if comparison_key(cell.value) == target_key:
    
                    translation = sheet.cell(
                        row=cell.row,
                        column=language_column
                    ).value
    
                    return (
                        ""
                        if translation is None
                        else str(translation)
                    )
    
    return None
    

    
def compare_with_excel(sheet, actual_value, english_key, description):
   
    expected_value = get_translation(
        sheet,
        english_key,
        LANGUAGE_NAME
    )
    
    if expected_value is None:
    
        # test.fail(
        #     "Excel translation not found for English string: "
        #     + str(english_key)
        # )
    
        return
    
    test.compare(
        comparison_key(actual_value),
        comparison_key(expected_value),
        description
    )
    
 
def main():
    
    
  
    
    workbook = load_workbook(
        EXCEL_PATH,
        read_only=True,
        data_only=True
    )
    
    if SHEET_NAME not in workbook.sheetnames:
    
        test.fail(
            "Excel sheet not found: " + SHEET_NAME
        )
    
        workbook.close()
        return
    
    sheet = workbook[SHEET_NAME]
    
  
    
    language_column = find_language_column(
        sheet,
        LANGUAGE_NAME
    )
    
    if language_column is None:
    
        test.fail(
            "Could not find language column for: "
            + LANGUAGE_NAME
        )
    
        workbook.close()
        return

    
    clickButton(
        waitForObject(
            names.numpad_pushButton_8_NumpadButton
        )
    )
    
    clickButton(
        waitForObject(
            names.numpad_pushButton_8_NumpadButton
        )
    )
    
    clickButton(
        waitForObject(
            names.numpad_pushButton_8_NumpadButton
        )
    )
    
    clickButton(
        waitForObject(
            names.numpad_pushButton_8_NumpadButton
        )
    )
    
    clickButton(
        waitForObject(
            names.mainFrame_okButton_QPushButton
        )
    )
 
    
    clickButton(
        waitForObject(
            names.homeScreen_settingsButton_QPushButton
        )
    )
    
    clickTab(
        waitForObject(
            names.settingsScreen_tbSystemInfo_TabWidget
        ),
        "INFO"
    )
    
    clickButton(
        waitForObject(
            names.gbSystemInformation_pbLanguageSettingsUpdate_QPushButton
        )
    )
    
  
    
    clickButton(
        waitForObject(
            names.languagesFrame_pbDenmark_QPushButton
        )
    )
    
    clickButton(
        waitForObject(
            names.buttonsFrame_acceptButton_QPushButton
        )
    )
  
    
    clickTab(
        waitForObject(
            names.settingsScreen_tbSystemInfo_TabWidget
        ),
        "TRÅDLØS FODKONTAKT"
    )
    
    clickTab(
        waitForObject(
            names.settingsScreen_tbSystemInfo_TabWidget
        ),
        "FORUDINDSTILLI\nNGER FOR BRUGER"
    )
    
    sendEvent(
        "QMoveEvent",
        waitForObject(names.o_ScreenSwitcher),
        0,
        -66,
        1285,
        233
    )
  
    
    clickButton(
        waitForObject(
            names.settingsScreen_addButton_QPushButton
        )
    )
    
   
    
    actualText1 = str(
        waitForObjectExists(
            names.centralWidget_lbOperation_QLabel
        ).text
    )
    
    compare_with_excel(
        sheet,
        actualText1,
        "Add Preset",
        "Add Preset Operation"
    )
 
    
    actualText2 = str(
        waitForObjectExists(
            names.centralWidget_pbName_QPushButton
        ).text
    )
    
    test.vp("Add Preset Popup")

    compare_with_excel(
        sheet,
        actualText2,
        "Preset Name",
        "Preset Name"
    )
    
  
    actualText3 = str(
        waitForObjectExists(
            names.averagePowerWidget_nameLabel_QLabel
        ).text
    )
    
    compare_with_excel(
        sheet,
        actualText3,
        "Average Power",
        "Average Power"
    )
 
    
    actualText4 = str(
        waitForObjectExists(
            names.averagePowerWidget_unitsOfMeasureLabel_QLabel
        ).text
    )
    
    compare_with_excel(
        sheet,
        actualText4,
        "Average Power Unit",
        "Average Power Unit"
    )
    
 
    
    actualText5 = str(
        waitForObjectExists(
            names.pulseEnergyWidget_nameLabel_QLabel
        ).text
    )
    
    compare_with_excel(
        sheet,
        actualText5,
        "Pulse Energy",
        "Pulse Energy"
    )
 
    mouseClick(
        waitForObject(
            names.centralWidget_frame_QFrame
        ),
        271,
        13,
        Qt.NoModifier,
        Qt.LeftButton
    )
    
    actualText6 = str(
        waitForObjectExists(
            names.pulseEnergyWidget_unitsOfMeasureLabel_QLabel
        ).text
    )
    
    compare_with_excel(
        sheet,
        actualText6,
        "Pulse Energy Unit",
        "Pulse Energy Unit"
    )
    
 
    actualText7 = str(
        waitForObjectExists(
            names.frequencyWidget_nameLabel_QLabel
        ).text
    )
    
    compare_with_excel(
        sheet,
        actualText7,
        "Frequency",
        "Frequency"
    )
    
 
    actualText8 = str(
        waitForObjectExists(
            names.frequencyWidget_unitsOfMeasureLabel_QLabel
        ).text
    )
    
    compare_with_excel(
        sheet,
        actualText8,
        "Frequency Unit",
        "Frequency Unit"
    )
 
    actualText9 = str(
        waitForObjectExists(
            names.modeButton_nameLabel_QLabel_3
        ).text
    )
    
    compare_with_excel(
        sheet,
        actualText9,
        "Regular Pulse",
        "Regular Pulse"
    )
    
   
    actualText10 = str(
        waitForObjectExists(
            names.peakPowerButton_nameLabel_QLabel_2
        ).text
    )
    
    compare_with_excel(
        sheet,
        actualText10,
        "Peak Power",
        "Peak Power"
    )
  
    
    actualText11 = str(
        waitForObjectExists(
            names.pulseModeWidget_saveButton_QPushButton
        ).text
    )
    
    compare_with_excel(
        sheet,
        actualText11,
        "Add",
        "Add Button"
    )
   
    
    mouseClick(
        waitForObject(
            names.pulseModeWidget_buttonsWidget_QWidget
        ),
        508,
        46,
        Qt.NoModifier,
        Qt.LeftButton
    )
    
    mouseClick(
        waitForObject(
            names.pulseModeWidget_buttonsWidget_QWidget
        ),
        411,
        1,
        Qt.NoModifier,
        Qt.LeftButton
    )
    
    clickButton(
        waitForObject(
            names.centralWidget_pbName_QPushButton
        )
    )
    
    type(
        waitForObject(
            names.keyboardPopup_wdEdit_WdLineEdit
        ),
        "preset one"
    )
    
    mouseClick(
        waitForObject(
            names.keyboardPopup_lbText_QLabel
        ),
        77,
        26,
        Qt.NoModifier,
        Qt.LeftButton
    )
    
    clickButton(
        waitForObject(
            names.pulseModeWidget_saveButton_QPushButton
        )
    )
  
    
    mouseClick(
        waitForObject(
            names.tabSettingsPresetsUserStonePresets_presetListView_PresetListView
        ),
        546,
        72,
        Qt.NoModifier,
        Qt.LeftButton
    )
  
    
    clickButton(
        waitForObject(
            names.settingsScreen_editButton_QPushButton
        )
    )
    
    mouseClick(
        waitForObject(
            names.tabSettingsPresetsUserStonePresets_presetListView_PresetListView
        ),
        494,
        66,
        Qt.NoModifier,
        Qt.LeftButton
    )
    
  
    
    actualEditOperation = str(
        waitForObjectExists(
            names.centralWidget_lbOperation_QLabel
        ).text
    )
    
    compare_with_excel(
        sheet,
        actualEditOperation,
        "Edit Preset",
        "Edit Preset Operation"
    )
  
    
    actualEditName = str(
        waitForObjectExists(
            names.centralWidget_pbName_QPushButton
        ).text
    )
    
    test.compare(
        comparison_key(actualEditName),
        comparison_key("preset one"),
        "Edit Preset Name"
    )
 
    
    actualEditAveragePower = str(
        waitForObjectExists(
            names.averagePowerWidget_nameLabel_QLabel
        ).text
    )
    
    compare_with_excel(
        sheet,
        actualEditAveragePower,
        "Average Power",
        "Edit - Average Power"
    )
    
  
    actualEditAveragePowerUnit = str(
        waitForObjectExists(
            names.averagePowerWidget_unitsOfMeasureLabel_QLabel
        ).text
    )
    
    compare_with_excel(
        sheet,
        actualEditAveragePowerUnit,
        "Average Power Unit",
        "Edit - Average Power Unit"
    )
    

    
    actualEditPulseEnergy = str(
        waitForObjectExists(
            names.pulseEnergyWidget_nameLabel_QLabel
        ).text
    )
    
    compare_with_excel(
        sheet,
        actualEditPulseEnergy,
        "Pulse Energy",
        "Edit - Pulse Energy"
    )

    
    actualEditPulseEnergyUnit = str(
        waitForObjectExists(
            names.pulseEnergyWidget_unitsOfMeasureLabel_QLabel
        ).text
    )
    
    compare_with_excel(
        sheet,
        actualEditPulseEnergyUnit,
        "Pulse Energy Unit",
        "Edit - Pulse Energy Unit"
    )
 
    
    actualEditFrequency = str(
        waitForObjectExists(
            names.frequencyWidget_nameLabel_QLabel
        ).text
    )
    
    compare_with_excel(
        sheet,
        actualEditFrequency,
        "Frequency",
        "Edit - Frequency"
    )
    
   
    actualEditFrequencyUnit = str(
        waitForObjectExists(
            names.frequencyWidget_unitsOfMeasureLabel_QLabel
        ).text
    )
    
    compare_with_excel(
        sheet,
        actualEditFrequencyUnit,
        "Frequency Unit",
        "Edit - Frequency Unit"
    )

    
    actualEditRegularPulse = str(
        waitForObjectExists(
            names.modeButton_nameLabel_QLabel_3
        ).text
    )
    
    compare_with_excel(
        sheet,
        actualEditRegularPulse,
        "Regular Pulse",
        "Edit - Regular Pulse"
    )
    
   
    
    actualEditPeakPower = str(
        waitForObjectExists(
            names.peakPowerButton_nameLabel_QLabel_2
        ).text
    )
    
    compare_with_excel(
        sheet,
        actualEditPeakPower,
        "Peak Power",
        "Edit - Peak Power"
    )
    
  
    actualUpdateButton = str(
        waitForObjectExists(
            names.pulseModeWidget_saveButton_QPushButton
        ).text
    )
    
    compare_with_excel(
        sheet,
        actualUpdateButton,
        "Update",
        "Update Button"
    )
    
   
    
    mouseClick(
        waitForObject(
            names.pulseModeWidget_buttonsWidget_QWidget
        ),
        531,
        22,
        Qt.NoModifier,
        Qt.LeftButton
    )
    
    mouseClick(
        waitForObject(
            names.pulseModeWidget_buttonsWidget_QWidget
        ),
        509,
        30,
        Qt.NoModifier,
        Qt.LeftButton
    )
    
    clickButton(
        waitForObject(
            names.pulseModeWidget_saveButton_QPushButton
        )
    )
    
   
    
    clickButton(
        waitForObject(
            names.settingsScreen_deletePresetButton_QPushButton
        )
    )
    
    mouseClick(
        waitForObject(
            names.tabSettingsPresetsUserStonePresets_presetListView_PresetListView
        ),
        553,
        80,
        Qt.NoModifier,
        Qt.LeftButton
    )
    
   
    
    actualDeleteTitle = str(
        waitForObjectExists(
            names.centralWidget_lbDialogTitle_QLabel
        ).text
    )
    
    compare_with_excel(
        sheet,
        actualDeleteTitle,
        "Delete Preset",
        "Delete Dialog Title"
    )
   
    
    actualDeleteName = str(
        waitForObjectExists(
            names.centralWidget_questionLabel_QLabel
        ).text
    )
    
    test.compare(
        comparison_key(actualDeleteName),
        comparison_key("preset one"),
        "Delete Preset Name"
    )
    
   
    
    actualInfo = str(
        waitForObjectExists(
            names.infoFrame_extInfoLabel_QLabel
        ).text
    )
    
    expectedInfo = "6 J     3 Hz     18 W"
    
    test.compare(
        comparison_key(actualInfo),
        comparison_key(expectedInfo),
        "Preset Information"
    )
    
   
    
    actualCancel = str(
        waitForObjectExists(
            names.buttonsFrame_rejectButton_QPushButton
        ).text
    )
    
    compare_with_excel(
        sheet,
        actualCancel,
        "Cancel",
        "Cancel Button"
    )
    
    
    
    actualDelete = str(
        waitForObjectExists(
            names.centralWidget_acceptButton_QPushButton
        ).text
    )
    
   

    compare_with_excel(
        sheet,
        actualDelete,
        "Delete",
        "Delete Button"
    )
    
    test.vp("Delete Preset Popup")

    clickButton(waitForObject(names.centralWidget_acceptButton_QPushButton))
    workbook.close()
    
