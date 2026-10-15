
# -*- coding: utf-8 -*-

import names
from openpyxl import load_workbook



EXCEL_PATH = (
    "/home/ezen/Test_Automation_Gemini/Localisation/"
    "suite_GFL_Danish/"
    "tst_GEM-LOC-Danish-PeakPower-SoftTissueCharacters/"
    "testdata/Gemini Split strings.xlsx"
)

EXCEL_SHEET_NAME = "Treatment Modes"

# Column B = English
ENGLISH_COLUMN = 2

# Column F = Danish
DANISH_COLUMN = 6



def normalize_text(value):

    if value is None:
        return ""

    return " ".join(
        str(value)
        .replace("\n", " ")
        .replace("\r", " ")
        .replace("\xa0", " ")
        .split()
    ).strip().casefold()




def load_excel_data():

    
    workbook = load_workbook(
        EXCEL_PATH,
        read_only=True,
        data_only=True
    )

    if EXCEL_SHEET_NAME not in workbook.sheetnames:

        test.fail(
            "Excel sheet not found: "
            + EXCEL_SHEET_NAME
        )

        workbook.close()
        return None, None

    sheet = workbook[
        EXCEL_SHEET_NAME
    ]

    danish_data = {}

    for row in sheet.iter_rows(
        min_col=ENGLISH_COLUMN,
        max_col=DANISH_COLUMN
    ):

        english_value = row[0].value
        danish_value = row[
            DANISH_COLUMN - ENGLISH_COLUMN
        ].value

        if english_value is None:
            continue

        english_key = normalize_text(
            english_value
        )

        if english_key == "":
            continue

        if danish_value is None:
            danish_value = ""

        danish_data[english_key] = str(
            danish_value
        ).strip()

    # test.log(
    #     "Excel loaded: {} | Sheet: {} | "
    #     "English/Danish entries: {}".format(
    #         EXCEL_PATH,
    #         EXCEL_SHEET_NAME,
    #         len(danish_data)
    #     )
    # )

    return workbook, danish_data




def get_danish_text(
    danish_data,
    english_text
):

    search_text = normalize_text(
        english_text
    )

    if search_text not in danish_data:

        test.fail(
            "English text not found in Excel Column B: "
            + str(english_text)
        )

        return None

    danish_value = danish_data[
        search_text
    ]

    if normalize_text(
        danish_value
    ) == "":

        test.fail(
            "Danish translation is empty for English text: "
            + str(english_text)
        )

        return None

    return danish_value




def verify_text(
    danish_data,
    object_name,
    english_text,
    description
):

    snooze(2)

    expected_text = get_danish_text(
        danish_data,
        english_text
    )

    # Do not continue with an empty Excel value.
    if expected_text is None:
        # test.fail(
        #     description
        #     + " skipped because Excel lookup failed"
        # )
        return

    try:

        obj = waitForObjectExists(
            getattr(names, object_name)
        )

    except Exception as e:

        test.fail(
            description
            + " | Object not found: "
            + object_name
            + " | "
            + str(e)
        )

        return

    actual_text = normalize_text(
        obj.text
    )

    expected_normalized = normalize_text(
        expected_text
    )

    # test.log(
    #     "{} | Expected Danish: [{}] | Actual: [{}]".format(
    #         description,
    #         expected_text,
    #         actual_text
    #     )
    # )

    test.compare(
        actual_text,
        expected_normalized,
        description
        + " - Danish text verification"
    )




def get_visible_tab_items(widget):

    result = []

    try:

        children = object.children(
            widget
        )

    except Exception:

        children = []

    for child in children:

        try:

            if not child.visible:
                continue

        except Exception:

            pass

        try:

            child_text = str(
                child.text
            ).strip()

        except Exception:

            child_text = ""

        if child_text:

            result.append(
                child_text
            )

        try:

            grandchildren = object.children(
                child
            )

        except Exception:

            grandchildren = []

        for grandchild in grandchildren:

            try:

                if not grandchild.visible:
                    continue

            except Exception:

                pass

            try:

                grandchild_text = str(
                    grandchild.text
                ).strip()

            except Exception:

                grandchild_text = ""

            if grandchild_text:

                result.append(
                    grandchild_text
                )

    return result



def verify_left_treatment_name(
    danish_data,
    english_text,
    description
):

    snooze(6)

    expected_text = get_danish_text(
        danish_data,
        english_text
    )

    if expected_text is None:

        # test.fail(
        #     description
        #     + " skipped because Excel lookup failed"
        # )

        return

    expected_normalized = normalize_text(
        expected_text
    )

    test.log(
        "{} | Expected Danish: [{}]".format(
            description,
            expected_text
        )
    )



    try:

        treatment_screen = waitForObjectExists(
            names.o_ScreenSwitcher
        )

    except Exception as e:

        test.fail(
            description
            + " | ScreenSwitcher not found | "
            + str(e)
        )

        return

 
    def search_object(obj):

        try:

            if not obj.visible:
                return None

        except Exception:

            pass

        try:

            obj_text = normalize_text(
                obj.text
            )

            if obj_text == expected_normalized:

                return obj

        except Exception:

            pass

        try:

            children = object.children(
                obj
            )

        except Exception:

            children = []

        for child in children:

            found = search_object(
                child
            )

            if found is not None:

                return found

        return None

    found_object = search_object(
        treatment_screen
    )

    if found_object is None:

        test.fail(
            "{} | Danish treatment text not found "
            "on screen: [{}]".format(
                description,
                expected_text
            )
        )

        return

    try:

        actual_text = normalize_text(
            found_object.text
        )

    except Exception:

        actual_text = ""

    test.log(
        "{} | Expected: [{}] | Actual: [{}]".format(
            description,
            expected_normalized,
            actual_text
        )
    )

    test.compare(
        actual_text,
        expected_normalized,
        description
        + " - Danish text verification"
    )



def verify_right_treatment_name(
    danish_data,
    english_text,
    description
):

 
    verify_left_treatment_name(
        danish_data,
        english_text,
        description
    )


def main():


    workbook, danish_data = load_excel_data()

    if workbook is None:
        return


 
    for i in range(4):

        clickButton(
            waitForObject(
                names.numpad_pushButton_8_NumpadButton
            )
        )

        snooze(1)


    clickButton(
        waitForObject(
            names.mainFrame_okButton_QPushButton
        )
    )

    



    clickButton(
        waitForObject(
            names.homeScreen_settingsButton_QPushButton
        )
    )

    snooze(2)


    clickTab(
        waitForObject(
            names.settingsScreen_tbSystemInfo_TabWidget
        ),
        "INFO"
    )

 

    clickButton(
        waitForObject(
            names.gbSystemInformation_pbLanguageSettingsUpdate_QPushButton
        )
    )

  

    

    clickButton(
        waitForObject(
            names.languagesFrame_pbDenmark_QPushButton
        )
    )

    

    clickButton(
        waitForObject(
            names.buttonsFrame_acceptButton_QPushButton
        )
    )

    

    clickButton(
        waitForObject(
            names.buttonsBar_homeButton_QPushButton
        )
    )

 


    clickButton(
        waitForObject(
            names.homeScreen_pbSoftTissueQuickStart_QToolButton
        )
    )

    


    mouseClick(
        waitForObjectItem(
            names.tabLeftMainPresets_leftPresetsView_PresetListView,
            "_1"
        ),
        238,
        58,
        Qt.NoModifier,
        Qt.LeftButton
    )

   


    mouseClick(
        waitForObjectItem(
            names.tabRightMainPresets_rightPresetsView_PresetListView,
            "_1"
        ),
        148,
        32,
        Qt.NoModifier,
        Qt.LeftButton
    )

    


    mouseClick(
        waitForObject(
            names.quickStartScreen_wdBottom_QWidget
        ),
        1167,
        9,
        Qt.NoModifier,
        Qt.LeftButton
    )

    snooze(2)


    sendEvent(
        "QMoveEvent",
        waitForObject(
            names.o_ScreenSwitcher
        ),
        0,
        -51,
        1286,
        61
    )

   




    clickButton(
        waitForObject(
            names.quickStartScreen_pbContinue_QPushButton
        )
    )

 


    clickButton(
        waitForObject(
            names.mainFrame_okButton_QPushButton_2
        )
    )

   


 
    clickButton(
        waitForObject(
            names.stateSwitch_readyButton_QPushButton
        )
    )





    verify_left_treatment_name(
        danish_data,
        "PROSTATE ENUCLEATION-(HIGH ENERGY)",
        "Left Treatment Name"
    )


    verify_text(
        danish_data,
        "averagePowerWidgetEmission_nameLabel_QLabel",
        "Average Power",
        "Left Average Power"
    )


    verify_text(
        danish_data,
        "pulseEnergyWidgetEmission_nameLabel_QLabel",
        "Pulse Energy",
        "Left Pulse Energy"
    )


    verify_text(
        danish_data,
        "frequencyWidgetEmission_nameLabel_QLabel",
        "Frequency",
        "Left Frequency"
    )

    # clickTab(waitForObject(names.leftPedalWidget_wdPresetTabs_TabWidget), "PROSTATAENUKLEATION-(HØJ ENERGI)")
    # test.compare(str(waitForObjectExists(names.modeButtonEmission_nameLabel_QLabel).text), "REGELMÆSSIG\nPULS")

    verify_text(
        danish_data,
        "modeButtonEmission_nameLabel_QLabel",
        "Pulse Mode",
        "Left Pulse Mode"
    )


    verify_text(
        danish_data,
        "peakPowerButtonEmission_nameLabel_QLabel",
        "Peak Power",
        "Left Peak Power"
    )


    verify_text(
        danish_data,
        "totalEnergyFrame_nameLabel_QLabel",
        "Total Energy",
        "Left Total Energy"
    )


    verify_text(
        danish_data,
        "totalTimeFrame_nameLabel_QLabel",
        "Total Lasing Time",
        "Left Total Time"
    )


    verify_text(
        danish_data,
        "averagePowerWidgetEmission_unitsOfMeasureLabel_QLabel",
        "W",
        "Left Power Unit"
    )


    verify_text(
        danish_data,
        "pulseEnergyWidgetEmission_unitsOfMeasureLabel_QLabel",
        "J",
        "Left Energy Unit"
    )


    verify_text(
        danish_data,
        "frequencyWidgetEmission_unitsOfMeasureLabel_QLabel",
        "Hz",
        "Left Frequency Unit"
    )


    verify_text(
        danish_data,
        "fiberFrame_nameLabel_QLabel",
        "Fiber",
        "Left Fiber"
    )


    verify_text(
        danish_data,
        "fiberUsesFrame_nameLabel_QLabel",
        "Fiber Uses Remaining",
        "Left Remaining Uses"
    )


    verify_text(
        danish_data,
        "bottomBar_aimingBeamButton_AimingBeamModeButton",
        "Aiming Beam",
        "Left Aiming Beam"
    )


    verify_text(
        danish_data,
        "stateSwitch_standbyButton_QPushButton",
        "Standby",
        "Left Standby"
    )


    verify_text(
        danish_data,
        "stateSwitch_readyButton_QPushButton",
        "Ready",
        "Left Ready"
    )


    # test.log(
    #     "Starting right-side Danish verification"
    # )


    verify_right_treatment_name(
        danish_data,
        "PROSTATE ENUCLEATION-(HIGH ENERGY)",
        "Right Treatment Name"
    )


    verify_text(
        danish_data,
        "averagePowerWidgetEmission_nameLabel_QLabel_2",
        "Average Power",
        "Right Average Power"
    )


    verify_text(
        danish_data,
        "pulseEnergyWidgetEmission_nameLabel_QLabel_2",
        "Pulse Energy",
        "Right Pulse Energy"
    )


    verify_text(
        danish_data,
        "frequencyWidgetEmission_nameLabel_QLabel_2",
        "Frequency",
        "Right Frequency"
    )


    verify_text(
        danish_data,
        "modeButtonEmission_nameLabel_QLabel_2",
        "Pulse Mode",
        "Right Pulse Mode"
    )


    verify_text(
        danish_data,
        "peakPowerButtonEmission_nameLabel_QLabel_2",
        "Peak Power",
        "Right Peak Power"
    )


    verify_text(
        danish_data,
        "averagePowerWidgetEmission_unitsOfMeasureLabel_QLabel_2",
        "W",
        "Right Power Unit"
    )


    verify_text(
        danish_data,
        "pulseEnergyWidgetEmission_unitsOfMeasureLabel_QLabel_2",
        "J",
        "Right Energy Unit"
    )


    verify_text(
        danish_data,
        "frequencyWidgetEmission_unitsOfMeasureLabel_QLabel_2",
        "Hz",
        "Right Frequency Unit"
    )


    # ========================================================
    # READY BUTTON
    # ========================================================

    clickButton(
        waitForObject(
            names.stateSwitch_readyButton_QPushButton
        )
    )

    


    # test.imagePresent(
    #     "Ready screen header image"
    # )


    clickButton(
        waitForObject(
            names.stateSwitch_readyButton_QPushButton
        )
    )



    test.vp("Readyscreen")

    

    workbook.close()

    # test.log(
    #     "Danish Ready Screen localization verification completed"
    # )

