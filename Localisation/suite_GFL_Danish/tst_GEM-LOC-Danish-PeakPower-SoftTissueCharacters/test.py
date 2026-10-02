
# -*- coding: utf-8 -*-

import names
import openpyxl
import re
import time


# ---------------------------------------------------------
# Normalize text for comparison
# ---------------------------------------------------------
def normalize(text):
    if text is None:
        return ""

    text = str(text)

    # Remove HTML tags
    text = re.sub(r"<br\s*/?>", " ", text, flags=re.IGNORECASE)
    text = re.sub(r"<[^>]+>", " ", text)

    # Remove extra spaces/new lines
    text = text.replace("\xa0", " ")
    text = " ".join(text.split())

    return text.strip().casefold()


# ---------------------------------------------------------
# Load Excel and create English -> Danish dictionary
#
# Excel:
# Column B = English
# Column F = Danish
# ---------------------------------------------------------
def load_danish_excel(excelpath):

   

    start_time = time.time()

    workbook = openpyxl.load_workbook(
        excelpath,
        data_only=True,
        read_only=True
    )

    load_time = time.time() - start_time

    # test.log(
    #     "Excel loading completed in "
    #     + str(round(load_time, 2))
    #     + " seconds"
    # )

    if "Treatment Modes" not in workbook.sheetnames:
        test.fail(
            "Worksheet 'Treatment Modes' does not exist in Excel file"
        )

    sheet = workbook["Treatment Modes"]

    # English Column B = 2
    # Danish Column F = 6
    english_column = 2
    danish_column = 6

    danish_data = {}

    start_time = time.time()

    for row in sheet.iter_rows(
        min_col=english_column,
        max_col=danish_column
    ):

        english_value = row[0].value
        danish_value = row[4].value

        if english_value is None:
            continue

        english_key = normalize(english_value)

        if english_key == "":
            continue

        danish_data[english_key] = (
            "" if danish_value is None else str(danish_value)
        )

    dictionary_time = time.time() - start_time

    test.log(
        "Excel data dictionary created in "
        + str(round(dictionary_time, 2))
        + " seconds"
    )

    return workbook, sheet, danish_data


# ---------------------------------------------------------
# Get Danish expected text using English text/key
# ---------------------------------------------------------
def get_danish_text(danish_data, english_text):

    key = normalize(english_text)

    if key not in danish_data:

        # test.fail(
        #     "English text not found in Excel: "
        #     + str(english_text)
        # )

        return ""

    expected = danish_data[key]

    
    return expected


# ---------------------------------------------------------
# Compare application text with Excel Danish text
# ---------------------------------------------------------
def verify_text(danish_data, object_name, english_text):

    actual_text = ""

    try:
        obj = waitForObjectExists(object_name)
        actual_text = str(obj.text).strip()
    except Exception as e:

        test.fail(
            "Unable to read application object: "
            + str(object_name)
            + " | Error: "
            + str(e)
        )

        return False

    expected_text = get_danish_text(
        danish_data,
        english_text
    )

    actual_normalized = normalize(actual_text)
    expected_normalized = normalize(expected_text)


    if actual_normalized == expected_normalized:

        test.log(
            "PASS - Text verified: "
            + str(english_text)
            + " -> "
            + actual_text
        )

        test.compare(
            actual_normalized,
            expected_normalized,
            "Danish text verification: " + str(english_text)
        )

        return True

    else:

        test.fail(
            "FAIL - Danish text mismatch for: "
            + str(english_text)
            + " | Actual: "
            + actual_text
            + " | Expected: "
            + str(expected_text)
        )

        return False


# ---------------------------------------------------------
# MAIN
# ---------------------------------------------------------
def main():

    excelpath = (
        "/home/ezen/Test_Automation_Gemini/Localisation/"
        "suite_GFL_Danish/tst_GEM-LOC-Danish-PeakPower-SoftTissueCharacters/"
        "testdata/Gemini Split strings.xlsx"
    )

   


    # -----------------------------------------------------
    # Login
    # -----------------------------------------------------

   
    clickButton(
        waitForObject(
            names.numpad_pushButton_8_NumpadButton
        )
    )

    clickButton(
        waitForObject(
            names.numpad_pushButton_8_NumpadButton
        )
    )

    clickButton(
        waitForObject(
            names.numpad_pushButton_8_NumpadButton
        )
    )

    clickButton(
        waitForObject(
            names.numpad_pushButton_8_NumpadButton
        )
    )

    snooze(6)

    clickButton(
        waitForObject(
            names.mainFrame_okButton_QPushButton
        )
    )


    # -----------------------------------------------------
    # Change language to Danish
    # -----------------------------------------------------

    

    clickButton(
        waitForObject(
            names.homeScreen_settingsButton_QPushButton
        )
    )

    clickTab(
        waitForObject(
            names.settingsScreen_tbSystemInfo_TabWidget
        ),
        "INFO"
    )

    clickButton(
        waitForObject(
            names.gbSystemInformation_pbLanguageSettingsUpdate_QPushButton
        )
    )

    

    clickButton(
        waitForObject(
            names.languagesFrame_pbDenmark_QPushButton
        )
    )

    clickButton(
        waitForObject(
            names.buttonsFrame_acceptButton_QPushButton
        )
    )

    clickButton(
        waitForObject(
            names.settingsScreen_pbBack_QPushButton
        )
    )


    # -----------------------------------------------------
    # Soft Tissue Quick Start
    # -----------------------------------------------------

    
    clickButton(
        waitForObject(
            names.homeScreen_pbSoftTissueQuickStart_QToolButton
        )
    )

    mouseClick(
        waitForObjectItem(
            names.tabLeftMainPresets_leftPresetsView_PresetListView,
            "_1"
        ),
        305,
        35,
        Qt.NoModifier,
        Qt.LeftButton
    )

    mouseClick(
        waitForObjectItem(
            names.tabRightMainPresets_rightPresetsView_PresetListView,
            "_1"
        ),
        299,
        49,
        Qt.NoModifier,
        Qt.LeftButton
    )

    clickButton(
        waitForObject(
            names.quickStartScreen_pbContinue_QPushButton
        )
    )

    clickButton(
        waitForObject(
            names.mainFrame_okButton_QPushButton_2
        )
    )


    # -----------------------------------------------------
    # Peak Power
    # -----------------------------------------------------

   

    mouseClick(
        waitForObject(
            names.peakPowerButton_nameLabel_QLabel
        ),
        34,
        34,
        Qt.NoModifier,
        Qt.LeftButton
    )

    snooze(2)


    # -----------------------------------------------------
    # NOW load Excel
    #
    # Excel loading happens only after application
    # navigation reaches Peak Power.
    # -----------------------------------------------------

    workbook, sheet, danish_data = load_danish_excel(
        excelpath
    )

   

    # -----------------------------------------------------
    # Peak Power verification
    #
    # NO ROW NUMBER
    # NO COLUMN NUMBER
    #
    # English text is used as lookup key.
    # -----------------------------------------------------

    verify_text(
        danish_data,
        names.centralWidget_lbTitle_QLabel,
        "SELECT PEAK POWER"
    )

    verify_text(
        danish_data,
        names.centralWidget_pw500_QPushButton,
        "High"
    )

    verify_text(
        danish_data,
        names.centralWidget_pw250_QPushButton,
        "Medium"
    )

    verify_text(
        danish_data,
        names.centralWidget_pw125_QPushButton,
        "Low"
    )

    verify_text(
        danish_data,
        names.buttonsFrame_rejectButton_QPushButton,
        "CANCEL"
    )


    # -----------------------------------------------------
    # Close Peak Power
    # -----------------------------------------------------

    

    mouseClick(
        waitForObject(
            names.centralWidget_mainWidget_QWidget
        ),
        460,
        549,
        Qt.NoModifier,
        Qt.LeftButton
    )

    clickButton(
        waitForObject(
            names.buttonsFrame_rejectButton_QPushButton
        )
    )


    # -----------------------------------------------------
    # Soft Tissue Assistant
    # -----------------------------------------------------

   
    clickButton(
        waitForObject(
            names.buttonsBar_homeButton_QPushButton_2
        )
    )

    clickButton(
        waitForObject(
            names.homeScreen_pbSoftTissueAssistant_QToolButton
        )
    )

    clickButton(
        waitForObject(
            names.gbSoftTissueProcedures_pbIncision_QPushButton
        )
    )

    clickButton(
        waitForObject(
            names.softTissueAssistantModeScreen_pbContinue_QPushButton
        )
    )

    clickButton(
        waitForObject(
            names.treatmentCharacteristicsBar_pbEdit_QPushButton
        )

    )

    snooze(2)


    # -----------------------------------------------------
    # Soft Tissue Characteristics verification
    #
    # Again:
    # NO ROW NUMBER
    # NO COLUMN NUMBER
    # -----------------------------------------------------

    
    verify_text(
        danish_data,
        names.centralWidget_titleLabel_QLabel,
        "SOFT TISSUE CHARACTERISTICS"
    )

    verify_text(
        danish_data,
        names.centralWidget_titleLabel_QLabel_2,
        "PROCEDURE"
    )

    verify_text(
        danish_data,
        names.centralWidget_valueLabel_QLabel,
        "Incision"
    )

    verify_text(
        danish_data,
        names.centralWidget_rejectButton_QPushButton,
        "EDIT"
    )

    verify_text(
        danish_data,
        names.centralWidget_acceptButton_QPushButton,
        "OK"
    )


    # -----------------------------------------------------
    # Close Excel
    # -----------------------------------------------------

    workbook.close()

