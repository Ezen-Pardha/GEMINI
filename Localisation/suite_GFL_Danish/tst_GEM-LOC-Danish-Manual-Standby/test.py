
# -*- coding: utf-8 -*-

import names
from openpyxl import load_workbook


# ============================================================
# EXCEL CONFIGURATION
# ============================================================

EXCEL_PATH = (
    "/home/ezen/Test_Automation_Gemini/Localisation/"
    "suite_GFL_Danish/"
    "tst_GEM-LOC-Danish-PeakPower-SoftTissueCharacters/"
    "testdata/Gemini Split strings.xlsx"
)

EXCEL_SHEET_NAME = "Treatment Modes"

# Column B = English
ENGLISH_COLUMN = 2

# Column F = Danish
DANISH_COLUMN = 6


# ============================================================
# TEXT NORMALIZATION
# ============================================================

def normalize_text(value):

    if value is None:
        return ""

    return " ".join(
        str(value)
        .replace("\n", " ")
        .replace("\r", " ")
        .split()
    ).strip().casefold()


# ============================================================
# GET DANISH TEXT FROM EXCEL
# ============================================================

def get_danish_text(sheet, english_text):

    search_text = normalize_text(
        english_text
    )

    for row in sheet.iter_rows(
        min_row=1,
        min_col=ENGLISH_COLUMN,
        max_col=DANISH_COLUMN,
        values_only=True
    ):

        english_value = row[0]

        if normalize_text(
            english_value
        ) == search_text:

            danish_value = row[
                DANISH_COLUMN - ENGLISH_COLUMN
            ]

            if danish_value is None:

                test.fail(
                    "Danish translation is empty for: "
                    + english_text
                )

                return ""

            return str(
                danish_value
            ).strip()

    test.fail(
        "English text not found in Excel Column B: "
        + english_text
    )

    return ""


# ============================================================
# VERIFY TEXT AGAINST EXCEL
# ============================================================

def verify_text(
    sheet,
    object_name,
    english_text,
    description
):

    # Wait before every verification
    snooze(6)

    # Get expected Danish value from Excel
    expected_text = get_danish_text(
        sheet,
        english_text
    )

    expected_normalized = normalize_text(
        expected_text
    )

    # Get application object
    obj = waitForObjectExists(
        getattr(names, object_name)
    )

    # Get actual application text
    actual_text = normalize_text(
        obj.text
    )

    # test.log(
    #     "{} | English Key: [{}] | "
    #     "Expected Danish: [{}] | Actual: [{}]".format(
    #         description,
    #         english_text,
    #         expected_text,
    #         actual_text
    #     )
    # )

    # Compare application text with Excel
    test.compare(
        actual_text,
        expected_normalized,
        description + " - Danish text verification"
    )


# ============================================================
# VERIFY LEFT MANUAL TAB
# ============================================================

def verify_left_manual_tab(
    sheet,
    english_text,
    description
):

    snooze(6)

    expected_text = get_danish_text(
        sheet,
        english_text
    )

    expected_normalized = normalize_text(
        expected_text
    )

    # IMPORTANT:
    # Use the recorded Squish object.
    # Do NOT construct the object name manually.

    obj = waitForObjectExists(
        names.mANUEL_TabItem
    )

    actual_text = normalize_text(
        obj.text
    )

    # test.log(
    #     "{} | English Key: [{}] | "
    #     "Expected Danish: [{}] | Actual: [{}]".format(
    #         description,
    #         english_text,
    #         expected_text,
    #         actual_text
    #     )
    # )

    test.compare(
        actual_text,
        expected_normalized,
        description + " - Danish text verification"
    )


# ============================================================
# VERIFY RIGHT MANUAL TAB
# ============================================================

def verify_right_manual_tab(
    sheet,
    english_text,
    description
):

    snooze(6)

    expected_text = get_danish_text(
        sheet,
        english_text
    )

    expected_normalized = normalize_text(
        expected_text
    )

    # IMPORTANT:
    # Use the recorded Squish object.
    # Do NOT construct the object name manually.

    obj = waitForObjectExists(
        names.mANUEL_TabItem_2
    )

    actual_text = normalize_text(
        obj.text
    )

    # test.log(
    #     "{} | English Key: [{}] | "
    #     "Expected Danish: [{}] | Actual: [{}]".format(
    #         description,
    #         english_text,
    #         expected_text,
    #         actual_text
    #     )
    # )

    test.compare(
        actual_text,
        expected_normalized,
        description + " - Danish text verification"
    )


# ============================================================
# MAIN
# ============================================================

def main():

    # ========================================================
    # LOAD EXCEL
    # ========================================================

    

    workbook = load_workbook(
        EXCEL_PATH,
        read_only=True,
        data_only=True
    )

    if EXCEL_SHEET_NAME not in workbook.sheetnames:

        test.fail(
            "Excel sheet not found: "
            + EXCEL_SHEET_NAME
        )

        workbook.close()
        return

    sheet = workbook[
        EXCEL_SHEET_NAME
    ]


    # test.log(
    #     "Sheet: {}".format(
    #         EXCEL_SHEET_NAME
    #     )
    # )


    # ========================================================
    # LOGIN
    # ========================================================

    doubleClick(
        waitForObject(
            names.numpad_pushButton_8_NumpadButton,
            100461
        ),
        96,
        56,
        Qt.NoModifier,
        Qt.LeftButton
    )

    snooze(1)

    doubleClick(
        waitForObject(
            names.numpad_pushButton_8_NumpadButton
        ),
        96,
        56,
        Qt.NoModifier,
        Qt.LeftButton
    )

    snooze(1)

    clickButton(
        waitForObject(
            names.mainFrame_okButton_QPushButton
        )
    )

    snooze(2)


    # ========================================================
    # OPEN SETTINGS
    # ========================================================

    clickButton(
        waitForObject(
            names.homeScreen_settingsButton_QPushButton
        )
    )

    snooze(2)


    # ========================================================
    # OPEN SYSTEM INFORMATION
    # ========================================================

    clickTab(
        waitForObject(
            names.settingsScreen_tbSystemInfo_TabWidget
        ),
        "INFO"
    )

    snooze(2)


    # ========================================================
    # OPEN LANGUAGE SETTINGS
    # ========================================================

    clickButton(
        waitForObject(
            names.gbSystemInformation_pbLanguageSettingsUpdate_QPushButton
        )
    )

    snooze(2)


    # ========================================================
    # SELECT DANISH
    # ========================================================

    
    clickButton(
        waitForObject(
            names.languagesFrame_pbDenmark_QPushButton
        )
    )

    snooze(1)


    # ========================================================
    # ACCEPT LANGUAGE
    # ========================================================

    clickButton(
        waitForObject(
            names.buttonsFrame_acceptButton_QPushButton
        )
    )

    snooze(3)


    # ========================================================
    # RETURN TO HOME
    # ========================================================

    clickButton(
        waitForObject(
            names.buttonsBar_homeButton_QPushButton
        )
    )

    snooze(2)


    # ========================================================
    # OPEN EXPERT MODE
    # ========================================================

    clickButton(
        waitForObject(
            names.gbExpert_expertButton_QPushButton
        )
    )

    snooze(3)





    # --------------------------------------------------------
    # MANUAL TAB
    # Excel:
    # MANUAL -> Danish translation
    # --------------------------------------------------------

    verify_left_manual_tab(
        sheet,
        "MANUAL",
        "Left Manual Tab"
    )


    # --------------------------------------------------------
    # AVERAGE POWER
    # --------------------------------------------------------

    verify_text(
        sheet,
        "averagePowerWidget_nameLabel_QLabel_2",
        "AVERAGE POWER",
        "Left Average Power"
    )


    # --------------------------------------------------------
    # POWER UNIT
    # --------------------------------------------------------

    verify_text(
        sheet,
        "averagePowerWidget_unitsOfMeasureLabel_QLabel_2",
        "W",
        "Left Power Unit"
    )


    # --------------------------------------------------------
    # PULSE ENERGY
    # --------------------------------------------------------

    verify_text(
        sheet,
        "pulseEnergyWidget_nameLabel_QLabel_2",
        "PULSE ENERGY",
        "Left Pulse Energy"
    )


    # --------------------------------------------------------
    # ENERGY UNIT
    # --------------------------------------------------------

    verify_text(
        sheet,
        "pulseEnergyWidget_unitsOfMeasureLabel_QLabel_2",
        "J",
        "Left Energy Unit"
    )


    # --------------------------------------------------------
    # FREQUENCY
    # --------------------------------------------------------

    verify_text(
        sheet,
        "frequencyWidget_nameLabel_QLabel_2",
        "FREQUENCY",
        "Left Frequency"
    )


    # --------------------------------------------------------
    # FREQUENCY UNIT
    # --------------------------------------------------------

    verify_text(
        sheet,
        "frequencyWidget_unitsOfMeasureLabel_QLabel_2",
        "Hz",
        "Left Frequency Unit"
    )


    # --------------------------------------------------------
    # REGULAR PULSE
    # --------------------------------------------------------

    verify_text(
        sheet,
        "modeButton_nameLabel_QLabel",
        "REGULAR PULSE",
        "Left Pulse Mode"
    )


    # --------------------------------------------------------
    # PEAK POWER
    # --------------------------------------------------------

    verify_text(
        sheet,
        "peakPowerButton_nameLabel_QLabel",
        "PEAK POWER",
        "Left Peak Power"
    )


    # --------------------------------------------------------
    # SAVE
    # --------------------------------------------------------

    verify_text(
        sheet,
        "pulseModeWidget_saveButton_QPushButton_2",
        "SAVE",
        "Left Save Button"
    )


    # --------------------------------------------------------
    # TOTAL ENERGY
    # --------------------------------------------------------

    verify_text(
        sheet,
        "totalEnergyFrame_nameLabel_QLabel",
        "TOTAL ENERGY",
        "Left Total Energy"
    )


    # --------------------------------------------------------
    # TOTAL LASER TIME
    # --------------------------------------------------------

    verify_text(
        sheet,
        "totalTimeFrame_nameLabel_QLabel",
        "TOTAL LASING TIME",
        "Left Total Lasing Time"
    )


    # ========================================================
    # RIGHT SIDE VERIFICATION
    # ========================================================

    # test.log(
    #     "Starting right-side Danish verification"
    # )


    # --------------------------------------------------------
    # MANUAL TAB
    # --------------------------------------------------------

    verify_right_manual_tab(
        sheet,
        "MANUAL",
        "Right Manual Tab"
    )


    # --------------------------------------------------------
    # AVERAGE POWER
    # --------------------------------------------------------

    verify_text(
        sheet,
        "averagePowerWidget_nameLabel_QLabel_3",
        "AVERAGE POWER",
        "Right Average Power"
    )


    # --------------------------------------------------------
    # POWER UNIT
    # --------------------------------------------------------

    verify_text(
        sheet,
        "averagePowerWidget_unitsOfMeasureLabel_QLabel_3",
        "W",
        "Right Power Unit"
    )


    # --------------------------------------------------------
    # PULSE ENERGY
    # --------------------------------------------------------

    verify_text(
        sheet,
        "pulseEnergyWidget_nameLabel_QLabel_3",
        "PULSE ENERGY",
        "Right Pulse Energy"
    )


    # --------------------------------------------------------
    # ENERGY UNIT
    # --------------------------------------------------------

    verify_text(
        sheet,
        "pulseEnergyWidget_unitsOfMeasureLabel_QLabel_3",
        "J",
        "Right Energy Unit"
    )


    # --------------------------------------------------------
    # FREQUENCY
    # --------------------------------------------------------

    verify_text(
        sheet,
        "frequencyWidget_nameLabel_QLabel_3",
        "FREQUENCY",
        "Right Frequency"
    )


    # --------------------------------------------------------
    # FREQUENCY UNIT
    # --------------------------------------------------------

    # verify_text(
    #     sheet,
    #     "frequencyWidget_unitsOfMeasureLabel_QLabel_3",
    #     "Hz",
    #     "Right Frequency Unit"
    # )


    # --------------------------------------------------------
    # SAVE
    # --------------------------------------------------------

    verify_text(
        sheet,
        "pulseModeWidget_saveButton_QPushButton_3",
        "SAVE",
        "Right Save Button"
    )


    # --------------------------------------------------------
    # PEAK POWER
    # --------------------------------------------------------

    verify_text(
        sheet,
        "peakPowerButton_nameLabel_QLabel_3",
        "PEAK POWER",
        "Right Peak Power"
    )


    # --------------------------------------------------------
    # REGULAR PULSE
    # --------------------------------------------------------

    verify_text(
        sheet,
        "modeButton_nameLabel_QLabel_2",
        "REGULAR PULSE",
        "Right Pulse Mode"
    )


    # ========================================================
    # READY BUTTON
    # ========================================================

    verify_text(
        sheet,
        "stateSwitch_readyButton_QPushButton",
        "READY",
        "Ready Button"
    )


    # ========================================================
    # STANDBY BUTTON
    # ========================================================

    verify_text(
        sheet,
        "stateSwitch_standbyButton_QPushButton",
        "STANDBY",
        "Standby Button"
    )


    # ========================================================
    # FIBER
    # ========================================================

    verify_text(
        sheet,
        "fiberFrame_nameLabel_QLabel",
        "FIBER",
        "Fiber"
    )


    # ========================================================
    # FIBER USES REMAINING
    # ========================================================

    verify_text(
        sheet,
        "fiberUsesFrame_nameLabel_QLabel",
        "FIBER USES REMAINING",
        "Fiber Uses Remaining"
    )


    # ========================================================
    # AIMING BEAM
    # ========================================================

    verify_text(
        sheet,
        "bottomBar_aimingBeamButton_AimingBeamModeButton",
        "AIMING BEAM",
        "Aiming Beam"
    )


    # ========================================================
    # VERIFICATION POINT
    # ========================================================

    # test.vp(
    #     "ExpertManualTreatmentModes"
    # )


    # ========================================================
    # CLOSE EXCEL
    # ========================================================

    workbook.close()

    test.log(
        "Danish Expert Manual Treatment Modes "
        "verification completed"
    )

