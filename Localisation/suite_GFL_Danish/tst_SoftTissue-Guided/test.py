
# -*- coding: utf-8 -*-

import names
import openpyxl
import re
import os

TESTCASE_DIR = os.path.dirname(os.path.abspath(__file__))

EXCEL_PATH = os.path.abspath(
    os.path.join(
        TESTCASE_DIR,
        "..",
        "shared",
        "testdata",
        "Gemini Split strings.xlsx"
    )
)

SHEET_NAME = "Soft Tissue-Guided"


def normalize(text):
    text = str(text or "")
    text = re.sub(r"<[^>]+>", "", text)
    text = text.replace("\u00a0", " ")
    text = " ".join(text.split())
    return text.casefold()


def get_all_texts(obj):
    texts = []

    try:
        text = str(obj.text).strip()

        if text:
            if "<html>" in text.lower():
                text = re.sub(r"<[^>]+>", "", text)
                text = " ".join(text.split())

            if text:
                texts.append(text)

    except Exception:
        pass

    try:
        children = object.children(obj)

        for child in children:
            texts.extend(get_all_texts(child))

    except Exception:
        pass

    return texts


def ignore_text(text):
    text = str(text).strip()

    if not text:
        return True

    if re.fullmatch(r"[-_=.]+", text):
        return True

    return False


def load_danish_strings():
    if not os.path.isfile(EXCEL_PATH):
        test.fail(
            "Excel file not found: {}".format(EXCEL_PATH)
        )
        return {}

    try:
        workbook = openpyxl.load_workbook(
            EXCEL_PATH,
            data_only=True,
            read_only=True
        )
    except Exception as error:
        test.fail(
            "Unable to load Danish Excel: {}".format(error)
        )
        return {}

    try:
        if SHEET_NAME not in workbook.sheetnames:
            test.fail(
                "Excel worksheet not found: {}".format(SHEET_NAME)
            )
            return {}

        sheet = workbook[SHEET_NAME]
        excel_strings = {}

        # Column B = English
        # Column F = Danish

        for row in sheet.iter_rows(
            min_row=2,
            values_only=True
        ):
            if len(row) < 6:
                continue

            english_value = row[1]
            danish_value = row[5]

            if danish_value is None:
                continue

            english = (
                str(english_value).strip()
                if english_value is not None
                else ""
            )
            danish = str(danish_value).strip()

            if ignore_text(danish):
                continue

            key = normalize(danish)

            if key not in excel_strings:
                excel_strings[key] = []

            excel_strings[key].append(english)

        return excel_strings

    finally:
        workbook.close()


def main():

    # LOAD DANISH EXCEL CATALOG

    excel_strings = load_danish_strings()

    if not excel_strings:
        test.fail("No Danish strings loaded from Excel.")
        return

    # LOGIN

    doubleClick(
        waitForObject(names.numpad_pushButton_8_NumpadButton),
        68, 52,
        Qt.NoModifier,
        Qt.LeftButton
    )

    doubleClick(
        waitForObject(names.numpad_pushButton_8_NumpadButton),
        116, 46,
        Qt.NoModifier,
        Qt.LeftButton
    )

    clickButton(
        waitForObject(names.mainFrame_okButton_QPushButton)
    )

    # CHANGE LANGUAGE TO DANISH

    clickButton(
        waitForObject(names.homeScreen_settingsButton_QPushButton)
    )

    clickTab(
        waitForObject(names.settingsScreen_tbSystemInfo_TabWidget),
        "INFO"
    )

    clickButton(
        waitForObject(
            names.gbSystemInformation_pbLanguageSettingsUpdate_QPushButton
        )
    )

    clickButton(
        waitForObject(names.languagesFrame_pbDenmark_QPushButton)
    )

    clickButton(
        waitForObject(names.buttonsFrame_acceptButton_QPushButton)
    )

    clickButton(
        waitForObject(names.settingsScreen_pbBack_QPushButton)
    )

    # OPEN SOFT TISSUE ASSISTANT

    clickButton(
        waitForObject(
            names.homeScreen_pbSoftTissueAssistant_QToolButton
        )
    )

    mouseClick(
        waitForObject(
            names.softTissueAssistantModeScreen_gbSoftTissueProcedures_QFrame
        ),
        90, 150,
        Qt.NoModifier,
        Qt.LeftButton
    )

    # COLLECT SCREEN TEXT

    soft_tissue = waitForObject(
        names.softTissueAssistantModeScreen_SoftTissueAssistantModeScreen
    )

    actual_texts = get_all_texts(soft_tissue)

    # REMOVE DUPLICATES

    unique_texts = []
    seen = set()

    for text in actual_texts:
        if ignore_text(text):
            continue

        key = normalize(text)

        if key not in seen:
            seen.add(key)
            unique_texts.append(text)

    # VALIDATE SCREEN STRINGS

    total_pass = 0
    total_fail = 0

    for text in unique_texts:
        key = normalize(text)

        if key in excel_strings:
            english_values = excel_strings[key]

            english = next(
                (value for value in english_values if value),
                ""
            )

            if english:
                test.passes(
                    "PASS: {} -> {}".format(english, text)
                )
            else:
                test.passes(
                    "PASS: Danish string found in Excel -> {}".format(
                        text
                    )
                )

            total_pass += 1

        else:
            test.fail(
                "FAIL: Screen string not found in Danish Excel: {}".format(
                    text
                )
            )
            total_fail += 1

    # FINAL SUMMARY

    test.log(
        "Danish Localization Summary | Total: {} | PASS: {} | FAIL: {}".format(
            len(unique_texts),
            total_pass,
            total_fail
        )
    )

    if total_fail == 0 and total_pass > 0:
        test.passes("DANISH LOCALIZATION VALIDATION PASSED")
    else:
        test.fail("DANISH LOCALIZATION VALIDATION FAILED")
