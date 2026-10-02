# -*- coding: utf-8 -*-

import names
import openpyxl
from openpyxl import load_workbook


# =========================================================
# EXCEL GLOBALS
# =========================================================

sheet = None
workbook = None


# =========================================================
# TEXT VERIFICATION
# =========================================================

def get_object_text(obj):
    # -----------------------------------------------------
    # Try text property
    # -----------------------------------------------------
    try:
        value = obj.text
        if value is not None:
            value = str(value).strip()
            if value:
                return value
    except:
        pass

    # -----------------------------------------------------
    # Try title property
    # -----------------------------------------------------
    try:
        value = obj.title
        if value is not None:
            value = str(value).strip()
            if value:
                return value
    except:
        pass

    # -----------------------------------------------------
    # Try windowTitle property
    # -----------------------------------------------------
    try:
        value = obj.windowTitle
        if value is not None:
            value = str(value).strip()
            if value:
                return value
    except:
        pass

    # -----------------------------------------------------
    # Try accessibleName property
    # -----------------------------------------------------
    try:
        value = obj.accessibleName
        if value is not None:
            value = str(value).strip()
            if value:
                return value
    except:
        pass

    return ""


# =========================================================
# VERIFY TEXT FROM EXCEL
# =========================================================

def verify_text(objectName, row):
    # Wait 6 seconds before every verification
    snooze(6)

    obj = waitForObject(objectName)

    # -----------------------------------------------------
    # Get actual UI text
    # -----------------------------------------------------
    actual = get_object_text(obj)

    # -----------------------------------------------------
    # Get expected text from Excel column B
    # -----------------------------------------------------
    if sheet is None:
        test.fail("Excel sheet is not initialized")
        return

    try:
        expected_value = sheet.cell(
            row=row,
            column=2
        ).value

        expected = str(expected_value or "")

    except Exception as e:
        test.fail(
            "Failed to read Excel row "
            + str(row)
            + ": "
            + str(e)
        )
        return

    # -----------------------------------------------------
    # Normalize spaces and line breaks
    # -----------------------------------------------------
    actual = " ".join(actual.split())
    expected = " ".join(expected.split())


    test.compare(
        actual,
        expected,
        "Text verification"
    )


# =========================================================
# VERIFICATION GROUP
# =========================================================

def verify_group(group_name):
    

    for objectName, row in verification_objects[group_name]:
        verify_text(
            objectName,
            row
        )



def verify_popup(group_name):
    test.log(
        "Popup verification started: "
        + group_name
    )

    verify_group(group_name)

    # Wait 6 seconds before clicking OK
    snooze(6)

    clickButton(
        waitForObject(
            names.centralWidget_OKButton_QPushButton_2
        )
    )



# =========================================================
# CLICK INCREMENT UNTIL POPUP OPENS
# =========================================================

def click_until_popup_opens():
    max_clicks = 50


    for i in range(max_clicks):

        # Wait 6 seconds before every increment click
        snooze(6)

        clickButton(
            waitForObject(
                names.pulseEnergyWidget_incrementButton_QPushButton
            )
        )

       

        # Give application time to update
        snooze(0.2)

        # -------------------------------------------------
        # Check popup
        # -------------------------------------------------
        try:
            central_widget = findObject({
                "name": "centralWidget",
                "type": "QFrame",
                "window": ":o_ScreenSwitcher"
            })

            if central_widget.visible:
                test.log(
                    "Popup opened after "
                    + str(i + 1)
                    + " clicks"
                )

                # Wait 6 seconds after popup opens
                snooze(6)

                return True

        except:
            pass

    # -----------------------------------------------------
    # Popup did not open
    # -----------------------------------------------------
    

    return False


# =========================================================
# ALL VERIFICATION OBJECTS
# =========================================================

verification_objects = {

    # -----------------------------------------------------
    # Main Screen
    # -----------------------------------------------------
    "main": [
        (names.centralWidget_lbTitle_QLabel, 2),
        (names.centralWidget_label_2_QLabel, 3),
        (names.centralWidget_CW_QPushButton, 4),
        (names.centralWidget_Square_QPushButton, 5),
        (names.centralWidget_label_QLabel, 6)
    ],

    # -----------------------------------------------------
    # Pulse Buttons
    # -----------------------------------------------------
    "pulse_buttons": [
        (names.centralWidget_MRP_QPushButton, 7),
        (names.centralWidget_Fragmentation_QPushButton, 8),
        (names.centralWidget_Cleaning_QPushButton, 9),
        (names.centralWidget_Enucleation_QPushButton, 10),
        (names.buttonsFrame_rejectButton_QPushButton, 11)
    ],

    # -----------------------------------------------------
    # Power Mode Update
    # -----------------------------------------------------
    "power_mode": [
        (names.centralWidget_TitleLabel_QLabel, 12),
        (names.centralWidget_MessageLabel_QLabel, 13),
        (names.centralWidget_CancelButton_QPushButton_2, 14),
        (names.centralWidget_DisableButton_QPushButton, 15)
    ],

    # -----------------------------------------------------
    # Standard Pulse Popup
    # -----------------------------------------------------
    "standard_pulse": [
        (names.centralWidget_TitleLabel_QLabel, 16),
        (names.centralWidget_MessageLabel_QLabel, 17),
        (names.centralWidget_OKButton_QPushButton_2, 18)
    ],

    # -----------------------------------------------------
    # CW Popup
    # -----------------------------------------------------
    "cw": [
        (names.centralWidget_TitleLabel_QLabel, 19),
        (names.centralWidget_MessageLabel_QLabel, 20),
        (names.centralWidget_OKButton_QPushButton_2, 18)
    ],

    # -----------------------------------------------------
    # MRP Popup
    # -----------------------------------------------------
    "mrp": [
        (names.centralWidget_TitleLabel_QLabel, 21),
        (names.centralWidget_MessageLabel_QLabel, 22),
        (names.centralWidget_OKButton_QPushButton_2, 18)
    ],

    # -----------------------------------------------------
    # Fragmentation Popup
    # -----------------------------------------------------
    "fragmentation": [
        (names.centralWidget_TitleLabel_QLabel, 23),
        (names.centralWidget_MessageLabel_QLabel, 24),
        (names.centralWidget_OKButton_QPushButton_2, 18)
    ],

    # -----------------------------------------------------
    # Cleaning Popup
    # -----------------------------------------------------
    "cleaning": [
        (names.centralWidget_TitleLabel_QLabel, 25),
        (names.centralWidget_MessageLabel_QLabel, 26),
        (names.centralWidget_OKButton_QPushButton_2, 18)
    ],

    # -----------------------------------------------------
    # Enucleation Popup
    # -----------------------------------------------------
    "enucleation": [
        (names.centralWidget_TitleLabel_QLabel, 27),
        (names.centralWidget_MessageLabel_QLabel, 28),
        (names.centralWidget_OKButton_QPushButton_2, 18)
    ]
}


# =========================================================
# MAIN
# =========================================================

def main():

    global sheet
    global workbook

    # =====================================================
    # EXCEL
    # =====================================================

    excel_file = (
        "/home/ezen/Test_Automation_Katana/"
        "suite_Gemini_Localization_Croatia/"
        "tst_GEM-LOC-CROATIA-Regularpulse/"
        "testdata/Rpulse.xlsx"
    )

   
    

    # -----------------------------------------------------
    # Load Excel workbook
    # -----------------------------------------------------

    try:
        workbook = load_workbook(
            excel_file,
            data_only=True
        )

        # Use the active sheet
        sheet = workbook.active

        

        
    except Exception as e:
        test.fail(
            "Failed to load Excel file: "
            + str(e)
        )
        return

    # =====================================================
    # LOGIN
    # =====================================================

    

    for i in range(4):

        snooze(6)

        clickButton(
            waitForObject(
                names.numpad_pushButton_8_NumpadButton
            )
        )

     
    snooze(6)

    clickButton(
        waitForObject(
            names.mainFrame_okButton_QPushButton
        )
    )

    # =====================================================
    # LANGUAGE SETTINGS
    # =====================================================

   

    snooze(6)

    clickButton(
        waitForObject(
            names.homeScreen_settingsButton_QPushButton
        )
    )

    snooze(6)

    clickTab(
        waitForObject(
            names.settingsScreen_tbSystemInfo_TabWidget
        ),
        "INFO"
    )

    snooze(6)

    clickButton(
        waitForObject(
            names.gbSystemInformation_pbLanguageSettingsUpdate_QPushButton
        )
    )

    snooze(6)

    clickButton(
        waitForObject(
            names.languagesFrame_pbCroatia_QPushButton
        )
    )

    snooze(6)

    clickButton(
        waitForObject(
            names.buttonsFrame_acceptButton_QPushButton
        )
    )

    snooze(6)

    clickButton(
        waitForObject(
            names.buttonsBar_homeButton_QPushButton
        )
    )

    # =====================================================
    # EXPERT MODE
    # =====================================================

    

    snooze(6)

    clickButton(
        waitForObject(
            names.gbExpert_expertButton_QPushButton
        )
    )

    # =====================================================
    # OPEN MODE SELECTION
    # =====================================================

    snooze(6)

    mouseClick(
        waitForObject(
            names.modeButton_nameLabel_QLabel
        ),
        98,
        24,
        Qt.NoModifier,
        Qt.LeftButton
    )

    # =====================================================
    # MAIN SCREEN VERIFICATION
    # Rows 2 - 6
    # =====================================================

    verify_group("main")

    # =====================================================
    # PULSE BUTTON VERIFICATION
    # Rows 7 - 11
    # =====================================================

    verify_group("pulse_buttons")

    # =====================================================
    # OPEN CW
    # =====================================================

    snooze(6)

    mouseClick(
        waitForObject(
            names.modeButton_nameLabel_QLabel
        ),
        69,
        21,
        Qt.NoModifier,
        Qt.LeftButton
    )

    snooze(6)

    clickButton(
        waitForObject(
            names.centralWidget_CW_QPushButton
        )
    )

    # =====================================================
    # POWER MODE UPDATE
    # Rows 12 - 15
    # =====================================================

    verify_group("power_mode")

    # =====================================================
    # CLOSE POWER MODE UPDATE
    # =====================================================

    snooze(6)

    clickButton(
        waitForObject(
            names.centralWidget_CancelButton_QPushButton_2
        )
    )

    # =====================================================
    # STANDARD PULSE
    # Rows 16 - 18
    # =====================================================

    snooze(6)

    clickButton(
        waitForObject(
            names.square_icon_QPushButton
        )
    )

    if click_until_popup_opens():
        verify_popup("standard_pulse")

    # =====================================================
    # CONTINUOUS WAVE
    # Rows 19 - 20 + OK row 18
    # =====================================================

    snooze(6)

    clickButton(
        waitForObject(
            names.cW_icon_QPushButton
        )
    )

    if click_until_popup_opens():
        verify_popup("cw")

    # =====================================================
    # MRP
    # Rows 21 - 22 + OK row 18
    # =====================================================

    snooze(6)

    clickButton(
        waitForObject(
            names.mRP_icon_QPushButton
        )
    )

    if click_until_popup_opens():
        verify_popup("mrp")

    # =====================================================
    # FRAGMENTATION
    # Rows 23 - 24 + OK row 18
    # =====================================================

    snooze(6)

    clickButton(
        waitForObject(
            names.fragmentation_icon_QPushButton
        )
    )

    if click_until_popup_opens():
        verify_popup("fragmentation")

    # =====================================================
    # CLEANING
    # Rows 25 - 26 + OK row 18
    # =====================================================

    snooze(6)

    clickButton(
        waitForObject(
            names.cleaning_icon_QPushButton
        )
    )

    if click_until_popup_opens():
        verify_popup("cleaning")

    # =====================================================
    # ENUCLEATION
    # Rows 27 - 28 + OK row 18
    # =====================================================

    snooze(6)

    clickButton(
        waitForObject(
            names.enucleation_icon_QPushButton
        )
    )

    if click_until_popup_opens():
        verify_popup("enucleation")

    # =====================================================
    # FINAL REJECT
    # =====================================================

    snooze(6)

    clickButton(
        waitForObject(
            names.buttonsFrame_rejectButton_QPushButton
        )
    )

    