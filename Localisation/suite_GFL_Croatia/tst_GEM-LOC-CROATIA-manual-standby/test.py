
# -*- coding: utf-8 -*-

import names
import openpyxl
import os
import re


TESTCASE_DIR = os.path.dirname(os.path.abspath(__file__))

EXCEL_PATH = os.path.abspath(
    os.path.join(
        TESTCASE_DIR,
        "..",
        "shared",
        "testdata",
        "Gemini Split strings.xlsx"
    )
)

SHEET_NAME = "Treatment Modes"


def normalize(text):

    if text is None:
        return ""

    text = str(text)
    text = re.sub(r"<[^>]+>", " ", text)
    text = " ".join(text.split())
    text = re.sub(r"\s*<\s*", " < ", text)
    text = re.sub(r"\s*>\s*", " > ", text)

    return text.casefold()


def get_all_texts(obj):

    texts = []

    try:
        if not obj.visible:
            return texts
    except Exception:
        pass

    try:
        value = str(obj.text).strip()

        if value:
            value = re.sub(r"<[^>]+>", " ", value)
            value = " ".join(value.split())

            if value:
                texts.append(value)

    except Exception:
        pass

    try:
        value = str(obj.title).strip()

        if value:
            value = re.sub(r"<[^>]+>", " ", value)
            value = " ".join(value.split())

            if value:
                texts.append(value)

    except Exception:
        pass

    try:
        for child in object.children(obj):
            texts.extend(get_all_texts(child))
    except Exception:
        pass

    return texts


def load_croatian_strings():

    if not os.path.isfile(EXCEL_PATH):
        test.fail("Excel file not found: " + EXCEL_PATH)
        return []

    workbook = None
    strings = []

    try:
        workbook = openpyxl.load_workbook(
            EXCEL_PATH,
            data_only=True,
            read_only=True
        )

        if SHEET_NAME not in workbook.sheetnames:
            test.fail(
                "Excel sheet not found: " + SHEET_NAME +
                ". Available sheets: " + ", ".join(workbook.sheetnames)
            )
            return []

        sheet = workbook[SHEET_NAME]

        for row in sheet.iter_rows(min_row=2, values_only=True):

            if len(row) < 4:
                continue

            english = row[1]
            croatian = row[3]

            if english is None or croatian is None:
                continue

            english = str(english).strip()
            croatian = str(croatian).strip()

            if not english or not croatian:
                continue

            strings.append({
                "english": english,
                "croatian": croatian,
                "normalized": normalize(croatian)
            })

    except Exception as e:
        test.fail("Error reading Croatian Excel strings: " + str(e))

    finally:
        if workbook is not None:
            workbook.close()

    return strings


def verify_screen(root_object, screen_name):

    excel_strings = load_croatian_strings()

    if not excel_strings:
        test.fail(
            "No Croatian strings loaded for screen: " + screen_name
        )
        return

    screen_texts = get_all_texts(root_object)
    unique_texts = []

    for text in screen_texts:
        text = text.strip()

        if text and text not in unique_texts:
            unique_texts.append(text)

    verified = []

    for screen_text in unique_texts:

        screen_normalized = normalize(screen_text)
        matches = []

        for item in excel_strings:

            excel_normalized = item["normalized"]

            # Exact match
            if screen_normalized == excel_normalized:
                matches.append(item)

            # Handle combined Squish text
            elif (
                len(excel_normalized) >= 4
                and excel_normalized in screen_normalized
            ):
                matches.append(item)

        matches.sort(
            key=lambda x: len(x["normalized"]),
            reverse=True
        )

        for item in matches:

            croatian = item["croatian"]

            if croatian in verified:
                continue

            verified.append(croatian)

            test.passes(
                item["english"] + " -> " + item["croatian"]
            )


def main():

    # LOGIN

    sendEvent(
        "QMoveEvent",
        waitForObject(names.o_ScreenSwitcher),
        62, 205, 692, 220
    )

    doubleClick(
        waitForObject(names.numpad_pushButton_8_NumpadButton),
        64, 50,
        Qt.NoModifier,
        Qt.LeftButton
    )

    clickButton(
        waitForObject(names.numpad_pushButton_8_NumpadButton)
    )

    clickButton(
        waitForObject(names.numpad_pushButton_8_NumpadButton)
    )

    clickButton(
        waitForObject(names.mainFrame_okButton_QPushButton)
    )

    # CHANGE LANGUAGE TO CROATIAN

    clickButton(
        waitForObject(names.homeScreen_settingsButton_QPushButton)
    )

    clickTab(
        waitForObject(names.settingsScreen_tbSystemInfo_TabWidget),
        "INFO"
    )

    clickButton(
        waitForObject(
            names.gbSystemInformation_pbLanguageSettingsUpdate_QPushButton
        )
    )

    clickButton(
        waitForObject(names.languagesFrame_pbCroatia_QPushButton)
    )

    clickButton(
        waitForObject(names.buttonsFrame_acceptButton_QPushButton)
    )

    clickButton(
        waitForObject(names.settingsScreen_pbBack_QPushButton)
    )

    # OPEN EXPERT SCREEN

    clickButton(
        waitForObject(names.gbExpert_expertButton_QPushButton)
    )

    test.vp("Verify Manual Screen")

    # VERIFY COMPLETE EXPERT SCREEN

    verify_screen(
        waitForObject(names.o_ScreenSwitcher),
        "Expert"
    )

    # CLICK AVERAGE POWER UNITS

    mouseClick(
        waitForObject(
            names.averagePowerWidget_unitsOfMeasureLabel_QLabel
        ),
        153, 75,
        Qt.NoModifier,
        Qt.LeftButton
    )

    verify_screen(
        waitForObject(names.o_ScreenSwitcher),
        "Expert after Average Power unit click"
    )

    # CLICK AVERAGE POWER NAME

    mouseClick(
        waitForObject(names.averagePowerWidget_nameLabel_QLabel),
        13, 64,
        Qt.NoModifier,
        Qt.LeftButton
    )

    verify_screen(
        waitForObject(names.o_ScreenSwitcher),
        "Expert after Average Power click"
    )

    # CLICK LEFT PEDAL PRESET TAB

    mouseClick(
        waitForObject(names.leftPedalWidget_wdPresetTabs_TabWidget),
        587, 18,
        Qt.NoModifier,
        Qt.LeftButton
    )

    verify_screen(
        waitForObject(names.o_ScreenSwitcher),
        "Expert after Left Pedal preset click"
    )
