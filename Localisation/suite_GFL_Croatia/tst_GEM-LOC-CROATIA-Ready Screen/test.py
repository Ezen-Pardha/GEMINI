# -*- coding: utf-8 -*-

import names
import openpyxl
from openpyxl import load_workbook


def verifyText(sheet, object_name, excel_row, description):

    
    
    # Get object
    obj = waitForObjectExists(getattr(names, object_name))

    # Get actual text from application
    actual_text = str(obj.text)

    # Get expected text from Excel
    expected_text = str(sheet.cell(row=excel_row, column=2).value)

    # Remove extra spaces and new lines
    actual_text = " ".join(actual_text.split())
    expected_text = " ".join(expected_text.split())

    
    # Compare
    test.compare(
        actual_text,
        expected_text,
        description + " text verification"
    )

    

   
    # test.imagePresent("Ready screen image 1")
    # test.log("image is present")
    # test.vp("VP1")
    # test.log("whole screen is verified")
    
 

def main():

    # ============================================================
    # EXCEL FILE
    # ============================================================

    excel_file = "/home/ezen/Test_Automation_Katana/suite_Gemini_Localization_Croatia/tst_GEM-LOC-CROATIA-Ready Screen/testdata/Ready Screen Croatia.xlsx"

    workbook = openpyxl.load_workbook(excel_file)
    sheet = workbook["Sheet1"]

   


    # ============================================================
    # LOGIN / NUMPAD
    # ============================================================

       

    for i in range(4):
        

        clickButton(
            waitForObject(
                names.numpad_pushButton_8_NumpadButton
            )
        )

        snooze(2)

    

    clickButton(
        waitForObject(
            names.mainFrame_okButton_QPushButton
        )
    )

   
    snooze(2)


    # ============================================================
    # OPEN SETTINGS
    # ============================================================

    

    clickButton(
        waitForObject(
            names.homeScreen_settingsButton_QPushButton
        )
    )

    snooze(2) 


    # ============================================================
    # OPEN SYSTEM INFORMATION
    # ============================================================

    

    clickTab(
        waitForObject(
            names.settingsScreen_tbSystemInfo_TabWidget
        ),
        "INFO"
    )

    snooze(2)


    # ============================================================
    # LANGUAGE SETTINGS
    # ============================================================

    

    clickButton(
        waitForObject(
            names.gbSystemInformation_pbLanguageSettingsUpdate_QPushButton
        )
    )

    snooze(2)

   

    clickButton(
        waitForObject(
            names.languagesFrame_pbCroatia_QPushButton
        )
    )

    snooze(2)

    clickButton(
        waitForObject(
            names.buttonsFrame_acceptButton_QPushButton
        )
    )

    
    snooze(2)


    # ============================================================
    # RETURN TO HOME
    # ============================================================

    

    clickButton(
        waitForObject(
            names.buttonsBar_homeButton_QPushButton
        )
    )

    snooze(2)


    # ============================================================
    # SOFT TISSUE QUICK START
    # ============================================================

    

    clickButton(
        waitForObject(
            names.homeScreen_pbSoftTissueQuickStart_QToolButton
        )
    )

    snooze(2)


    # ============================================================
    # SELECT LEFT TREATMENT
    # ============================================================

    

    mouseClick(
        waitForObjectItem(
            names.tabLeftMainPresets_leftPresetsView_PresetListView,
            "_1"
        ),
        369,
        68,
        Qt.NoModifier,
        Qt.LeftButton
    )

    snooze(2)


    # ============================================================
    # SELECT RIGHT TREATMENT
    # ============================================================

    

    mouseClick(
        waitForObjectItem(
            names.tabRightMainPresets_rightPresetsView_PresetListView,
            "_1"
        ),
        154,
        32,
        Qt.NoModifier,
        Qt.LeftButton
    )

    snooze(2)


    # ============================================================
    # CONTINUE
    # ============================================================

    

    clickButton(
        waitForObject(
            names.quickStartScreen_pbContinue_QPushButton
        )
    )

    snooze(2)

    clickButton(
        waitForObject(
            names.mainFrame_okButton_QPushButton_2
        )
    )

    
    snooze(2)
  


    # ============================================================
    # READY STATE
    # ============================================================

    

    clickButton(
        waitForObject(
            names.stateSwitch_readyButton_QPushButton
        )
    )

    snooze(2)

    


    # ============================================================
    # LEFT SIDE OBJECTS
    # ============================================================

    left_text_objects = [

        (
            "eNUKLEACIJA_PROSTATE_VELIKA_ENERGIJA_TabItem",
            2,
            "Left Treatment Name"
        ),

        (
            "averagePowerWidgetEmission_nameLabel_QLabel",
            3,
            "Left Average Power"
        ),

        (
            "pulseEnergyWidgetEmission_nameLabel_QLabel",
            4,
            "Left Pulse Energy"
        ),

        (
            "frequencyWidgetEmission_nameLabel_QLabel",
            5,
            "Left Frequency"
        ),

        (
            "modeButtonEmission_nameLabel_QLabel",
            6,
            "Left Pulse Mode"
        ),

        (
            "peakPowerButtonEmission_nameLabel_QLabel",
            7,
            "Left Peak Power"
        ),

        (
            "totalEnergyFrame_nameLabel_QLabel",
            8,
            "Left Total Energy"
        ),

        (
            "totalTimeFrame_nameLabel_QLabel",
            9,
            "Left Total Time"
        ),

        (
            "averagePowerWidgetEmission_unitsOfMeasureLabel_QLabel",
            10,
            "Left Power Unit"
        ),

        (
            "pulseEnergyWidgetEmission_unitsOfMeasureLabel_QLabel",
            11,
            "Left Energy Unit"
        ),

        (
            "frequencyWidgetEmission_unitsOfMeasureLabel_QLabel",
            12,
            "Left Frequency Unit"
        ),

        (
            "fiberFrame_nameLabel_QLabel",
            13,
            "Left Fiber"
        ),

        (
            "valueFrame_valueLabel_QLabel",
            14,
            "Left Fiber Value"
        ),

        (
            "bottomBar_aimingBeamButton_AimingBeamModeButton",
            15,
            "Left Aiming Beam"
        ),

        (
            "stateSwitch_standbyButton_QPushButton",
            16,
            "Left Standby"
        ),

        (
            "stateSwitch_readyButton_QPushButton",
            17,
            "Left Ready"
        )
    ]


    # ============================================================
    # VERIFY LEFT SIDE
    # ============================================================

   
    for object_name, excel_row, description in left_text_objects:

        verifyText(
            sheet,
            object_name,
            excel_row,
            description
        )

    


    # ============================================================
    # RIGHT SIDE OBJECTS
    # ============================================================

    right_text_objects = [

        (
            "eNUKLEACIJA_PROSTATE_VELIKA_ENERGIJA_TabItem_2",
            2,
            "Right Treatment Name"
        ),

        (
            "averagePowerWidgetEmission_nameLabel_QLabel_2",
            3,
            "Right Average Power"
        ),

        (
            "pulseEnergyWidgetEmission_nameLabel_QLabel_2",
            4,
            "Right Pulse Energy"
        ),

        (
            "frequencyWidgetEmission_nameLabel_QLabel_2",
            5,
            "Right Frequency"
        ),

        (
            "modeButtonEmission_nameLabel_QLabel_2",
            6,
            "Right Pulse Mode"
        ),

        (
            "peakPowerButtonEmission_nameLabel_QLabel_2",
            7,
            "Right Peak Power"
        ),

        (
            "averagePowerWidgetEmission_unitsOfMeasureLabel_QLabel_2",
            10,
            "Right Power Unit"
        ),

        (
            "pulseEnergyWidgetEmission_unitsOfMeasureLabel_QLabel_2",
            11,
            "Right Energy Unit"
        ),

        (
            "frequencyWidgetEmission_unitsOfMeasureLabel_QLabel_2",
            12,
            "Right Frequency Unit"
        )
    ]


    # ============================================================
    # VERIFY RIGHT SIDE
    # ============================================================

   

    for object_name, excel_row, description in right_text_objects:

        verifyText(
            sheet,
            object_name,
            excel_row,
            description
        )

    


    # ============================================================
    # FINISH
    # ============================================================

    

    # Click READY button again if required by your original flow
    clickButton(
        waitForObject(
            names.stateSwitch_readyButton_QPushButton
        )
    )

    snooze(2)

   
    

   
    
    
    
