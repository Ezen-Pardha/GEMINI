
# -*- coding: utf-8 -*-

import names
import openpyxl
import re
import os

TESTCASE_DIR = os.path.dirname(os.path.abspath(__file__))

EXCEL_PATH = os.path.abspath(
    os.path.join(
        TESTCASE_DIR,
        "..",
        "shared",
        "testdata",
        "Gemini Split strings.xlsx"
    )
)

SHEET_NAME = "Treatment Modes"

# Zero-based Excel column indexes
ENGLISH_COLUMN = 1       # Column B
CROATIAN_COLUMN = 3      # Column D


def normalize(text):
    if text is None:
        return ""

    text = str(text)
    text = re.sub(r"<br\s*/?>", " ", text, flags=re.IGNORECASE)
    text = re.sub(r"<[^>]+>", " ", text)
    text = text.replace("\xa0", " ")
    text = text.replace("\u200b", "")
    text = " ".join(text.split())
    text = re.sub(r"\s*<\s*", "<", text)
    text = re.sub(r"\s*>\s*", ">", text)

    return text.strip().casefold()


def get_object_text(obj):
    texts = []

    for attribute in ("text", "title", "windowTitle"):
        try:
            value = getattr(obj, attribute)

            if callable(value):
                value = value()

            if value is None:
                continue

            value = str(value).strip()
            value = re.sub(
                r"<br\s*/?>", " ", value, flags=re.IGNORECASE
            )
            value = re.sub(r"<[^>]+>", " ", value)
            value = " ".join(value.split())

            if value:
                texts.append(value)

        except Exception:
            pass

    return texts


def get_screen_candidates(obj, depth=0, visited=None):
    candidates = []

    if obj is None or depth > 25:
        return candidates

    if visited is None:
        visited = set()

    try:
        object_id = id(obj)

        if object_id in visited:
            return candidates

        visited.add(object_id)
    except Exception:
        pass

    try:
        if not obj.visible:
            return candidates
    except Exception:
        pass

    candidates.extend(get_object_text(obj))

    try:
        children = object.children(obj)
    except Exception:
        children = []

    child_texts = []

    for child in children:
        child_values = get_object_text(child)

        for value in child_values:
            candidates.append(value)
            child_texts.append(value)

    # Combine adjacent child labels for UI strings split across widgets.
    for start in range(len(child_texts)):
        combined = ""

        for end in range(start, len(child_texts)):
            combined = (
                combined + " " + child_texts[end]
                if combined else child_texts[end]
            )
            candidates.append(combined)

    for child in children:
        candidates.extend(
            get_screen_candidates(child, depth + 1, visited)
        )

    return candidates


def load_croatian_strings():
    strings = []
    workbook = None

    if not os.path.isfile(EXCEL_PATH):
        test.fail("Excel file not found: {}".format(EXCEL_PATH))
        return strings

    try:
        workbook = openpyxl.load_workbook(
            EXCEL_PATH,
            data_only=True,
            read_only=True
        )

        if SHEET_NAME not in workbook.sheetnames:
            test.fail(
                "Worksheet '{}' not found. Available sheets: {}".format(
                    SHEET_NAME,
                    ", ".join(workbook.sheetnames)
                )
            )
            return strings

        sheet = workbook[SHEET_NAME]

        for row in sheet.iter_rows(min_row=2, values_only=True):
            if len(row) <= max(ENGLISH_COLUMN, CROATIAN_COLUMN):
                continue

            english = row[ENGLISH_COLUMN]
            croatian = row[CROATIAN_COLUMN]

            if english is None or croatian is None:
                continue

            english = str(english).strip()
            croatian = str(croatian).strip()

            if not normalize(english) or not normalize(croatian):
                continue

            strings.append({
                "english": english,
                "croatian": croatian,
                "normalized": normalize(croatian)
            })

    except Exception as error:
        test.fail(
            "Unable to load Croatian translations: {}".format(error)
        )

    finally:
        if workbook is not None:
            workbook.close()

    test.log(
        "Worksheet '{}': {} translations loaded".format(
            SHEET_NAME, len(strings)
        )
    )

    return strings


def verify_screen(screen, screen_name):
    excel_strings = load_croatian_strings()

    if not excel_strings:
        test.fail("No Croatian strings loaded from Excel.")
        return

    candidates = get_screen_candidates(screen)

    try:
        parent = object.parent(screen)

        if parent is not None:
            candidates.extend(get_screen_candidates(parent))
    except Exception:
        pass

    normalized_candidates = set(
        normalize(text) for text in candidates if normalize(text)
    )

    test.log(
        "{}: {} unique screen text candidates collected".format(
            screen_name, len(normalized_candidates)
        )
    )

    matched = []
    missing = []

    for item in excel_strings:
        expected = item["normalized"]

        found = any(
            actual == expected
            or (len(expected) >= 4 and expected in actual)
            for actual in normalized_candidates
        )

        if found:
            matched.append(item)
            test.passes(
                "{} -> {}".format(
                    item["english"], item["croatian"]
                )
            )
        else:
            missing.append(item)
            test.fail(
                "Missing on screen: {} -> {}".format(
                    item["english"], item["croatian"]
                )
            )

    test.log("========== LOCALIZATION RESULTS ==========")
    test.log("Worksheet: {}".format(SHEET_NAME))
    test.log("Total Excel strings: {}".format(len(excel_strings)))
    test.log("Matched strings: {}".format(len(matched)))
    test.log("Missing strings: {}".format(len(missing)))

    if missing:
        test.fail(
            "{} validation: {} matched, {} missing.".format(
                screen_name, len(matched), len(missing)
            )
        )
    else:
        test.passes(
            "All {} Croatian strings found on {}.".format(
                len(matched), screen_name
            )
        )


def main():

    # LOGIN
    sendEvent(
        "QMoveEvent",
        waitForObject(names.o_ScreenSwitcher),
        28, 208, 534, 218
    )

    clickButton(
        waitForObject(names.numpad_pushButton_8_NumpadButton)
    )
    clickButton(
        waitForObject(names.numpad_pushButton_8_NumpadButton)
    )
    clickButton(
        waitForObject(names.numpad_pushButton_8_NumpadButton)
    )
    clickButton(
        waitForObject(names.numpad_pushButton_8_NumpadButton)
    )

    clickButton(
        waitForObject(names.mainFrame_okButton_QPushButton)
    )

    # CHANGE LANGUAGE TO CROATIAN
    clickButton(
        waitForObject(names.homeScreen_settingsButton_QPushButton)
    )

    clickTab(
        waitForObject(names.settingsScreen_tbSystemInfo_TabWidget),
        "INFO"
    )

    clickButton(
        waitForObject(
            names.gbSystemInformation_pbLanguageSettingsUpdate_QPushButton
        )
    )

    clickButton(
        waitForObject(names.languagesFrame_pbCroatia_QPushButton)
    )

    clickButton(
        waitForObject(names.buttonsFrame_acceptButton_QPushButton)
    )

    clickButton(
        waitForObject(names.settingsScreen_pbBack_QPushButton)
    )

    # OPEN SOFT TISSUE QUICK START
    clickButton(
        waitForObject(names.homeScreen_pbSoftTissueQuickStart_QToolButton)
    )

    mouseClick(
        waitForObjectItem(
            names.tabLeftMainPresets_leftPresetsView_PresetListView,
            "_1"
        ),
        163, 46, Qt.NoModifier, Qt.LeftButton
    )

    mouseClick(
        waitForObjectItem(
            names.tabRightMainPresets_rightPresetsView_PresetListView,
            "_1"
        ),
        109, 31, Qt.NoModifier, Qt.LeftButton
    )

    clickButton(
        waitForObject(names.quickStartScreen_pbContinue_QPushButton)
    )

    clickButton(
        waitForObject(names.mainFrame_okButton_QPushButton_2)
    )

    mouseClick(
        waitForObject(names.averagePowerWidget_nameLabel_QLabel),
        7, 42, Qt.NoModifier, Qt.LeftButton
    )

    mouseClick(
        waitForObject(names.averagePowerWidget_unitsOfMeasureLabel_QLabel),
        109, 40, Qt.NoModifier, Qt.LeftButton
    )

    # VERIFY CROATIAN STRINGS FROM THE SPECIFIED WORKSHEET
    screen = waitForObject(
        names.treatmentScreen_bottomBar_QFrame
    )

    verify_screen(
        screen,
        "Soft Tissue Quick Start Treatment Screen"
    )


def cleanup():
    pass
