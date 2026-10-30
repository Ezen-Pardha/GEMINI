
# -*- coding: utf-8 -*-

import names
import openpyxl
import re
import os

TESTCASE_DIR = os.path.dirname(os.path.abspath(__file__))

EXCEL_PATH = os.path.abspath(
    os.path.join(
        TESTCASE_DIR,
        "..",
        "shared",
        "testdata",
        "Gemini Split strings.xlsx"
    )
)

SHEET_NAME = "Stone-Guided"
ENGLISH_COLUMN = 2
CROATIAN_COLUMN = 4


def normalize(text):
    if text is None:
        return ""

    text = str(text)
    text = re.sub(r"<br\s*/?>", " ", text, flags=re.IGNORECASE)
    text = re.sub(r"<[^>]+>", " ", text)
    text = text.replace("\xa0", " ")
    text = text.replace("\u200b", "")
    text = text.replace("\u202f", " ")
    text = re.sub(r"\s+", " ", text)
    text = re.sub(r"\s*<\s*", "<", text)
    text = re.sub(r"\s*>\s*", ">", text)

    return text.strip().casefold()


def load_croatian_strings():
    translations = []
    workbook = None

    if not os.path.isfile(EXCEL_PATH):
        test.fail("Excel file not found: {}".format(EXCEL_PATH))
        return translations

    try:
        workbook = openpyxl.load_workbook(
            EXCEL_PATH,
            data_only=True,
            read_only=True
        )

        if SHEET_NAME not in workbook.sheetnames:
            test.fail(
                "Worksheet '{}' not found. Available sheets: {}".format(
                    SHEET_NAME,
                    ", ".join(workbook.sheetnames)
                )
            )
            return translations

        sheet = workbook[SHEET_NAME]

        for row in sheet.iter_rows(min_row=2, values_only=True):
            if len(row) < CROATIAN_COLUMN:
                continue

            trace_id = row[0]
            english = row[ENGLISH_COLUMN - 1]
            croatian = row[CROATIAN_COLUMN - 1]

            if english is None or croatian is None:
                continue

            english = str(english).strip()
            croatian = str(croatian).strip()

            if not normalize(english) or not normalize(croatian):
                continue

            translations.append({
                "trace_id": (
                    str(trace_id).strip()
                    if trace_id is not None else ""
                ),
                "english": english,
                "croatian": croatian,
                "normalized": normalize(croatian)
            })

    except Exception as error:
        test.fail(
            "Unable to load Croatian translations: {}".format(error)
        )

    finally:
        if workbook is not None:
            workbook.close()

    test.log(
        "Croatian Excel translations loaded: {}".format(
            len(translations)
        )
    )

    return translations


def get_object_text(obj):
    texts = []

    for attribute in (
        "text",
        "title",
        "windowTitle",
        "accessibleName",
        "placeholderText",
        "toolTip"
    ):
        try:
            value = getattr(obj, attribute)

            if callable(value):
                value = value()

            if value is not None:
                value = str(value).strip()

                if value and normalize(value) and value not in texts:
                    texts.append(value)

        except Exception:
            pass

    return texts


def collect_screen_text(obj, collected=None, visited=None, depth=0):
    if collected is None:
        collected = []

    if visited is None:
        visited = set()

    if obj is None or depth > 30:
        return collected

    try:
        object_id = id(obj)

        if object_id in visited:
            return collected

        visited.add(object_id)

    except Exception:
        pass

    try:
        if not object.exists(obj):
            return collected
    except Exception:
        pass

    try:
        if not obj.visible:
            return collected
    except Exception:
        pass

    for value in get_object_text(obj):
        if value not in collected:
            collected.append(value)

    try:
        children = object.children(obj)

        for child in children:
            collect_screen_text(
                child,
                collected,
                visited,
                depth + 1
            )

    except Exception:
        pass

    return collected


def build_screen_candidates(screen_texts):
    candidates = set()

    for text_value in screen_texts:
        normalized = normalize(text_value)

        if normalized:
            candidates.add(normalized)

        for line in re.split(r"[\r\n]+", str(text_value)):
            normalized_line = normalize(line)

            if normalized_line:
                candidates.add(normalized_line)

    return candidates


def verify_screen(root_object, screen_name):
    test.log("Starting verification: {}".format(screen_name))

    excel_strings = load_croatian_strings()

    if not excel_strings:
        test.fail("No Croatian translations loaded.")
        return

    screen_texts = collect_screen_text(root_object)
    screen_candidates = build_screen_candidates(screen_texts)

    test.log(
        "Visible screen text items collected: {}".format(
            len(screen_texts)
        )
    )

    for text_value in screen_texts:
        test.log("VISIBLE TEXT: {!r}".format(text_value))

    matched = []
    missing = []

    for item in excel_strings:
        expected = item["normalized"]

        found = any(
            candidate == expected
            or (len(expected) >= 4 and expected in candidate)
            for candidate in screen_candidates
        )

        if found:
            matched.append(item)
        else:
            missing.append(item)

    test.log("========== {} RESULTS ==========".format(screen_name))
    test.log("Excel translations loaded: {}".format(len(excel_strings)))
    test.log("Matched translations: {}".format(len(matched)))
    test.log("Missing translations: {}".format(len(missing)))

    for item in matched:
        test.log(
            "PASS | {} | {}".format(
                item["english"],
                item["croatian"]
            )
        )
        test.passes(
            "{} -> {}".format(
                item["english"],
                item["croatian"]
            )
        )

    for item in missing:
        test.log(
            "NOT FOUND | {} | {}".format(
                item["english"],
                item["croatian"]
            )
        )

    test.log("========================================")

    if not matched:
        test.fail(
            "No Croatian translations matched on {}.".format(
                screen_name
            )
        )
    elif missing:
        test.log(
            "Some Excel strings may not belong to the current screen."
        )
    else:
        test.passes(
            "All loaded Croatian translations matched."
        )


def main():
    test.log("STEP 1: Starting Croatian localization test")

    test.log("STEP 2: Loading Excel translations")
    translations = load_croatian_strings()

    if not translations:
        test.fail("Cannot continue without Excel translations.")
        return

    test.log("STEP 3: Logging in")

    sendEvent(
        "QMoveEvent",
        waitForObject(names.o_ScreenSwitcher),
        73,
        169,
        688,
        186
    )

    doubleClick(
        waitForObject(names.numpad_pushButton_8_NumpadButton),
        79,
        62,
        Qt.NoModifier,
        Qt.LeftButton
    )

    clickButton(
        waitForObject(names.numpad_pushButton_8_NumpadButton)
    )

    clickButton(
        waitForObject(names.numpad_pushButton_8_NumpadButton)
    )

    clickButton(
        waitForObject(names.mainFrame_okButton_QPushButton)
    )

    test.log("STEP 4: Selecting Croatian language")

    clickButton(
        waitForObject(names.homeScreen_settingsButton_QPushButton)
    )

    clickTab(
        waitForObject(names.settingsScreen_tbSystemInfo_TabWidget),
        "INFO"
    )

    clickButton(
        waitForObject(
            names.gbSystemInformation_pbLanguageSettingsUpdate_QPushButton
        )
    )

    clickButton(
        waitForObject(names.languagesFrame_pbCroatia_QPushButton)
    )

    clickButton(
        waitForObject(names.buttonsFrame_acceptButton_QPushButton)
    )

    clickButton(
        waitForObject(names.settingsScreen_pbBack_QPushButton)
    )

    test.log("STEP 5: Opening Stone Assistant")

    clickButton(
        waitForObject(names.homeScreen_pbStoneAssistant_QToolButton)
    )

    clickButton(
        waitForObject(
            names.locationFrame_kidneyPcnlLocationButton_QPushButton
        )
    )

    clickButton(
        waitForObject(
            names.flowRateFrame_lowFlowRateButton_QPushButton
        )
    )

    clickButton(
        waitForObject(
            names.hardnessFrame_mediumHardnessButton_QPushButton
        )
    )

    test.log("STEP 6: Verifying Stone Guided screen")

    stone_screen = waitForObject(
        names.stoneAssistantModeScreen_StoneAssistantModeScreen
    )

    verify_screen(stone_screen, "Stone Guided")

    test.log("STEP 7: Croatian Stone Guided verification completed")


def cleanup():
    pass
