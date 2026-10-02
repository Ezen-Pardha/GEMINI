
# -*- coding: utf-8 -*-

import names
import openpyxl
import re


# =========================================================
# EXCEL CONFIGURATION
# =========================================================

EXCEL_FILE = (
    "/home/ezen/Test_Automation_Gemini/Localisation/"
    "suite_GFL_Croatia/"
    "tst_GEM-LOC-CROATIA-PeakPower-SoftTissueCharacters/"
    "testdata/Gemini Split strings.xlsx"
)

SHEET_NAME = "Treatment Modes"

# Column B = English
ENGLISH_COLUMN = 2

# Column D = Croatian
CROATIAN_COLUMN = 4


# =========================================================
# TEXT NORMALIZATION
# =========================================================

def normalize_text(value):

    if value is None:
        return ""

    value = str(value)

    # Remove HTML tags
    value = re.sub(r"<[^>]*>", "", value)

    # Replace line breaks with spaces
    value = value.replace("\n", " ")
    value = value.replace("\r", " ")

    # Normalize spaces
    value = " ".join(value.split())

    return value.strip().casefold()


# =========================================================
# GET CROATIAN TEXT FROM EXCEL
# =========================================================

def get_croatian_text(sheet, english_text):

    search_text = normalize_text(
        english_text
    )

    for row in sheet.iter_rows(
        min_row=1,
        min_col=ENGLISH_COLUMN,
        max_col=CROATIAN_COLUMN
    ):

        # Column B = English
        english_value = row[0].value

        if normalize_text(
            english_value
        ) == search_text:

            # Column D = Croatian
            croatian_value = row[
                CROATIAN_COLUMN - ENGLISH_COLUMN
            ].value

            if croatian_value is None:

                test.fail(
                    "Croatian translation is empty for: "
                    + english_text
                )

                return ""

            return str(
                croatian_value
            ).strip()

    # test.fail(
    #     "English text not found in Excel Column B: "
    #     + english_text
    # )

    return ""


# =========================================================
# VERIFY APPLICATION TEXT AGAINST CROATIAN EXCEL
# =========================================================

def verify_text(
    sheet,
    object_name,
    english_text,
    description
):

    # Wait before every verification
    snooze(6)

    obj = waitForObjectExists(
        getattr(names, object_name)
    )

    actual_text = str(
        obj.text
    ).strip()

    expected_text = get_croatian_text(
        sheet,
        english_text
    )

    actual_normalized = normalize_text(
        actual_text
    )

    expected_normalized = normalize_text(
        expected_text
    )

    # test.log(
    #     "{} | Expected Croatian: [{}] | Actual: [{}]".format(
    #         description,
    #         expected_text,
    #         actual_text
    #     )
    # )

    test.compare(
        actual_normalized,
        expected_normalized,
        description + " - Croatian text verification"
    )


# =========================================================
# MAIN
# =========================================================

def main():

    # =====================================================
    # LOAD EXCEL
    # =====================================================

   
    workbook = openpyxl.load_workbook(
        EXCEL_FILE,
        data_only=True,
        read_only=True
    )

    if SHEET_NAME not in workbook.sheetnames:

        test.fail(
            "Excel sheet not found: "
            + SHEET_NAME
        )

        workbook.close()
        return

    sheet = workbook[
        SHEET_NAME
    ]

    


    # =====================================================
    # LOGIN
    # =====================================================

    clickButton(
        waitForObject(
            names.numpad_pushButton_8_NumpadButton
        )
    )

    clickButton(
        waitForObject(
            names.numpad_pushButton_8_NumpadButton
        )
    )

    clickButton(
        waitForObject(
            names.numpad_pushButton_8_NumpadButton
        )
    )

    clickButton(
        waitForObject(
            names.numpad_pushButton_8_NumpadButton
        )
    )

    clickButton(
        waitForObject(
            names.mainFrame_okButton_QPushButton
        )
    )


    # =====================================================
    # OPEN SETTINGS
    # =====================================================

    clickButton(
        waitForObject(
            names.homeScreen_settingsButton_QPushButton
        )
    )


    # =====================================================
    # OPEN SYSTEM INFORMATION
    # =====================================================

    clickTab(
        waitForObject(
            names.settingsScreen_tbSystemInfo_TabWidget
        ),
        "INFO"
    )


    # =====================================================
    # OPEN LANGUAGE SETTINGS
    # =====================================================

    clickButton(
        waitForObject(
            names.gbSystemInformation_pbLanguageSettingsUpdate_QPushButton
        )
    )


    # =====================================================
    # SELECT CROATIAN LANGUAGE
    # =====================================================

    clickButton(
        waitForObject(
            names.languagesFrame_pbCroatia_QPushButton
        )
    )


    # =====================================================
    # ACCEPT LANGUAGE
    # =====================================================

    clickButton(
        waitForObject(
            names.buttonsFrame_acceptButton_QPushButton
        )
    )


    # =====================================================
    # BACK TO HOME
    # =====================================================

    clickButton(
        waitForObject(
            names.settingsScreen_pbBack_QPushButton
        )
    )


    # =====================================================
    # NAVIGATE TO SOFT TISSUE QUICK START
    # =====================================================

    clickButton(
        waitForObject(
            names.homeScreen_pbSoftTissueQuickStart_QToolButton
        )
    )


    # =====================================================
    # SELECT LEFT PRESET
    # =====================================================

    mouseClick(
        waitForObjectItem(
            names.tabLeftMainPresets_leftPresetsView_PresetListView,
            "_1"
        ),
        289,
        45,
        Qt.NoModifier,
        Qt.LeftButton
    )


    # =====================================================
    # SELECT RIGHT PRESET
    # =====================================================

    mouseClick(
        waitForObjectItem(
            names.tabRightMainPresets_rightPresetsView_PresetListView,
            "_1"
        ),
        173,
        22,
        Qt.NoModifier,
        Qt.LeftButton
    )


    # =====================================================
    # CONTINUE
    # =====================================================

    clickButton(
        waitForObject(
            names.quickStartScreen_pbContinue_QPushButton
        )
    )

    clickButton(
        waitForObject(
            names.mainFrame_okButton_QPushButton_2
        )
    )


    # =====================================================
    # OPEN PEAK POWER
    # =====================================================

    mouseClick(
        waitForObject(
            names.peakPowerButton_iconLabel_QLabel
        ),
        27,
        38,
        Qt.NoModifier,
        Qt.LeftButton
    )


    # =====================================================
    # VERIFY PEAK POWER - CROATIAN
    # =====================================================

    verify_text(
        sheet,
        "centralWidget_lbTitle_QLabel",
        "Peak Power",
        "Peak Power title"
    )

    verify_text(
        sheet,
        "centralWidget_pw500_QPushButton",
        "High",
        "High button"
    )

    verify_text(
        sheet,
        "centralWidget_pw250_QPushButton",
        "Medium",
        "Medium button"
    )

    verify_text(
        sheet,
        "centralWidget_pw125_QPushButton",
        "Low",
        "Low button"
    )

    verify_text(
        sheet,
        "buttonsFrame_rejectButton_QPushButton",
        "Cancel",
        "Cancel button"
    )


    # =====================================================
    # EXIT PEAK POWER
    # =====================================================

    clickButton(
        waitForObject(
            names.buttonsFrame_rejectButton_QPushButton
        )
    )

    clickButton(
        waitForObject(
            names.buttonsBar_homeButton_QPushButton_2
        )
    )


    # =====================================================
    # SOFT TISSUE ASSISTANT
    # =====================================================

    clickButton(
        waitForObject(
            names.homeScreen_pbSoftTissueAssistant_QToolButton
        )
    )


    # =====================================================
    # SELECT INCISION
    # =====================================================

    clickButton(
        waitForObject(
            names.gbSoftTissueProcedures_pbIncision_QPushButton
        )
    )


    # =====================================================
    # CONTINUE
    # =====================================================

    clickButton(
        waitForObject(
            names.softTissueAssistantModeScreen_pbContinue_QPushButton
        )
    )


    # =====================================================
    # OPEN TREATMENT CHARACTERISTICS
    # =====================================================

    clickButton(
        waitForObject(
            names.treatmentCharacteristicsBar_pbEdit_QPushButton
        )
    )

    mouseClick(
        waitForObject(
            names.centralWidget_QFrame
        ),
        495,
        191,
        Qt.NoModifier,
        Qt.LeftButton
    )

    mouseClick(
        waitForObject(
            names.centralWidget_titleLabel_QLabel
        ),
        0,
        24,
        Qt.NoModifier,
        Qt.LeftButton
    )


    # =====================================================
    # VERIFY SOFT TISSUE CHARACTERISTICS - CROATIAN
    # =====================================================

    verify_text(
        sheet,
        "centralWidget_titleLabel_QLabel",
        "Soft Tissue Characteristics",
        "Soft Tissue Characteristics title"
    )

    verify_text(
        sheet,
        "centralWidget_titleLabel_QLabel_2",
        "Procedure",
        "Procedure label"
    )

    verify_text(
        sheet,
        "centralWidget_valueLabel_QLabel",
        "Incision",
        "Procedure value"
    )

    verify_text(
        sheet,
        "centralWidget_rejectButton_QPushButton",
        "Edit",
        "Edit button"
    )

    verify_text(
        sheet,
        "centralWidget_acceptButton_QPushButton",
        "Accept",
        "Accept button"
    )


    # =====================================================
    # VERIFICATION POINT
    # =====================================================

    test.vp(
        "SoftTissueCharacters"
    )


    # =====================================================
    # CLOSE EXCEL
    # =====================================================

    workbook.close()

    test.log(
        "Croatian Peak Power and Soft Tissue "
        "Characteristics verification completed"
    )

