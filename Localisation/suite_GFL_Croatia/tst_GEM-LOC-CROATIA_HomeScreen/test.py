
# -*- coding: utf-8 -*-

import names
import openpyxl
import os
import re


TESTCASE_DIR = os.path.dirname(os.path.abspath(__file__))

EXCEL_PATH = os.path.abspath(
    os.path.join(
        TESTCASE_DIR,
        "..",
        "shared",
        "testdata",
        "Gemini Split strings.xlsx"
    )
)

SHEET_NAME = "Home Screen"


def normalize(text):
    if text is None:
        return ""

    text = str(text).strip()
    text = re.sub(r"<[^>]+>", "", text)
    text = text.replace("\xa0", " ")
    text = " ".join(text.split())

    return text.casefold()


def get_all_texts(obj):
    texts = []

    try:
        text = str(obj.text).strip()

        if text:
            texts.append(text)

    except Exception:
        pass

    try:
        children = object.children(obj)

        for child in children:
            texts.extend(get_all_texts(child))

    except Exception:
        pass

    return texts


def main():

    # Login
    doubleClick(
        waitForObject(names.numpad_pushButton_8_NumpadButton),
        68, 52,
        Qt.NoModifier,
        Qt.LeftButton
    )

    doubleClick(
        waitForObject(names.numpad_pushButton_8_NumpadButton),
        116, 46,
        Qt.NoModifier,
        Qt.LeftButton
    )

    clickButton(
        waitForObject(names.mainFrame_okButton_QPushButton)
    )

    # Change language to Croatian
    clickButton(
        waitForObject(names.homeScreen_settingsButton_QPushButton)
    )

    clickTab(
        waitForObject(names.settingsScreen_tbSystemInfo_TabWidget),
        "INFO"
    )

    clickButton(
        waitForObject(
            names.gbSystemInformation_pbLanguageSettingsUpdate_QPushButton
        )
    )

    clickButton(
        waitForObject(names.languagesFrame_pbCroatia_QPushButton)
    )

    clickButton(
        waitForObject(names.buttonsFrame_acceptButton_QPushButton)
    )

    clickButton(
        waitForObject(names.settingsScreen_pbBack_QPushButton)
    )

    test.vp("Verify Home Screen1")

    # Home screen navigation
    mouseClick(
        waitForObject(names.homeScreen_HomeScreen),
        637, 179,
        Qt.NoModifier,
        Qt.LeftButton
    )

    mouseClick(
        waitForObject(names.homeScreen_gbExpert_QFrame),
        169, 60,
        Qt.NoModifier,
        Qt.LeftButton
    )

    mouseClick(
        waitForObject(names.topBar_buttonsBar_ButtonsBar),
        371, 25,
        Qt.NoModifier,
        Qt.LeftButton
    )

    # Check Excel file
    if not os.path.isfile(EXCEL_PATH):
        test.fail("Excel file not found: " + EXCEL_PATH)
        return

    workbook = None

    try:
        # Read Excel workbook
        workbook = openpyxl.load_workbook(
            EXCEL_PATH,
            data_only=True,
            read_only=True
        )

        # Check and select the specified worksheet
        if SHEET_NAME not in workbook.sheetnames:
            test.fail(
                "Excel sheet not found: " + SHEET_NAME +
                ". Available sheets: " + ", ".join(workbook.sheetnames)
            )
            return

        sheet = workbook[SHEET_NAME]

        # Get all text from Home Screen
        home = waitForObject(names.homeScreen_HomeScreen)
        actual_texts = get_all_texts(home)

        normalized_actual_texts = {
            normalize(text)
            for text in actual_texts
            if normalize(text)
        }

        # Verify English -> Croatian
        for row in sheet.iter_rows(min_row=2, values_only=True):

            if len(row) < 3:
                continue

            english_value = row[0]
            croatian_value = row[3]

            if english_value is None or croatian_value is None:
                continue

            english = str(english_value).strip()
            croatian = str(croatian_value).strip()

            expected = normalize(croatian)

            if not expected:
                continue

            if expected in normalized_actual_texts:
                test.passes(
                    "PASS: " + english + " -> " + croatian
                )
            else:
                test.fail(
                    "FAIL: " + english + " -> " + croatian
                )

    except Exception as e:
        test.fail("Localization validation error: " + str(e))

    finally:
        if workbook is not None:
            workbook.close()
