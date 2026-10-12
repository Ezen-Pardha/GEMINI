# :::writing{variant="document" id="58321" title="Modified Preset Test Script - Excel Driven Localization"}
# -*- coding: utf-8 -*-

import names
import openpyxl
import re


EXCEL_PATH = "/home/ezen/Test_Automation_Katana/suite_Gemini_Localization_Croatia/tst_GEM-LOC-CROATIA-Userpresets-Add-Edit-Delete/testdata/Gemini strings.xlsx"

 
def normalize_localization(text):

    if text is None:
        return ""

    text = str(text)

    # Remove HTML tags
    text = re.sub(r"<[^>]+>", "", text)

    # Remove escaped/newline characters
    text = text.replace("\\n", "")
    text = text.replace("\n", "")
    text = text.replace("\r", "")

    # Remove all whitespace so:
    # PROSJEČNA
    # SNAGA
    # becomes PROSJEČNASNAGA
    text = re.sub(r"\s+", "", text)

    return text.casefold()




def load_croatian_strings():

    
    

    workbook = openpyxl.load_workbook(
        EXCEL_PATH,
        data_only=True
    )

    sheet = workbook.active

    excel_strings = []

    for row_number in range(2, sheet.max_row + 1):

        english = sheet.cell(
            row=row_number,
            column=2
        ).value

        croatian = sheet.cell(
            row=row_number,
            column=9
        ).value

        if english is not None and croatian is not None:

            excel_strings.append({
                "english": str(english),
                "croatian": str(croatian),
                "row": row_number
            })

    workbook.close()

    

    return excel_strings




def verify_excel_text(object_name, excel_strings, description=""):

    actual_text = str(
        waitForObjectExists(object_name).text
    )

    actual_normalized = normalize_localization(actual_text)

    matches = []

    for item in excel_strings:

        excel_croatian = item["croatian"]

        excel_normalized = normalize_localization(
            excel_croatian
        )

        if actual_normalized == excel_normalized:

            matches.append(item)

    if len(matches) == 0:

        test.fail(
            "No matching Croatian translation found in Excel "
            "for UI text: {}".format(actual_text)
        )

        return

    # Use the first exact match
    match = matches[0]

    

    

    if description:

        test.compare(
            actual_normalized,
            normalize_localization(match["croatian"]),
            "{} | English: {} | Excel Row: {}".format(
                description,
                match["english"],
                match["row"]
            )
        )

    else:

        test.compare(
            actual_normalized,
            normalize_localization(match["croatian"]),
            "Localization verified | English: {} | Excel Row: {}".format(
                match["english"],
                match["row"]
            )
        )



def main():

    excel_strings = load_croatian_strings()


  

    clickButton(
        waitForObject(
            names.numpad_pushButton_8_NumpadButton
        )
    )

    clickButton(
        waitForObject(
            names.numpad_pushButton_8_NumpadButton
        )
    )

    clickButton(
        waitForObject(
            names.numpad_pushButton_8_NumpadButton
        )
    )

    clickButton(
        waitForObject(
            names.numpad_pushButton_8_NumpadButton
        )
    )

    clickButton(
        waitForObject(
            names.mainFrame_okButton_QPushButton
        )
    )



    clickButton(
        waitForObject(
            names.homeScreen_settingsButton_QPushButton
        )
    )




    clickTab(
        waitForObject(
            names.settingsScreen_tbSystemInfo_TabWidget
        ),
        "INFO"
    )

    clickButton(
        waitForObject(
            names.gbSystemInformation_pbLanguageSettingsUpdate_QPushButton
        )
    )

    mouseClick(
        waitForObject(
            names.centralWidget_languagesFrame_QFrame
        ),
        747,
        10,
        Qt.NoModifier,
        Qt.LeftButton
    )

    clickButton(
        waitForObject(
            names.languagesFrame_pbCroatia_QPushButton
        )
    )

    clickButton(
        waitForObject(
            names.buttonsFrame_acceptButton_QPushButton
        )
    )


 

    clickTab(
        waitForObject(
            names.settingsScreen_tbSystemInfo_TabWidget
        ),
        "ZADANE KORISN\nIČKE POSTAVKE"
    )

    clickButton(
        waitForObject(
            names.settingsScreen_addButton_QPushButton
        )
    )




    verify_excel_text(
        names.centralWidget_lbOperation_QLabel,
        excel_strings,
        "Add Default Preset title"
    )

    verify_excel_text(
        names.centralWidget_pbName_QPushButton,
        excel_strings,
        "Preset name field"
    )


    verify_excel_text(
        names.centralWidget_pbName_QPushButton,
        excel_strings,
        "Preset name field - second verification"
    )

    verify_excel_text(
        names.averagePowerWidget_nameLabel_QLabel,
        excel_strings,
        "Average Power label"
    )

    verify_excel_text(
        names.pulseEnergyWidget_nameLabel_QLabel,
        excel_strings,
        "Pulse Energy label"
    )

    verify_excel_text(
        names.frequencyWidget_nameLabel_QLabel,
        excel_strings,
        "Frequency label"
    )




    test.compare(
        str(
            waitForObjectExists(
                names.averagePowerWidget_unitsOfMeasureLabel_QLabel
            ).text
        ),
        "W",
        "Average Power unit"
    )

    test.compare(
        str(
            waitForObjectExists(
                names.pulseEnergyWidget_unitsOfMeasureLabel_QLabel
            ).text
        ),
        "J",
        "Pulse Energy unit"
    )

    test.compare(
        str(
            waitForObjectExists(
                names.frequencyWidget_unitsOfMeasureLabel_QLabel
            ).text
        ),
        "Hz",
        "Frequency unit"
    )


 

    verify_excel_text(
        names.modeButton_nameLabel_QLabel_3,
        excel_strings,
        "Standard Pulse label"
    )

    verify_excel_text(
        names.peakPowerButton_nameLabel_QLabel_2,
        excel_strings,
        "Peak Power label"
    )

    verify_excel_text(
        names.pulseModeWidget_saveButton_QPushButton,
        excel_strings,
        "Add button"
    )




    mouseClick(
        waitForObject(
            names.averagePowerWidget_unitsOfMeasureLabel_QLabel
        ),
        161,
        76,
        Qt.NoModifier,
        Qt.LeftButton
    )

    clickButton(
        waitForObject(
            names.centralWidget_pbName_QPushButton
        )
    )

    sendEvent(
        "QMoveEvent",
        waitForObject(names.o_ScreenSwitcher),
        92,
        83,
        1137,
        110
    )

    mouseDrag(
        waitForObject(
            names.keyboardPopup_KeyboardPopup
        ),
        952,
        49,
        7,
        34,
        1,
        Qt.LeftButton
    )

    type(
        waitForObject(
            names.keyboardPopup_wdEdit_WdLineEdit
        ),
        "Preset One"
    )

    mouseDrag(
        waitForObject(
            names.keyboardPopup_wdEdit_WdLineEdit
        ),
        450,
        21,
        -151,
        10,
        1,
        Qt.LeftButton
    )

    type(
        waitForObject(
            names.keyboardPopup_wdEdit_WdLineEdit
        ),
        "<Del>"
    )

    clickButton(
        waitForObject(
            names.keyboardPopup_p_KeyButton
        )
    )

    clickButton(
        waitForObject(
            names.keyboardPopup_r_KeyButton
        )
    )

    clickButton(
        waitForObject(
            names.keyboardPopup_e_KeyButton
        )
    )

    clickButton(
        waitForObject(
            names.keyboardPopup_s_KeyButton
        )
    )

    clickButton(
        waitForObject(
            names.keyboardPopup_e_KeyButton
        )
    )

    clickButton(
        waitForObject(
            names.keyboardPopup_t_KeyButton
        )
    )

    clickButton(
        waitForObject(
            names.keyboardPopup_KeyButton_2
        )
    )

    clickButton(
        waitForObject(
            names.keyboardPopup_o_KeyButton
        )
    )

    clickButton(
        waitForObject(
            names.keyboardPopup_n_KeyButton
        )
    )

    clickButton(
        waitForObject(
            names.keyboardPopup_e_KeyButton
        )
    )

    mouseClick(
        waitForObject(
            names.keyboardPopup_lbText_QLabel
        ),
        157,
        25,
        Qt.NoModifier,
        Qt.LeftButton
    )




    clickButton(
        waitForObject(
            names.pulseModeWidget_saveButton_QPushButton
        )
    )

    sendEvent(
        "QMoveEvent",
        waitForObject(names.o_ScreenSwitcher),
        95,
        59,
        793,
        84
    )




    verify_excel_text(
        names.tabBarPresetsForSettings_KAMEN_TabItem,
        excel_strings,
        "Stone preset tab"
    )

    verify_excel_text(
        names.tabBarPresetsForSettings_MEKO_TKIVO_TabItem,
        excel_strings,
        "Soft Tissue preset tab"
    )


 

    clickButton(
        waitForObject(
            names.settingsScreen_editButton_QPushButton
        )
    )

    mouseClick(
        waitForObject(
            names.tabSettingsPresetsUserStonePresets_presetListView_PresetListView
        ),
        462,
        80,
        Qt.NoModifier,
        Qt.LeftButton
    )




    verify_excel_text(
        names.centralWidget_lbOperation_QLabel,
        excel_strings,
        "Edit Default Preset title"
    )


  

    test.compare(
        str(
            waitForObjectExists(
                names.centralWidget_pbName_QPushButton
            ).text
        ),
        "preset one",
        "Preset name"
    )


 

    verify_excel_text(
        names.pulseEnergyWidget_nameLabel_QLabel,
        excel_strings,
        "Pulse Energy label - Edit"
    )

    verify_excel_text(
        names.frequencyWidget_nameLabel_QLabel,
        excel_strings,
        "Frequency label - Edit"
    )




    test.compare(
        str(
            waitForObjectExists(
                names.averagePowerWidget_unitsOfMeasureLabel_QLabel
            ).text
        ),
        "W",
        "Average Power unit - Edit"
    )

    test.compare(
        str(
            waitForObjectExists(
                names.pulseEnergyWidget_unitsOfMeasureLabel_QLabel
            ).text
        ),
        "J",
        "Pulse Energy unit - Edit"
    )

    test.compare(
        str(
            waitForObjectExists(
                names.frequencyWidget_unitsOfMeasureLabel_QLabel
            ).text
        ),
        "Hz",
        "Frequency unit - Edit"
    )



    verify_excel_text(
        names.modeButton_nameLabel_QLabel_3,
        excel_strings,
        "Standard Pulse label - Edit"
    )

    verify_excel_text(
        names.peakPowerButton_nameLabel_QLabel_2,
        excel_strings,
        "Peak Power label - Edit"
    )

    verify_excel_text(
        names.pulseModeWidget_saveButton_QPushButton,
        excel_strings,
        "Update button"
    )

    verify_excel_text(
        names.averagePowerWidget_nameLabel_QLabel,
        excel_strings,
        "Average Power label - Edit"
    )




    mouseClick(
        waitForObject(
            names.centralWidget_QFrame
        ),
        352,
        14,
        Qt.NoModifier,
        Qt.LeftButton
    )

 



    clickButton(
        waitForObject(
            names.pulseModeWidget_saveButton_QPushButton
        )
    )


    
    test.vp("Preset Screen")
    
    clickButton(
        waitForObject(
            names.settingsScreen_deletePresetButton_QPushButton
        )
    )

    mouseClick(
        waitForObject(
            names.tabSettingsPresetsUserStonePresets_presetListView_PresetListView
        ),
        494,
        79,
        Qt.NoModifier,
        Qt.LeftButton
    )




    verify_excel_text(
        names.centralWidget_lbDialogTitle_QLabel,
        excel_strings,
        "Delete preset dialog title"
    )


    # Preset name is dynamic test data
    test.compare(
        str(
            waitForObjectExists(
                names.centralWidget_questionLabel_QLabel
            ).text
        ),
        "preset one",
        "Preset name in delete dialog"
    )




    verify_excel_text(
        names.buttonsFrame_rejectButton_QPushButton,
        excel_strings,
        "Cancel button"
    )

    verify_excel_text(
        names.centralWidget_acceptButton_QPushButton,
        excel_strings,
        "Delete button"
    )


    test.compare(
        str(
            waitForObjectExists(
                names.infoFrame_extInfoLabel_QLabel
            ).text
        ),
        "6 J     3 Hz     18 W",
        "Preset parameter values"
    )


  

    clickButton(
        waitForObject(
            names.centralWidget_acceptButton_QPushButton
        )
    )



    verify_excel_text(
        names.settingsScreen_pbBack_QPushButton,
        excel_strings,
        "Back button"
    )
