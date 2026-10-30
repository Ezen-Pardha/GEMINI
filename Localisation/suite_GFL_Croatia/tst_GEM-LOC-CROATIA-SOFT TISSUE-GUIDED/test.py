
# -*- coding: utf-8 -*-

import names
import openpyxl
import re
import os

TESTCASE_DIR = os.path.dirname(os.path.abspath(__file__))

EXCEL_PATH = os.path.abspath(
    os.path.join(
        TESTCASE_DIR,
        "..",
        "shared",
        "testdata",
        "Gemini Split strings.xlsx"
    )
)

SHEET_NAME = "Stone-Guided"
ENGLISH_COLUMN = 1
CROATIAN_COLUMN = 4


def normalize(text):
    if text is None:
        return ""

    text = str(text)
    text = re.sub(r"<br\s*/?>", " ", text, flags=re.IGNORECASE)
    text = re.sub(r"<[^>]+>", "", text)
    text = text.replace("\xa0", " ")
    text = text.replace("\u200b", "")
    text = text.replace("\u202f", " ")
    text = text.replace("\u2013", "-")
    text = text.replace("\u2014", "-")
    text = re.sub(r"\s+", " ", text)

    return text.strip().casefold()


def load_croatian_strings():
    translations = {}
    workbook = None

    if not os.path.isfile(EXCEL_PATH):
        test.fail("Excel file not found: {}".format(EXCEL_PATH))
        return translations

    try:
        workbook = openpyxl.load_workbook(
            EXCEL_PATH,
            data_only=True,
            read_only=True
        )

        if SHEET_NAME not in workbook.sheetnames:
            test.fail(
                "Worksheet '{}' not found. Available sheets: {}".format(
                    SHEET_NAME,
                    ", ".join(workbook.sheetnames)
                )
            )
            return translations

        sheet = workbook[SHEET_NAME]

        for row in sheet.iter_rows(min_row=2, values_only=True):
            if len(row) < CROATIAN_COLUMN:
                continue

            english = row[ENGLISH_COLUMN - 1]
            croatian = row[CROATIAN_COLUMN - 1]

            if english is None or croatian is None:
                continue

            english = str(english).strip()
            croatian = str(croatian).strip()

            if normalize(english) and normalize(croatian):
                translations[english] = croatian

    except Exception as error:
        test.fail(
            "Unable to read Croatian Excel translations: {}".format(error)
        )

    finally:
        if workbook is not None:
            workbook.close()

    test.log(
        "Croatian Excel translations loaded: {}".format(
            len(translations)
        )
    )

    return translations


def get_object_text(obj):
    texts = []

    for attribute in (
        "text",
        "windowTitle",
        "placeholderText",
        "accessibleName",
        "displayText",
        "toolTip"
    ):
        try:
            value = getattr(obj, attribute)

            if callable(value):
                value = value()

            if value is not None:
                value = str(value).strip()

                if value and normalize(value):
                    texts.append(value)

        except Exception:
            pass

    return texts


def collect_screen_text(obj, collected=None, visited=None, depth=0):
    if collected is None:
        collected = []

    if visited is None:
        visited = set()

    if obj is None or depth > 30:
        return collected

    try:
        object_id = id(obj)

        if object_id in visited:
            return collected

        visited.add(object_id)

    except Exception:
        pass

    try:
        if not object.exists(obj):
            return collected
    except Exception:
        pass

    try:
        if not obj.visible:
            return collected
    except Exception:
        pass

    for value in get_object_text(obj):
        if value not in collected:
            collected.append(value)

    try:
        children = object.children(obj)

        for child in children:
            collect_screen_text(
                child,
                collected,
                visited,
                depth + 1
            )

    except Exception:
        pass

    return collected


def build_screen_candidates(screen_texts):
    candidates = set()

    for text_value in screen_texts:
        cleaned = normalize(text_value)

        if cleaned:
            candidates.add(cleaned)

        for line in re.split(r"[\r\n]+", str(text_value)):
            cleaned_line = normalize(line)

            if cleaned_line:
                candidates.add(cleaned_line)

    return candidates


def validate_croatian_strings(translations, screen_candidates):
    matched_strings = []
    missing_strings = []

    for english, croatian in translations.items():
        expected = normalize(croatian)

        if not expected:
            continue

        found = False

        for actual in screen_candidates:
            if actual == expected:
                found = True
                break

            if len(expected) >= 4 and expected in actual:
                found = True
                break

        if found:
            matched_strings.append((english, croatian))
        else:
            missing_strings.append((english, croatian))

    test.log("========== CROATIAN LOCALIZATION RESULTS ==========")
    test.log("Total Excel strings: {}".format(len(translations)))
    test.log("Matched strings: {}".format(len(matched_strings)))
    test.log("Missing strings: {}".format(len(missing_strings)))

    for english, croatian in matched_strings:
        test.log("PASS | {} | {}".format(english, croatian))

    for english, croatian in missing_strings:
        test.log("FAIL | {} | {}".format(english, croatian))

    test.log("===================================================")

    if missing_strings:
        test.fail(
            "Croatian localization validation completed: "
            "{} matched, {} missing.".format(
                len(matched_strings),
                len(missing_strings)
            )
        )
    else:
        test.passes(
            "All {} Croatian Excel strings were found.".format(
                len(matched_strings)
            )
        )


def main():
    croatian_translations = load_croatian_strings()

    if not croatian_translations:
        test.fail("No Croatian translations were loaded.")
        return

    # Login
    clickButton(waitForObject(names.numpad_pushButton_8_NumpadButton))
    clickButton(waitForObject(names.numpad_pushButton_8_NumpadButton))
    clickButton(waitForObject(names.numpad_pushButton_8_NumpadButton))
    clickButton(waitForObject(names.numpad_pushButton_8_NumpadButton))
    clickButton(waitForObject(names.mainFrame_okButton_QPushButton))

    # The application must already be configured for Croatian.
    # Open Stone Assistant.
    clickButton(
        waitForObject(names.homeScreen_pbStoneAssistant_QToolButton)
    )

    # Select Stone Guided options.
    clickButton(
        waitForObject(
            names.locationFrame_kidneyPcnlLocationButton_QPushButton
        )
    )

    clickButton(
        waitForObject(
            names.flowRateFrame_mediumFlowRateButton_QPushButton
        )
    )

    clickButton(
        waitForObject(
            names.hardnessFrame_largeHardnessButton_QPushButton
        )
    )

    stone_screen = waitForObject(
        names.stoneAssistantModeScreen_StoneAssistantModeScreen
    )

    test.log("Screen: Stone Guided")

    screen_texts = collect_screen_text(stone_screen)

    test.log(
        "Screen text items collected: {}".format(len(screen_texts))
    )

    for text_value in screen_texts:
        test.log("VISIBLE TEXT: {!r}".format(text_value))

    screen_candidates = build_screen_candidates(screen_texts)

    test.log(
        "Screen text candidates collected: {}".format(
            len(screen_candidates)
        )
    )

    validate_croatian_strings(
        croatian_translations,
        screen_candidates
    )


def cleanup():
    pass
