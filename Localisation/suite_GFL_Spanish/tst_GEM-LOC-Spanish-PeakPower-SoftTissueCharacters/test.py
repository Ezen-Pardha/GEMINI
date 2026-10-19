# -*- coding: utf-8 -*-

import names
import re
import unicodedata
from openpyxl import load_workbook
import os 

TESTCASE_DIR = os.path.dirname(os.path.abspath(__file__))
 
EXCEL_PATH = os.path.abspath(
    os.path.join( TESTCASE_DIR,
        "..",
        "shared",
        "testdata",
        "Gemini Split strings.xlsx"
    )
)


def normalize(text):

    if text is None:
        return ""

    text = str(text)

    text = re.sub(
        r"<br\s*/?>",
        "",
        text,
        flags=re.IGNORECASE
    )

    text = re.sub(
        r"<[^>]+>",
        "",
        text
    )

    try:
        text = unicodedata.normalize(
            "NFKC",
            text
        )
    except Exception:
        pass

    text = (
        text
        .replace("\xa0", "")
        .replace("\u2007", "")
        .replace("\u202f", "")
        .replace("\n", "")
        .replace("\r", "")
        .replace("\t", "")
        .replace("™", "")
        .replace("®", "")
        .replace("©", "")
    )

    text = re.sub(
        r"\s+",
        "",
        text
    )

    return text.casefold()


def load_spanish_excel(EXCEL_PATH):

    try:
        workbook = load_workbook(
            EXCEL_PATH,
            data_only=True,
            read_only=True
        )

    except Exception as e:
        test.fail(
            "Unable to load Spanish localization Excel: "
            + str(e)
        )
        return None, None, {}

    sheet_name = "Treatment Modes"

    if sheet_name not in workbook.sheetnames:
        test.fail(
            "Worksheet '"
            + sheet_name
            + "' does not exist in Excel file"
        )
        workbook.close()
        return None, None, {}

    sheet = workbook[sheet_name]

    english_column = 2
    spanish_column = 24
    spanish_index = spanish_column - english_column

    spanish_data = {}

    for row in sheet.iter_rows(
        min_row=1,
        min_col=english_column,
        max_col=spanish_column,
        values_only=True
    ):

        english_value = row[0]

        if english_value is None:
            continue

        english_key = normalize(english_value)

        if english_key == "":
            continue

        if spanish_index >= len(row):
            continue

        spanish_value = row[spanish_index]

        if spanish_value is None:
            spanish_value = ""
        else:
            spanish_value = str(spanish_value).strip()

        if english_key in spanish_data:

            existing_values = spanish_data[english_key].split("|||")

            duplicate_found = any(
                normalize(existing_translation) == normalize(spanish_value)
                for existing_translation in existing_values
            )

            if not duplicate_found:
                spanish_data[english_key] += "|||" + spanish_value

        else:
            spanish_data[english_key] = spanish_value

    return workbook, sheet, spanish_data


def get_spanish_text(spanish_data, english_text):

    key = normalize(english_text)

    if key not in spanish_data:
        return ""

    return spanish_data[key]


def get_object_text(obj):

    try:
        value = obj.text
        if value is not None:
            return str(value)
    except Exception:
        pass

    try:
        value = obj.plainText
        if value is not None:
            return str(value)
    except Exception:
        pass

    try:
        value = obj.title
        if value is not None:
            return str(value)
    except Exception:
        pass

    try:
        value = obj.windowTitle
        if value is not None:
            return str(value)
    except Exception:
        pass

    try:
        value = obj.accessibleName
        if value is not None:
            return str(value)
    except Exception:
        pass

    try:
        value = obj.property("text")
        if value is not None:
            return str(value)
    except Exception:
        pass

    return ""


def verify_text(spanish_data, object_name, english_text):

    try:
        obj = waitForObjectExists(object_name)

    except Exception as e:
        test.fail(
            "Unable to find application object: "
            + str(object_name)
            + " | Error: "
            + str(e)
        )
        return False

    actual_text = get_object_text(obj)

    if actual_text is None:
        actual_text = ""

    expected_text = get_spanish_text(
        spanish_data,
        english_text
    )

    if expected_text == "":
        test.fail(
            "English text not found in Excel Column B: "
            + str(english_text)
        )
        return False

    expected_values = expected_text.split("|||")
    actual_normalized = normalize(actual_text)

    for value in expected_values:

        expected_normalized = normalize(value)

        if actual_normalized == expected_normalized:
            test.compare(
                actual_normalized,
                expected_normalized,
                "Spanish text verification: "
                + str(english_text)
            )
            return True

    test.fail(
        "Spanish text mismatch for: "
        + str(english_text)
        + " | Actual: "
        + actual_text
        + " | Expected: "
        + expected_text
    )

    return False


def main():


    workbook = None

    try:

        clickButton(
            waitForObject(
                names.numpad_pushButton_8_NumpadButton
            )
        )

        clickButton(
            waitForObject(
                names.numpad_pushButton_8_NumpadButton
            )
        )

        clickButton(
            waitForObject(
                names.numpad_pushButton_8_NumpadButton
            )
        )

        clickButton(
            waitForObject(
                names.numpad_pushButton_8_NumpadButton
            )
        )

        snooze(6)

        clickButton(
            waitForObject(
                names.mainFrame_okButton_QPushButton
            )
        )

        clickButton(
            waitForObject(
                names.homeScreen_settingsButton_QPushButton
            )
        )

        clickTab(
            waitForObject(
                names.settingsScreen_tbSystemInfo_TabWidget
            ),
            "INFO"
        )

        clickButton(
            waitForObject(
                names.gbSystemInformation_pbLanguageSettingsUpdate_QPushButton
            )
        )

        clickButton(
            waitForObject(
                names.languagesFrame_pbSpain_QPushButton
            )
        )

        clickButton(
            waitForObject(
                names.buttonsFrame_acceptButton_QPushButton
            )
        )

        clickButton(
            waitForObject(
                names.settingsScreen_pbBack_QPushButton
            )
        )

        clickButton(
            waitForObject(
                names.homeScreen_pbSoftTissueQuickStart_QToolButton
            )
        )

        mouseClick(
            waitForObjectItem(
                names.tabLeftMainPresets_leftPresetsView_PresetListView,
                "_1"
            ),
            305,
            35,
            Qt.NoModifier,
            Qt.LeftButton
        )

        mouseClick(
            waitForObjectItem(
                names.tabRightMainPresets_rightPresetsView_PresetListView,
                "_1"
            ),
            299,
            49,
            Qt.NoModifier,
            Qt.LeftButton
        )

        clickButton(
            waitForObject(
                names.quickStartScreen_pbContinue_QPushButton
            )
        )

        clickButton(
            waitForObject(
                names.mainFrame_okButton_QPushButton_2
            )
        )

        mouseClick(
            waitForObject(
                names.peakPowerButton_nameLabel_QLabel
            ),
            34,
            34,
            Qt.NoModifier,
            Qt.LeftButton
        )

        snooze(6)

        workbook, sheet, spanish_data = load_spanish_excel(
            EXCEL_PATH
        )

        if workbook is None:
            return

        verify_text(
            spanish_data,
            names.centralWidget_lbTitle_QLabel,
            "SELECT PEAK POWER"
        )

        snooze(6)

        verify_text(
            spanish_data,
            names.centralWidget_pw500_QPushButton,
            "High"
        )

        snooze(6)

        verify_text(
            spanish_data,
            names.centralWidget_pw250_QPushButton,
            "Medium"
        )

        snooze(6)

        verify_text(
            spanish_data,
            names.centralWidget_pw125_QPushButton,
            "Low"
        )

        snooze(6)

        verify_text(
            spanish_data,
            names.buttonsFrame_rejectButton_QPushButton,
            "CANCEL"
        )

        mouseClick(
            waitForObject(
                names.centralWidget_mainWidget_QWidget
            ),
            460,
            549,
            Qt.NoModifier,
            Qt.LeftButton
        )

        clickButton(
            waitForObject(
                names.buttonsFrame_rejectButton_QPushButton
            )
        )

        clickButton(
            waitForObject(
                names.buttonsBar_homeButton_QPushButton_2
            )
        )

        clickButton(
            waitForObject(
                names.homeScreen_pbSoftTissueAssistant_QToolButton
            )
        )

        clickButton(
            waitForObject(
                names.gbSoftTissueProcedures_pbIncision_QPushButton
            )
        )

        clickButton(
            waitForObject(
                names.softTissueAssistantModeScreen_pbContinue_QPushButton
            )
        )

        clickButton(
            waitForObject(
                names.treatmentCharacteristicsBar_pbEdit_QPushButton
            )
        )

        snooze(6)

        verify_text(
            spanish_data,
            names.centralWidget_titleLabel_QLabel,
            "SOFT TISSUE CHARACTERISTICS"
        )

        snooze(6)

        verify_text(
            spanish_data,
            names.centralWidget_titleLabel_QLabel_2,
            "PROCEDURE"
        )

        snooze(6)

        verify_text(
            spanish_data,
            names.centralWidget_valueLabel_QLabel,
            "Incision"
        )

        snooze(6)

        verify_text(
            spanish_data,
            names.centralWidget_rejectButton_QPushButton,
            "EDIT"
        )

        snooze(6)

        verify_text(
            spanish_data,
            names.centralWidget_acceptButton_QPushButton,
            "OK"
        )

    except Exception as e:
        test.fail(
            "Spanish Peak Power / Soft Tissue localization test failed: "
            + str(e)
        )

    finally:
        if workbook is not None:
            try:
                workbook.close()
            except Exception:
                pass