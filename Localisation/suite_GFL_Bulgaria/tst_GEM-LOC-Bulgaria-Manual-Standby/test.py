
# -*- coding: utf-8 -*-

import names
from openpyxl import load_workbook
import os

TESTCASE_DIR = os.path.dirname(os.path.abspath(__file__))
 
EXCEL_PATH = os.path.abspath(
    os.path.join( TESTCASE_DIR,
        "..",
        "shared",
        "testdata",
        "Gemini Split strings.xlsx"
    )
)

# EXCEL_PATH = (
#     "/home/ezen/Test_Automation_Gemini/Localisation/"
#     "suite_GFL_Bulgaria/"
#     "tst_GEM-LOC-Bulgaria-PeakPower-SoftTissueCharacters/"
#     "testdata/Gemini Split strings.xlsx"
# )

EXCEL_SHEET_NAME = "Treatment Modes"

# Column B = English
ENGLISH_COLUMN = 2

# Column C = Bulgarian
BULGARIAN_COLUMN = 3


#

def normalize_text(value):

    if value is None:
        return ""

    return " ".join(
        str(value)
        .replace("\n", " ")
        .replace("\r", " ")
        .split()
    ).strip().casefold()


def load_bulgarian_data(sheet):

    bulgarian_data = {}

    for row in sheet.iter_rows(
        min_col=ENGLISH_COLUMN,
        max_col=BULGARIAN_COLUMN,
        values_only=True
    ):

        english_value = row[0]

        bulgarian_value = row[
            BULGARIAN_COLUMN - ENGLISH_COLUMN
        ]

        if english_value is None:
            continue

        english_key = normalize_text(
            english_value
        )

        if english_key == "":
            continue

        if bulgarian_value is None:

            test.fail(
                "Bulgarian translation is empty in Excel "
                "for English text: "
                + str(english_value)
            )

            continue

        bulgarian_text = str(
            bulgarian_value
        ).strip()

        if bulgarian_text == "":

            test.fail(
                "Bulgarian translation is blank in Excel "
                "for English text: "
                + str(english_value)
            )

            continue

        bulgarian_data[
            english_key
        ] = bulgarian_text

    return bulgarian_data


def get_bulgarian_text(
    bulgarian_data,
    english_text
):

    search_text = normalize_text(
        english_text
    )

    if search_text not in bulgarian_data:

        test.fail(
            "English text not found in Excel Column B: "
            + str(english_text)
        )

        return ""

    return bulgarian_data[
        search_text
    ]


def get_actual_text(
    obj,
    description
):

    actual_value = None

    try:

        actual_value = obj.text

    except Exception:

        actual_value = None


    if actual_value is None:

        try:

            actual_value = obj.property(
                "text"
            )

        except Exception:

            actual_value = None


    if actual_value is None:

        try:

            actual_value = obj.windowTitle

        except Exception:

            actual_value = None


    if actual_value is None:

        try:

            actual_value = obj.accessibleName

        except Exception:

            actual_value = None


    actual_text = normalize_text(
        actual_value
    )


    if actual_text == "":

        test.fail(
            "Application text is empty or None for: "
            + description
        )

        return None

    return actual_text


def verify_text(
    bulgarian_data,
    object_name,
    english_text,
    description
):

    snooze(6)


    expected_text = get_bulgarian_text(
        bulgarian_data,
        english_text
    )

    expected_normalized = normalize_text(
        expected_text
    )


    # --------------------------------------------------------
    # Validate expected Excel value
    # --------------------------------------------------------

    if expected_normalized == "":

        test.fail(
            "Expected Bulgarian text is empty for: "
            + description
        )

        return


    # --------------------------------------------------------
    # Get application object
    # --------------------------------------------------------

    obj = waitForObjectExists(
        getattr(
            names,
            object_name
        )
    )


    actual_text = get_actual_text(
        obj,
        description
    )


    if actual_text is None:
        return


    test.compare(
        actual_text,
        expected_normalized,
        description
        + " - Bulgarian text verification"
    )


def verify_left_manual_tab(
    bulgarian_data,
    english_text,
    description
):

    snooze(6)


    expected_text = get_bulgarian_text(
        bulgarian_data,
        english_text
    )

    expected_normalized = normalize_text(
        expected_text
    )


    if expected_normalized == "":

        test.fail(
            "Expected Bulgarian text is empty for: "
            + description
        )

        return


    manual_tab_object = getattr(
        names,
        "mANUEL_TabItem"
    ).copy()


    if "text" in manual_tab_object:

        del manual_tab_object["text"]


    obj = waitForObjectExists(
        manual_tab_object
    )


    actual_text = get_actual_text(
        obj,
        description
    )


    if actual_text is None:
        return


    test.compare(
        actual_text,
        expected_normalized,
        description
        + " - Bulgarian text verification"
    )


def verify_right_manual_tab(
    bulgarian_data,
    english_text,
    description
):

    snooze(6)



    expected_text = get_bulgarian_text(
        bulgarian_data,
        english_text
    )

    expected_normalized = normalize_text(
        expected_text
    )


    if expected_normalized == "":

        test.fail(
            "Expected Bulgarian text is empty for: "
            + description
        )

        return


    manual_tab_object = getattr(
        names,
        "mANUEL_TabItem_2"
    ).copy()


    if "text" in manual_tab_object:

        del manual_tab_object["text"]


    obj = waitForObjectExists(
        manual_tab_object
    )


    actual_text = get_actual_text(
        obj,
        description
    )


    if actual_text is None:
        return


    test.compare(
        actual_text,
        expected_normalized,
        description
        + " - Bulgarian text verification"
    )


def main():

    workbook = load_workbook(
        EXCEL_PATH,
        read_only=True,
        data_only=True
    )


    if EXCEL_SHEET_NAME not in workbook.sheetnames:

        test.fail(
            "Excel sheet not found: "
            + EXCEL_SHEET_NAME
        )

        workbook.close()

        return


    sheet = workbook[
        EXCEL_SHEET_NAME
    ]


    bulgarian_data = load_bulgarian_data(
        sheet
    )


    doubleClick(
        waitForObject(
            names.numpad_pushButton_8_NumpadButton,
            100461
        ),
        96,
        56,
        Qt.NoModifier,
        Qt.LeftButton
    )

    snooze(1)


    doubleClick(
        waitForObject(
            names.numpad_pushButton_8_NumpadButton
        ),
        96,
        56,
        Qt.NoModifier,
        Qt.LeftButton
    )

    snooze(1)


    clickButton(
        waitForObject(
            names.mainFrame_okButton_QPushButton
        )
    )

    snooze(2)


    clickButton(
        waitForObject(
            names.homeScreen_settingsButton_QPushButton
        )
    )

    snooze(2)


   
    clickTab(
        waitForObject(
            names.settingsScreen_tbSystemInfo_TabWidget
        ),
        "INFO"
    )

    snooze(2)


    clickButton(
        waitForObject(
            names.gbSystemInformation_pbLanguageSettingsUpdate_QPushButton
        )
    )

    snooze(2)


    clickButton(
        waitForObject(
            names.languagesFrame_pbBulgaria_QPushButton
        )
    )

    snooze(1)


    clickButton(
        waitForObject(
            names.buttonsFrame_acceptButton_QPushButton
        )
    )

    snooze(3)


    clickButton(
        waitForObject(
            names.buttonsBar_homeButton_QPushButton
        )
    )

    snooze(2)


    clickButton(
        waitForObject(
            names.gbExpert_expertButton_QPushButton
        )
    )

    snooze(3)



    
    verify_left_manual_tab(
        bulgarian_data,
        "MANUAL",
        "Left Manual Tab"
    )


    verify_text(
        bulgarian_data,
        "averagePowerWidget_nameLabel_QLabel_2",
        "AVERAGE POWER",
        "Left Average Power"
    )


   
    verify_text(
        bulgarian_data,
        "averagePowerWidget_unitsOfMeasureLabel_QLabel_2",
        "W",
        "Left Power Unit"
    )


    verify_text(
        bulgarian_data,
        "pulseEnergyWidget_nameLabel_QLabel_2",
        "PULSE ENERGY",
        "Left Pulse Energy"
    )


    verify_text(
        bulgarian_data,
        "pulseEnergyWidget_unitsOfMeasureLabel_QLabel_2",
        "J",
        "Left Energy Unit"
    )


    verify_text(
        bulgarian_data,
        "frequencyWidget_nameLabel_QLabel_2",
        "FREQUENCY",
        "Left Frequency"
    )


    verify_text(
        bulgarian_data,
        "frequencyWidget_unitsOfMeasureLabel_QLabel_2",
        "Hz",
        "Left Frequency Unit"
    )


    verify_text(
        bulgarian_data,
        "modeButton_nameLabel_QLabel",
        "REGULAR PULSE",
        "Left Pulse Mode"
    )


    verify_text(
        bulgarian_data,
        "peakPowerButton_nameLabel_QLabel",
        "PEAK POWER",
        "Left Peak Power"
    )


    verify_text(
        bulgarian_data,
        "pulseModeWidget_saveButton_QPushButton_2",
        "SAVE",
        "Left Save Button"
    )


    verify_text(
        bulgarian_data,
        "totalEnergyFrame_nameLabel_QLabel",
        "TOTAL ENERGY",
        "Left Total Energy"
    )


    verify_text(
        bulgarian_data,
        "totalTimeFrame_nameLabel_QLabel",
        "TOTAL LASING TIME",
        "Left Total Lasing Time"
    )


  
    verify_right_manual_tab(
        bulgarian_data,
        "MANUAL",
        "Right Manual Tab"
    )


    verify_text(
        bulgarian_data,
        "averagePowerWidget_nameLabel_QLabel_3",
        "AVERAGE POWER",
        "Right Average Power"
    )


    verify_text(
        bulgarian_data,
        "averagePowerWidget_unitsOfMeasureLabel_QLabel_3",
        "W",
        "Right Power Unit"
    )


    verify_text(
        bulgarian_data,
        "pulseEnergyWidget_nameLabel_QLabel_3",
        "PULSE ENERGY",
        "Right Pulse Energy"
    )


    verify_text(
        bulgarian_data,
        "pulseEnergyWidget_unitsOfMeasureLabel_QLabel_3",
        "J",
        "Right Energy Unit"
    )


    verify_text(
        bulgarian_data,
        "frequencyWidget_nameLabel_QLabel_3",
        "FREQUENCY",
        "Right Frequency"
    )


    verify_text(
        bulgarian_data,
        "pulseModeWidget_saveButton_QPushButton_3",
        "SAVE",
        "Right Save Button"
    )


    verify_text(
        bulgarian_data,
        "peakPowerButton_nameLabel_QLabel_3",
        "PEAK POWER",
        "Right Peak Power"
    )


  
    verify_text(
        bulgarian_data,
        "modeButton_nameLabel_QLabel_2",
        "REGULAR PULSE",
        "Right Pulse Mode"
    )


    verify_text(
        bulgarian_data,
        "stateSwitch_readyButton_QPushButton",
        "READY",
        "Ready Button"
    )


    verify_text(
        bulgarian_data,
        "stateSwitch_standbyButton_QPushButton",
        "STANDBY",
        "Standby Button"
    )


    verify_text(
        bulgarian_data,
        "fiberFrame_nameLabel_QLabel",
        "FIBER",
        "Fiber"
    )


    verify_text(
        bulgarian_data,
        "fiberUsesFrame_nameLabel_QLabel",
        "FIBER USES REMAINING",
        "Fiber Uses Remaining"
    )


    verify_text(
        bulgarian_data,
        "bottomBar_aimingBeamButton_AimingBeamModeButton",
        "AIMING BEAM",
        "Aiming Beam"
    )


    workbook.close()

