
# -*- coding: utf-8 -*-

import names
import openpyxl
import re
import unicodedata
import html
import os

TESTCASE_DIR = os.path.dirname(os.path.abspath(__file__))
 
EXCEL_PATH = os.path.abspath(
    os.path.join( TESTCASE_DIR,
        "..",
        "shared",
        "testdata",
        "Gemini Split strings.xlsx"
    )
)

def normalize(text):
    if text is None:
        return ""

    text = html.unescape(str(text))
    text = re.sub(r"<br\s*/?>", " ", text, flags=re.IGNORECASE)
    text = re.sub(r"<[^>]+>", " ", text)
    text = unicodedata.normalize("NFKC", text)

    text = text.replace("\xa0", " ")
    text = text.replace("\u2007", " ")
    text = text.replace("\u202f", " ")
    text = text.replace("\r", " ")
    text = text.replace("\n", " ")
    text = text.replace("\t", " ")

    text = text.replace("™", "")
    text = text.replace("®", "")
    text = text.replace("©", "")

    return " ".join(text.split()).strip().casefold()


def comparison_key(text):
    return re.sub(r"\s+", "", normalize(text))


def load_bulgarian_excel(EXCEL_PATH):
    workbook = openpyxl.load_workbook(
        EXCEL_PATH,
        data_only=True,
        read_only=True
    )

    if "Treatment Modes" not in workbook.sheetnames:
        workbook.close()
        test.fail("Worksheet 'Treatment Modes' does not exist in Excel file")
        return None, {}

    sheet = workbook["Treatment Modes"]

    english_column = 2
    bulgarian_column = 3

    bulgarian_data = {}

    for row in sheet.iter_rows(
        min_col=english_column,
        max_col=bulgarian_column
    ):
        english_value = row[0].value
        bulgarian_value = row[1].value

        if english_value is None:
            continue

        english_key = comparison_key(english_value)

        if not english_key:
            continue

        if bulgarian_value is None:
            continue

        translated_value = str(bulgarian_value).strip()

        if translated_value:
            bulgarian_data[english_key] = translated_value

    return workbook, bulgarian_data


def get_bulgarian_text(bulgarian_data, english_text):
    key = comparison_key(english_text)
    return bulgarian_data.get(key, "")


def verify_text(bulgarian_data, object_name, english_text):
    try:
        obj = waitForObjectExists(object_name)
        actual_text = str(obj.text).strip()

    except Exception as e:
        test.fail(
            "Unable to read application object: "
            + str(object_name)
            + " | Error: "
            + str(e)
        )
        return False

    expected_text = get_bulgarian_text(
        bulgarian_data,
        english_text
    )

    if not expected_text:
        test.fail(
            "English text not found in Excel or Bulgarian translation is empty: "
            + str(english_text)
        )
        return False

    actual_normalized = comparison_key(actual_text)
    expected_normalized = comparison_key(expected_text)

    test.compare(
        actual_normalized,
        expected_normalized,
        "Bulgarian text verification: "
        + str(english_text)
        
        
    )

    return actual_normalized == expected_normalized


def main():


    workbook = None

    try:
        clickButton(waitForObject(names.numpad_pushButton_8_NumpadButton))
        clickButton(waitForObject(names.numpad_pushButton_8_NumpadButton))
        clickButton(waitForObject(names.numpad_pushButton_8_NumpadButton))
        clickButton(waitForObject(names.numpad_pushButton_8_NumpadButton))

        clickButton(waitForObject(names.mainFrame_okButton_QPushButton))
        clickButton(waitForObject(names.homeScreen_settingsButton_QPushButton))
        clickTab(
            waitForObject(names.settingsScreen_tbSystemInfo_TabWidget),
            "INFO"
        )
        clickButton(
            waitForObject(
                names.gbSystemInformation_pbLanguageSettingsUpdate_QPushButton
            )
        )
        clickButton(
            waitForObject(names.languagesFrame_pbBulgaria_QPushButton)
        )
        clickButton(
            waitForObject(names.buttonsFrame_acceptButton_QPushButton)
        )
        clickButton(waitForObject(names.buttonsBar_homeButton_QPushButton))
        clickButton(waitForObject(names.homeScreen_pbStoneAssistant_QToolButton))
        clickButton(
            waitForObject(names.locationFrame_kidneyPcnlLocationButton_QPushButton)
        )
        clickButton(
            waitForObject(names.flowRateFrame_lowFlowRateButton_QPushButton)
        )
        clickButton(
            waitForObject(names.hardnessFrame_smallHardnessButton_QPushButton)
        )
        clickButton(
            waitForObject(names.stoneAssistantModeScreen_continueButton_QPushButton)
        )
        clickButton(
            waitForObject(names.treatmentCharacteristicsBar_pbEdit_QPushButton_2)
        )

        workbook, bulgarian_data = load_bulgarian_excel(EXCEL_PATH)

        if workbook is None:
            return

        verify_text(
            bulgarian_data,
            names.centralWidget_titleLabel_QLabel,
            "STONE CHARACTERISTICS"
        )

        verify_text(
            bulgarian_data,
            names.centralWidget_titleLabel_QLabel_2,
            "LOCATION"
        )

        verify_text(
            bulgarian_data,
            names.centralWidget_valueLabel_QLabel,
            "KIDNEY (PCNL)"
        )

        verify_text(
            bulgarian_data,
            names.centralWidget_titleLabel_QLabel_3,
            "FLOW RATE"
        )

        verify_text(
            bulgarian_data,
            names.centralWidget_valueLabel_QLabel_2,
            "LOW"
        )

        verify_text(
            bulgarian_data,
            names.centralWidget_titleLabel_QLabel_4,
            "HARDNESS"
        )

        verify_text(
            bulgarian_data,
            names.centralWidget_descriptionLabel_QLabel,
            "Hounsfield Scale < 500"
        )

        verify_text(
            bulgarian_data,
            names.centralWidget_valueLabel_QLabel_3,
            "SOFT"
        )

        verify_text(
            bulgarian_data,
            names.centralWidget_rejectButton_QPushButton,
            "EDIT"
        )

        verify_text(
            bulgarian_data,
            names.centralWidget_acceptButton_QPushButton,
            "OK"
        )

    finally:
        if workbook is not None:
            workbook.close()
