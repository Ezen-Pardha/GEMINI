
# -*- coding: utf-8 -*-

import names
from openpyxl import load_workbook
import re
import os

TESTCASE_DIR = os.path.dirname(os.path.abspath(__file__))
 
EXCEL_PATH = os.path.abspath(
    os.path.join( TESTCASE_DIR,
        "..",
        "shared",
        "testdata",
        "Gemini Split strings.xlsx"
    )
)

def normalize(text):

    if text is None:
        return ""

    text = str(text)

    text = re.sub(
        r"<br\s*/?>",
        " ",
        text,
        flags=re.IGNORECASE
    )

    text = re.sub(
        r"<[^>]+>",
        " ",
        text
    )

    text = text.replace("\xa0", " ")

    text = " ".join(text.split())

    return text.strip().casefold()


def load_bulgarian_excel(EXCEL_PATH):

    workbook = load_workbook(
        EXCEL_PATH,
        read_only=True,
        data_only=True
    )

    if "Treatment Modes" not in workbook.sheetnames:

        workbook.close()

        test.fail(
            "Worksheet 'Treatment Modes' does not exist "
            "in Excel file"
        )

        return None, {}

    sheet = workbook["Treatment Modes"]

    english_column = 2
    bulgarian_column = 3

    bulgarian_data = {}

    for row in sheet.iter_rows(
        min_col=english_column,
        max_col=bulgarian_column
    ):

        english_value = row[0].value
        bulgarian_value = row[1].value

        if english_value is None:
            continue

        english_key = normalize(
            english_value
        )

        if english_key == "":
            continue

        if bulgarian_value is None:
            bulgarian_value = ""

        bulgarian_data[english_key] = str(
            bulgarian_value
        )

    return workbook, bulgarian_data


def get_bulgarian_text(
    bulgarian_data,
    english_text
):

    key = normalize(
        english_text
    )

    if key not in bulgarian_data:

        # test.fail(
        #     "English text not found in Excel: "
        #     + str(english_text)
        # )

        return ""

    return bulgarian_data[key]


def verify_excel_string(
    object_name,
    bulgarian_data,
    english_text,
    description
):

    snooze(6)

    try:

        actual = str(
            waitForObjectExists(
                object_name
            ).text
        ).strip()

    except Exception as e:

        test.fail(
            "Unable to read application object: "
            + str(object_name)
            + " | Error: "
            + str(e)
        )

        return

    expected = get_bulgarian_text(
        bulgarian_data,
        english_text
    )

    actual_normalized = normalize(
        actual
    )

    expected_normalized = normalize(
        expected
    )

    test.compare(
        actual_normalized,
        expected_normalized,
        description
    )


def main():

    workbook = None

    # EXCEL_PATH = (
    #     "/home/ezen/Test_Automation_Gemini/Localisation/"
    #     "suite_GFL_Bulgaria/"
    #     "tst_GEM-LOC-Bulgaria-Userpresets-Add-Edit-Delete/"
    #     "testdata/Gemini Split strings.xlsx"
    # )


    workbook, bulgarian_data = (
        load_bulgarian_excel(
            EXCEL_PATH
        )
    )

    if workbook is None:
        return


    clickButton(
        waitForObject(
            names.numpad_pushButton_8_NumpadButton
        )
    )

    clickButton(
        waitForObject(
            names.numpad_pushButton_8_NumpadButton
        )
    )

    clickButton(
        waitForObject(
            names.numpad_pushButton_8_NumpadButton
        )
    )

    clickButton(
        waitForObject(
            names.numpad_pushButton_8_NumpadButton
        )
    )

    clickButton(
        waitForObject(
            names.mainFrame_okButton_QPushButton
        )
    )


    clickButton(
        waitForObject(
            names.homeScreen_settingsButton_QPushButton
        )
    )

    clickTab(
        waitForObject(
            names.settingsScreen_tbSystemInfo_TabWidget
        ),
        "INFO"
    )

    clickButton(
        waitForObject(
            names.gbSystemInformation_pbLanguageSettingsUpdate_QPushButton
        )
    )


    clickButton(
        waitForObject(
            names.languagesFrame_pbBulgaria_QPushButton
        )
    )

    clickButton(
        waitForObject(
            names.buttonsFrame_acceptButton_QPushButton
        )
    )

    snooze(6)


    snooze(3)


    tabWidget = waitForObject(
        names.settingsScreen_tbSystemInfo_TabWidget
    )

    tabWidget.setCurrentIndex(2)

    snooze(3)


    tabWidget.setCurrentIndex(2)

    snooze(3)


    sendEvent(
        "QMoveEvent",
        waitForObject(
            names.o_ScreenSwitcher
        ),
        0,
        -66,
        1285,
        233
    )


    clickButton(
        waitForObject(
            names.settingsScreen_addButton_QPushButton
        )
    )

    verify_excel_string(
        names.centralWidget_lbOperation_QLabel,
        bulgarian_data,
        "ADD PRESET",
        "Add Preset Operation"
    )


    verify_excel_string(
        names.centralWidget_pbName_QPushButton,
        bulgarian_data,
        "PRESET NAME",
        "Preset Name"
    )


    verify_excel_string(
        names.averagePowerWidget_nameLabel_QLabel,
        bulgarian_data,
        "AVERAGE POWER",
        "Average Power"
    )


    verify_excel_string(
        names.averagePowerWidget_unitsOfMeasureLabel_QLabel,
        bulgarian_data,
        "W",
        "Average Power Unit"
    )


    verify_excel_string(
        names.pulseEnergyWidget_nameLabel_QLabel,
        bulgarian_data,
        "PULSE ENERGY",
        "Pulse Energy"
    )


    mouseClick(
        waitForObject(
            names.centralWidget_frame_QFrame
        ),
        271,
        13,
        Qt.NoModifier,
        Qt.LeftButton
    )

    verify_excel_string(
        names.pulseEnergyWidget_unitsOfMeasureLabel_QLabel,
        bulgarian_data,
        "J",
        "Pulse Energy Unit"
    )

    verify_excel_string(
        names.frequencyWidget_nameLabel_QLabel,
        bulgarian_data,
        "FREQUENCY",
        "Frequency"
    )


    verify_excel_string(
        names.frequencyWidget_unitsOfMeasureLabel_QLabel,
        bulgarian_data,
        "Hz",
        "Frequency Unit"
    )


    verify_excel_string(
        names.modeButton_nameLabel_QLabel_3,
        bulgarian_data,
        "REGULAR PULSE",
        "Regular Pulse"
    )



    verify_excel_string(
        names.peakPowerButton_nameLabel_QLabel_2,
        bulgarian_data,
        "PEAK POWER",
        "Peak Power"
    )


    verify_excel_string(
        names.pulseModeWidget_saveButton_QPushButton,
        bulgarian_data,
        "ADD",
        "Add Button"
    )


    mouseClick(
        waitForObject(
            names.pulseModeWidget_buttonsWidget_QWidget
        ),
        508,
        46,
        Qt.NoModifier,
        Qt.LeftButton
    )

    mouseClick(
        waitForObject(
            names.pulseModeWidget_buttonsWidget_QWidget
        ),
        411,
        1,
        Qt.NoModifier,
        Qt.LeftButton
    )

    clickButton(
        waitForObject(
            names.centralWidget_pbName_QPushButton
        )
    )

    type(
        waitForObject(
            names.keyboardPopup_wdEdit_WdLineEdit
        ),
        "preset one"
    )

    mouseClick(
        waitForObject(
            names.keyboardPopup_lbText_QLabel
        ),
        77,
        26,
        Qt.NoModifier,
        Qt.LeftButton
    )

    clickButton(
        waitForObject(
            names.pulseModeWidget_saveButton_QPushButton
        )
    )

    snooze(6)


    mouseClick(
        waitForObject(
            names.tabSettingsPresetsUserStonePresets_presetListView_PresetListView
        ),
        546,
        72,
        Qt.NoModifier,
        Qt.LeftButton
    )


    clickButton(
        waitForObject(
            names.settingsScreen_editButton_QPushButton
        )
    )

    mouseClick(
        waitForObject(
            names.tabSettingsPresetsUserStonePresets_presetListView_PresetListView
        ),
        494,
        66,
        Qt.NoModifier,
        Qt.LeftButton
    )

    snooze(6)


    verify_excel_string(
        names.centralWidget_lbOperation_QLabel,
        bulgarian_data,
        "EDIT PRESET",
        "Edit Preset Operation"
    )


    snooze(6)

    actualEditName = str(
        waitForObjectExists(
            names.centralWidget_pbName_QPushButton
        ).text
    )

    test.compare(
        actualEditName,
        "preset one",
        "Edit Preset Name"
    )


    verify_excel_string(
        names.averagePowerWidget_nameLabel_QLabel,
        bulgarian_data,
        "AVERAGE POWER",
        "Edit - Average Power"
    )


    verify_excel_string(
        names.averagePowerWidget_unitsOfMeasureLabel_QLabel,
        bulgarian_data,
        "W",
        "Edit - Average Power Unit"
    )


    verify_excel_string(
        names.pulseEnergyWidget_nameLabel_QLabel,
        bulgarian_data,
        "PULSE ENERGY",
        "Edit - Pulse Energy"
    )


    verify_excel_string(
        names.pulseEnergyWidget_unitsOfMeasureLabel_QLabel,
        bulgarian_data,
        "J",
        "Edit - Pulse Energy Unit"
    )


    verify_excel_string(
        names.frequencyWidget_nameLabel_QLabel,
        bulgarian_data,
        "FREQUENCY",
        "Edit - Frequency"
    )


    verify_excel_string(
        names.frequencyWidget_unitsOfMeasureLabel_QLabel,
        bulgarian_data,
        "Hz",
        "Edit - Frequency Unit"
    )


    verify_excel_string(
        names.modeButton_nameLabel_QLabel_3,
        bulgarian_data,
        "REGULAR PULSE",
        "Edit - Regular Pulse"
    )


    verify_excel_string(
        names.peakPowerButton_nameLabel_QLabel_2,
        bulgarian_data,
        "PEAK POWER",
        "Edit - Peak Power"
    )


    verify_excel_string(
        names.pulseModeWidget_saveButton_QPushButton,
        bulgarian_data,
        "UPDATE",
        "Update Button"
    )


    mouseClick(
        waitForObject(
            names.pulseModeWidget_buttonsWidget_QWidget
        ),
        531,
        22,
        Qt.NoModifier,
        Qt.LeftButton
    )

    mouseClick(
        waitForObject(
            names.pulseModeWidget_buttonsWidget_QWidget
        ),
        509,
        30,
        Qt.NoModifier,
        Qt.LeftButton
    )

    clickButton(
        waitForObject(
            names.pulseModeWidget_saveButton_QPushButton
        )
    )

    snooze(6)


    clickButton(
        waitForObject(
            names.settingsScreen_deletePresetButton_QPushButton
        )
    )

    mouseClick(
        waitForObject(
            names.tabSettingsPresetsUserStonePresets_presetListView_PresetListView
        ),
        553,
        80,
        Qt.NoModifier,
        Qt.LeftButton
    )

    snooze(6)


    verify_excel_string(
        names.centralWidget_lbDialogTitle_QLabel,
        bulgarian_data,
        "DELETE PRESET",
        "Delete Dialog Title"
    )


    snooze(6)

    actualDeleteName = str(
        waitForObjectExists(
            names.centralWidget_questionLabel_QLabel
        ).text
    )

    test.compare(
        actualDeleteName,
        "preset one",
        "Delete Preset Name"
    )


    snooze(6)

    actualInfo = str(
        waitForObjectExists(
            names.infoFrame_extInfoLabel_QLabel
        ).text
    )

    expectedInfo = "6 J     3 Hz     18 W"

    test.compare(
        actualInfo,
        expectedInfo,
        "Preset Information"
    )


 

    verify_excel_string(
        names.buttonsFrame_rejectButton_QPushButton,
        bulgarian_data,
        "CANCEL",
        "Cancel Button"
    )




    verify_excel_string(
        names.centralWidget_acceptButton_QPushButton,
        bulgarian_data,
        "DELETE",
        "Delete Button"
    ) 

    
   
    if workbook is not None:
        
        clickButton(waitForObject(names.centralWidget_acceptButton_QPushButton))
        workbook.close()

