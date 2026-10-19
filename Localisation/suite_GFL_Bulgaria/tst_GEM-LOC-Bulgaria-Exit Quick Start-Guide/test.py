
# -*- coding: utf-8 -*-

import names
import openpyxl
import re
import os

TESTCASE_DIR = os.path.dirname(os.path.abspath(__file__))
 
EXCEL_PATH = os.path.abspath(
    os.path.join( TESTCASE_DIR,
        "..",
        "shared",
        "testdata",
        "Gemini Split strings.xlsx"
    )
)

ENGLISH_COLUMN = 2
BULGARIAN_COLUMN = 3

SHEET_NAME = "Soft tissue"




def normalize(text):

    if text is None:
        return ""

    text = str(text)

    # Remove HTML tags
    text = re.sub(
        r"<[^>]+>",
        " ",
        text
    )

    # Normalize spaces and new lines
    text = " ".join(
        text.split()
    )

    # Normalize angle brackets
    text = re.sub(
        r"\s*<\s*",
        " < ",
        text
    )

    text = re.sub(
        r"\s*>\s*",
        " > ",
        text
    )

    return text.strip().casefold()




def load_excel():

    try:

        workbook = openpyxl.load_workbook(
            EXCEL_PATH,
            data_only=True,
            read_only=True
        )

        if SHEET_NAME not in workbook.sheetnames:

            test.fail(
                "Excel sheet not found: "
                + SHEET_NAME
            )

            workbook.close()

            return None, None

        sheet = workbook[SHEET_NAME]

        return workbook, sheet

    except Exception as e:

        test.fail(
            "Unable to load Bulgarian localization Excel: "
            + str(e)
        )

        return None, None



def get_bulgarian_values(
    sheet,
    english_text
):

    expected_english = normalize(
        english_text
    )

    bulgarian_values = []

    for row in sheet.iter_rows(
        min_col=ENGLISH_COLUMN,
        max_col=BULGARIAN_COLUMN,
        values_only=True
    ):

        english_value = row[0]

        if english_value is None:
            continue

        if normalize(
            english_value
        ) == expected_english:

            bulgarian_value = row[
                BULGARIAN_COLUMN - ENGLISH_COLUMN
            ]

            if bulgarian_value is not None:

                bulgarian_value = str(
                    bulgarian_value
                ).strip()

                if bulgarian_value:

                    bulgarian_values.append(
                        bulgarian_value
                    )

    return bulgarian_values




def verify_text(
    sheet,
    object_name,
    english_text,
    description
):

    # Wait before every verification
    snooze(6)

    try:

        obj = waitForObjectExists(
            object_name
        )

        actual_value = obj.text

        actual = ""

        if actual_value is not None:

            actual = str(
                actual_value
            ).strip()

    except Exception as e:

        test.fail(
            description
            + " | Unable to read object: "
            + str(e)
        )

        return

   

    expected_values = get_bulgarian_values(
        sheet,
        english_text
    )


    if not expected_values:

        test.fail(
            description
            + " | English string not found in Column B: "
            + english_text
        )

        return

    actual_normalized = normalize(
        actual
    )

 

    for expected in expected_values:

        expected_normalized = normalize(
            expected
        )

        if actual_normalized == expected_normalized:

            test.compare(
                actual_normalized,
                expected_normalized,
                description
            )

            return


    test.fail(
        description
        + " | Expected Bulgarian value from Column C: "
        + " / ".join(
            expected_values
        )
        + " | Actual: "
        + actual
    )



def main():

    workbook, sheet = load_excel()

    if workbook is None or sheet is None:
        return


    clickButton(
        waitForObject(
            names.numpad_pushButton_8_NumpadButton
        )
    )

    doubleClick(
        waitForObject(
            names.numpad_pushButton_8_NumpadButton
        ),
        94,
        77,
        Qt.NoModifier,
        Qt.LeftButton
    )

    doubleClick(
        waitForObject(
            names.numpad_pushButton_8_NumpadButton
        ),
        94,
        77,
        Qt.NoModifier,
        Qt.LeftButton
    )

    clickButton(
        waitForObject(
            names.mainFrame_okButton_QPushButton
        )
    )


    
    clickButton(
        waitForObject(
            names.homeScreen_settingsButton_QPushButton
        )
    )

    clickTab(
        waitForObject(
            names.settingsScreen_tbSystemInfo_TabWidget
        ),
        "INFO"
    )

    clickButton(
        waitForObject(
            names.gbSystemInformation_pbLanguageSettingsUpdate_QPushButton
        )
    )

    # Select BULGARIAN
    clickButton(
        waitForObject(
            names.languagesFrame_pbBulgaria_QPushButton
        )
    )

    clickButton(
        waitForObject(
            names.buttonsFrame_acceptButton_QPushButton
        )
    )

    clickButton(
        waitForObject(
            names.buttonsBar_homeButton_QPushButton
        )
    )



    clickButton(
        waitForObject(
            names.homeScreen_pbSoftTissueQuickStart_QToolButton
        )
    )

    mouseClick(
        waitForObjectItem(
            names.tabLeftMainPresets_leftPresetsView_PresetListView,
            "_1"
        ),
        377,
        65,
        Qt.NoModifier,
        Qt.LeftButton
    )

    mouseClick(
        waitForObjectItem(
            names.tabRightMainPresets_rightPresetsView_PresetListView,
            "_1"
        ),
        150,
        37,
        Qt.NoModifier,
        Qt.LeftButton
    )

    sendEvent(
        "QMoveEvent",
        waitForObject(
            names.o_ScreenSwitcher
        ),
        0,
        -50,
        1288,
        292
    )

    clickButton(
        waitForObject(
            names.quickStartScreen_pbContinue_QPushButton
        )
    )

    clickButton(
        waitForObject(
            names.mainFrame_okButton_QPushButton_2
        )
    )



    mouseClick(
        waitForObject(
            names.modeButton_nameLabel_QLabel
        ),
        115,
        15,
        Qt.NoModifier,
        Qt.LeftButton
    )



    verify_text(
        sheet,
        names.centralWidget_TitleLabel_QLabel_2,
        "EXIT QUICK START MODE?",
        "Quick Start Exit Title"
    )



    verify_text(
        sheet,
        names.centralWidget_MessageLabel_QLabel_2,
        "Please confirm that you want to exit quick start mode.",
        "Quick Start Exit Message"
    )


    verify_text(
        sheet,
        names.centralWidget_CancelButton_QPushButton_3,
        "CANCEL",
        "Quick Start Cancel Button"
    )


    verify_text(
        sheet,
        names.centralWidget_DisableButton_QPushButton_2,
        "CONFIRM",
        "Quick Start Confirm Button"
    )


    clickButton(
        waitForObject(
            names.centralWidget_DisableButton_QPushButton_2
        )
    )


    clickButton(
        waitForObject(
            names.buttonsBar_homeButton_QPushButton_2
        )
    )


    clickButton(
        waitForObject(
            names.homeScreen_pbSoftTissueAssistant_QToolButton
        )
    )

    clickButton(
        waitForObject(
            names.gbSoftTissueProcedures_pbIncision_QPushButton
        )
    )

    clickButton(
        waitForObject(
            names.softTissueAssistantModeScreen_pbContinue_QPushButton
        )
    )


    mouseDrag(
        waitForObject(
            names.indexPicker_valuePicker_IntervalSlider
        ),
        195,
        15,
        211,
        29,
        1,
        Qt.LeftButton
    )


    verify_text(
        sheet,
        names.centralWidget_TitleLabel_QLabel_2,
        "EXIT GUIDED MODE?",
        "Guided Mode Exit Title"
    )


    
    verify_text(
        sheet,
        names.centralWidget_MessageLabel_QLabel_2,
        "Please confirm that you want to exit guided mode.",
        "Guided Mode Exit Message"
    )


    verify_text(
        sheet,
        names.centralWidget_CancelButton_QPushButton_3,
        "CANCEL",
        "Guided Mode Cancel Button"
    )


    verify_text(
        sheet,
        names.centralWidget_DisableButton_QPushButton_2,
        "CONFIRM",
        "Guided Mode Confirm Button"
    )


    clickButton(
        waitForObject(
            names.centralWidget_DisableButton_QPushButton_2
        )
    )


    workbook.close()


