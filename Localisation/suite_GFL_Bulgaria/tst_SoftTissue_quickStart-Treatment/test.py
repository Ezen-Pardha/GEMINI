
# -*- coding: utf-8 -*-

import names
import openpyxl
import re
import os
TESTCASE_DIR = os.path.dirname(os.path.abspath(__file__))
 
EXCEL_PATH = os.path.abspath(
    os.path.join( TESTCASE_DIR,
        "..",
        "shared",
        "testdata",
        "Gemini Split strings.xlsx"
    )
)


# EXCEL_PATH = "/home/ntc/Test_Automation_Gemini/Localisation/suite_GFL_Bulgaria/tst_Home Screen/testdata/Gemini Split strings.xlsx"
SHEET_NAME = "Treatment Modes"

ENGLISH_COLUMN = 1
Bulgaria_COLUMN = 2


def normalize(text):
    if text is None:
        return ""

    text = str(text)
    text = re.sub(r"<br\s*/?>", " ", text, flags=re.IGNORECASE)
    text = re.sub(r"<[^>]+>", " ", text)
    text = text.replace("\xa0", " ")
    text = " ".join(text.split())
    text = re.sub(r"\s*<\s*", "<", text)
    text = re.sub(r"\s*>\s*", ">", text)

    return text.strip().casefold()


def get_object_text(obj):
    # Avoid bool() conversion of Squish properties.
    try:
        visibility = obj.visible

        if visibility is not None:
            try:
                if visibility == False:
                    return ""
            except Exception:
                pass
    except Exception:
        pass

    for property_name in ("text", "title"):
        try:
            value = getattr(obj, property_name)

            if callable(value):
                value = value()

            if value is None:
                continue

            text = str(value).strip()

            if not text:
                continue

            text = re.sub(
                r"<br\s*/?>", " ", text, flags=re.IGNORECASE
            )
            text = re.sub(r"<[^>]+>", " ", text)
            text = text.replace("\xa0", " ")
            text = " ".join(text.split())

            if text:
                return text.strip()

        except Exception:
            pass

    return ""


def get_screen_candidates(obj, visited=None):
    if visited is None:
        visited = set()

    candidates = []

    try:
        object_id = id(obj)

        if object_id in visited:
            return candidates

        visited.add(object_id)
    except Exception:
        pass

    own_text = get_object_text(obj)

    if own_text:
        candidates.append(own_text)

    try:
        children = object.children(obj)
    except Exception:
        children = []

    child_texts = []

    for child in children:
        child_text = get_object_text(child)

        if child_text:
            child_texts.append(child_text)
            candidates.append(child_text)

    # Collect combined text for labels split across widgets.
    for start in range(len(child_texts)):
        combined = ""

        for end in range(start, len(child_texts)):
            if combined:
                combined += " " + child_texts[end]
            else:
                combined = child_texts[end]

            candidates.append(combined)

    for child in children:
        candidates.extend(
            get_screen_candidates(child, visited)
        )

    return candidates


def is_dynamic_numeric_value(text):
    if text is None:
        return False

    text = str(text).strip()

    number = r"[-+]?\d+(?:[.,]\d+)?"
    unit = (
        r"(?:J|W|Hz|kHz|MHz|GHz|"
        r"µm|μm|um|nm|mm|cm|m|"
        r"min|s|ms|h|%|°C|°)"
    )

    pattern = (
        r"\s*(?:"
        + number
        + r"\s*(?:"
        + unit
        + r")?\s*)+\s*"
    )

    return re.fullmatch(
        pattern, text, flags=re.IGNORECASE
    ) is not None


def load_Bulgaria_strings():
    try:
        workbook = openpyxl.load_workbook(
            EXCEL_PATH,
            data_only=True,
            read_only=True
        )
    except Exception as error:
        test.fail(
            "FAIL: Unable to load Bulgaria Excel: " + str(error)
        )
        return {}

    try:
        if SHEET_NAME not in workbook.sheetnames:
            test.fail(
                "FAIL: Excel worksheet not found: " + SHEET_NAME
            )
            return {}

        sheet = workbook[SHEET_NAME]
        catalog = {}

        for row in sheet.iter_rows(
            min_row=2,
            values_only=True
        ):
            if len(row) <= Bulgaria_COLUMN:
                continue

            trace_id = row[0]
            english = row[ENGLISH_COLUMN]
            Bulgaria = row[Bulgaria_COLUMN]

            if english is None or Bulgaria is None:
                continue

            english = str(english).strip()
            Bulgaria = str(Bulgaria).strip()

            if not english or not Bulgaria:
                continue

            key = normalize(Bulgaria)

            if key:
                catalog.setdefault(
                    key,
                    {
                        "english": english,
                        "Bulgaria": Bulgaria,
                        "trace_id": (
                            "" if trace_id is None
                            else str(trace_id).strip()
                        )
                    }
                )

        return catalog

    finally:
        workbook.close()


def main():

    # LOGIN

    sendEvent(
        "QMoveEvent",
        waitForObject(names.o_ScreenSwitcher),
        28, 208, 534, 218
    )

    for _ in range(4):
        clickButton(
            waitForObject(
                names.numpad_pushButton_8_NumpadButton
            )
        )

    clickButton(
        waitForObject(names.mainFrame_okButton_QPushButton)
    )

    # CHANGE LANGUAGE TO Bulgaria

    clickButton(
        waitForObject(
            names.homeScreen_settingsButton_QPushButton
        )
    )

    clickTab(
        waitForObject(
            names.settingsScreen_tbSystemInfo_TabWidget
        ),
        "INFO"
    )

    clickButton(
        waitForObject(
            names.gbSystemInformation_pbLanguageSettingsUpdate_QPushButton
        )
    )

    clickButton(
        waitForObject(names.languagesFrame_pbBulgaria_QPushButton)
    )

    clickButton(
        waitForObject(names.buttonsFrame_acceptButton_QPushButton)
    )

    clickButton(
        waitForObject(names.settingsScreen_pbBack_QPushButton)
    )

    # OPEN SOFT TISSUE QUICK START

    clickButton(
        waitForObject(
            names.homeScreen_pbSoftTissueQuickStart_QToolButton
        )
    )

    mouseClick(
        waitForObjectItem(
            names.tabLeftMainPresets_leftPresetsView_PresetListView,
            "_1", 44166
        ),
        467, 28, Qt.NoModifier, Qt.LeftButton
    )

    mouseClick(
        waitForObjectItem(
            names.tabRightMainPresets_rightPresetsView_PresetListView,
            "_1"
        ),
        121, 80, Qt.NoModifier, Qt.LeftButton
    )

    clickButton(
        waitForObject(
            names.quickStartScreen_pbContinue_QPushButton
        )
    )

    mouseClick(
        waitForObject(
            names.recommendedFiberWarningDialog_mainFrame_QFrame
        ),
        539, 43, Qt.NoModifier, Qt.LeftButton
    )

    clickButton(
        waitForObject(names.mainFrame_okButton_QPushButton_2)
    )

    mouseClick(
        waitForObject(
            names.leftPedalWidget_wdPresetTabs_TabWidget
        ),
        497, 44, Qt.NoModifier, Qt.LeftButton
    )

    mouseClick(
        waitForObject(
            names.rightPedalWidget_wdPresetTabs_TabWidget
        ),
        98, 44, Qt.NoModifier, Qt.LeftButton
    )

    mouseClick(
        waitForObject(names.treatmentScreen_bottomBar_QFrame),
        890, 65, Qt.NoModifier, Qt.LeftButton
    )

    snooze(1)

    # LOAD Bulgaria EXCEL CATALOG

    excel_strings = load_Bulgaria_strings()

    if not excel_strings:
        test.fail("FAIL: No Bulgaria strings loaded from Excel.")
        return

    # COLLECT TREATMENT SCREEN TEXT

    try:
        screen = waitForObject(
            names.treatmentScreen_bottomBar_QFrame
        )

        screen_candidates = get_screen_candidates(screen)

        try:
            parent = object.parent(screen)

            if parent is not None:
                screen_candidates.extend(
                    get_screen_candidates(parent)
                )
        except Exception:
            pass

    except Exception as error:
        test.fail(
            "FAIL: Unable to collect Treatment screen strings: "
            + str(error)
        )
        return

    # CHECK RECORDED OBJECTS FROM THE REFERENCE SCRIPT

    recorded_objects = [
        names.topBar_buttonsBar_ButtonsBar,
        names.averagePowerWidget_progressFrame_QFrame,
        names.averagePowerWidget_unitsOfMeasureLabel_QLabel_2,
        names.averagePowerWidget_progressFrame_QFrame_2,
        names.pulseModeWidget_frequencyWidget_ParameterSettingsWidget,
        names.pulseModeWidget_buttonsWidget_QWidget_2,
        names.bottomBar_lbPedalStatusIcon_QLabel,
        names.treatmentScreen_bottomBar_QFrame
    ]

    for object_name in recorded_objects:
        try:
            obj = waitForObject(object_name)
            screen_candidates.extend(
                get_screen_candidates(obj)
            )
        except Exception:
            pass

    # SEPARATE INDIVIDUAL AND COMBINED CANDIDATES

    individual_candidates = []
    combined_candidates = []

    seen_individual = set()
    seen_combined = set()

    for candidate in screen_candidates:
        if candidate is None:
            continue

        text = str(candidate).strip()
        key = normalize(text)

        if not key:
            continue

        if is_dynamic_numeric_value(text):
            continue

        if not re.search(r"[^\W\d_]", text, re.UNICODE):
            continue

        if " " in text:
            if key not in seen_combined:
                seen_combined.add(key)
                combined_candidates.append(text)
        else:
            if key not in seen_individual:
                seen_individual.add(key)
                individual_candidates.append(text)

    test.log(
        "Screen candidates: " + str(len(screen_candidates))
    )
    test.log(
        "Individual candidates: " + str(len(individual_candidates))
    )
    test.log(
        "Combined candidates: " + str(len(combined_candidates))
    )
    test.log(
        "Bulgaria Excel entries: " + str(len(excel_strings))
    )

    # MATCH INDIVIDUAL STRINGS FIRST

    matches = {}
    missing_strings = {}

    for candidate in individual_candidates:
        key = normalize(candidate)

        if key in excel_strings:
            matches[key] = excel_strings[key]
        else:
            missing_strings[key] = candidate

    # MATCH COMBINED CANDIDATES AGAINST EXCEL

    for candidate in combined_candidates:
        key = normalize(candidate)

        if key in excel_strings:
            matches[key] = excel_strings[key]
            missing_strings.pop(key, None)

    # REPORT MATCHED STRINGS

    pass_count = 0

    for key in sorted(matches):
        data = matches[key]

        test.passes(
            "PASS: "
            + data["english"]
            + " -> "
            + data["Bulgaria"]
        )

        pass_count += 1

    # REPORT UNMATCHED INDIVIDUAL STRINGS

    fail_count = 0

    for key in sorted(missing_strings):
        test.fail(
            "FAIL: Visible screen string missing from Bulgaria Excel: "
            + missing_strings[key]
        )

        fail_count += 1

    # FINAL RESULT

    if pass_count == 0 and fail_count == 0:
        test.fail(
            "Bulgaria SOFT TISSUE QUICK START LOCALIZATION "
            "VALIDATION FAILED"
            " | No strings matched or were reported."
            " | Check screen collection and Excel column mapping."
        )

    elif fail_count == 0:
        test.passes(
            "Bulgaria SOFT TISSUE QUICK START LOCALIZATION "
            "VALIDATION PASSED"
            + " | Matched: " + str(pass_count)
            + " | Missing: 0"
        )

    else:
        test.fail(
            "Bulgaria SOFT TISSUE QUICK START LOCALIZATION "
            "VALIDATION FAILED"
            + " | Matched: " + str(pass_count)
            + " | Missing: " + str(fail_count)
        )
