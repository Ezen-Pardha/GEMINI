# -*- coding: utf-8 -*-

import names
import re
import unicodedata
from openpyxl import load_workbook
import os 
TESTCASE_DIR = os.path.dirname(os.path.abspath(__file__))
 
EXCEL_PATH = os.path.abspath(
    os.path.join( TESTCASE_DIR,
        "..",
        "shared",
        "testdata",
        "Gemini Split strings.xlsx"
    )
)


EXCEL_SHEET_NAME = "Treatment Modes"
ENGLISH_COLUMN = 2
Romanian_COLUMN = 21


def normalize_text(value):

    if value is None:
        return ""

    text = str(value)

    text = re.sub(
        r"<br\s*/?>",
        " ",
        text,
        flags=re.IGNORECASE
    )

    text = re.sub(
        r"<[^>]*>",
        " ",
        text
    )

    text = text.replace("\n", " ")
    text = text.replace("\r", " ")
    text = text.replace("\t", " ")
    text = text.replace("\xa0", " ")
    text = text.replace("\u2007", " ")
    text = text.replace("\u202f", " ")

    text = text.replace("™", "")
    text = text.replace("®", "")
    text = text.replace("©", "")

    try:
        text = unicodedata.normalize(
            "NFKC",
            text
        )
    except Exception:
        pass

    text = re.sub(
        r"\s+",
        " ",
        text
    ).strip()

    text = re.sub(
        r"\s+([.,;:!?])",
        r"\1",
        text
    )

    return text.casefold()


def comparison_key(value):

    normalized = normalize_text(value)

    return re.sub(
        r"\s+",
        "",
        normalized
    )


def load_excel_data():

    workbook = None

    try:

        workbook = load_workbook(
            EXCEL_PATH,
            read_only=True,
            data_only=True
        )

        if EXCEL_SHEET_NAME not in workbook.sheetnames:

            test.fail(
                "Excel sheet not found: "
                + EXCEL_SHEET_NAME
                + " | Available sheets: "
                + ", ".join(workbook.sheetnames)
            )

            workbook.close()
            return None, None

        sheet = workbook[EXCEL_SHEET_NAME]

        Romanian_data = {}

        for row in sheet.iter_rows(
            min_row=2,
            min_col=Romanian_COLUMN,
            max_col=Romanian_COLUMN,
            values_only=True
        ):

            Romanian_value = row[0]

            if Romanian_value is None:
                continue

            Romanian_value = str(
                Romanian_value
            ).strip()

            Romanian_key = comparison_key(
                Romanian_value
            )

            if Romanian_key == "":
                continue

            if Romanian_key not in Romanian_data:
                Romanian_data[Romanian_key] = Romanian_value

        if not Romanian_data:

            test.fail(
                "No Romanian localization data found in Excel Column U."
            )

            workbook.close()
            return None, None

        return workbook, Romanian_data

    except Exception as e:

        test.fail(
            "Unable to load/read Romanian localization Excel: "
            + str(e)
        )

        if workbook is not None:

            try:
                workbook.close()
            except Exception:
                pass

        return None, None


def get_Romanian_excel_text(
        Romanian_data,
        actual_text
        ):

    actual_key = comparison_key(
        actual_text
    )

    if actual_key == "":
        return None

    if actual_key not in Romanian_data:

        test.fail(
            "Romanian text not found in Excel Column U: "
            + str(actual_text)
        )

        return None

    return Romanian_data[actual_key]


def get_object_text(obj):

    properties = [
        "text",
        "plainText",
        "title",
        "windowTitle",
        "accessibleName"
    ]

    for property_name in properties:

        try:

            value = getattr(
                obj,
                property_name
            )

            if value is not None:

                value = str(value)

                if value.strip() != "":
                    return value

        except Exception:
            pass

    try:

        value = obj.property("text")

        if value is not None:

            value = str(value)

            if value.strip() != "":
                return value

    except Exception:
        pass

    return ""


def verify_text(
        Romanian_data,
        object_name,
        description
        ):

    snooze(2)

    try:

        obj = waitForObjectExists(
            getattr(
                names,
                object_name
            )
        )

    except Exception as e:

        test.fail(
            description
            + " | Object not found: "
            + object_name
            + " | "
            + str(e)
        )

        return False

    actual_text = get_object_text(obj)

    if comparison_key(actual_text) == "":

        test.fail(
            description
            + " | Object contains no readable text"
        )

        return False

    expected_text = get_Romanian_excel_text(
        Romanian_data,
        actual_text
    )

    if expected_text is None:

        test.fail(
            description
            + " skipped because Romanian text was not found in Excel"
        )

        return False

    actual_normalized = comparison_key(
        actual_text
    )

    expected_normalized = comparison_key(
        expected_text
    )

    test.compare(
        actual_normalized,
        expected_normalized,
        description + " - Romanian text verification"
    )

    if actual_normalized == expected_normalized:
        return True

    test.fail(
        description
        + " | Romanian text mismatch"
        + " | Actual: "
        + str(actual_text)
        + " | Expected: "
        + str(expected_text)
    )

    return False


def get_Romanian_treatment_text(english_text):

    workbook = None

    try:

        workbook = load_workbook(
            EXCEL_PATH,
            read_only=True,
            data_only=True
        )

        if EXCEL_SHEET_NAME not in workbook.sheetnames:

            test.fail(
                "Excel sheet not found: "
                + EXCEL_SHEET_NAME
            )

            return None

        sheet = workbook[EXCEL_SHEET_NAME]

        english_key = comparison_key(
            english_text
        )

        for row in sheet.iter_rows(
            min_row=2,
            min_col=ENGLISH_COLUMN,
            max_col=Romanian_COLUMN,
            values_only=True
        ):

            english_value = row[0]

            if english_value is None:
                continue

            if comparison_key(english_value) != english_key:
                continue

            Romanian_index = (
                Romanian_COLUMN - ENGLISH_COLUMN
            )

            if Romanian_index >= len(row):
                continue

            Romanian_value = row[Romanian_index]

            if Romanian_value is None:
                continue

            Romanian_value = str(
                Romanian_value
            ).strip()

            if comparison_key(Romanian_value) == "":
                continue

            return Romanian_value

        test.fail(
            "English treatment text not found in Excel Column B: "
            + str(english_text)
        )

        return None

    except Exception as e:

        test.fail(
            "Unable to find Romanian treatment translation: "
            + str(e)
        )

        return None

    finally:

        if workbook is not None:

            try:
                workbook.close()
            except Exception:
                pass


def find_text_recursively(
        parent_object,
        expected_normalized
        ):

    try:

        if not parent_object.visible:
            return None

    except Exception:
        pass

    current_text = get_object_text(
        parent_object
    )

    if comparison_key(current_text) == expected_normalized:
        return parent_object

    try:
        children = object.children(parent_object)
    except Exception:
        children = []

    for child in children:

        found = find_text_recursively(
            child,
            expected_normalized
        )

        if found is not None:
            return found

    return None


def verify_treatment_name(
        Romanian_data,
        english_text,
        description
        ):

    snooze(6)

    expected_text = get_Romanian_treatment_text(
        english_text
    )

    if expected_text is None:

        test.fail(
            description
            + " | Romanian treatment translation not found in Excel"
        )

        return False

    expected_normalized = comparison_key(
        expected_text
    )

    try:

        screen_switcher = waitForObjectExists(
            names.o_ScreenSwitcher
        )

    except Exception as e:

        test.fail(
            description
            + " | ScreenSwitcher not found | "
            + str(e)
        )

        return False

    found_object = find_text_recursively(
        screen_switcher,
        expected_normalized
    )

    if found_object is None:

        test.fail(
            description
            + " | Romanian treatment text not found on screen: ["
            + str(expected_text)
            + "]"
        )

        return False

    actual_text = get_object_text(
        found_object
    )

    actual_normalized = comparison_key(
        actual_text
    )

    test.compare(
        actual_normalized,
        expected_normalized,
        description + " - Romanian treatment verification"
    )

    if actual_normalized == expected_normalized:
        return True

    test.fail(
        description
        + " | Romanian treatment mismatch"
        + " | Actual: "
        + str(actual_text)
        + " | Expected: "
        + str(expected_text)
    )

    return False


def verify_left_treatment_name(
        Romanian_data,
        english_text,
        description
        ):

    return verify_treatment_name(
        Romanian_data,
        english_text,
        description
    )


def verify_right_treatment_name(
        Romanian_data,
        english_text,
        description
        ):

    return verify_treatment_name(
        Romanian_data,
        english_text,
        description
    )


def main():

    workbook = None
    Romanian_data = None

    result = load_excel_data()

    if result is None:

        test.fail(
            "Romanian Excel data could not be loaded"
        )

        return

    try:

        workbook, Romanian_data = result

    except Exception as e:

        test.fail(
            "Unable to unpack Romanian Excel data: "
            + str(e)
        )

        return

    if workbook is None or Romanian_data is None:

        test.fail(
            "Romanian Excel data could not be loaded"
        )

        return

    try:

        for i in range(4):

            clickButton(
                waitForObject(
                    names.numpad_pushButton_8_NumpadButton
                )
            )

            snooze(1)

        clickButton(
            waitForObject(
                names.mainFrame_okButton_QPushButton
            )
        )

        clickButton(
            waitForObject(
                names.homeScreen_settingsButton_QPushButton
            )
        )

        clickTab(
            waitForObject(
                names.settingsScreen_tbSystemInfo_TabWidget
            ),
            "INFO"
        )

        clickButton(
            waitForObject(
                names.gbSystemInformation_pbLanguageSettingsUpdate_QPushButton
            )
        )

        clickButton(
            waitForObject(
                names.languagesFrame_pbRomania_QPushButton
            )
        )

        clickButton(
            waitForObject(
                names.buttonsFrame_acceptButton_QPushButton
            )
        )

        clickButton(
            waitForObject(
                names.buttonsBar_homeButton_QPushButton
            )
        )

        clickButton(
            waitForObject(
                names.homeScreen_pbSoftTissueQuickStart_QToolButton
            )
        )

        mouseClick(
            waitForObjectItem(
                names.tabLeftMainPresets_leftPresetsView_PresetListView,
                "_1"
            ),
            238,
            58,
            Qt.NoModifier,
            Qt.LeftButton
        )

        mouseClick(
            waitForObjectItem(
                names.tabRightMainPresets_rightPresetsView_PresetListView,
                "_1"
            ),
            148,
            32,
            Qt.NoModifier,
            Qt.LeftButton
        )

        mouseClick(
            waitForObject(
                names.quickStartScreen_wdBottom_QWidget
            ),
            1167,
            9,
            Qt.NoModifier,
            Qt.LeftButton
        )

        sendEvent(
            "QMoveEvent",
            waitForObject(
                names.o_ScreenSwitcher
            ),
            0,
            -51,
            1286,
            61
        )

        clickButton(
            waitForObject(
                names.quickStartScreen_pbContinue_QPushButton
            )
        )

        clickButton(
            waitForObject(
                names.mainFrame_okButton_QPushButton_2
            )
        )

        clickButton(
            waitForObject(
                names.stateSwitch_readyButton_QPushButton
            )
        )

        verify_left_treatment_name(
            Romanian_data,
            "PROSTATE ENUCLEATION-(HIGH ENERGY)",
            "Left Treatment Name"
        )

        verify_text(
            Romanian_data,
            "averagePowerWidgetEmission_nameLabel_QLabel",
            "Left Average Power"
        )

        verify_text(
            Romanian_data,
            "pulseEnergyWidgetEmission_nameLabel_QLabel",
            "Left Pulse Energy"
        )

        verify_text(
            Romanian_data,
            "frequencyWidgetEmission_nameLabel_QLabel",
            "Left Frequency"
        )

        verify_text(
            Romanian_data,
            "modeButtonEmission_nameLabel_QLabel",
            "Left Pulse Mode"
        )

        verify_text(
            Romanian_data,
            "peakPowerButtonEmission_nameLabel_QLabel",
            "Left Peak Power"
        )

        verify_text(
            Romanian_data,
            "totalEnergyFrame_nameLabel_QLabel",
            "Left Total Energy"
        )

        verify_text(
            Romanian_data,
            "totalTimeFrame_nameLabel_QLabel",
            "Left Total Lasing Time"
        )

        verify_text(
            Romanian_data,
            "averagePowerWidgetEmission_unitsOfMeasureLabel_QLabel",
            "Left Power Unit"
        )

        verify_text(
            Romanian_data,
            "pulseEnergyWidgetEmission_unitsOfMeasureLabel_QLabel",
            "Left Energy Unit"
        )

        verify_text(
            Romanian_data,
            "frequencyWidgetEmission_unitsOfMeasureLabel_QLabel",
            "Left Frequency Unit"
        )

        verify_text(
            Romanian_data,
            "fiberFrame_nameLabel_QLabel",
            "Left Fiber"
        )

        verify_text(
            Romanian_data,
            "fiberUsesFrame_nameLabel_QLabel",
            "Left Fiber Uses Remaining"
        )

        verify_text(
            Romanian_data,
            "bottomBar_aimingBeamButton_AimingBeamModeButton",
            "Left Aiming Beam"
        )

        verify_text(
            Romanian_data,
            "stateSwitch_standbyButton_QPushButton",
            "Left Standby"
        )

        verify_text(
            Romanian_data,
            "stateSwitch_readyButton_QPushButton",
            "Left Ready"
        )

        verify_right_treatment_name(
            Romanian_data,
            "PROSTATE ENUCLEATION-(HIGH ENERGY)",
            "Right Treatment Name"
        )

        verify_text(
            Romanian_data,
            "averagePowerWidgetEmission_nameLabel_QLabel_2",
            "Right Average Power"
        )

        verify_text(
            Romanian_data,
            "pulseEnergyWidgetEmission_nameLabel_QLabel_2",
            "Right Pulse Energy"
        )

        verify_text(
            Romanian_data,
            "frequencyWidgetEmission_nameLabel_QLabel_2",
            "Right Frequency"
        )

        verify_text(
            Romanian_data,
            "modeButtonEmission_nameLabel_QLabel_2",
            "Right Pulse Mode"
        )

        verify_text(
            Romanian_data,
            "peakPowerButtonEmission_nameLabel_QLabel_2",
            "Right Peak Power"
        )

        verify_text(
            Romanian_data,
            "averagePowerWidgetEmission_unitsOfMeasureLabel_QLabel_2",
            "Right Power Unit"
        )

        verify_text(
            Romanian_data,
            "pulseEnergyWidgetEmission_unitsOfMeasureLabel_QLabel_2",
            "Right Energy Unit"
        )

        verify_text(
            Romanian_data,
            "frequencyWidgetEmission_unitsOfMeasureLabel_QLabel_2",
            "Right Frequency Unit"
        )

        clickButton(
            waitForObject(
                names.stateSwitch_readyButton_QPushButton
            )
        )

        test.vp("Readyscreen")

        clickButton(
            waitForObject(
                names.stateSwitch_readyButton_QPushButton
            )
        )

    except Exception as e:

        test.fail(
            "Romanian Ready Screen localization test failed: "
            + str(e)
        )

    finally:

        if workbook is not None:

            try:
                workbook.close()
            except Exception:
                pass