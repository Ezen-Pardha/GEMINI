
# -*- coding: utf-8 -*-

import names
import openpyxl
import re
import os

TESTCASE_DIR = os.path.dirname(os.path.abspath(__file__))

EXCEL_PATH = os.path.abspath(
    os.path.join(
        TESTCASE_DIR,
        "..",
        "shared",
        "testdata",
        "Gemini Split strings.xlsx"
    )
)


def normalize(text):
    text = str(text or "")

    # Remove HTML tags
    text = re.sub(r"<[^>]+>", "", text)

    # Normalize non-breaking spaces
    text = text.replace("\u00a0", " ")

    # Normalize tabs, newlines, and repeated whitespace
    text = " ".join(text.split())

    # Case-insensitive comparison
    return text.casefold()


def get_all_texts(obj):
    texts = []

    try:
        text = str(obj.text).strip()

        if text:
            if "<html>" in text.lower():
                text = re.sub(r"<[^>]+>", "", text)
                text = " ".join(text.split())

            if text:
                texts.append(text)

    except:
        pass

    try:
        children = object.children(obj)

        for child in children:
            texts.extend(get_all_texts(child))

    except:
        pass

    return texts


def ignore_text(text):
    text = str(text).strip()

    if not text:
        return True

    # Ignore separators and placeholders
    if re.fullmatch(r"[-_=.]+", text):
        return True

    return False


def main():

    # Load Excel using the relative path
    workbook = openpyxl.load_workbook(
        EXCEL_PATH,
        data_only=True,
        read_only=True
    )

    try:
        sheet = workbook["Soft Tissue-Guided"]

        # Load all France translations from Excel.
        # Column B = English
        # Column J = French (zero-based index 9)

        excel_strings = {}

        for row in sheet.iter_rows(
            min_row=2,
            values_only=True
        ):

            if len(row) < 10:
                continue

            english_value = row[1]
            France_value = row[9]

            if France_value is None:
                continue

            english = (
                str(english_value).strip()
                if english_value is not None
                else ""
            )
            France = str(France_value).strip()

            if ignore_text(France):
                continue

            key = normalize(France)

            if key not in excel_strings:
                excel_strings[key] = []

            excel_strings[key].append(english)

        # Login

        doubleClick(
            waitForObject(names.numpad_pushButton_8_NumpadButton),
            68, 52,
            Qt.NoModifier,
            Qt.LeftButton
        )

        doubleClick(
            waitForObject(names.numpad_pushButton_8_NumpadButton),
            116, 46,
            Qt.NoModifier,
            Qt.LeftButton
        )

        clickButton(
            waitForObject(names.mainFrame_okButton_QPushButton)
        )

        # Settings

        clickButton(
            waitForObject(names.homeScreen_settingsButton_QPushButton)
        )

        clickTab(
            waitForObject(names.settingsScreen_tbSystemInfo_TabWidget),
            "INFO"
        )

        clickButton(
            waitForObject(
                names.gbSystemInformation_pbLanguageSettingsUpdate_QPushButton
            )
        )

        clickButton(
            waitForObject(names.languagesFrame_pbFrance_QPushButton)
        )

        clickButton(
            waitForObject(names.buttonsFrame_acceptButton_QPushButton)
        )

        clickButton(
            waitForObject(names.settingsScreen_pbBack_QPushButton)
        )

        # Open Soft Tissue Assistant

        clickButton(
            waitForObject(
                names.homeScreen_pbSoftTissueAssistant_QToolButton
            )
        )

        mouseClick(
            waitForObject(
                names.softTissueAssistantModeScreen_gbSoftTissueProcedures_QFrame
            ),
            90, 150,
            Qt.NoModifier,
            Qt.LeftButton
        )

        # Read all visible/available object text from the screen

        soft_tissue = waitForObject(
            names.softTissueAssistantModeScreen_SoftTissueAssistantModeScreen
        )

        actual_texts = get_all_texts(soft_tissue)

        # Remove duplicates while preserving the original order

        unique_texts = []
        seen = set()

        for text in actual_texts:

            if ignore_text(text):
                continue

            key = normalize(text)

            if key not in seen:
                seen.add(key)
                unique_texts.append(text)

        # Verify every collected screen string against Excel

        total_pass = 0
        total_fail = 0

        for text in unique_texts:

            key = normalize(text)

            if key in excel_strings:

                english_values = excel_strings[key]
                english = next(
                    (value for value in english_values if value),
                    ""
                )

                if english:
                    test.passes(
                        "PASS: {} -> {}".format(english, text)
                    )
                else:
                    test.passes(
                        "PASS: France string found in Excel -> {}".format(
                            text
                        )
                    )

                total_pass += 1

            else:

                test.fail(
                    "FAIL: Screen string not found in France Excel: {}".format(
                        text
                    )
                )

                total_fail += 1

        # Final summary

        test.log(
            "Localization Summary | Total: {} | PASS: {} | FAIL: {}".format(
                len(unique_texts),
                total_pass,
                total_fail
            )
        )

    finally:
        workbook.close()
