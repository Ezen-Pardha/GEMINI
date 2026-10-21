# -*- coding: utf-8 -*-

import names
import openpyxl
import re
import unicodedata
import os 

TESTCASE_DIR = os.path.dirname(os.path.abspath(__file__))
 
EXCEL_PATH = os.path.abspath(
    os.path.join( TESTCASE_DIR,
        "..",
        "shared",
        "testdata",
        "Gemini Split strings.xlsx"
    )
)


def normalize(text):
    if text is None:
        return ""

    text = str(text)

    text = re.sub(
        r"<br\s*/?>",
        "",
        text,
        flags=re.IGNORECASE
    )

    text = re.sub(
        r"<[^>]+>",
        "",
        text
    )

    text = text.replace("\xa0", "")
    text = text.replace("\u2007", "")
    text = text.replace("\u202f", "")

    text = text.replace("™", "")
    text = text.replace("®", "")
    text = text.replace("©", "")

    try:
        text = unicodedata.normalize("NFKC", text)
    except Exception:
        pass

    text = re.sub(r"\s+", "", text)

    return text.casefold()


def load_french_excel(EXCEL_PATH):
    workbook = None

    try:
        workbook = openpyxl.load_workbook(
            EXCEL_PATH,
            data_only=True,
            read_only=True
        )

        sheet_name = "Treatment Modes"

        if sheet_name not in workbook.sheetnames:
            test.fail(
                "Worksheet '" + sheet_name
                + "' does not exist in Excel file. Available sheets: "
                + ", ".join(workbook.sheetnames)
            )

            workbook.close()
            return None, None, {}

        sheet = workbook[sheet_name]

        english_column = 2
        french_column = 10

        french_data = {}

        for row in sheet.iter_rows(
            min_row=2,
            min_col=english_column,
            max_col=french_column
        ):
            english_value = row[0].value
            french_value = row[french_column - english_column].value

            if english_value is None:
                continue

            english_key = normalize(english_value)

            if not english_key:
                continue

            french_value = (
                "" if french_value is None
                else str(french_value).strip()
            )

            if english_key in french_data:
                existing_values = french_data[english_key].split("|||")

                if not any(
                    normalize(value) == normalize(french_value)
                    for value in existing_values
                ):
                    french_data[english_key] += "|||" + french_value

            else:
                french_data[english_key] = french_value

        if not french_data:
            test.fail(
                "No English/French localization data found in Excel."
            )
            workbook.close()
            return None, None, {}

        return workbook, sheet, french_data

    except Exception as e:
        test.fail(
            "Unable to load French localization Excel: " + str(e)
        )

        if workbook is not None:
            try:
                workbook.close()
            except Exception:
                pass

        return None, None, {}


def get_french_text(french_data, english_text):
    key = normalize(english_text)

    if key not in french_data:
        return ""

    return french_data[key]


def get_object_text(obj):
    for property_name in (
        "text",
        "plainText",
        "title",
        "windowTitle",
        "accessibleName"
    ):
        try:
            value = getattr(obj, property_name)

            if value is not None:
                value = str(value)

                if value.strip():
                    return value

        except Exception:
            pass

    return ""


def verify_text(french_data, object_name, english_text):
    try:
        obj = waitForObjectExists(object_name)

    except Exception as e:
        test.fail(
            "Unable to find application object: "
            + str(object_name)
            + " | Error: "
            + str(e)
        )
        return False

    actual_text = get_object_text(obj)

    if not actual_text.strip():
        test.fail(
            "Application object contains no readable text: "
            + str(english_text)
        )
        return False

    expected_text = get_french_text(
        french_data,
        english_text
    )

    if not expected_text:
        test.fail(
            "English text not found in Excel Column B: "
            + str(english_text)
        )
        return False

    expected_values = expected_text.split("|||")
    actual_normalized = normalize(actual_text)

    for value in expected_values:
        if actual_normalized == normalize(value):
            test.compare(
                actual_normalized,
                normalize(value),
                "French text verification: " + str(english_text)
            )
            return True

    test.fail(
        "French text mismatch for: "
        + str(english_text)
        + " | Actual: "
        + str(actual_text)
        + " | Expected: "
        + str(expected_text)
    )

    return False


def main():
    

    workbook = None

    try:
        for i in range(4):
            clickButton(
                waitForObject(
                    names.numpad_pushButton_8_NumpadButton
                )
            )

        snooze(6)

        clickButton(
            waitForObject(
                names.mainFrame_okButton_QPushButton
            )
        )

        clickButton(
            waitForObject(
                names.homeScreen_settingsButton_QPushButton
            )
        )

        clickTab(
            waitForObject(
                names.settingsScreen_tbSystemInfo_TabWidget
            ),
            "INFO"
        )

        clickButton(
            waitForObject(
                names.gbSystemInformation_pbLanguageSettingsUpdate_QPushButton
            )
        )

        clickButton(
            waitForObject(
                names.languagesFrame_pbFrance_QPushButton
            )
        )

        clickButton(
            waitForObject(
                names.buttonsFrame_acceptButton_QPushButton
            )
        )

        clickButton(
            waitForObject(
                names.settingsScreen_pbBack_QPushButton
            )
        )

        clickButton(
            waitForObject(
                names.homeScreen_pbSoftTissueQuickStart_QToolButton
            )
        )

        mouseClick(
            waitForObjectItem(
                names.tabLeftMainPresets_leftPresetsView_PresetListView,
                "_1"
            ),
            305,
            35,
            Qt.NoModifier,
            Qt.LeftButton
        )

        mouseClick(
            waitForObjectItem(
                names.tabRightMainPresets_rightPresetsView_PresetListView,
                "_1"
            ),
            299,
            49,
            Qt.NoModifier,
            Qt.LeftButton
        )

        clickButton(
            waitForObject(
                names.quickStartScreen_pbContinue_QPushButton
            )
        )

        clickButton(
            waitForObject(
                names.mainFrame_okButton_QPushButton_2
            )
        )

        mouseClick(
            waitForObject(
                names.peakPowerButton_nameLabel_QLabel
            ),
            34,
            34,
            Qt.NoModifier,
            Qt.LeftButton
        )

        snooze(6)

        workbook, sheet, french_data = load_french_excel(
            EXCEL_PATH
        )

        if workbook is None or not french_data:
            test.fail("French Excel data could not be loaded.")
            return

        snooze(6)

        verify_text(
            french_data,
            names.centralWidget_lbTitle_QLabel,
            "SELECT PEAK POWER"
        )

        snooze(2)

        verify_text(
            french_data,
            names.centralWidget_pw500_QPushButton,
            "High"
        )

        snooze(2)

        verify_text(
            french_data,
            names.centralWidget_pw250_QPushButton,
            "Medium"
        )

        snooze(2)

        verify_text(
            french_data,
            names.centralWidget_pw125_QPushButton,
            "Low"
        )

        snooze(2)

        verify_text(
            french_data,
            names.buttonsFrame_rejectButton_QPushButton,
            "CANCEL"
        )

        mouseClick(
            waitForObject(
                names.centralWidget_mainWidget_QWidget
            ),
            460,
            549,
            Qt.NoModifier,
            Qt.LeftButton
        )

        clickButton(
            waitForObject(
                names.buttonsFrame_rejectButton_QPushButton
            )
        )

        clickButton(
            waitForObject(
                names.buttonsBar_homeButton_QPushButton_2
            )
        )

        clickButton(
            waitForObject(
                names.homeScreen_pbSoftTissueAssistant_QToolButton
            )
        )

        clickButton(
            waitForObject(
                names.gbSoftTissueProcedures_pbIncision_QPushButton
            )
        )

        clickButton(
            waitForObject(
                names.softTissueAssistantModeScreen_pbContinue_QPushButton
            )
        )

        clickButton(
            waitForObject(
                names.treatmentCharacteristicsBar_pbEdit_QPushButton
            )
        )

        snooze(6)

        verify_text(
            french_data,
            names.centralWidget_titleLabel_QLabel,
            "SOFT TISSUE CHARACTERISTICS"
        )

        snooze(2)

        verify_text(
            french_data,
            names.centralWidget_titleLabel_QLabel_2,
            "PROCEDURE"
        )

        snooze(2)

        verify_text(
            french_data,
            names.centralWidget_valueLabel_QLabel,
            "Incision"
        )

        snooze(2)

        verify_text(
            french_data,
            names.centralWidget_rejectButton_QPushButton,
            "EDIT"
        )

        snooze(2)

        test.vp("Soft tissue chararectristics")

        verify_text(
            french_data,
            names.centralWidget_acceptButton_QPushButton,
            "OK"
        )

    except Exception as e:
        test.fail(
            "French localization test failed: " + str(e)
        )

    finally:
        if workbook is not None:
            try:
                workbook.close()
            except Exception:
                pass
