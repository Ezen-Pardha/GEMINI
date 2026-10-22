# -*- coding: utf-8 -*-

import names
import re
import unicodedata
from openpyxl import load_workbook
import os 

TESTCASE_DIR = os.path.dirname(os.path.abspath(__file__))
 
EXCEL_PATH = os.path.abspath(
    os.path.join( TESTCASE_DIR,
        "..",
        "shared",
        "testdata",
        "Gemini Split strings.xlsx"
    )
)



EXCEL_SHEET_NAME = "Treatment Modes"
ENGLISH_COLUMN = 2
FRENCH_COLUMN = 10


def normalize_text(value):
    if value is None:
        return ""

    text = str(value)

    text = re.sub(r"<[^>]*>", " ", text)

    text = text.replace("\n", " ")
    text = text.replace("\r", " ")
    text = text.replace("\t", " ")
    text = text.replace("\xa0", " ")
    text = text.replace("\u2007", " ")
    text = text.replace("\u202f", " ")

    text = text.replace("™", "")
    text = text.replace("®", "")
    text = text.replace("©", "")

    try:
        text = unicodedata.normalize("NFKC", text)
    except Exception:
        pass

    text = re.sub(r"\s+", " ", text).strip()
    text = re.sub(r"\s+([.,;:!?])", r"\1", text)

    return text.casefold()


def comparison_key(value):
    normalized = normalize_text(value)
    return re.sub(r"\s+", "", normalized)


def load_excel_data():
    workbook = None

    try:
        workbook = load_workbook(
            EXCEL_PATH,
            read_only=True,
            data_only=True
        )

        if EXCEL_SHEET_NAME not in workbook.sheetnames:
            test.fail(
                "Excel sheet not found: "
                + EXCEL_SHEET_NAME
                + " | Available sheets: "
                + ", ".join(workbook.sheetnames)
            )

            workbook.close()
            return None, None

        sheet = workbook[EXCEL_SHEET_NAME]
        french_data = {}

        for row in sheet.iter_rows(
            min_row=2,
            min_col=FRENCH_COLUMN,
            max_col=FRENCH_COLUMN
        ):
            french_value = row[0].value

            if french_value is None:
                continue

            french_value = str(french_value).strip()
            french_key = comparison_key(french_value)

            if not french_key:
                continue

            if french_key not in french_data:
                french_data[french_key] = french_value

        if not french_data:
            test.fail(
                "No French localization data found in Excel Column J."
            )

            workbook.close()
            return None, None

        return workbook, french_data

    except Exception as e:
        test.fail(
            "Unable to load/read French localization Excel: "
            + str(e)
        )

        if workbook is not None:
            try:
                workbook.close()
            except Exception:
                pass

        return None, None


def get_french_excel_text(french_data, actual_text):
    actual_key = comparison_key(actual_text)

    if not actual_key:
        return None

    return french_data.get(actual_key)


def get_object_text(obj):
    for property_name in (
        "text",
        "title",
        "windowTitle",
        "accessibleName"
    ):
        try:
            value = getattr(obj, property_name)

            if value is not None:
                value = str(value)

                if value.strip():
                    return value

        except Exception:
            pass

    return ""


def verify_text(french_data, object_name, description):
    snooze(2)

    try:
        obj = waitForObjectExists(
            getattr(names, object_name)
        )

    except Exception as e:
        test.fail(
            description
            + " | Object not found: "
            + object_name
            + " | "
            + str(e)
        )
        return False

    actual_text = get_object_text(obj)

    if not comparison_key(actual_text):
        test.fail(
            description + " | Object contains no readable text"
        )
        return False

    expected_text = get_french_excel_text(
        french_data,
        actual_text
    )

    if expected_text is None:
        test.fail(
            description
            + " | French text not found in Excel Column J: "
            + str(actual_text)
        )
        return False

    actual_normalized = comparison_key(actual_text)
    expected_normalized = comparison_key(expected_text)

    test.compare(
        actual_normalized,
        expected_normalized,
        description + " - French text verification"
    )

    if actual_normalized == expected_normalized:
        return True

    test.fail(
        description
        + " | French text mismatch"
        + " | Actual: "
        + str(actual_text)
        + " | Expected: "
        + str(expected_text)
    )

    return False


def find_text_recursively(parent_object, expected_normalized):
    try:
        if not parent_object.visible:
            return None
    except Exception:
        pass

    current_text = get_object_text(parent_object)

    if comparison_key(current_text) == expected_normalized:
        return parent_object

    try:
        children = object.children(parent_object)
    except Exception:
        children = []

    for child in children:
        found = find_text_recursively(
            child,
            expected_normalized
        )

        if found is not None:
            return found

    return None


def verify_treatment_name(french_data, english_text, description):
    snooze(6)
    workbook = None
    expected_text = None

    try:
        workbook = load_workbook(
            EXCEL_PATH,
            read_only=True,
            data_only=True
        )

        if EXCEL_SHEET_NAME not in workbook.sheetnames:
            test.fail(
                "Excel sheet not found: "
                + EXCEL_SHEET_NAME
            )
            return False

        sheet = workbook[EXCEL_SHEET_NAME]
        english_key = comparison_key(english_text)

        for row in sheet.iter_rows(
            min_row=2,
            min_col=ENGLISH_COLUMN,
            max_col=FRENCH_COLUMN
        ):
            english_value = row[0].value

            if english_value is None:
                continue

            if comparison_key(english_value) != english_key:
                continue

            french_value = row[
                FRENCH_COLUMN - ENGLISH_COLUMN
            ].value

            if french_value is None:
                continue

            french_value = str(french_value).strip()

            if comparison_key(french_value):
                expected_text = french_value
                break

        if expected_text is None:
            test.fail(
                description
                + " | English treatment text not found in Excel Column B: "
                + str(english_text)
            )
            return False

    except Exception as e:
        test.fail(
            description
            + " | Unable to find treatment translation: "
            + str(e)
        )
        return False

    finally:
        if workbook is not None:
            try:
                workbook.close()
            except Exception:
                pass

    expected_normalized = comparison_key(expected_text)

    try:
        screen_switcher = waitForObjectExists(
            names.o_ScreenSwitcher
        )

    except Exception as e:
        test.fail(
            description
            + " | ScreenSwitcher not found | "
            + str(e)
        )
        return False

    found_object = find_text_recursively(
        screen_switcher,
        expected_normalized
    )

    if found_object is None:
        test.fail(
            description
            + " | French treatment text not found on screen: "
            + str(expected_text)
        )
        return False

    actual_text = get_object_text(found_object)
    actual_normalized = comparison_key(actual_text)

    test.compare(
        actual_normalized,
        expected_normalized,
        description + " - French treatment verification"
    )

    if actual_normalized == expected_normalized:
        return True

    test.fail(
        description
        + " | French treatment mismatch"
        + " | Actual: "
        + str(actual_text)
        + " | Expected: "
        + str(expected_text)
    )

    return False


def verify_left_treatment_name(
    french_data,
    english_text,
    description
):
    return verify_treatment_name(
        french_data,
        english_text,
        description
    )


def verify_right_treatment_name(
    french_data,
    english_text,
    description
):
    return verify_treatment_name(
        french_data,
        english_text,
        description
    )


def main():
    workbook = None
    french_data = None

    workbook, french_data = load_excel_data()

    if workbook is None or french_data is None:
        test.fail("French Excel data could not be loaded")
        return

    try:
        for i in range(4):
            clickButton(
                waitForObject(
                    names.numpad_pushButton_8_NumpadButton
                )
            )
            snooze(1)

        clickButton(
            waitForObject(
                names.mainFrame_okButton_QPushButton
            )
        )

        clickButton(
            waitForObject(
                names.homeScreen_settingsButton_QPushButton
            )
        )

        clickTab(
            waitForObject(
                names.settingsScreen_tbSystemInfo_TabWidget
            ),
            "INFO"
        )

        clickButton(
            waitForObject(
                names.gbSystemInformation_pbLanguageSettingsUpdate_QPushButton
            )
        )

        clickButton(
            waitForObject(
                names.languagesFrame_pbFrance_QPushButton
            )
        )

        clickButton(
            waitForObject(
                names.buttonsFrame_acceptButton_QPushButton
            )
        )

        clickButton(
            waitForObject(
                names.buttonsBar_homeButton_QPushButton
            )
        )

        clickButton(
            waitForObject(
                names.homeScreen_pbSoftTissueQuickStart_QToolButton
            )
        )

        mouseClick(
            waitForObjectItem(
                names.tabLeftMainPresets_leftPresetsView_PresetListView,
                "_1"
            ),
            238,
            58,
            Qt.NoModifier,
            Qt.LeftButton
        )

        mouseClick(
            waitForObjectItem(
                names.tabRightMainPresets_rightPresetsView_PresetListView,
                "_1"
            ),
            148,
            32,
            Qt.NoModifier,
            Qt.LeftButton
        )

        mouseClick(
            waitForObject(
                names.quickStartScreen_wdBottom_QWidget
            ),
            1167,
            9,
            Qt.NoModifier,
            Qt.LeftButton
        )

        sendEvent(
            "QMoveEvent",
            waitForObject(
                names.o_ScreenSwitcher
            ),
            0,
            -51,
            1286,
            61
        )

        clickButton(
            waitForObject(
                names.quickStartScreen_pbContinue_QPushButton
            )
        )

        clickButton(
            waitForObject(
                names.mainFrame_okButton_QPushButton_2
            )
        )

        clickButton(
            waitForObject(
                names.stateSwitch_readyButton_QPushButton
            )
        )

        verify_left_treatment_name(
            french_data,
            "PROSTATE ENUCLEATION-(HIGH ENERGY)",
            "Left Treatment Name"
        )

        verify_text(
            french_data,
            "averagePowerWidgetEmission_nameLabel_QLabel",
            "Left Average Power"
        )

        verify_text(
            french_data,
            "pulseEnergyWidgetEmission_nameLabel_QLabel",
            "Left Pulse Energy"
        )

        verify_text(
            french_data,
            "frequencyWidgetEmission_nameLabel_QLabel",
            "Left Frequency"
        )

        verify_text(
            french_data,
            "modeButtonEmission_nameLabel_QLabel",
            "Left Pulse Mode"
        )

        verify_text(
            french_data,
            "peakPowerButtonEmission_nameLabel_QLabel",
            "Left Peak Power"
        )

        verify_text(
            french_data,
            "totalEnergyFrame_nameLabel_QLabel",
            "Left Total Energy"
        )

        verify_text(
            french_data,
            "totalTimeFrame_nameLabel_QLabel",
            "Left Total Lasing Time"
        )

        verify_text(
            french_data,
            "averagePowerWidgetEmission_unitsOfMeasureLabel_QLabel",
            "Left Power Unit"
        )

        verify_text(
            french_data,
            "pulseEnergyWidgetEmission_unitsOfMeasureLabel_QLabel",
            "Left Energy Unit"
        )

        verify_text(
            french_data,
            "frequencyWidgetEmission_unitsOfMeasureLabel_QLabel",
            "Left Frequency Unit"
        )

        verify_text(
            french_data,
            "fiberFrame_nameLabel_QLabel",
            "Left Fiber"
        )

        verify_text(
            french_data,
            "fiberUsesFrame_nameLabel_QLabel",
            "Left Fiber Uses Remaining"
        )

        verify_text(
            french_data,
            "bottomBar_aimingBeamButton_AimingBeamModeButton",
            "Left Aiming Beam"
        )

        verify_text(
            french_data,
            "stateSwitch_standbyButton_QPushButton",
            "Left Standby"
        )

        verify_text(
            french_data,
            "stateSwitch_readyButton_QPushButton",
            "Left Ready"
        )

        verify_right_treatment_name(
            french_data,
            "PROSTATE ENUCLEATION-(HIGH ENERGY)",
            "Right Treatment Name"
        )

        verify_text(
            french_data,
            "averagePowerWidgetEmission_nameLabel_QLabel_2",
            "Right Average Power"
        )

        verify_text(
            french_data,
            "pulseEnergyWidgetEmission_nameLabel_QLabel_2",
            "Right Pulse Energy"
        )

        verify_text(
            french_data,
            "frequencyWidgetEmission_nameLabel_QLabel_2",
            "Right Frequency"
        )

        verify_text(
            french_data,
            "modeButtonEmission_nameLabel_QLabel_2",
            "Right Pulse Mode"
        )

        verify_text(
            french_data,
            "peakPowerButtonEmission_nameLabel_QLabel_2",
            "Right Peak Power"
        )

        verify_text(
            french_data,
            "averagePowerWidgetEmission_unitsOfMeasureLabel_QLabel_2",
            "Right Power Unit"
        )

        verify_text(
            french_data,
            "pulseEnergyWidgetEmission_unitsOfMeasureLabel_QLabel_2",
            "Right Energy Unit"
        )

        verify_text(
            french_data,
            "frequencyWidgetEmission_unitsOfMeasureLabel_QLabel_2",
            "Right Frequency Unit"
        )

        clickButton(
            waitForObject(
                names.stateSwitch_readyButton_QPushButton
            )
        )

       
        test.vp("Readyscreen")

        clickButton(
            waitForObject(
                names.stateSwitch_readyButton_QPushButton
            )
        )

    except Exception as e:
        test.fail(
            "French Ready Screen test failed: " + str(e)
        )

    finally:
        if workbook is not None:
            try:
                workbook.close()
            except Exception:
                pass
