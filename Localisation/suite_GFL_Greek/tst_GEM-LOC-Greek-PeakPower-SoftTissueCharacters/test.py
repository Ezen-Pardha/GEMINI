# -*- coding: utf-8 -*-

import names
import openpyxl
import re
import unicodedata


def normalize(text):

    if text is None:
        return ""

    text = str(text)

    text = re.sub(
        r"<br\s*/?>",
        "",
        text,
        flags=re.IGNORECASE
    )

    text = re.sub(
        r"<[^>]+>",
        "",
        text
    )

    try:
        text = unicodedata.normalize(
            "NFKC",
            text
        )
    except Exception:
        pass

    text = (
        text
        .replace("\xa0", "")
        .replace("\u2007", "")
        .replace("\u202f", "")
        .replace("™", "")
        .replace("®", "")
        .replace("©", "")
    )

    text = re.sub(
        r"\s+",
        "",
        text
    )

    return text.casefold()


def load_greek_excel(excelpath):

    try:
        workbook = openpyxl.load_workbook(
            excelpath,
            data_only=True,
            read_only=True
        )

    except Exception as e:
        test.fail(
            "Unable to load Greek localization Excel: "
            + str(e)
        )
        return None, None, {}

    sheet_name = "Treatment Modes"

    if sheet_name not in workbook.sheetnames:

        workbook.close()

        test.fail(
            "Worksheet '"
            + sheet_name
            + "' does not exist in Excel file"
        )

        return None, None, {}

    sheet = workbook[sheet_name]

    english_column = 2
    greek_column = 12

    greek_data = {}

    greek_index = (
        greek_column
        - english_column
    )

    for row in sheet.iter_rows(
        min_col=english_column,
        max_col=greek_column,
        values_only=True
    ):

        english_value = row[0]

        if english_value is None:
            continue

        english_key = normalize(
            english_value
        )

        if english_key == "":
            continue

        if greek_index >= len(row):
            continue

        greek_value = row[greek_index]

        greek_value = (
            ""
            if greek_value is None
            else str(greek_value).strip()
        )

        if english_key in greek_data:

            existing_values = greek_data[
                english_key
            ].split("|||")

            duplicate_found = False

            for existing_translation in existing_values:

                if normalize(
                    existing_translation
                ) == normalize(greek_value):

                    duplicate_found = True
                    break

            if not duplicate_found:

                greek_data[english_key] = (
                    greek_data[english_key]
                    + "|||"
                    + greek_value
                )

        else:

            greek_data[english_key] = greek_value

    return workbook, sheet, greek_data


def get_greek_text(greek_data, english_text):

    key = normalize(
        english_text
    )

    if key not in greek_data:
        return ""

    return greek_data[key]


def get_object_text(obj):

    try:
        value = obj.text

        if value is not None:
            return str(value)

    except Exception:
        pass

    try:
        value = obj.plainText

        if value is not None:
            return str(value)

    except Exception:
        pass

    try:
        value = obj.title

        if value is not None:
            return str(value)

    except Exception:
        pass

    try:
        value = obj.windowTitle

        if value is not None:
            return str(value)

    except Exception:
        pass

    try:
        value = obj.accessibleName

        if value is not None:
            return str(value)

    except Exception:
        pass

    try:
        value = obj.property("text")

        if value is not None:
            return str(value)

    except Exception:
        pass

    return ""


def verify_text(greek_data, object_name, english_text):

    try:
        obj = waitForObjectExists(
            object_name
        )

    except Exception as e:

        test.fail(
            "Unable to find application object: "
            + str(object_name)
            + " | Error: "
            + str(e)
        )

        return False

    actual_text = get_object_text(
        obj
    )

    if actual_text is None:
        actual_text = ""

    expected_text = get_greek_text(
        greek_data,
        english_text
    )

    if expected_text == "":

        test.fail(
            "English text not found in Excel Column B: "
            + str(english_text)
        )

        return False

    expected_values = expected_text.split(
        "|||"
    )

    actual_normalized = normalize(
        actual_text
    )

    matched_expected = None

    for value in expected_values:

        expected_normalized = normalize(
            value
        )

        if actual_normalized == expected_normalized:

            matched_expected = expected_normalized
            break

    if matched_expected is not None:

        test.compare(
            actual_normalized,
            matched_expected,
            "Greek text verification: "
            + str(english_text)
        )

        return True

    test.fail(
        "Greek text mismatch for: "
        + str(english_text)
        + " | Actual: "
        + actual_text
        + " | Expected: "
        + expected_text
    )

    return False


def main():

    excelpath = (
        "/home/ezen/Test_Automation_Gemini/Localisation/"
        "suite_GFL_Greek/"
        "tst_GEM-LOC-Greek-PeakPower-SoftTissueCharacters/"
        "testdata/Gemini Split strings.xlsx"
    )

    workbook = None

    clickButton(
        waitForObject(
            names.numpad_pushButton_8_NumpadButton
        )
    )

    clickButton(
        waitForObject(
            names.numpad_pushButton_8_NumpadButton
        )
    )

    clickButton(
        waitForObject(
            names.numpad_pushButton_8_NumpadButton
        )
    )

    clickButton(
        waitForObject(
            names.numpad_pushButton_8_NumpadButton
        )
    )

    snooze(6)

    clickButton(
        waitForObject(
            names.mainFrame_okButton_QPushButton
        )
    )

    clickButton(
        waitForObject(
            names.homeScreen_settingsButton_QPushButton
        )
    )

    clickTab(
        waitForObject(
            names.settingsScreen_tbSystemInfo_TabWidget
        ),
        "INFO"
    )

    clickButton(
        waitForObject(
            names.gbSystemInformation_pbLanguageSettingsUpdate_QPushButton
        )
    )

    clickButton(
        waitForObject(
            names.languagesFrame_pbGreece_QPushButton
        )
    )

    clickButton(
        waitForObject(
            names.buttonsFrame_acceptButton_QPushButton
        )
    )

    clickButton(
        waitForObject(
            names.settingsScreen_pbBack_QPushButton
        )
    )

    clickButton(
        waitForObject(
            names.homeScreen_pbSoftTissueQuickStart_QToolButton
        )
    )

    mouseClick(
        waitForObjectItem(
            names.tabLeftMainPresets_leftPresetsView_PresetListView,
            "_1"
        ),
        305,
        35,
        Qt.NoModifier,
        Qt.LeftButton
    )

    mouseClick(
        waitForObjectItem(
            names.tabRightMainPresets_rightPresetsView_PresetListView,
            "_1"
        ),
        299,
        49,
        Qt.NoModifier,
        Qt.LeftButton
    )

    clickButton(
        waitForObject(
            names.quickStartScreen_pbContinue_QPushButton
        )
    )

    clickButton(
        waitForObject(
            names.mainFrame_okButton_QPushButton_2
        )
    )

    mouseClick(
        waitForObject(
            names.peakPowerButton_nameLabel_QLabel
        ),
        34,
        34,
        Qt.NoModifier,
        Qt.LeftButton
    )

    snooze(6)

    workbook, sheet, greek_data = load_greek_excel(
        excelpath
    )

    if workbook is None:
        return

    try:

        snooze(6)

        verify_text(
            greek_data,
            names.centralWidget_lbTitle_QLabel,
            "SELECT PEAK POWER"
        )

        snooze(6)

        verify_text(
            greek_data,
            names.centralWidget_pw500_QPushButton,
            "High"
        )

        snooze(6)

        verify_text(
            greek_data,
            names.centralWidget_pw250_QPushButton,
            "Medium"
        )

        snooze(6)

        verify_text(
            greek_data,
            names.centralWidget_pw125_QPushButton,
            "Low"
        )

        snooze(6)
        test.vp("PeakPower")
        verify_text(
            greek_data,
            names.buttonsFrame_rejectButton_QPushButton,
            "CANCEL"
        )

        mouseClick(
            waitForObject(
                names.centralWidget_mainWidget_QWidget
            ),
            460,
            549,
            Qt.NoModifier,
            Qt.LeftButton
        )

        clickButton(
            waitForObject(
                names.buttonsFrame_rejectButton_QPushButton
            )
        )

        clickButton(
            waitForObject(
                names.buttonsBar_homeButton_QPushButton_2
            )
        )

        clickButton(
            waitForObject(
                names.homeScreen_pbSoftTissueAssistant_QToolButton
            )
        )

        clickButton(
            waitForObject(
                names.gbSoftTissueProcedures_pbIncision_QPushButton
            )
        )

        clickButton(
            waitForObject(
                names.softTissueAssistantModeScreen_pbContinue_QPushButton
            )
        )

        clickButton(
            waitForObject(
                names.treatmentCharacteristicsBar_pbEdit_QPushButton
            )
        )

        snooze(6)

        verify_text(
            greek_data,
            names.centralWidget_titleLabel_QLabel,
            "SOFT TISSUE CHARACTERISTICS"
        )

        snooze(6)

        verify_text(
            greek_data,
            names.centralWidget_titleLabel_QLabel_2,
            "PROCEDURE"
        )

        snooze(6)

        verify_text(
            greek_data,
            names.centralWidget_valueLabel_QLabel,
            "Incision"
        )

        snooze(6)

        verify_text(
            greek_data,
            names.centralWidget_rejectButton_QPushButton,
            "EDIT"
        )

        snooze(6)

        verify_text(
            greek_data,
            names.centralWidget_acceptButton_QPushButton,
            "OK"
        )

    finally:

        try:
            workbook.close()
        except Exception:
            pass