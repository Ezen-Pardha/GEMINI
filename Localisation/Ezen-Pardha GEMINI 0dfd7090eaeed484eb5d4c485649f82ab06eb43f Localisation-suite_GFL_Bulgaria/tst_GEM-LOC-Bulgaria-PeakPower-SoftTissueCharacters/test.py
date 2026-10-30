
# -*- coding: utf-8 -*-

import names
import openpyxl
import re
import os

TESTCASE_DIR = os.path.dirname(os.path.abspath(__file__))
 
EXCEL_PATH = os.path.abspath(
    os.path.join( TESTCASE_DIR,
        "..",
        "shared",
        "testdata",
        "Gemini Split strings.xlsx"
    )
)


def normalize(text):

    if text is None:
        return ""

    text = str(text)

    # Remove HTML line breaks/tags
    text = re.sub(
        r"<br\s*/?>",
        " ",
        text,
        flags=re.IGNORECASE
    )

    text = re.sub(
        r"<[^>]+>",
        " ",
        text
    )

    # Replace non-breaking spaces
    text = text.replace("\xa0", " ")

    # Remove extra spaces/newlines
    text = " ".join(text.split())

    return text.strip().casefold()


def load_bulgarian_excel(EXCEL_PATH ):

    workbook = openpyxl.load_workbook(
        EXCEL_PATH ,
        data_only=True,
        read_only=True
    )

    if "Treatment Modes" not in workbook.sheetnames:

        workbook.close()

        test.fail(
            "Worksheet 'Treatment Modes' does not exist "
            "in Excel file"
        )

        return None, None, {}

    sheet = workbook["Treatment Modes"]

    english_column = 2
    bulgarian_column = 3

    bulgarian_data = {}

    for row in sheet.iter_rows(
        min_col=english_column,
        max_col=bulgarian_column
    ):

        english_value = row[0].value
        bulgarian_value = row[1].value

        if english_value is None:
            continue

        english_key = normalize(
            english_value
        )

        if english_key == "":
            continue

        bulgarian_data[english_key] = (
            ""
            if bulgarian_value is None
            else str(bulgarian_value)
        )

    return workbook, sheet, bulgarian_data


def get_bulgarian_text(
    bulgarian_data,
    english_text
):

    key = normalize(
        english_text
    )

    if key not in bulgarian_data:

        # test.fail(
        #     "English text not found in Excel: "
        #     + str(english_text)
        # )

        return ""

    return bulgarian_data[key]


def verify_text(
    bulgarian_data,
    object_name,
    english_text
):

    actual_text = ""

    try:

        obj = waitForObjectExists(
            object_name
        )

        actual_text = str(
            obj.text
        ).strip()

    except Exception as e:

        test.fail(
            "Unable to read application object: "
            + str(object_name)
            + " | Error: "
            + str(e)
        )

        return False

    expected_text = get_bulgarian_text(
        bulgarian_data,
        english_text
    )

    actual_normalized = normalize(
        actual_text
    )

    expected_normalized = normalize(
        expected_text
    )

    test.compare(
        actual_normalized,
        expected_normalized,
        "Bulgarian text verification: "
        + str(english_text)
    )

    if actual_normalized == expected_normalized:
        return True

    # test.fail(
    #     "Bulgarian text mismatch for: "
    #     + str(english_text)
    #     + " | Actual: "
    #     + actual_text
    #     + " | Expected: "
    #     + str(expected_text)
    # )

    return False

def main():

    workbook = None

  
   

    clickButton(
        names.numpad_pushButton_8_NumpadButton
    )

    clickButton(
        names.numpad_pushButton_8_NumpadButton
    )

    clickButton(
        names.numpad_pushButton_8_NumpadButton
    )

    clickButton(
        names.numpad_pushButton_8_NumpadButton
    )

    snooze(2)

    clickButton(
        names.mainFrame_okButton_QPushButton
    )

 

    snooze(3) 

    clickButton(waitForObject(names.homeScreen_settingsButton_QPushButton))

    
    snooze(2)

    clickTab(waitForObject(names.settingsScreen_tbSystemInfo_TabWidget), "INFO")
    

    snooze(2)

    clickButton(waitForObject(names.gbSystemInformation_pbLanguageSettingsUpdate_QPushButton))
    clickButton(waitForObject(names.languagesFrame_pbBulgaria_QPushButton))
    clickButton(waitForObject(names.buttonsFrame_acceptButton_QPushButton))

    

    snooze(5)

    clickButton(
        names.settingsScreen_pbBack_QPushButton
    )

    snooze(5)

   
    clickButton(
        names.homeScreen_pbSoftTissueQuickStart_QToolButton
    )

    snooze(3)

    mouseClick(waitForObjectItem(names.tabLeftMainPresets_leftPresetsView_PresetListView, "_1"), 336, 37, Qt.NoModifier, Qt.LeftButton)
    mouseClick(waitForObjectItem(names.tabRightMainPresets_rightPresetsView_PresetListView, "_1"), 241, 78, Qt.NoModifier, Qt.LeftButton)
    
    # mouseClick(waitForObject(names.pulseModeWidget_peakPowerButton_ParameterSettingsButton), 4, 34, Qt.NoModifier, Qt.LeftButton)
    # test.compare(str(waitForObjectExists(names.centralWidget_lbTitle_QLabel).text), "ИЗБОР НА ПИКОВА МОЩНОСТ")
    # test.compare(str(waitForObjectExists(names.centralWidget_pw500_QPushButton).text), "Високо")
    # test.compare(str(waitForObjectExists(names.centralWidget_pw250_QPushButton).text), "Среден")
    # test.compare(str(waitForObjectExists(names.centralWidget_pw125_QPushButton).text), "Ниско")
    # test.compare(str(waitForObjectExists(names.centralWidget_pw125_QPushButton).text), "Ниско")
    # test.compare(str(waitForObjectExists(names.buttonsFrame_rejectButton_QPushButton).text), "ОТКАЗ")

    
    
    snooze(2)


    clickButton(
        names.quickStartScreen_pbContinue_QPushButton
    )

    snooze(3)

    clickButton(
        names.mainFrame_okButton_QPushButton_2
    )

    snooze(5)
    mouseClick(waitForObject(names.pulseModeWidget_peakPowerButton_ParameterSettingsButton), 4, 34, Qt.NoModifier, Qt.LeftButton)
    snooze(6)


    workbook, sheet, bulgarian_data = (
        load_bulgarian_excel(
            EXCEL_PATH 
        )
    )

    if workbook is None:
        return

    snooze(6)

    verify_text(
        bulgarian_data,
        names.centralWidget_lbTitle_QLabel,
        "SELECT PEAK POWER"
    )

    snooze(6)

    verify_text(
        bulgarian_data,
        names.centralWidget_pw500_QPushButton,
        "High"
    )

    snooze(6)

    verify_text(
        bulgarian_data,
        names.centralWidget_pw250_QPushButton,
        "Medium"
    )

    snooze(6)

    verify_text(
        bulgarian_data,
        names.centralWidget_pw125_QPushButton,
        "Low"
    )

    snooze(6)

    verify_text(
        bulgarian_data,
        names.buttonsFrame_rejectButton_QPushButton,
        "CANCEL"
    )

  
    snooze(2)
    
    clickButton(waitForObject(names.buttonsFrame_rejectButton_QPushButton))
    clickButton(waitForObject(names.buttonsBar_homeButton_QPushButton_2))
    clickButton(waitForObject(names.homeScreen_pbSoftTissueAssistant_QToolButton))
    clickButton(waitForObject(names.gbSoftTissueProcedures_pbIncision_QPushButton))
    clickButton(waitForObject(names.softTissueAssistantModeScreen_pbContinue_QPushButton))
    clickButton(waitForObject(names.treatmentCharacteristicsBar_pbEdit_QPushButton))
    # test.compare(str(waitForObjectExists(names.centralWidget_titleLabel_QLabel).text), "SOFT TISSUE CHARACTERISTICS")
    # test.compare(str(waitForObjectExists(names.centralWidget_titleLabel_QLabel_2).text), "ПРОЦЕДУРА")
    # test.compare(str(waitForObjectExists(names.centralWidget_valueLabel_QLabel).text), "Incision")
    # test.compare(str(waitForObjectExists(names.centralWidget_rejectButton_QPushButton).text), "РЕДАКТИРАНЕ")
    # test.compare(str(waitForObjectExists(names.centralWidget_acceptButton_QPushButton).text), "OK")

    # mouseClick(
    #     names.centralWidget_mainWidget_QWidget,
    #     460,
    #     549,
    #     MouseButton.LeftButton
    # )
    #
    # snooze(2)
    #
    # clickButton(
    #     names.buttonsFrame_rejectButton_QPushButton
    # )
    #
    # snooze(5)
    #
    # # ========================================================
    # # RETURN HOME
    # # ========================================================
    #
    # clickButton(
    #     names.buttonsBar_homeButton_QPushButton_2
    # )
    #
    # snooze(5)
    #
    # # ========================================================
    # # SOFT TISSUE ASSISTANT
    # # ========================================================
    #
    # clickButton(
    #     names.homeScreen_pbSoftTissueAssistant_QToolButton
    # )
    #
    # snooze(4)
    #
    # clickButton(
    #     names.gbSoftTissueProcedures_pbIncision_QPushButton
    # )
    #
    # snooze(3)
    #
    # clickButton(
    #     names.softTissueAssistantModeScreen_pbContinue_QPushButton
    # )
    #
    # snooze(4)
    #
    # clickButton(
    #     names.treatmentCharacteristicsBar_pbEdit_QPushButton
    # )

    snooze(6)

   
    verify_text(
        bulgarian_data,
        names.centralWidget_titleLabel_QLabel,
        "SOFT TISSUE CHARACTERISTICS"
    )

    snooze(6)

    verify_text(
        bulgarian_data,
        names.centralWidget_titleLabel_QLabel_2,
        "PROCEDURE"
    )

    snooze(6)

    verify_text(
        bulgarian_data,
        names.centralWidget_valueLabel_QLabel,
        "Incision"
    )

    snooze(6)

    verify_text(
        bulgarian_data,
        names.centralWidget_rejectButton_QPushButton,
        "EDIT"
    )

    snooze(6)

    verify_text(
        bulgarian_data,
        names.centralWidget_acceptButton_QPushButton,
        "OK"
    )

 
    if workbook is not None:
        workbook.close()

