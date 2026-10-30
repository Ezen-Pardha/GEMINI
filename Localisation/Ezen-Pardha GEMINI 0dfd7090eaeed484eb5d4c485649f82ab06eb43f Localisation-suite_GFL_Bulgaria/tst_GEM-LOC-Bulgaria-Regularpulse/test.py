    # -*- coding: utf-8 -*-
    
import names
import re
import unicodedata
from openpyxl import load_workbook
import os

TESTCASE_DIR = os.path.dirname(os.path.abspath(__file__))
 
EXCEL_PATH = os.path.abspath(
    os.path.join( TESTCASE_DIR,
        "..",
        "shared",
        "testdata",
        "Gemini Split strings.xlsx"
    )
)





ENGLISH_COLUMN = 2


workbook = None
sheet = None
BULGARIAN_COLUMN = None


    
verification_objects = {
    
    
    "main": [
    
        (
            names.centralWidget_lbTitle_QLabel,
            "Mode Selection Title"
        ),
    
        (
            names.centralWidget_label_2_QLabel,
            "Standard Mode Label"
        ),
    
        (
            names.centralWidget_Square_QPushButton,
            "Standard Pulse Button"
        ),
    
        (
            names.centralWidget_CW_QPushButton,
            "Continuous Wave Button"
        ),
    
        (
            names.centralWidget_MRP_QPushButton,
            "MRP Button"
        ),
    
        (
            names.centralWidget_label_QLabel,
            "Expert Mode Label"
        ),
    
        (
            names.centralWidget_Fragmentation_QPushButton,
            "Fragmentation Pulse Button"
        ),
    
        (
            names.centralWidget_Cleaning_QPushButton,
            "Cleaning Pulse Button"
        ),
    
        (
            names.centralWidget_Enucleation_QPushButton,
            "Enucleation Pulse Button"
        ),
    
        (
            names.buttonsFrame_rejectButton_QPushButton,
            "Cancel Button"
        )
    ],
    
    
    "standard_pulse": [
    
        (
            names.centralWidget_TitleLabel_QLabel_3,
            "Standard Pulse Popup Title"
        ),
    
        (
            names.centralWidget_MessageLabel_QLabel_3,
            "Standard Pulse Popup Message"
        ),
    
        (
            names.centralWidget_OKButton_QPushButton,
            "Standard Pulse Popup OK Button"
        )
    ],
    
     
    "cw": [
    
        (
            names.centralWidget_TitleLabel_QLabel_3,
            "Continuous Wave Popup Title"
        ),
    
        (
            names.centralWidget_MessageLabel_QLabel_3,
            "Continuous Wave Popup Message"
        ),
    
        (
            names.centralWidget_OKButton_QPushButton,
            "Continuous Wave Popup OK Button"
        )
    ],
    
    
    "mrp": [
    
        (
            names.centralWidget_TitleLabel_QLabel_3,
            "MRP Popup Title"
        ),
    
        (
            names.centralWidget_MessageLabel_QLabel_3,
            "MRP Popup Message"
        ),
    
        (
            names.centralWidget_OKButton_QPushButton,
            "MRP Popup OK Button"
        )
    ],
    
    
    "fragmentation": [
    
        (
            names.centralWidget_TitleLabel_QLabel_3,
            "Fragmentation Popup Title"
        ),
    
        (
            names.centralWidget_MessageLabel_QLabel_3,
            "Fragmentation Popup Message"
        ),
    
        (
            names.centralWidget_OKButton_QPushButton,
            "Fragmentation Popup OK Button"
        )
    ],
    
    
    "cleaning": [
    
        (
            names.centralWidget_TitleLabel_QLabel_3,
            "Cleaning Popup Title"
        ),
    
        (
            names.centralWidget_MessageLabel_QLabel_3,
            "Cleaning Popup Message"
        ),
    
        (
            names.centralWidget_OKButton_QPushButton,
            "Cleaning Popup OK Button"
        )
    ],
    
    
    
    "enucleation": [
    
        (
            names.centralWidget_TitleLabel_QLabel_3,
            "Enucleation Popup Title"
        ),
    
        (
            names.centralWidget_MessageLabel_QLabel_3,
            "Enucleation Popup Message"
        ),
    
        (
            names.centralWidget_OKButton_QPushButton,
            "Enucleation Popup OK Button"
        )
    ]
    
    
    }
    
def get_object_text(obj):
    
    
    try:
        value = obj.text
    
        if value is not None:
            value = str(value)
    
            if value.strip():
                return value
    
    except:
        pass
    
    
    try:
        value = obj.title
    
        if value is not None:
            value = str(value)
    
            if value.strip():
                return value
    
    except:
        pass
    
    
    try:
        value = obj.windowTitle
    
        if value is not None:
            value = str(value)
    
            if value.strip():
                return value
    
    except:
        pass
    
    
    try:
        value = obj.accessibleName
    
        if value is not None:
            value = str(value)
    
            if value.strip():
                return value
    
    except:
        pass
    
    
    return ""
    
  
    
def normalize_text(text):
    
    
    if text is None:
        return ""
    
    text = str(text)
    
    
    # Remove HTML/XML tags
    
    text = re.sub(
        r"<[^>]*>",
        " ",
        text
    )
    
    
    # Normalize special spaces
    
    text = text.replace(
        "\u00A0",
        " "
    )
    
    text = text.replace(
        "\u2007",
        " "
    )
    
    text = text.replace(
        "\u202F",
        " "
    )
    
    
    # Normalize line breaks and tabs
    
    text = re.sub(
        r"[\r\n\t]+",
        " ",
        text
    )
    
    
    # Remove trademark/copyright symbols
    
    text = text.replace(
        "™",
        ""
    )
    
    text = text.replace(
        "®",
        ""
    )
    
    text = text.replace(
        "©",
        ""
    )
    
    
    # Normalize Unicode
    
    try:
    
        text = unicodedata.normalize(
            "NFKC",
            text
        )
    
    except:
        pass
    
    
    # Normalize spaces
    
    text = re.sub(
        r"\s+",
        " ",
        text
    )
    
    
    # Remove spaces before punctuation
    
    text = re.sub(
        r"\s+([.,;:!?])",
        r"\1",
        text
    )
    
    
    return text.strip().casefold()
    
    
  
    
def find_bulgarian_column():
    
    
    global sheet
    
    if sheet is None:
        return None
    
    
    possible_headers = [
    
        "bulgarian",
        "bulgaria",
        "български",
        "български език",
        "българия"
    
    ]
    
    
    # Search first 10 rows for Bulgarian header
    
    max_rows = min(
        sheet.max_row,
        10
    )
    
    max_columns = sheet.max_column
    
    
    for row_number in range(
        1,
        max_rows + 1
    ):
    
        for column_number in range(
            1,
            max_columns + 1
        ):
    
            try:
    
                value = sheet.cell(
                    row=row_number,
                    column=column_number
                ).value
    
            except:
    
                continue
    
    
            if value is None:
                continue
    
    
            normalized_value = normalize_text(
                value
            )
    
    
            for header in possible_headers:
    
                normalized_header = normalize_text(
                    header
                )
    
                if normalized_header == normalized_value:
    
                    return column_number
    
    
    return None

    
  
def find_bulgarian_text(actual_text):
    
  
    global sheet
    global BULGARIAN_COLUMN
    
    
    if sheet is None:
        return None
    
    
    if BULGARIAN_COLUMN is None:
        return None
    
    
    actual_normalized = normalize_text(
        actual_text
    )
    
    
    if not actual_normalized:
        return None
    
    
    # Search complete Bulgarian column
    
    for row in sheet.iter_rows(
        min_row=1,
        min_col=BULGARIAN_COLUMN,
        max_col=BULGARIAN_COLUMN
    ):
    
        bulgarian_value = row[0].value
    
    
        if bulgarian_value is None:
            continue
    
    
        excel_normalized = normalize_text(
            bulgarian_value
        )
    
    
        if not excel_normalized:
            continue
    
    
        if actual_normalized == excel_normalized:
    
            return bulgarian_value
    
    
    return None
  
 
def verify_text(object_name, description):
    
  
    # Wait before every verification
    
    snooze(6)
    
    
    # Get object
    
    try:
    
        obj = waitForObject(
            object_name
        )
    
    except Exception as e:
    
        test.fail(
            description
            + " | Object not found: "
            + str(e)
        )
    
        return False
    
    
    # Get actual UI text
    
    actual_text = get_object_text(
        obj
    )
    
    
    if not actual_text:
    
        test.fail(
            description
            + " | UI text is empty"
        )
    
        return False
    
    
    # Search actual Bulgarian text in Excel
    
    matched_bulgarian = find_bulgarian_text(
        actual_text
    )
    
    
    # Match found
    
    if matched_bulgarian is not None:
    
        actual_normalized = normalize_text(
            actual_text
        )
    
        expected_normalized = normalize_text(
            matched_bulgarian
        )
    
    
        test.compare(
            actual_normalized,
            expected_normalized,
            description
        )
    
        return True
    
    
    # Match not found
    
    test.fail(
        description
        + " | Bulgarian text not found in Excel: "
        + str(actual_text)
    )
    
    return False
    
    
    
def verify_group(group_name):
    
    
    if group_name not in verification_objects:
    
        test.fail(
            "Verification group not found: "
            + str(group_name)
        )
    
        return False
    
    
    for object_name, description in verification_objects[group_name]:
    
        verify_text(
            object_name,
            description
        )
    
    
    return True
    
    
    
def verify_popup(group_name):
    
    
    verify_group(
        group_name
    )
    
    
    snooze(6)
    
    
    clickButton(
        waitForObject(
            names.centralWidget_OKButton_QPushButton
        )
    )
    
   
    
def main():
    
    
    global workbook
    global sheet
    global BULGARIAN_COLUMN
    
    
    # ========================================================
    # LOAD EXCEL
    # ========================================================
    
    try:
    
        workbook = load_workbook(
            EXCEL_PATH,
            read_only=True,
            data_only=True
        )
    
    
        # Support both possible sheet-name capitalization
    
        if "Treatment Modes" in workbook.sheetnames:
    
            sheet = workbook[
                "Treatment Modes"
            ]
    
        elif "Treatment modes" in workbook.sheetnames:
    
            sheet = workbook[
                "Treatment modes"
            ]
    
        else:
    
            test.fail(
                "Treatment Modes sheet not found in Gemini Split strings.xlsx"
            )
    
            return
    
    
    except Exception as e:
    
        test.fail(
            "Failed to load Gemini Split strings.xlsx: "
            + str(e)
        )
    
        return
    
    
    # ========================================================
    # FIND BULGARIAN COLUMN
    # ========================================================
    
    BULGARIAN_COLUMN = find_bulgarian_column()
    
    
    if BULGARIAN_COLUMN is None:
    
        test.fail(
            "Bulgarian column was not found in Gemini Split strings.xlsx"
        )
    
    
        try:
    
            workbook.close()
    
        except:
    
            pass
    
    
        return
    
    
  
    
    doubleClick(
        waitForObject(
            names.numpad_pushButton_8_NumpadButton
        ),
        91,
        77,
        Qt.NoModifier,
        Qt.LeftButton
    )
    
    
    snooze(6)
    
    
    clickButton(
        waitForObject(
            names.numpad_pushButton_8_NumpadButton
        )
    )
    
    
    snooze(6)
    
    
    clickButton(
        waitForObject(
            names.numpad_pushButton_8_NumpadButton
        )
    )
    
    
    snooze(6)
    
    
    clickButton(
        waitForObject(
            names.mainFrame_okButton_QPushButton
        )
    )
    
    
    
    snooze(6)
    
    
    clickButton(
        waitForObject(
            names.homeScreen_settingsButton_QPushButton
        )
    )
    
    
    snooze(6)
    
    
    clickTab(
        waitForObject(
            names.settingsScreen_tbSystemInfo_TabWidget
        ),
        "INFO"
    )
    
    
    
    snooze(6)
    
    
    clickButton(
        waitForObject(
            names.gbSystemInformation_pbLanguageSettingsUpdate_QPushButton
        )
    )
    
    
    snooze(6)
    
    
    clickButton(
        waitForObject(
            names.languagesFrame_pbBulgaria_QPushButton
        )
    )
    
    
    snooze(6)
    
    
    clickButton(
        waitForObject(
            names.buttonsFrame_acceptButton_QPushButton
        )
    )
    
    
    snooze(6)
    
    
    clickButton(
        waitForObject(
            names.buttonsBar_homeButton_QPushButton
        )
    )
    
    
    snooze(6)
    
    
    clickButton(
        waitForObject(
            names.gbExpert_expertButton_QPushButton
        )
    )
    
    
    snooze(6)
    
    
    mouseClick(
        waitForObject(
            names.modeButton_nameLabel_QLabel
        ),
        73,
        25,
        Qt.NoModifier,
        Qt.LeftButton
    )
    
    
    verify_group(
        "main"
    )
    
    
    snooze(6)
    
    
    clickButton(
        waitForObject(
            names.square_icon_QPushButton
        )
    )
    
    
    verify_popup(
        "standard_pulse"
    )
    
    
    snooze(6)
    
    
    clickButton(
        waitForObject(
            names.cW_icon_QPushButton
        )
    )
    
    
    verify_popup(
        "cw"
    )
    
    
    snooze(6)
    
    
    clickButton(
        waitForObject(
            names.mRP_icon_QPushButton
        )
    )
    
    
    verify_popup(
        "mrp"
    )
    
    
    snooze(6)
    
    
    clickButton(
        waitForObject(
            names.fragmentation_icon_QPushButton
        )
    )
    
    
    verify_popup(
        "fragmentation"
    )
    
    
    snooze(6)
    
    
    clickButton(
        waitForObject(
            names.cleaning_icon_QPushButton
        )
    )
    
    
    verify_popup(
        "cleaning"
    )
    
    
    snooze(6)
    
    
    clickButton(
        waitForObject(
            names.enucleation_icon_QPushButton
        )
    )
    
    
    verify_popup(
        "enucleation"
    )
    
    
    snooze(6)
    
    
    clickButton(
        waitForObject(
            names.buttonsFrame_rejectButton_QPushButton
        )
    )
    
    
    try:
    
        workbook.close()
    
    except:
    
        pass
    
