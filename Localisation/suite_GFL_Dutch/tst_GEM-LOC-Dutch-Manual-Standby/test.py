# -*- coding: utf-8 -*-

import names
from openpyxl import load_workbook
import os

TESTCASE_DIR = os.path.dirname(os.path.abspath(__file__))
 
EXCEL_PATH = os.path.abspath(
    os.path.join( TESTCASE_DIR,
        "..",
        "shared",
        "testdata",
        "Gemini Split strings.xlsx"
    )
)


EXCEL_SHEET_NAME = "Treatment Modes"

# Column B = English

ENGLISH_COLUMN = 2

# Column G = Dutch

DUTCH_COLUMN = 7


def normalize_text(value):


    if value is None:
        return ""
    
    return " ".join(
        str(value)
        .replace("\n", " ")
        .replace("\r", " ")
        .split()
    ).strip().casefold()
    
    

def get_dutch_text(sheet, english_text):


    search_text = normalize_text(
        english_text
    )
    
    for row in sheet.iter_rows(
        min_row=1,
        min_col=ENGLISH_COLUMN,
        max_col=DUTCH_COLUMN,
        values_only=True
    ):
    
        english_value = row[0]
    
        if normalize_text(
            english_value
        ) == search_text:
    
            dutch_value = row[
                DUTCH_COLUMN - ENGLISH_COLUMN
            ]
    
            # Do not allow empty Excel translations
            if dutch_value is None:
    
                test.fail(
                    "Dutch translation is empty in Excel "
                    "for English text: "
                    + english_text
                )
    
                return ""
    
            dutch_text = str(
                dutch_value
            ).strip()
    
            if dutch_text == "":
    
                test.fail(
                    "Dutch translation is blank in Excel "
                    "for English text: "
                    + english_text
                )
    
                return ""
    
            return dutch_text
    
    test.fail(
        "English text not found in Excel Column B: "
        + english_text
    )
    
    return ""



def get_actual_text(obj, description):

    actual_value = None


    try:
        actual_value = obj.text
    except Exception:
        actual_value = None



    if actual_value is None:
    
        try:
            actual_value = obj.property(
                "text"
            )
        except Exception:
            actual_value = None



    if actual_value is None:
    
        try:
            actual_value = obj.windowTitle
        except Exception:
            actual_value = None
    
    
    
    if actual_value is None:
    
        try:
            actual_value = obj.accessibleName
        except Exception:
            actual_value = None
    
    
    
    actual_text = normalize_text(
        actual_value
    )
    
    
    
    if actual_text == "":
    
        test.fail(
            "Application text is empty or None for: "
            + description
        )
    
        return None
    
    
    return actual_text
    
    
    
def verify_text(
        sheet,
        object_name,
        english_text,
        description
        ):
    
    
    # Wait before every verification
    snooze(6)
    
    
    
    expected_text = get_dutch_text(
        sheet,
        english_text
    )
    
    
    expected_normalized = normalize_text(
        expected_text
    )
    
    
    
    if expected_normalized == "":
    
        test.fail(
            "Expected Dutch text is empty for: "
            + description
        )
    
        return
    
    
    
    obj = waitForObjectExists(
        getattr(
            names,
            object_name
        )
    )
    
    
    
    actual_text = get_actual_text(
        obj,
        description
    )
    
    
    
    if actual_text is None:
    
        return
    
    
    
    test.compare(
        actual_text,
        expected_normalized,
        description + " - Dutch text verification"
    )
    
    
    
def verify_left_manual_tab(
        sheet,
        english_text,
        description
        ):
    
    
    snooze(6)
    
    
    
    expected_text = get_dutch_text(
        sheet,
        english_text
    )
    
    expected_normalized = normalize_text(
        expected_text
    )
    
    
    if expected_normalized == "":
    
        test.fail(
            "Expected Dutch text is empty for: "
            + description
        )
    
        return
    
    
    
    manual_tab_object = getattr(
        names,
        "mANUEL_TabItem"
    ).copy()
    
    
    
    if "text" in manual_tab_object:
    
        del manual_tab_object["text"]
    
    
    
    obj = waitForObjectExists(
        manual_tab_object
    )
    
    
    
    actual_text = get_actual_text(
        obj,
        description
    )
    
    
    if actual_text is None:
    
        return
    
    
    
    test.compare(
        actual_text,
        expected_normalized,
        description + " - Dutch text verification"
    )
    
    
    
    
def verify_right_manual_tab(
        sheet,
        english_text,
        description
        ):
    
    
    snooze(6)
    
    
    
    expected_text = get_dutch_text(
        sheet,
        english_text
    )
    
    expected_normalized = normalize_text(
        expected_text
    )
    
    
    if expected_normalized == "":
    
        test.fail(
            "Expected Dutch text is empty for: "
            + description
        )
    
        return
    
    
    
    manual_tab_object = getattr(
        names,
        "mANUEL_TabItem_2"
    ).copy()
    
    
    
    if "text" in manual_tab_object:
    
        del manual_tab_object["text"]
    
    
    
    obj = waitForObjectExists(
        manual_tab_object
    )
    
    
    
    actual_text = get_actual_text(
        obj,
        description
    )
    
    
    if actual_text is None:
    
        return
    
    
    
    test.compare(
        actual_text,
        expected_normalized,
        description + " - Dutch text verification"
    )
    
    
    
def main():
    
    
    workbook = load_workbook(
        EXCEL_PATH,
        read_only=True,
        data_only=True
    )
    
    
    
    if EXCEL_SHEET_NAME not in workbook.sheetnames:
    
        test.fail(
            "Excel sheet not found: "
            + EXCEL_SHEET_NAME
        )
    
        workbook.close()
    
        return
    
    
    sheet = workbook[
        EXCEL_SHEET_NAME
    ]
    
    
    
    doubleClick(
        waitForObject(
            names.numpad_pushButton_8_NumpadButton,
            100461
        ),
        96,
        56,
        Qt.NoModifier,
        Qt.LeftButton
    )
    
    snooze(1)
    
    
    doubleClick(
        waitForObject(
            names.numpad_pushButton_8_NumpadButton
        ),
        96,
        56,
        Qt.NoModifier,
        Qt.LeftButton
    )
    
    snooze(1)
    
    
    clickButton(
        waitForObject(
            names.mainFrame_okButton_QPushButton
        )
    )
    
    snooze(2)
    
    
    
    clickButton(
        waitForObject(
            names.homeScreen_settingsButton_QPushButton
        )
    )
    
    snooze(2)
    
    
    
    clickTab(
        waitForObject(
            names.settingsScreen_tbSystemInfo_TabWidget
        ),
        "INFO"
    )
    
    snooze(2)
    
    
    
    clickButton(
        waitForObject(
            names.gbSystemInformation_pbLanguageSettingsUpdate_QPushButton
        )
    )
    
    snooze(2)
    
    
    
    clickButton(
        waitForObject(
            names.languagesFrame_pbDutch_QPushButton
        )
    )
    
    snooze(1)
    
    
    
    clickButton(
        waitForObject(
            names.buttonsFrame_acceptButton_QPushButton
        )
    )
    
    snooze(3)
    
    
    
    clickButton(
        waitForObject(
            names.buttonsBar_homeButton_QPushButton
        )
    )
    
    snooze(2)
    
    
    
    clickButton(
        waitForObject(
            names.gbExpert_expertButton_QPushButton
        )
    )
    
    snooze(3)
    
    test.vp("Manul Standby screen")

    verify_left_manual_tab(
        sheet,
        "MANUAL",
        "Left Manual Tab"
    )
    
    
    
    verify_text(
        sheet,
        "averagePowerWidget_nameLabel_QLabel_2",
        "AVERAGE POWER",
        "Left Average Power"
    )
    
    
    
    verify_text(
        sheet,
        "averagePowerWidget_unitsOfMeasureLabel_QLabel_2",
        "W",
        "Left Power Unit"
    )
    
    
    
    verify_text(
        sheet,
        "pulseEnergyWidget_nameLabel_QLabel_2",
        "PULSE ENERGY",
        "Left Pulse Energy"
    )
    
    
    
    verify_text(
        sheet,
        "pulseEnergyWidget_unitsOfMeasureLabel_QLabel_2",
        "J",
        "Left Energy Unit"
    )
    
    
    
    verify_text(
        sheet,
        "frequencyWidget_nameLabel_QLabel_2",
        "FREQUENCY",
        "Left Frequency"
    )
    
    
    
    verify_text(
        sheet,
        "frequencyWidget_unitsOfMeasureLabel_QLabel_2",
        "Hz",
        "Left Frequency Unit"
    )
    
    
    
    verify_text(
        sheet,
        "modeButton_nameLabel_QLabel",
        "REGULAR PULSE",
        "Left Pulse Mode"
    )
    
    
    
    verify_text(
        sheet,
        "peakPowerButton_nameLabel_QLabel",
        "PEAK POWER",
        "Left Peak Power"
    )
    
    
    
    verify_text(
        sheet,
        "pulseModeWidget_saveButton_QPushButton_2",
        "SAVE",
        "Left Save Button"
    )
    
    
    
    verify_text(
        sheet,
        "totalEnergyFrame_nameLabel_QLabel",
        "TOTAL ENERGY",
        "Left Total Energy"
    )
    
    
    
    verify_text(
        sheet,
        "totalTimeFrame_nameLabel_QLabel",
        "TOTAL LASING TIME",
        "Left Total Lasing Time"
    )
    
    
    
    verify_right_manual_tab(
        sheet,
        "MANUAL",
        "Right Manual Tab"
    )
    
    
    
    verify_text(
        sheet,
        "averagePowerWidget_nameLabel_QLabel_3",
        "AVERAGE POWER",
        "Right Average Power"
    )
    
    
    
    verify_text(
        sheet,
        "averagePowerWidget_unitsOfMeasureLabel_QLabel_3",
        "W",
        "Right Power Unit"
    )
    
    
    
    verify_text(
        sheet,
        "pulseEnergyWidget_nameLabel_QLabel_3",
        "PULSE ENERGY",
        "Right Pulse Energy"
    )
    
    
    
    verify_text(
        sheet,
        "pulseEnergyWidget_unitsOfMeasureLabel_QLabel_3",
        "J",
        "Right Energy Unit"
    )
    
    
    
    verify_text(
        sheet,
        "frequencyWidget_nameLabel_QLabel_3",
        "FREQUENCY",
        "Right Frequency"
    )
    
    
    
    verify_text(
        sheet,
        "pulseModeWidget_saveButton_QPushButton_3",
        "SAVE",
        "Right Save Button"
    )
    
    
    
    verify_text(
        sheet,
        "peakPowerButton_nameLabel_QLabel_3",
        "PEAK POWER",
        "Right Peak Power"
    )
    
    
    
    verify_text(
        sheet,
        "modeButton_nameLabel_QLabel_2",
        "REGULAR PULSE",
        "Right Pulse Mode"
    )
    
    
    
    verify_text(
        sheet,
        "stateSwitch_readyButton_QPushButton",
        "READY",
        "Ready Button"
    )
    
    
    
    verify_text(
        sheet,
        "stateSwitch_standbyButton_QPushButton",
        "STANDBY",
        "Standby Button"
    )
    
    
    
    verify_text(
        sheet,
        "fiberFrame_nameLabel_QLabel",
        "FIBER",
        "Fiber"
    )
    
    
    
    verify_text(
        sheet,
        "fiberUsesFrame_nameLabel_QLabel",
        "FIBER USES REMAINING",
        "Fiber Uses Remaining"
    )
    
    
    
    verify_text(
        sheet,
        "bottomBar_aimingBeamButton_AimingBeamModeButton",
        "AIMING BEAM",
        "Aiming Beam"
    )
    
    workbook.close()
    
    
