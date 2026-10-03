
# -*- coding: utf-8 -*-

import names
from openpyxl import load_workbook


def main():

    # ============================================================
    # EXCEL FILE
    # ============================================================

    excelpath = (
        "/home/ezen/Test_Automation_Gemini/Localisation/"
        "suite_GFL_Dutch/"
        "tst_GEM-LOC-Dutch-PeakPower-SoftTissueCharacters/"
        "testdata/Gemini Split strings.xlsx"
    )

    test.log("Loading Excel file...")

    workbook = load_workbook(
        excelpath,
        read_only=True,
        data_only=True
    )

    if "Treatment Modes" not in workbook.sheetnames:
        test.fail(
            "Worksheet 'Treatment Modes' does not exist in Excel file"
        )
        workbook.close()
        return

    sheet = workbook["Treatment Modes"]

    test.log(
        "Excel sheet loaded: Treatment Modes"
    )


    # ============================================================
    # LOGIN
    # ============================================================

    test.log("Entering login PIN")

    clickButton(
        waitForObject(
            names.numpad_pushButton_8_NumpadButton
        )
    )

    clickButton(
        waitForObject(
            names.numpad_pushButton_8_NumpadButton
        )
    )

    clickButton(
        waitForObject(
            names.numpad_pushButton_8_NumpadButton
        )
    )

    clickButton(
        waitForObject(
            names.numpad_pushButton_8_NumpadButton
        )
    )

    snooze(6)

    clickButton(
        waitForObject(
            names.mainFrame_okButton_QPushButton
        )
    )


    # ============================================================
    # OPEN SETTINGS
    # ============================================================

    test.log("Opening Settings")

    clickButton(
        waitForObject(
            names.homeScreen_settingsButton_QPushButton
        )
    )

    clickTab(
        waitForObject(
            names.settingsScreen_tbSystemInfo_TabWidget
        ),
        "INFO"
    )

    clickButton(
        waitForObject(
            names.gbSystemInformation_pbLanguageSettingsUpdate_QPushButton
        )
    )


    # ============================================================
    # CHANGE LANGUAGE TO DUTCH
    # ============================================================

    test.log("Selecting Dutch language")

    clickButton(
        waitForObject(
            names.languagesFrame_pbDutch_QPushButton
        )
    )

    snooze(6)

    clickButton(
        waitForObject(
            names.buttonsFrame_acceptButton_QPushButton
        )
    )

    snooze(6)


    # ============================================================
    # OPEN USER PRESETS
    # ============================================================

    test.log("Opening User Presets")

    # ------------------------------------------------------------
    # NOTE:
    # The tab text below must match the Dutch text displayed by
    # your application.
    #
    # If Squish identifies the tab by object instead of text,
    # this can be changed to clickTab without the text.
    # ------------------------------------------------------------

    clickTab(
        waitForObject(
            names.settingsScreen_tbSystemInfo_TabWidget
        ),
        "INFO"
    )

    # ------------------------------------------------------------
    # If the Dutch application displays a translated tab name,
    # replace the following text with the exact Dutch tab name
    # from the application.
    # ------------------------------------------------------------

    clickTab(
        waitForObject(
            names.settingsScreen_tbSystemInfo_TabWidget
        ),
        "WIRELESS FOOTSWITCH"
    )

    clickTab(
        waitForObject(
            names.settingsScreen_tbSystemInfo_TabWidget
        ),
        "USER PRESETS"
    )

    sendEvent(
        "QMoveEvent",
        waitForObject(
            names.o_ScreenSwitcher
        ),
        0,
        -66,
        1285,
        233
    )


    # ============================================================
    # ADD PRESET
    # ============================================================

    test.log("Clicking Add Preset")

    clickButton(
        waitForObject(
            names.settingsScreen_addButton_QPushButton
        )
    )

    snooze(6)


    # ============================================================
    # 1. ADD PRESET OPERATION
    # ============================================================

    snooze(6)

    actualText1 = str(
        waitForObjectExists(
            names.centralWidget_lbOperation_QLabel
        ).text
    )

    expectedText1 = str(
        sheet.cell(
            row=67,
            column=7
        ).value
    )

    test.compare(
        actualText1,
        expectedText1,
        "Dutch - Add Preset Operation"
    )


    # ============================================================
    # 2. PRESET NAME
    # ============================================================

    snooze(6)

    actualText2 = str(
        waitForObjectExists(
            names.centralWidget_pbName_QPushButton
        ).text
    )

    expectedText2 = str(
        sheet.cell(
            row=50,
            column=7
        ).value
    )

    test.compare(
        actualText2,
        expectedText2,
        "Dutch - Preset Name"
    )


    # ============================================================
    # 3. AVERAGE POWER
    # ============================================================

    snooze(6)

    actualText3 = str(
        waitForObjectExists(
            names.averagePowerWidget_nameLabel_QLabel
        ).text
    )

    expectedText3 = str(
        sheet.cell(
            row=6,
            column=7
        ).value
    )

    test.compare(
        actualText3,
        expectedText3,
        "Dutch - Average Power"
    )


    # ============================================================
    # 4. AVERAGE POWER UNIT
    # ============================================================

    snooze(6)

    actualText4 = str(
        waitForObjectExists(
            names.averagePowerWidget_unitsOfMeasureLabel_QLabel
        ).text
    )

    expectedText4 = str(
        sheet.cell(
            row=7,
            column=7
        ).value
    )

    test.compare(
        actualText4,
        expectedText4,
        "Dutch - Average Power Unit"
    )


    # ============================================================
    # 5. PULSE ENERGY
    # ============================================================

    snooze(6)

    actualText5 = str(
        waitForObjectExists(
            names.pulseEnergyWidget_nameLabel_QLabel
        ).text
    )

    expectedText5 = str(
        sheet.cell(
            row=8,
            column=7
        ).value
    )

    test.compare(
        actualText5,
        expectedText5,
        "Dutch - Pulse Energy"
    )


    # ============================================================
    # 6. PULSE ENERGY UNIT
    # ============================================================

    mouseClick(
        waitForObject(
            names.centralWidget_frame_QFrame
        ),
        271,
        13,
        Qt.NoModifier,
        Qt.LeftButton
    )

    snooze(6)

    actualText6 = str(
        waitForObjectExists(
            names.pulseEnergyWidget_unitsOfMeasureLabel_QLabel
        ).text
    )

    expectedText6 = str(
        sheet.cell(
            row=9,
            column=7
        ).value
    )

    test.compare(
        actualText6,
        expectedText6,
        "Dutch - Pulse Energy Unit"
    )


    # ============================================================
    # 7. FREQUENCY
    # ============================================================

    snooze(6)

    actualText7 = str(
        waitForObjectExists(
            names.frequencyWidget_nameLabel_QLabel
        ).text
    )

    expectedText7 = str(
        sheet.cell(
            row=10,
            column=7
        ).value
    )

    test.compare(
        actualText7,
        expectedText7,
        "Dutch - Frequency"
    )


    # ============================================================
    # 8. FREQUENCY UNIT
    # ============================================================

    snooze(6)

    actualText8 = str(
        waitForObjectExists(
            names.frequencyWidget_unitsOfMeasureLabel_QLabel
        ).text
    )

    expectedText8 = str(
        sheet.cell(
            row=11,
            column=7
        ).value
    )

    test.compare(
        actualText8,
        expectedText8,
        "Dutch - Frequency Unit"
    )


    # ============================================================
    # 9. REGULAR PULSE
    # ============================================================

    snooze(6)

    actualText9 = str(
        waitForObjectExists(
            names.modeButton_nameLabel_QLabel_3
        ).text
    )

    expectedText9 = str(
        sheet.cell(
            row=33,
            column=7
        ).value
    )

    test.compare(
        actualText9,
        expectedText9,
        "Dutch - Regular Pulse"
    )


    # ============================================================
    # 10. PEAK POWER
    # ============================================================

    snooze(6)

    actualText10 = str(
        waitForObjectExists(
            names.peakPowerButton_nameLabel_QLabel_2
        ).text
    )

    expectedText10 = str(
        sheet.cell(
            row=12,
            column=7
        ).value
    )

    test.compare(
        actualText10,
        expectedText10,
        "Dutch - Peak Power"
    )


    # ============================================================
    # 11. ADD BUTTON
    # ============================================================

    snooze(6)

    actualText11 = str(
        waitForObjectExists(
            names.pulseModeWidget_saveButton_QPushButton
        ).text
    )

    expectedText11 = str(
        sheet.cell(
            row=68,
            column=7
        ).value
    )

    test.compare(
        actualText11,
        expectedText11,
        "Dutch - Add Button"
    )


    # ============================================================
    # ENTER PRESET NAME
    # ============================================================

    mouseClick(
        waitForObject(
            names.pulseModeWidget_buttonsWidget_QWidget
        ),
        508,
        46,
        Qt.NoModifier,
        Qt.LeftButton
    )

    mouseClick(
        waitForObject(
            names.pulseModeWidget_buttonsWidget_QWidget
        ),
        411,
        1,
        Qt.NoModifier,
        Qt.LeftButton
    )

    clickButton(
        waitForObject(
            names.centralWidget_pbName_QPushButton
        )
    )

    type(
        waitForObject(
            names.keyboardPopup_wdEdit_WdLineEdit
        ),
        "preset one"
    )

    mouseClick(
        waitForObject(
            names.keyboardPopup_lbText_QLabel
        ),
        77,
        26,
        Qt.NoModifier,
        Qt.LeftButton
    )

    clickButton(
        waitForObject(
            names.pulseModeWidget_saveButton_QPushButton
        )
    )

    snooze(6)


    # ============================================================
    # SELECT PRESET
    # ============================================================

    mouseClick(
        waitForObject(
            names.tabSettingsPresetsUserStonePresets_presetListView_PresetListView
        ),
        546,
        72,
        Qt.NoModifier,
        Qt.LeftButton
    )

    snooze(6)


    # ============================================================
    # EDIT PRESET
    # ============================================================

    clickButton(
        waitForObject(
            names.settingsScreen_editButton_QPushButton
        )
    )

    mouseClick(
        waitForObject(
            names.tabSettingsPresetsUserStonePresets_presetListView_PresetListView
        ),
        494,
        66,
        Qt.NoModifier,
        Qt.LeftButton
    )

    snooze(6)


    # ============================================================
    # EDIT OPERATION
    # ============================================================

    snooze(6)

    actualEditOperation = str(
        waitForObjectExists(
            names.centralWidget_lbOperation_QLabel
        ).text
    )

    expectedEditOperation = str(
        sheet.cell(
            row=73,
            column=7
        ).value
    )

    test.compare(
        actualEditOperation,
        expectedEditOperation,
        "Dutch - Edit Preset Operation"
    )


    # ============================================================
    # EDIT PRESET NAME
    # ============================================================

    snooze(6)

    actualEditName = str(
        waitForObjectExists(
            names.centralWidget_pbName_QPushButton
        ).text
    )

    test.compare(
        actualEditName,
        "preset one",
        "Edit Preset Name"
    )


    # ============================================================
    # EDIT - AVERAGE POWER
    # ============================================================

    snooze(6)

    actualEditAveragePower = str(
        waitForObjectExists(
            names.averagePowerWidget_nameLabel_QLabel
        ).text
    )

    expectedEditAveragePower = str(
        sheet.cell(
            row=6,
            column=7
        ).value
    )

    test.compare(
        actualEditAveragePower,
        expectedEditAveragePower,
        "Dutch - Edit Average Power"
    )


    # ============================================================
    # EDIT - AVERAGE POWER UNIT
    # ============================================================

    snooze(6)

    actualEditAveragePowerUnit = str(
        waitForObjectExists(
            names.averagePowerWidget_unitsOfMeasureLabel_QLabel
        ).text
    )

    expectedEditAveragePowerUnit = str(
        sheet.cell(
            row=7,
            column=7
        ).value
    )

    test.compare(
        actualEditAveragePowerUnit,
        expectedEditAveragePowerUnit,
        "Dutch - Edit Average Power Unit"
    )


    # ============================================================
    # EDIT - PULSE ENERGY
    # ============================================================

    snooze(6)

    actualEditPulseEnergy = str(
        waitForObjectExists(
            names.pulseEnergyWidget_nameLabel_QLabel
        ).text
    )

    expectedEditPulseEnergy = str(
        sheet.cell(
            row=8,
            column=7
        ).value
    )

    test.compare(
        actualEditPulseEnergy,
        expectedEditPulseEnergy,
        "Dutch - Edit Pulse Energy"
    )


    # ============================================================
    # EDIT - PULSE ENERGY UNIT
    # ============================================================

    snooze(6)

    actualEditPulseEnergyUnit = str(
        waitForObjectExists(
            names.pulseEnergyWidget_unitsOfMeasureLabel_QLabel
        ).text
    )

    expectedEditPulseEnergyUnit = str(
        sheet.cell(
            row=9,
            column=7
        ).value
    )

    test.compare(
        actualEditPulseEnergyUnit,
        expectedEditPulseEnergyUnit,
        "Dutch - Edit Pulse Energy Unit"
    )


    # ============================================================
    # EDIT - FREQUENCY
    # ============================================================

    snooze(6)

    actualEditFrequency = str(
        waitForObjectExists(
            names.frequencyWidget_nameLabel_QLabel
        ).text
    )

    expectedEditFrequency = str(
        sheet.cell(
            row=10,
            column=7
        ).value
    )

    test.compare(
        actualEditFrequency,
        expectedEditFrequency,
        "Dutch - Edit Frequency"
    )


    # ============================================================
    # EDIT - FREQUENCY UNIT
    # ============================================================

    snooze(6)

    actualEditFrequencyUnit = str(
        waitForObjectExists(
            names.frequencyWidget_unitsOfMeasureLabel_QLabel
        ).text
    )

    expectedEditFrequencyUnit = str(
        sheet.cell(
            row=11,
            column=7
        ).value
    )

    test.compare(
        actualEditFrequencyUnit,
        expectedEditFrequencyUnit,
        "Dutch - Edit Frequency Unit"
    )


    # ============================================================
    # EDIT - REGULAR PULSE
    # ============================================================

    snooze(6)

    actualEditRegularPulse = str(
        waitForObjectExists(
            names.modeButton_nameLabel_QLabel_3
        ).text
    )

    expectedEditRegularPulse = str(
        sheet.cell(
            row=33,
            column=7
        ).value
    )

    test.compare(
        actualEditRegularPulse,
        expectedEditRegularPulse,
        "Dutch - Edit Regular Pulse"
    )


    # ============================================================
    # EDIT - PEAK POWER
    # ============================================================

    snooze(6)

    actualEditPeakPower = str(
        waitForObjectExists(
            names.peakPowerButton_nameLabel_QLabel_2
        ).text
    )

    expectedEditPeakPower = str(
        sheet.cell(
            row=12,
            column=7
        ).value
    )

    test.compare(
        actualEditPeakPower,
        expectedEditPeakPower,
        "Dutch - Edit Peak Power"
    )


    # ============================================================
    # UPDATE BUTTON
    # ============================================================

    snooze(6)

    actualUpdateButton = str(
        waitForObjectExists(
            names.pulseModeWidget_saveButton_QPushButton
        ).text
    )

    expectedUpdateButton = str(
        sheet.cell(
            row=69,
            column=7
        ).value
    )

    test.compare(
        actualUpdateButton,
        expectedUpdateButton,
        "Dutch - Update Button"
    )


    # ============================================================
    # SAVE EDITED PRESET
    # ============================================================

    mouseClick(
        waitForObject(
            names.pulseModeWidget_buttonsWidget_QWidget
        ),
        531,
        22,
        Qt.NoModifier,
        Qt.LeftButton
    )

    mouseClick(
        waitForObject(
            names.pulseModeWidget_buttonsWidget_QWidget
        ),
        509,
        30,
        Qt.NoModifier,
        Qt.LeftButton
    )

    clickButton(
        waitForObject(
            names.pulseModeWidget_saveButton_QPushButton
        )
    )

    snooze(6)


    # ============================================================
    # DELETE PRESET
    # ============================================================

    clickButton(
        waitForObject(
            names.settingsScreen_deletePresetButton_QPushButton
        )
    )

    mouseClick(
        waitForObject(
            names.tabSettingsPresetsUserStonePresets_presetListView_PresetListView
        ),
        553,
        80,
        Qt.NoModifier,
        Qt.LeftButton
    )

    snooze(6)


    # ============================================================
    # DELETE DIALOG TITLE
    # ============================================================

    snooze(6)

    actualDeleteTitle = str(
        waitForObjectExists(
            names.centralWidget_lbDialogTitle_QLabel
        ).text
    )

    expectedDeleteTitle = str(
        sheet.cell(
            row=74,
            column=7
        ).value
    )

    test.compare(
        actualDeleteTitle,
        expectedDeleteTitle,
        "Dutch - Delete Dialog Title"
    )


    # ============================================================
    # DELETE PRESET NAME
    # ============================================================

    snooze(6)

    actualDeleteName = str(
        waitForObjectExists(
            names.centralWidget_questionLabel_QLabel
        ).text
    )

    test.compare(
        actualDeleteName,
        "preset one",
        "Delete Preset Name"
    )


    # ============================================================
    # PRESET INFORMATION
    # ============================================================

    snooze(6)

    actualInfo = str(
        waitForObjectExists(
            names.infoFrame_extInfoLabel_QLabel
        ).text
    )

    expectedInfo = "6 J     3 Hz     18 W"

    test.compare(
        actualInfo,
        expectedInfo,
        "Preset Information"
    )


    # ============================================================
    # CANCEL BUTTON
    # ============================================================

    snooze(6)

    actualCancel = str(
        waitForObjectExists(
            names.buttonsFrame_rejectButton_QPushButton
        ).text
    )

    expectedCancel = str(
        sheet.cell(
            row=71,
            column=7
        ).value
    )

    test.compare(
        actualCancel,
        expectedCancel,
        "Dutch - Cancel Button"
    )


    # ============================================================
    # DELETE BUTTON
    # ============================================================

    snooze(6)

    actualDelete = str(
        waitForObjectExists(
            names.centralWidget_acceptButton_QPushButton
        ).text
    )

    expectedDelete = str(
        sheet.cell(
            row=72,
            column=7
        ).value
    )

    test.compare(
        actualDelete,
        expectedDelete,
        "Dutch - Delete Button"
    )


    # ============================================================
    # CLOSE EXCEL
    # ============================================================

    workbook.close()

    test.log(
        "Dutch Peak Power - Soft Tissue User Presets "
        "localization verification completed"
    )

