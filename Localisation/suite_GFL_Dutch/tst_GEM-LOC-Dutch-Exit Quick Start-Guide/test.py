# -*- coding: utf-8 -*-

import names
import openpyxl
import re

# =========================================================

# Excel Configuration

# =========================================================

EXCEL_PATH = (
"/home/ezen/Test_Automation_Gemini/Localisation/"
"suite_GFL_Dutch/"
"tst_GEM-LOC-Dutch-PeakPower-SoftTissueCharacters/"
"testdata/Gemini Split strings.xlsx"
)

# Excel:

# Column B = English

# Column G = Dutch

ENGLISH_COLUMN = 2
DUTCH_COLUMN = 7

SHEET_NAME = "Soft tissue"

# =========================================================

# Normalize Text

# =========================================================

def normalize(text):

    
    if text is None:
        return ""
    
    text = str(text)
    
    # Remove HTML tags
    text = re.sub(
        r"<[^>]+>",
        " ",
        text
    )
    
    # Normalize spaces and new lines
    text = " ".join(
        text.split()
    )
    
    # Normalize angle brackets
    text = re.sub(
        r"\s*<\s*",
        " < ",
        text
    )
    
    text = re.sub(
        r"\s*>\s*",
        " > ",
        text
    )
    
    return text.strip().casefold()
    

# =========================================================

# Load Excel

# =========================================================

def load_excel():

    
    try:
    
        test.log(
            "Loading Dutch localization Excel..."
        )
    
        workbook = openpyxl.load_workbook(
            EXCEL_PATH,
            data_only=True,
            read_only=True
        )
    
        if SHEET_NAME not in workbook.sheetnames:
    
            test.fail(
                "Excel sheet not found: "
                + SHEET_NAME
            )
    
            workbook.close()
    
            return None, None
    
        sheet = workbook[SHEET_NAME]
    
        test.log(
            "Excel sheet loaded successfully: "
            + SHEET_NAME
        )
    
        return workbook, sheet
    
    except Exception as e:
    
        test.fail(
            "Unable to load Dutch localization Excel: "
            + str(e)
        )
    
        return None, None


# =========================================================

# Get Dutch Values

#

# Search English text in Column B.

# Return ALL matching Dutch values from Column G.

#

# No Excel row numbers are hard-coded.

# =========================================================

def get_dutch_values(
sheet,
english_text
):
    
    expected_english = normalize(
        english_text
    )
    
    dutch_values = []
    
    for row in range(
        2,
        sheet.max_row + 1
    ):
    
        english_value = sheet.cell(
            row=row,
            column=ENGLISH_COLUMN
        ).value
    
        if english_value is None:
            continue
    
        if normalize(
            english_value
        ) == expected_english:
    
            dutch_value = sheet.cell(
                row=row,
                column=DUTCH_COLUMN
            ).value
    
            if dutch_value is not None:
    
                dutch_value = str(
                    dutch_value
                ).strip()
    
                if dutch_value:
    
                    dutch_values.append(
                        dutch_value
                    )
    
    return dutch_values


# =========================================================

# Verify Object Text Against Excel

# =========================================================

def verify_text(
sheet,
object_name,
english_text,
description
):

    
    # Wait before every verification
    snooze(6)
    
    try:
    
        obj = waitForObjectExists(
            object_name
        )
    
        actual_value = obj.text
    
        actual = ""
    
        if actual_value is not None:
    
            actual = str(
                actual_value
            ).strip()
    
    except Exception as e:
    
        test.fail(
            description
            + " | Unable to read object: "
            + str(e)
        )
    
        return
    
    
# =====================================================
# Get Dutch translation from Excel
# =====================================================

    expected_values = get_dutch_values(
        sheet,
        english_text
    )


# =====================================================
# English string not found
# =====================================================

    if not expected_values:
    
        test.fail(
            description
            + " | English string not found in Column B: "
            + english_text
        )
    
        return
    
    
    actual_normalized = normalize(
        actual
    )


# =====================================================
# Compare against all matching Dutch values
# =====================================================

    for expected in expected_values:
    
        expected_normalized = normalize(
            expected
        )
    
        if actual_normalized == expected_normalized:
    
            test.compare(
                actual_normalized,
                expected_normalized,
                description
            )
    
            test.log(
                description
                + " | PASS"
                + " | English: "
                + english_text
                + " | Dutch: "
                + expected
            )
    
            return


# =====================================================
# Verification failed
# =====================================================

    test.fail(
        description
        + " | Expected Dutch value from Column G: "
        + " / ".join(
            expected_values
        )
        + " | Actual: "
        + actual
    )
    

# =========================================================

# Get Visible centralWidget

# =========================================================

def get_central_widget():


    try:
    
        central_widget = findObject(
            {
                "name": "centralWidget",
                "type": "QFrame",
                "window": ":o_ScreenSwitcher"
            }
        )
    
        for i in range(20):
    
            try:
    
                if central_widget.visible:
    
                    test.log(
                        "centralWidget is visible"
                    )
    
                    return central_widget
    
            except:
    
                pass
    
            snooze(0.5)
    
        test.log(
            "centralWidget was found but is not visible"
        )
    
        return None
    
    except Exception as e:
    
        test.log(
            "centralWidget not found: "
            + str(e)
        )
    
        return None


# =========================================================

# Main

# =========================================================

def main():


    workbook, sheet = load_excel()
    
    if workbook is None or sheet is None:
        return


# =====================================================
# LOGIN
# =====================================================

    clickButton(
        waitForObject(
            names.numpad_pushButton_8_NumpadButton
        )
    )
    
    doubleClick(
        waitForObject(
            names.numpad_pushButton_8_NumpadButton
        ),
        94,
        77,
        Qt.NoModifier,
        Qt.LeftButton
    )
    
    doubleClick(
        waitForObject(
            names.numpad_pushButton_8_NumpadButton
        ),
        94,
        77,
        Qt.NoModifier,
        Qt.LeftButton
    )
    
    clickButton(
        waitForObject(
            names.mainFrame_okButton_QPushButton
        )
    )
    
    
    # =====================================================
    # CHANGE LANGUAGE TO DUTCH
    # =====================================================
    
    clickButton(
        waitForObject(
            names.homeScreen_settingsButton_QPushButton
        )
    )
    
    clickTab(
        waitForObject(
            names.settingsScreen_tbSystemInfo_TabWidget
        ),
        "INFO"
    )
    
    clickButton(
        waitForObject(
            names.gbSystemInformation_pbLanguageSettingsUpdate_QPushButton
        )
    )
    
    # Select DUTCH
    clickButton(
        waitForObject(
            names.languagesFrame_pbDutch_QPushButton
        )
    )
    
    clickButton(
        waitForObject(
            names.buttonsFrame_acceptButton_QPushButton
        )
    )
    
    clickButton(
        waitForObject(
            names.buttonsBar_homeButton_QPushButton
        )
    )
    
    
    # =====================================================
    # SOFT TISSUE QUICK START
    # =====================================================
    
    clickButton(
        waitForObject(
            names.homeScreen_pbSoftTissueQuickStart_QToolButton
        )
    )
    
    mouseClick(
        waitForObjectItem(
            names.tabLeftMainPresets_leftPresetsView_PresetListView,
            "_1"
        ),
        377,
        65,
        Qt.NoModifier,
        Qt.LeftButton
    )
    
    mouseClick(
        waitForObjectItem(
            names.tabRightMainPresets_rightPresetsView_PresetListView,
            "_1"
        ),
        150,
        37,
        Qt.NoModifier,
        Qt.LeftButton
    )
    
    sendEvent(
        "QMoveEvent",
        waitForObject(
            names.o_ScreenSwitcher
        ),
        0,
        -50,
        1288,
        292
    )
    
    clickButton(
        waitForObject(
            names.quickStartScreen_pbContinue_QPushButton
        )
    )
    
    clickButton(
        waitForObject(
            names.mainFrame_okButton_QPushButton_2
        )
    )
    
    
    # =====================================================
    # QUICK START EXIT POPUP
    # =====================================================
    
    mouseClick(
        waitForObject(
            names.modeButton_nameLabel_QLabel
        ),
        115,
        15,
        Qt.NoModifier,
        Qt.LeftButton
    )
    
    
    # =====================================================
    # QUICK START EXIT TITLE
    # English lookup:
    # EXIT QUICK START MODE?
    # =====================================================
    
    verify_text(
        sheet,
        names.centralWidget_TitleLabel_QLabel_2,
        "EXIT QUICK START MODE?",
        "Quick Start Exit Title"
    )
    
    
    # =====================================================
    # QUICK START EXIT MESSAGE
    # =====================================================
    
    verify_text(
        sheet,
        names.centralWidget_MessageLabel_QLabel_2,
        "Please confirm that you want to exit quick start mode.",
        "Quick Start Exit Message"
    )
    
    
    # =====================================================
    # QUICK START CANCEL
    # =====================================================
    
    verify_text(
        sheet,
        names.centralWidget_CancelButton_QPushButton_3,
        "CANCEL",
        "Quick Start Cancel Button"
    )
    
    
    # =====================================================
    # QUICK START CONFIRM
    # =====================================================
    
    verify_text(
        sheet,
        names.centralWidget_DisableButton_QPushButton_2,
        "CONFIRM",
        "Quick Start Confirm Button"
    )
    
    
    # =====================================================
    # CONFIRM EXIT
    # =====================================================
    
    clickButton(
        waitForObject(
            names.centralWidget_DisableButton_QPushButton_2
        )
    )
    
    
    # =====================================================
    # RETURN HOME
    # =====================================================
    
    clickButton(
        waitForObject(
            names.buttonsBar_homeButton_QPushButton_2
        )
    )
    
    
    # =====================================================
    # SOFT TISSUE ASSISTANT
    # =====================================================
    
    clickButton(
        waitForObject(
            names.homeScreen_pbSoftTissueAssistant_QToolButton
        )
    )
    
    clickButton(
        waitForObject(
            names.gbSoftTissueProcedures_pbIncision_QPushButton
        )
    )
    
    clickButton(
        waitForObject(
            names.softTissueAssistantModeScreen_pbContinue_QPushButton
        )
    )
    
    
    # =====================================================
    # GUIDED MODE EXIT POPUP
    # =====================================================
    
    mouseDrag(
        waitForObject(
            names.indexPicker_valuePicker_IntervalSlider
        ),
        195,
        15,
        211,
        29,
        1,
        Qt.LeftButton
    )
    
    
    # =====================================================
    # GUIDED MODE EXIT TITLE
    # =====================================================
    
    verify_text(
        sheet,
        names.centralWidget_TitleLabel_QLabel_2,
        "EXIT GUIDED MODE?",
        "Guided Mode Exit Title"
    )
    
    
    # =====================================================
    # GUIDED MODE EXIT MESSAGE
    # =====================================================
    
    verify_text(
        sheet,
        names.centralWidget_MessageLabel_QLabel_2,
        "Please confirm that you want to exit guided mode.",
        "Guided Mode Exit Message"
    )
    
    
    # =====================================================
    # GUIDED MODE CANCEL
    # =====================================================
    
    verify_text(
        sheet,
        names.centralWidget_CancelButton_QPushButton_3,
        "CANCEL",
        "Guided Mode Cancel Button"
    )
    
    
    # =====================================================
    # GUIDED MODE CONFIRM
    # =====================================================
    
    verify_text(
        sheet,
        names.centralWidget_DisableButton_QPushButton_2,
        "CONFIRM",
        "Guided Mode Confirm Button"
    )
    
    
    # =====================================================
    # CLOSE GUIDED MODE POPUP
    # =====================================================
    
    clickButton(
        waitForObject(
            names.centralWidget_DisableButton_QPushButton_2
        )
    )
    
    
    # =====================================================
    # CLOSE EXCEL
    # =====================================================
    
    workbook.close()
    
