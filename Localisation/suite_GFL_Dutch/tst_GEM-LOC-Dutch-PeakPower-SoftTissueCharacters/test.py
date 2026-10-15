# -*- coding: utf-8 -*-

import names
import openpyxl
import re
import time


def normalize(text):
    if text is None:
        return ""

    text = str(text)

    # Remove HTML <br> tags
    text = re.sub(
        r"<br\s*/?>",
        " ",
        text,
        flags=re.IGNORECASE
    )

    # Remove remaining HTML tags
    text = re.sub(
        r"<[^>]+>",
        " ",
        text
    )

    # Replace non-breaking spaces
    text = text.replace("\xa0", " ")

    # Remove extra spaces / new lines
    text = " ".join(text.split())

    return text.strip().casefold()


def load_dutch_excel(excelpath):

    start_time = time.time()

    workbook = openpyxl.load_workbook(
        excelpath,
        data_only=True,
        read_only=True
    )

    load_time = time.time() - start_time

    # test.log(
    #     "Excel loading completed in "
    #     + str(round(load_time, 2))
    #     + " seconds"
    # )

    if "Treatment Modes" not in workbook.sheetnames:

        # test.fail(
        #     "Worksheet 'Treatment Modes' does not exist "
        #     "in Excel file"
        # )

        workbook.close()
        return None, None, {}

    sheet = workbook["Treatment Modes"]

    english_column = 2
    dutch_column = 7

    dutch_data = {}

    start_time = time.time()

    for row in sheet.iter_rows(
        min_col=english_column,
        max_col=dutch_column
    ):

        english_value = row[0].value
        dutch_value = row[5].value

        if english_value is None:
            continue

        english_key = normalize(
            english_value
        )

        if english_key == "":
            continue

        dutch_data[english_key] = (
            ""
            if dutch_value is None
            else str(dutch_value)
        )

    dictionary_time = time.time() - start_time

   

    return workbook, sheet, dutch_data


def get_dutch_text(dutch_data, english_text):

    key = normalize(
        english_text
    )

    if key not in dutch_data:

        # test.fail(
        #     "English text not found in Excel: "
        #     + str(english_text)
        # )

        return ""

    expected = dutch_data[key]

    return expected


def verify_text(
    dutch_data,
    object_name,
    english_text
):

    actual_text = ""

    try:

        obj = waitForObjectExists(
            object_name
        )

        actual_text = str(
            obj.text
        ).strip()

    except Exception as e:

        test.fail(
            "Unable to read application object: "
            + str(object_name)
            + " | Error: "
            + str(e)
        )

        return False

    expected_text = get_dutch_text(
        dutch_data,
        english_text
    )

    actual_normalized = normalize(
        actual_text
    )

    expected_normalized = normalize(
        expected_text
    )

    if actual_normalized == expected_normalized:

        # test.log(
        #     "PASS - Dutch text verified: "
        #     + str(english_text)
        #     + " -> "
        #     + actual_text
        # )

        test.compare(
            actual_normalized,
            expected_normalized,
            "Dutch text verification: "
            + str(english_text)
        )

        return True

    else:

        test.fail(
            "FAIL - Dutch text mismatch for: "
            + str(english_text)
            + " | Actual: "
            + actual_text
            + " | Expected: "
            + str(expected_text)
        )

        return False


def main():

    excelpath = (
        "/home/ezen/Test_Automation_Gemini/Localisation/"
        "suite_GFL_Dutch/"
        "tst_GEM-LOC-Dutch-PeakPower-SoftTissueCharacters/"
        "testdata/Gemini Split strings.xlsx"
    )



    clickButton(
        waitForObject(
            names.numpad_pushButton_8_NumpadButton
        )
    )

    clickButton(
        waitForObject(
            names.numpad_pushButton_8_NumpadButton
        )
    )

    clickButton(
        waitForObject(
            names.numpad_pushButton_8_NumpadButton
        )
    )

    clickButton(
        waitForObject(
            names.numpad_pushButton_8_NumpadButton
        )
    )

    snooze(6)

    clickButton(
        waitForObject(
            names.mainFrame_okButton_QPushButton
        )
    )



    clickButton(
        waitForObject(
            names.homeScreen_settingsButton_QPushButton
        )
    )

    clickTab(
        waitForObject(
            names.settingsScreen_tbSystemInfo_TabWidget
        ),
        "INFO"
    )

    clickButton(
        waitForObject(
            names.gbSystemInformation_pbLanguageSettingsUpdate_QPushButton
        )
    )


    clickButton(
        waitForObject(
            names.languagesFrame_pbDutch_QPushButton
        )
    )

    clickButton(
        waitForObject(
            names.buttonsFrame_acceptButton_QPushButton
        )
    )

    clickButton(
        waitForObject(
            names.settingsScreen_pbBack_QPushButton
        )
    )



    clickButton(
        waitForObject(
            names.homeScreen_pbSoftTissueQuickStart_QToolButton
        )
    )

    mouseClick(
        waitForObjectItem(
            names.tabLeftMainPresets_leftPresetsView_PresetListView,
            "_1"
        ),
        305,
        35,
        Qt.NoModifier,
        Qt.LeftButton
    )

    mouseClick(
        waitForObjectItem(
            names.tabRightMainPresets_rightPresetsView_PresetListView,
            "_1"
        ),
        299,
        49,
        Qt.NoModifier,
        Qt.LeftButton
    )

    clickButton(
        waitForObject(
            names.quickStartScreen_pbContinue_QPushButton
        )
    )

    clickButton(
        waitForObject(
            names.mainFrame_okButton_QPushButton_2
        )
    )



    mouseClick(
        waitForObject(
            names.peakPowerButton_nameLabel_QLabel
        ),
        34,
        34,
        Qt.NoModifier,
        Qt.LeftButton
    )

    snooze(6)



    workbook, sheet, dutch_data = load_dutch_excel(
        excelpath
    )



    snooze(6)

    verify_text(
        dutch_data,
        names.centralWidget_lbTitle_QLabel,
        "SELECT PEAK POWER"
    )

    snooze(6)

    verify_text(
        dutch_data,
        names.centralWidget_pw500_QPushButton,
        "High"
    )

    snooze(6)

    verify_text(
        dutch_data,
        names.centralWidget_pw250_QPushButton,
        "Medium"
    )

    snooze(6)

    verify_text(
        dutch_data,
        names.centralWidget_pw125_QPushButton,
        "Low"
    )

    snooze(6)

    test.vp("Power Modes")

    verify_text(
        dutch_data,
        names.buttonsFrame_rejectButton_QPushButton,
        "CANCEL"
    )



    mouseClick(
        waitForObject(
            names.centralWidget_mainWidget_QWidget
        ),
        460,
        549,
        Qt.NoModifier,
        Qt.LeftButton
    )

    clickButton(
        waitForObject(
            names.buttonsFrame_rejectButton_QPushButton
        )
    )



    clickButton(
        waitForObject(
            names.buttonsBar_homeButton_QPushButton_2
        )
    )

    clickButton(
        waitForObject(
            names.homeScreen_pbSoftTissueAssistant_QToolButton
        )
    )

    clickButton(
        waitForObject(
            names.gbSoftTissueProcedures_pbIncision_QPushButton
        )
    )

    clickButton(
        waitForObject(
            names.softTissueAssistantModeScreen_pbContinue_QPushButton
        )
    )

    clickButton(
        waitForObject(
            names.treatmentCharacteristicsBar_pbEdit_QPushButton
        )
    )

    snooze(6)




    verify_text(
        dutch_data,
        names.centralWidget_titleLabel_QLabel,
        "SOFT TISSUE CHARACTERISTICS"
    )

    snooze(6)

    verify_text(
        dutch_data,
        names.centralWidget_titleLabel_QLabel_2,
        "PROCEDURE"
    )

    snooze(6)

    verify_text(
        dutch_data,
        names.centralWidget_valueLabel_QLabel,
        "Incision"
    )

    snooze(6)

    verify_text(
        dutch_data,
        names.centralWidget_rejectButton_QPushButton,
        "EDIT"
    )

    snooze(6)

    test.vp("Soft tissue characteristics") 

    verify_text(
        dutch_data,
        names.centralWidget_acceptButton_QPushButton,
        "OK"
    )


    workbook.close()
