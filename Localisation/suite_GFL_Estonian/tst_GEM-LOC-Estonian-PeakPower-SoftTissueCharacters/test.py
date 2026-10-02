
# -*- coding: utf-8 -*-

import names
import openpyxl
import re
import time


# ---------------------------------------------------------
# Normalize text for comparison
# ---------------------------------------------------------
def normalize(text):

    if text is None:
        return ""

    text = str(text)

    # Remove HTML <br> tags
    text = re.sub(
        r"<br\s*/?>",
        " ",
        text,
        flags=re.IGNORECASE
    )

    # Remove remaining HTML tags
    text = re.sub(
        r"<[^>]+>",
        " ",
        text
    )

    # Replace non-breaking spaces
    text = text.replace(
        "\xa0",
        " "
    )

    # Remove extra spaces / new lines
    text = " ".join(
        text.split()
    )

    return text.strip().casefold()


# ---------------------------------------------------------
# Load Excel and create:
# English -> Estonian dictionary
#
# Excel:
# Column B = English
# Column H = Estonian
# Sheet   = Treatment modes
# ---------------------------------------------------------
def load_estonian_excel(excelpath):

    start_time = time.time()

    workbook = openpyxl.load_workbook(
        excelpath,
        data_only=True,
        read_only=True
    )

    if "Treatment Modes" not in workbook.sheetnames:

        workbook.close()

        test.fail(
            "Worksheet 'Treatment Modes' does not exist "
            "in Excel file"
        )

        return None, None, {}

    sheet = workbook["Treatment Modes"]

    # -----------------------------------------------------
    # Excel columns
    # -----------------------------------------------------
    # English  = B = 2
    # Estonian = H = 8
    # -----------------------------------------------------

    english_column = 2
    estonian_column = 8

    estonian_data = {}

    for row in sheet.iter_rows(
        min_col=english_column,
        max_col=estonian_column
    ):

        english_value = row[0].value
        estonian_value = row[6].value

        if english_value is None:
            continue

        english_key = normalize(
            english_value
        )

        if english_key == "":
            continue

        estonian_data[english_key] = (
            ""
            if estonian_value is None
            else str(estonian_value)
        )

    return workbook, sheet, estonian_data


# ---------------------------------------------------------
# Get Estonian expected text using English text
# ---------------------------------------------------------
def get_estonian_text(
    estonian_data,
    english_text
):

    key = normalize(
        english_text
    )

    if key not in estonian_data:

        test.fail(
            "English text not found in Excel: "
            + str(english_text)
        )

        return ""

    expected = estonian_data[key]

    return expected


# ---------------------------------------------------------
# Compare application text with Excel Estonian text
# ---------------------------------------------------------
def verify_text(
    estonian_data,
    object_name,
    english_text
):

    actual_text = ""

    try:

        obj = waitForObjectExists(
            object_name
        )

        actual_text = str(
            obj.text
        ).strip()

    except Exception as e:

        test.fail(
            "Unable to read application object: "
            + str(object_name)
            + " | Error: "
            + str(e)
        )

        return False

    expected_text = get_estonian_text(
        estonian_data,
        english_text
    )

    actual_normalized = normalize(
        actual_text
    )

    expected_normalized = normalize(
        expected_text
    )

    if actual_normalized == expected_normalized:

        test.compare(
            actual_normalized,
            expected_normalized,
            "Estonian text verification: "
            + str(english_text)
        )

        return True

    else:

        test.fail(
            "Estonian text mismatch for: "
            + str(english_text)
            + " | Actual: "
            + actual_text
            + " | Expected: "
            + str(expected_text)
        )

        return False


# ---------------------------------------------------------
# MAIN
# ---------------------------------------------------------
def main():

    excelpath = (
        "/home/ezen/Test_Automation_Gemini/Localisation/"
        "suite_GFL_Estonian/"
        "tst_GEM-LOC-Estonian-PeakPower-SoftTissueCharacters/"
        "testdata/Gemini Split strings.xlsx"
    )


    # =====================================================
    # LOGIN
    # =====================================================

    clickButton(
        waitForObject(
            names.numpad_pushButton_8_NumpadButton
        )
    )

    clickButton(
        waitForObject(
            names.numpad_pushButton_8_NumpadButton
        )
    )

    clickButton(
        waitForObject(
            names.numpad_pushButton_8_NumpadButton
        )
    )

    clickButton(
        waitForObject(
            names.numpad_pushButton_8_NumpadButton
        )
    )

    snooze(6)

    clickButton(
        waitForObject(
            names.mainFrame_okButton_QPushButton
        )
    )


    # =====================================================
    # CHANGE LANGUAGE TO ESTONIAN
    # =====================================================

    clickButton(
        waitForObject(
            names.homeScreen_settingsButton_QPushButton
        )
    )

    clickTab(
        waitForObject(
            names.settingsScreen_tbSystemInfo_TabWidget
        ),
        "INFO"
    )

    clickButton(
        waitForObject(
            names.gbSystemInformation_pbLanguageSettingsUpdate_QPushButton
        )
    )


    # -----------------------------------------------------
    # ESTONIAN LANGUAGE BUTTON
    # -----------------------------------------------------

    clickButton(
        waitForObject(
            names.languagesFrame_pbEstonia_QPushButton
        )
    )


    # -----------------------------------------------------
    # ACCEPT LANGUAGE
    # -----------------------------------------------------

    clickButton(
        waitForObject(
            names.buttonsFrame_acceptButton_QPushButton
        )
    )

    clickButton(
        waitForObject(
            names.settingsScreen_pbBack_QPushButton
        )
    )


    # =====================================================
    # SOFT TISSUE QUICK START
    # =====================================================

    clickButton(
        waitForObject(
            names.homeScreen_pbSoftTissueQuickStart_QToolButton
        )
    )

    mouseClick(
        waitForObjectItem(
            names.tabLeftMainPresets_leftPresetsView_PresetListView,
            "_1"
        ),
        305,
        35,
        Qt.NoModifier,
        Qt.LeftButton
    )

    mouseClick(
        waitForObjectItem(
            names.tabRightMainPresets_rightPresetsView_PresetListView,
            "_1"
        ),
        299,
        49,
        Qt.NoModifier,
        Qt.LeftButton
    )

    clickButton(
        waitForObject(
            names.quickStartScreen_pbContinue_QPushButton
        )
    )

    clickButton(
        waitForObject(
            names.mainFrame_okButton_QPushButton_2
        )
    )


    # =====================================================
    # PEAK POWER
    # =====================================================

    # mouseClick(waitForObject(names.peakPowerButton_nameLabel_QLabel_2), 52, 19, Qt.NoModifier, Qt.LeftButton)
    # test.compare(str(waitForObjectExists(names.centralWidget_lbTitle_QLabel).text), "VALI TIPUVÕIMSUS")
    # test.compare(str(waitForObjectExists(names.centralWidget_pw500_QPushButton).text), "Suur")
    # test.compare(str(waitForObjectExists(names.centralWidget_pw250_QPushButton).text), "Keskmise suurusega")
    # test.compare(str(waitForObjectExists(names.centralWidget_pw125_QPushButton).text), "Väike")
    # test.compare(str(waitForObjectExists(names.buttonsFrame_rejectButton_QPushButton).text), "TÜHISTA")

    mouseClick(
        waitForObject(
            names.peakPowerButton_nameLabel_QLabel_2
        ),
        34,
        34,
        Qt.NoModifier,
        Qt.LeftButton
    )

    snooze(6)


    # =====================================================
    # LOAD EXCEL
    #
    # Excel is loaded only after reaching Peak Power.
    # =====================================================

    workbook, sheet, estonian_data = load_estonian_excel(
        excelpath
    )

    if workbook is None:
        return


    # =====================================================
    # PEAK POWER ESTONIAN VERIFICATION
    #
    # NO ROW NUMBER
    # NO COLUMN NUMBER
    #
    # English text is used as the lookup key.
    # =====================================================

    snooze(6)

    verify_text(
        estonian_data,
        names.centralWidget_lbTitle_QLabel,
        "SELECT PEAK POWER"
    )

    snooze(6)

    verify_text(
        estonian_data,
        names.centralWidget_pw500_QPushButton,
        "High"
    )

    snooze(6)

    verify_text(
        estonian_data,
        names.centralWidget_pw250_QPushButton,
        "Medium"
    )

    snooze(6)

    verify_text(
        estonian_data,
        names.centralWidget_pw125_QPushButton,
        "Low"
    )

    snooze(6)

    verify_text(
        estonian_data,
        names.buttonsFrame_rejectButton_QPushButton,
        "CANCEL"
    )


    # =====================================================
    # CLOSE PEAK POWER
    # =====================================================

    

    mouseClick(
        waitForObject(
            names.centralWidget_mainWidget_QWidget
        ),
        460,
        549,
        Qt.NoModifier,
        Qt.LeftButton
    )

    clickButton(
        waitForObject(
            names.buttonsFrame_rejectButton_QPushButton
        )
    )


    # =====================================================
    # SOFT TISSUE ASSISTANT
    # =====================================================

    clickButton(
        waitForObject(
            names.buttonsBar_homeButton_QPushButton_2
        )
    )

    clickButton(
        waitForObject(
            names.homeScreen_pbSoftTissueAssistant_QToolButton
        )
    )

    clickButton(
        waitForObject(
            names.gbSoftTissueProcedures_pbIncision_QPushButton
        )
    )

    clickButton(
        waitForObject(
            names.softTissueAssistantModeScreen_pbContinue_QPushButton
        )
    )

    
    

    clickButton(
        waitForObject(
            names.treatmentCharacteristicsBar_pbEdit_QPushButton
        )
    )

    snooze(6)


    # =====================================================
    # SOFT TISSUE CHARACTERISTICS
    # ESTONIAN VERIFICATION
    #
    # NO ROW NUMBER
    # NO COLUMN NUMBER
    # =====================================================

    snooze(6)

    # test.compare(str(waitForObjectExists(names.centralWidget_titleLabel_QLabel).text), "SOFT TISSUE CHARACTERISTICS")
    # test.compare(str(waitForObjectExists(names.centralWidget_titleLabel_QLabel_2).text), "PROTSEDUUR")
    # test.compare(str(waitForObjectExists(names.centralWidget_valueLabel_QLabel).text), "Incision")
    # test.compare(str(waitForObjectExists(names.centralWidget_rejectButton_QPushButton).text), "REDIGEERI")
    # test.compare(str(waitForObjectExists(names.centralWidget_acceptButton_QPushButton).text), "OK")

    verify_text(
        estonian_data,
        names.centralWidget_titleLabel_QLabel,
        "SOFT TISSUE CHARACTERISTICS"
    )

    snooze(6)

    verify_text(
        estonian_data,
        names.centralWidget_titleLabel_QLabel_2,
        "PROCEDURE"
    )

    snooze(6)

    verify_text(
        estonian_data,
        names.centralWidget_valueLabel_QLabel,
        "Incision"
    )

    snooze(6)

    verify_text(
        estonian_data,
        names.centralWidget_rejectButton_QPushButton,
        "EDIT"
    )

    snooze(6)

    verify_text(
        estonian_data,
        names.centralWidget_acceptButton_QPushButton,
        "OK"
    )


    # =====================================================
    # CLOSE EXCEL
    # =====================================================

    try:

        workbook.close()

    except:

        pass

