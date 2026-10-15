# -*- coding: utf-8 -*-

import names
from openpyxl import load_workbook


EXCEL_PATH = (
"/home/ezen/Test_Automation_Gemini/Localisation/"
"suite_GFL_Estonian/"
"tst_GEM-LOC-Estonian-PeakPower-SoftTissueCharacters/"
"testdata/Gemini Split strings.xlsx"
)

EXCEL_SHEET_NAME = "Treatment Modes"

# Column B = English

ENGLISH_COLUMN = 2

# Column H = Estonian

ESTONIAN_COLUMN = 8


def normalize_text(value):

    if value is None:
        return ""
    
    return (
        "".join(str(value).split())
        .casefold()
    )
    

def load_excel_data(workbook):


    if EXCEL_SHEET_NAME not in workbook.sheetnames:
    
        test.fail(
            "Excel sheet not found: "
            + EXCEL_SHEET_NAME
        )
    
        return None
    
    sheet = workbook[EXCEL_SHEET_NAME]
    
    translation_data = {}
    
    for row in sheet.iter_rows(
        min_row=1,
        min_col=ENGLISH_COLUMN,
        max_col=ESTONIAN_COLUMN,
        values_only=True
    ):
    
        english_value = row[0]
    
        if english_value is None:
            continue
    
        english_key = normalize_text(
            english_value
        )
    
        if english_key == "":
            continue
    
        estonian_value = row[
            ESTONIAN_COLUMN - ENGLISH_COLUMN
        ]
    
        if estonian_value is None:
            translation_data[english_key] = ""
        else:
            translation_data[english_key] = str(
                estonian_value
            )
    
    return translation_data
    

def get_estonian_text(
translation_data,
english_text
):


    search_text = normalize_text(
        english_text
    )
    
    if search_text not in translation_data:
    
        test.fail(
            "English text not found in Excel Column B: "
            + english_text
        )
    
        return ""
    
    estonian_value = translation_data[
        search_text
    ]
    
    if estonian_value is None:
    
        test.fail(
            "Estonian translation is empty in Excel "
            "for English text: "
            + english_text
        )
    
        return ""
    
    estonian_text = str(
        estonian_value
    ).strip()
    
    if estonian_text == "":
    
        test.fail(
            "Estonian translation is blank in Excel "
            "for English text: "
            + english_text
        )
    
        return ""
    
    return estonian_text
    
    
def get_actual_text(
    obj,
    description
    ):


    actual_value = None

    try:
        actual_value = obj.text
    except Exception:
        actual_value = None


    if actual_value is None:
    
        try:
            actual_value = obj.property(
                "text"
            )
        except Exception:
            actual_value = None
    

    if actual_value is None:
    
        try:
            actual_value = obj.windowTitle
        except Exception:
            actual_value = None
    

    if actual_value is None:
    
        try:
            actual_value = obj.accessibleName
        except Exception:
            actual_value = None
    

    actual_text = normalize_text(
        actual_value
    )


    if actual_text == "":
    
        test.fail(
            "Application text is empty or None for: "
            + description
        )
    
        return None
    
    return actual_text
    
    

def verify_text(
translation_data,
object_name,
english_text,
description
):


    snooze(2)


    expected_text = get_estonian_text(
        translation_data,
        english_text
    )
    
    expected_normalized = normalize_text(
        expected_text
    )


    if expected_normalized == "":
    
        test.fail(
            "Expected Estonian text is empty for: "
            + description
        )
    
        return
    

    try:
    
        obj = waitForObjectExists(
            getattr(
                names,
                object_name
            )
        )
    
    except Exception as error:
    
        test.fail(
            "Object not found for "
            + description
            + ": "
            + str(error)
        )
    
        return
    

    actual_text = get_actual_text(
        obj,
        description
    )
    
    
    if actual_text is None:
        return
    

    test.compare(
        actual_text,
        expected_normalized,
        description
        + " - Estonian text verification"
    )


def verify_left_manual_tab(
translation_data,
english_text,
description
):


    snooze(2)


    expected_text = get_estonian_text(
        translation_data,
        english_text
    )
    
    expected_normalized = normalize_text(
        expected_text
    )
    
    
    if expected_normalized == "":
    
        test.fail(
            "Expected Estonian text is empty for: "
            + description
        )
    
        return
    

    try:
    
        manual_tab_object = getattr(
            names,
            "mANUEL_TabItem"
        ).copy()
    
    except Exception as error:
    
        test.fail(
            "Manual tab object not found: "
            + str(error)
        )
    
        return
    

    if "text" in manual_tab_object:
    
        del manual_tab_object["text"]


    try:
    
        obj = waitForObjectExists(
            manual_tab_object
        )
    
    except Exception as error:
    
        test.fail(
            "Left Manual tab not found: "
            + str(error)
        )
    
        return
    

    actual_text = get_actual_text(
        obj,
        description
    )
    
    
    if actual_text is None:
        return
    
    
    test.compare(
        actual_text,
        expected_normalized,
        description
        + " - Estonian text verification"
    )
    
    
def verify_right_manual_tab(
    translation_data,
    english_text,
    description
    ):
    
    
    snooze(2)
    
    
    expected_text = get_estonian_text(
        translation_data,
        english_text
    )
    
    expected_normalized = normalize_text(
        expected_text
    )
    
    
    if expected_normalized == "":
    
        test.fail(
            "Expected Estonian text is empty for: "
        + description
    )

    return

    try:
    
        manual_tab_object = getattr(
            names,
            "mANUEL_TabItem_2"
        ).copy()
    
    except Exception as error:
    
        test.fail(
            "Right Manual tab object not found: "
            + str(error)
        )
    
        return
    
    
    if "text" in manual_tab_object:
    
        del manual_tab_object["text"]
    
    
    try:
    
        obj = waitForObjectExists(
            manual_tab_object
        )
    
    except Exception as error:
    
        test.fail(
            "Right Manual tab not found: "
            + str(error)
        )
    
        return
    
    
    actual_text = get_actual_text(
        obj,
        description
    )
    
    
    if actual_text is None:
        return
    
    
    test.compare(
        actual_text,
        expected_normalized,
        description
        + " - Estonian text verification"
    )
    
    
def find_text_recursively(
    parent,
    expected_text
    ):
    
    
    expected_normalized = normalize_text(
        expected_text
    )
    
    try:
    
        children = object.children(
            parent
        )
    
    except Exception:
    
        return None
    
    
    for child in children:
    
        try:
    
            child_text = get_actual_text(
                child,
                expected_text
            )
    
            if child_text is not None:
    
                if child_text == expected_normalized:
    
                    return child
    
        except Exception:
            pass
    
    
        result = find_text_recursively(
            child,
            expected_text
        )
    
        if result is not None:
    
            return result
    
    return None
    
    
    
def find_localized_button(
    translation_data,
    english_text
    ):
    
    
    expected_text = get_estonian_text(
        translation_data,
        english_text
    )
    
    if normalize_text(expected_text) == "":
    
        return None
    
    
    expected_normalized = normalize_text(
        expected_text
    )
    
    
    try:
    
        root = waitForObject(
            {
                "type": "ScreenSwitcher",
                "unnamed": 1,
                "visible": 1
            },
            10000
        )
    
    except Exception:
    
        root = None
    
    
    if root is None:
    
        return None
    
    
    obj = find_text_recursively(
        root,
        expected_normalized
    )
    
    return obj
    
    
def verify_localized_button(
    translation_data,
    english_text,
    description
    ):
    
    
    snooze(2)
    
    
    expected_text = get_estonian_text(
        translation_data,
        english_text
    )
    
    expected_normalized = normalize_text(
        expected_text
    )
    
    
    if expected_normalized == "":
    
        test.fail(
            "Expected Estonian text is empty for: "
            + description
        )
    
        return
    
    
    obj = find_localized_button(
        translation_data,
        english_text
    )
    
    
    if obj is None:
    
        test.fail(
            "Could not find localized button for: "
            + description
        )
    
        return
    
    
    actual_text = get_actual_text(
        obj,
        description
    )
    
    
    if actual_text is None:
        return
    
    
    test.compare(
        actual_text,
        expected_normalized,
        description
        + " - Estonian text verification"
    )
    
  
    
def main():
    
    
    workbook = None
    
    
    try:
  
        workbook = load_workbook(
            EXCEL_PATH,
            read_only=True,
            data_only=True
        )
    
    
        translation_data = load_excel_data(
            workbook
        )
    
    
        if translation_data is None:
    
            return
    
    
        if len(translation_data) == 0:
    
            test.fail(
                "No English/Estonian translation data "
                "was found in Excel sheet: "
                + EXCEL_SHEET_NAME
            )
    
            return
    
    
        doubleClick(
            waitForObject(
                names.numpad_pushButton_8_NumpadButton,
                100461
            ),
            96,
            56,
            Qt.NoModifier,
            Qt.LeftButton
        )
    
        snooze(1)
    
    
        doubleClick(
            waitForObject(
                names.numpad_pushButton_8_NumpadButton
            ),
            96,
            56,
            Qt.NoModifier,
            Qt.LeftButton
        )
    
        snooze(1)
    
    
        clickButton(
            waitForObject(
                names.mainFrame_okButton_QPushButton
            )
        )
    
        snooze(2)
    
        clickButton(
            waitForObject(
                names.homeScreen_settingsButton_QPushButton
            )
        )
    
        snooze(2)
    
    
        clickTab(
            waitForObject(
                names.settingsScreen_tbSystemInfo_TabWidget
            ),
            "INFO"
        )
    
        snooze(2)
    
    
        clickButton(
            waitForObject(
                names.gbSystemInformation_pbLanguageSettingsUpdate_QPushButton
            )
        )
    
        snooze(2)
    
    
        clickButton(
            waitForObject(
                names.languagesFrame_pbEstonia_QPushButton
            )
        )
    
        snooze(1)
    
    
        clickButton(
            waitForObject(
                names.buttonsFrame_acceptButton_QPushButton
            )
        )
    
        snooze(3)
    
    
        clickButton(
            waitForObject(
                names.buttonsBar_homeButton_QPushButton
            )
        )
    
        snooze(2)
    
    
        clickButton(waitForObject(names.gbExpert_expertButton_QPushButton))

        # clickButton(
        #     waitForObject(
        #         names.gbExpert_expertButton_QPushButton
        #     )
        # )
    
        snooze(3)
    
        verify_text(
            translation_data,
            "averagePowerWidget_nameLabel_QLabel",
            "AVERAGE POWER",
            "Left Average Power"
        )
    
    
        verify_text(
            translation_data,
            "averagePowerWidget_unitsOfMeasureLabel_QLabel_2",
            "W",
            "Left Power Unit"
        )
    
    
        verify_text(
            translation_data,
            "pulseEnergyWidget_nameLabel_QLabel_2",
            "PULSE ENERGY",
            "Left Pulse Energy"
        )
    
    
        verify_text(
            translation_data,
            "pulseEnergyWidget_unitsOfMeasureLabel_QLabel_2",
            "J",
            "Left Energy Unit"
        )
    
    
        verify_text(
            translation_data,
            "frequencyWidget_nameLabel_QLabel_2",
            "FREQUENCY",
            "Left Frequency"
        )
    
    
        verify_text(
            translation_data,
            "frequencyWidget_unitsOfMeasureLabel_QLabel_2",
            "Hz",
            "Left Frequency Unit"
        )
    
    
        verify_text(
            translation_data,
            "modeButton_nameLabel_QLabel",
            "REGULAR PULSE",
            "Left Pulse Mode"
        )
    
    
        verify_text(
            translation_data,
            "peakPowerButton_nameLabel_QLabel_2",
            "PEAK POWER",
            "Left Peak Power"
        )
    
        verify_text(
            translation_data,
            "pulseModeWidget_saveButton_QPushButton_2",
            "SAVE",
            "Left Save Button"
        )
    
    
        verify_text(
            translation_data,
            "totalEnergyFrame_nameLabel_QLabel",
            "TOTAL ENERGY",
            "Left Total Energy"
        )
    
    
        verify_text(
            translation_data,
            "totalTimeFrame_nameLabel_QLabel",
            "TOTAL LASING TIME",
            "Left Total Lasing Time"
        )
    
 
        verify_right_manual_tab(
            translation_data,
            "MANUAL",
            "Right Manual Tab"
        )
    
    
        verify_text(
            translation_data,
            "averagePowerWidget_nameLabel_QLabel_3",
            "AVERAGE POWER",
            "Right Average Power"
        )
    
    
        verify_text(
            translation_data,
            "averagePowerWidget_unitsOfMeasureLabel_QLabel",
            "W",
            "Right Power Unit"
        )
    
    
        verify_text(
            translation_data,
            "pulseEnergyWidget_nameLabel_QLabel_3",
            "PULSE ENERGY",
            "Right Pulse Energy"
        )
    
    
        verify_text(
            translation_data,
            "pulseEnergyWidget_unitsOfMeasureLabel_QLabel_3",
            "J",
            "Right Energy Unit"
        )
    
    
        verify_text(
            translation_data,
            "frequencyWidget_nameLabel_QLabel_3",
            "FREQUENCY",
            "Right Frequency"
        )
    
        verify_text(
            translation_data,
            "pulseModeWidget_saveButton_QPushButton_3",
            "SAVE",
            "Right Save Button"
        )
    
    
        verify_text(
            translation_data,
            "peakPowerButton_nameLabel_QLabel_3",
            "PEAK POWER",
            "Right Peak Power"
        )
    
    
        verify_text(
            translation_data,
            "modeButton_nameLabel_QLabel_3",
            "REGULAR PULSE",
            "Right Pulse Mode"
        )
    
    
        verify_text(
            translation_data,
            "stateSwitch_standbyButton_QPushButton",
            "STANDBY",
            "Standby Button"
        )
    
        verify_text(
            translation_data,
            "fiberFrame_nameLabel_QLabel",
            "FIBER",
            "Fiber"
        )
    
  
        verify_text(
            translation_data,
            "fiberUsesFrame_nameLabel_QLabel",
            "FIBER USES REMAINING",
            "Fiber Uses Remaining"
        )
    
        test.vp("Manual Standby Screen")

        verify_text(
            translation_data,
            "bottomBar_aimingBeamButton_AimingBeamModeButton",
            "AIMING BEAM",
            "Aiming Beam"
        )
    
    
    finally:
    
    
        if workbook is not None:
    
            workbook.close()
    
