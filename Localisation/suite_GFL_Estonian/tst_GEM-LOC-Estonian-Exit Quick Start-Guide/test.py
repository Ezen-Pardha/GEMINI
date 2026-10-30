# -*- coding: utf-8 -*-

import names
import openpyxl
import re
import os 

TESTCASE_DIR = os.path.dirname(os.path.abspath(__file__))
 
EXCEL_PATH = os.path.abspath(
    os.path.join( TESTCASE_DIR,
        "..",
        "shared",
        "testdata",
        "Gemini Split strings.xlsx"
    )
)

ENGLISH_COLUMN = 2
ESTONIAN_COLUMN = 8

SHEET_NAME = "Soft tissue"


def normalize(text):


    if text is None:
        return ""
    
    text = str(text)
    
    # Remove HTML tags
    text = re.sub(
        r"<[^>]+>",
        "",
        text
    )
    
    # Remove all whitespace:
    # spaces
    # multiple spaces
    # new lines
    # tabs
    # carriage returns
    #
    # Example:
    #
    # KESKMINE VÕIMSUS
    # KESKMINE\nVÕIMSUS
    # KESKMINE   VÕIMSUS
    #
    # All become:
    #
    # KESKMINEVÕIMSUS
    
    text = re.sub(
        r"\s+",
        "",
        text
    )
    
    # Case-insensitive comparison
    return text.casefold()
    


def load_excel():


    try:
    
        workbook = openpyxl.load_workbook(
            EXCEL_PATH,
            data_only=True,
            read_only=True
        )
    
        if SHEET_NAME not in workbook.sheetnames:
    
            test.fail(
                "Excel sheet not found: "
                + SHEET_NAME
            )
    
            workbook.close()
    
            return None, None
    
        sheet = workbook[SHEET_NAME]
    
        return workbook, sheet
    
    except Exception as e:
    
        test.fail(
            "Unable to load Estonian localization Excel: "
            + str(e)
        )
    
        return None, None
    

def get_estonian_values(
sheet,
english_text
):


    expected_english = normalize(
        english_text
    )
    
    estonian_values = []
    
    for row in range(
        2,
        sheet.max_row + 1
    ):
    
        english_value = sheet.cell(
            row=row,
            column=ENGLISH_COLUMN
        ).value
    
        if english_value is None:
            continue
    
        # Compare concatenated English strings
        if normalize(
            english_value
        ) == expected_english:
    
            estonian_value = sheet.cell(
                row=row,
                column=ESTONIAN_COLUMN
            ).value
    
            if estonian_value is not None:
    
                estonian_value = str(
                    estonian_value
                ).strip()
    
                if estonian_value:
    
                    estonian_values.append(
                        estonian_value
                    )
    
    return estonian_values
    

def verify_text(
sheet,
object_name,
english_text,
description
):

    # Wait before verification
    snooze(6)
    
    try:
    
        obj = waitForObjectExists(
            object_name
        )
    
        actual_value = obj.text
    
        if actual_value is None:
            actual = ""
        else:
            actual = str(
                actual_value
            )
    
    except Exception as e:
    
        test.fail(
            description
            + " | Unable to read object: "
            + str(e)
        )
    
        return
    

    expected_values = get_estonian_values(
        sheet,
        english_text
    )


    if not expected_values:
    
        test.fail(
            description
            + " | English string not found in Column B: "
            + english_text
        )
    
        return


    actual_normalized = normalize(
        actual
    )


    for expected in expected_values:
    
        expected_normalized = normalize(
            expected
        )
    
        if actual_normalized == expected_normalized:
    
            test.compare(
                actual_normalized,
                expected_normalized,
                description
            )
    
            return
    

    test.fail(
        description
        + " | Expected Estonian value from Column H: "
        + " / ".join(
            expected_values
        )
        + " | Actual: "
        + actual
    )


def get_central_widget():


    try:
    
        central_widget = findObject(
            {
                "name": "centralWidget",
                "type": "QFrame",
                "window": ":o_ScreenSwitcher"
            }
        )
    
        for i in range(20):
    
            try:
    
                if central_widget.visible:
                    return central_widget
    
            except:
                pass
    
            snooze(0.5)
    
        return None
    
    except Exception:
    
        return None
    

def main():

    workbook, sheet = load_excel()
    
    if workbook is None or sheet is None:
        return


    clickButton(
        waitForObject(
            names.numpad_pushButton_8_NumpadButton
        )
    )
    
    doubleClick(
        waitForObject(
            names.numpad_pushButton_8_NumpadButton
        ),
        94,
        77,
        Qt.NoModifier,
        Qt.LeftButton
    )
    
    doubleClick(
        waitForObject(
            names.numpad_pushButton_8_NumpadButton
        ),
        94,
        77,
        Qt.NoModifier,
        Qt.LeftButton
    )
    
    clickButton(
        waitForObject(
            names.mainFrame_okButton_QPushButton
        )
    )
    
    
    clickButton(
        waitForObject(
            names.homeScreen_settingsButton_QPushButton
        )
    )
    
    clickTab(
        waitForObject(
            names.settingsScreen_tbSystemInfo_TabWidget
        ),
        "INFO"
    )
    
    clickButton(
        waitForObject(
            names.gbSystemInformation_pbLanguageSettingsUpdate_QPushButton
        )
    )
    
    clickButton(
        waitForObject(
            names.languagesFrame_pbEstonia_QPushButton
        )
    )
    
    clickButton(
        waitForObject(
            names.buttonsFrame_acceptButton_QPushButton
        )
    )
    
    
    clickButton(waitForObject(names.buttonsBar_homeButton_QPushButton))
    
    clickButton(
        waitForObject(
            names.homeScreen_pbSoftTissueQuickStart_QToolButton
        )
    )
    
    mouseClick(
        waitForObjectItem(
            names.tabLeftMainPresets_leftPresetsView_PresetListView,
            "_1"
        ),
        377,
        65,
        Qt.NoModifier,
        Qt.LeftButton
    )
    
    mouseClick(
        waitForObjectItem(
            names.tabRightMainPresets_rightPresetsView_PresetListView,
            "_1"
        ),
        150,
        37,
        Qt.NoModifier,
        Qt.LeftButton
    )
    
    sendEvent(
        "QMoveEvent",
        waitForObject(
            names.o_ScreenSwitcher
        ),
        0,
        -50,
        1288,
        292
    )
    
    clickButton(
        waitForObject(
            names.quickStartScreen_pbContinue_QPushButton
        )
    )
    
    clickButton(
        waitForObject(
            names.mainFrame_okButton_QPushButton_2
        )
    )
    
    mouseClick(
        waitForObject(
            names.modeButton_nameLabel_QLabel
        ),
        115,
        15,
        Qt.NoModifier,
        Qt.LeftButton
    )

    # test.compare(str(waitForObjectExists(names.centralWidget_TitleLabel_QLabel_2).text), "KAS VÄLJUDA KIIRKÄIVITUSE REŽIIMIST?")
    # test.compare(str(waitForObjectExists(names.centralWidget_MessageLabel_QLabel_2).text), "Kinnitage, et soovite kiirolekurežiimist väljuda.")
    # test.compare(str(waitForObjectExists(names.centralWidget_CancelButton_QPushButton_3).text), "TÜHISTA")
    # test.compare(str(waitForObjectExists(names.centralWidget_DisableButton_QPushButton_2).text), "CONFIRM")

    verify_text(
        sheet,
        names.centralWidget_TitleLabel_QLabel_2,
        "EXIT QUICK START MODE?",
        "Quick Start Exit Title"
    )
    
    
    verify_text(
        sheet,
        names.centralWidget_MessageLabel_QLabel_2,
        "Please confirm you want to exit quick state mode.",
        "Quick Start Exit Message"
    )
    
    
    verify_text(
        sheet,
        names.centralWidget_CancelButton_QPushButton_3,
        "CANCEL",
        "Quick Start Cancel Button"
    )
    
    test.vp("Quick Start Exit Mode")

    verify_text(
        sheet,
        names.centralWidget_DisableButton_QPushButton_2,
        "CONFIRM",
        "Quick Start Confirm Button"
    )
    
    # clickButton(waitForObject(names.centralWidget_DisableButton_QPushButton_2))
    # clickButton(waitForObject(names.buttonsBar_homeButton_QPushButton_2))
    # clickButton(waitForObject(names.homeScreen_pbSoftTissueAssistant_QToolButton))
    # clickButton(waitForObject(names.gbSoftTissueProcedures_pbIncision_QPushButton))
    # clickButton(waitForObject(names.softTissueAssistantModeScreen_pbContinue_QPushButton))
    # mouseClick(waitForObject(names.peakPowerButton_nameLabel_QLabel), 36, 22, Qt.NoModifier, Qt.LeftButton)
    # clickButton(waitForObject(names.buttonsFrame_rejectButton_QPushButton))
    # clickButton(waitForObject(names.pulseEnergyWidget_incrementButton_QPushButton))
    # clickButton(waitForObject(names.pulseEnergyWidget_incrementButton_QPushButton))
    # doubleClick(waitForObject(names.pulseEnergyWidget_incrementButton_QPushButton), 30, 31, Qt.NoModifier, Qt.LeftButton)
    # doubleClick(waitForObject(names.pulseEnergyWidget_incrementButton_QPushButton), 30, 31, Qt.NoModifier, Qt.LeftButton)
    # doubleClick(waitForObject(names.pulseEnergyWidget_incrementButton_QPushButton), 30, 31, Qt.NoModifier, Qt.LeftButton)
    # clickButton(waitForObject(names.pulseEnergyWidget_incrementButton_QPushButton))
    # clickButton(waitForObject(names.pulseEnergyWidget_incrementButton_QPushButton))
    # clickButton(waitForObject(names.pulseEnergyWidget_incrementButton_QPushButton))
    # clickButton(waitForObject(names.pulseEnergyWidget_incrementButton_QPushButton))
    # test.compare(str(waitForObjectExists(names.centralWidget_TitleLabel_QLabel_2).text), "KAS VÄLJUDA JUHENDATUD REŽIIMIST?")
    # test.compare(str(waitForObjectExists(names.centralWidget_MessageLabel_QLabel_2).text), "Kinnitage, et soovite juhendatud režiimist väljuda.")
    # test.compare(str(waitForObjectExists(names.centralWidget_CancelButton_QPushButton_3).text), "TÜHISTA")
    # test.compare(str(waitForObjectExists(names.centralWidget_DisableButton_QPushButton_2).text), "CONFIRM")

    clickButton(waitForObject(names.centralWidget_CancelButton_QPushButton_3))

   
    
    clickButton(waitForObject(names.buttonsBar_homeButton_QPushButton_2))
   
    clickButton(
        waitForObject(
            names.homeScreen_pbSoftTissueAssistant_QToolButton
        )
    )
    
    clickButton(
        waitForObject(
            names.gbSoftTissueProcedures_pbIncision_QPushButton
        )
    )
    
    clickButton(
        waitForObject(
            names.softTissueAssistantModeScreen_pbContinue_QPushButton
        )
    )
    
    
    
    for i in range(30):

        clickButton(
            waitForObject(
                names.pulseEnergyWidget_incrementButton_QPushButton
            )
        )
    
        snooze(0.2)
    
        if i >= 15:
            try:
                waitForObjectExists(
                    names.somePopup_QDialog,
                    0.1
                )
                break
            except:
                pass

                # test.compare(str(waitForObjectExists(names.centralWidget_TitleLabel_QLabel_2).text), "KAS VÄLJUDA JUHENDATUD REŽIIMIST?")
                # test.compare(str(waitForObjectExists(names.centralWidget_MessageLabel_QLabel_2).text), "Kinnitage, et soovite juhendatud režiimist väljuda.")
                # test.compare(str(waitForObjectExists(names.centralWidget_CancelButton_QPushButton_3).text), "TÜHISTA")
                # test.compare(str(waitForObjectExists(names.centralWidget_DisableButton_QPushButton_2).text), "CONFIRM")

    verify_text(
        sheet,
        names.centralWidget_TitleLabel_QLabel_2,
        "EXIT GUIDED MODE?",
        "Guided Mode Exit Title"
    )
     
    
    
    verify_text(
        sheet,
        names.centralWidget_MessageLabel_QLabel_2,
        "Please confirm you want to exit guided mode.",
        "Guided Mode Exit Message"
    )
    
    
    
    verify_text(
        sheet,
        names.centralWidget_CancelButton_QPushButton_3,
        "CANCEL",
        "Guided Mode Cancel Button"
    )
    
    test.vp("Guided Exit Mode")

    verify_text(
        sheet,
        names.centralWidget_DisableButton_QPushButton_2,
        "CONFIRM",
        "Guided Mode Confirm Button"
    )
    
    
    clickButton(
        waitForObject(
            names.centralWidget_DisableButton_QPushButton_2
        )
    )
    
    
    workbook.close()
    
