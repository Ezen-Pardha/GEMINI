# -*- coding: utf-8 -*-
import names
import re
import unicodedata
from openpyxl import load_workbook
EXCEL_PATH = (
    "/home/ezen/Test_Automation_Gemini/Localisation/"
    "suite_GFL_Estonian/"
    "tst_GEM-LOC-Estonian-Userpresets-Add-Edit-Delete/"
    "testdata/Gemini Split strings.xlsx"
)
def normalize_text(value):
    if value is None:
        return ""
    text = unicodedata.normalize("NFKC", str(value))
    text = re.sub(r"<[^>]*>", " ", text)
    text = text.replace("\u00A0", " ").replace("\u2007", " ").replace("\u202F", " ")
    text = text.replace("™", "").replace("®", "").replace("©", "")
    text = re.sub(r"[\r\n\t]+", " ", text)
    return text.strip().casefold()


def comparison_key(value):
    return re.sub(r"\s+", "", normalize_text(value))


def find_column_by_header(sheet, accepted_headers, default_column):
    accepted = {normalize_text(header) for header in accepted_headers}
    for row in sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 15)):
        for cell in row:
            if cell.value is not None and normalize_text(cell.value) in accepted:
                return cell.column
    return default_column


def build_translation_map(sheet, english_column, estonian_column):
    translations = {}
    for row in sheet.iter_rows(min_row=1, max_row=sheet.max_row,
                               min_col=min(english_column, estonian_column),
                               max_col=max(english_column, estonian_column)):
        english_value = row[english_column - min(english_column, estonian_column)].value
        estonian_value = row[estonian_column - min(english_column, estonian_column)].value
        english_key = comparison_key(english_value)
        estonian_text = "" if estonian_value is None else str(estonian_value).strip()
        if english_key and estonian_text and english_key not in translations:
            translations[english_key] = estonian_text
    return translations


def get_excel_value(translation_map, english_text):
    key = comparison_key(english_text)
    if key in translation_map:
        return translation_map[key]
    # Try a normalized lookup in case the source has punctuation or spacing differences.
    for source_key, translated_value in translation_map.items():
        if source_key.replace("-", "") == key.replace("-", ""):
            return translated_value
    return ""


def verify_excel_string(object_name, translation_map, english_text, description):
    snooze(2)
    try:
        actual = get_object_text(waitForObjectExists(object_name))
    except Exception as e:
        test.fail(description + " | Object not found: " + str(e))
        return False
    expected = get_excel_value(translation_map, english_text)
    if not actual.strip():
        test.fail(description + " | UI text is empty")
        return False
    if not expected.strip():
        test.fail(description + " | English key not found or Estonian translation empty in Excel: " + str(english_text))
        return False
    test.compare(comparison_key(actual), comparison_key(expected), description)
    return True


def get_object_text(obj):
    for attribute in ("text", "title", "windowTitle", "accessibleName"):
        try:
            value = getattr(obj, attribute)
            if value is not None and str(value).strip():
                return str(value)
        except Exception:
            pass
    return ""


def main():
    workbook = load_workbook(EXCEL_PATH, read_only=True, data_only=True)
    matching_sheet = next((name for name in workbook.sheetnames
                           if normalize_text(name) == normalize_text("Treatment Modes")), None)
    if matching_sheet is None:
        workbook.close()
        test.fail("Excel sheet 'Treatment Modes' not found. Available sheets: " + ", ".join(workbook.sheetnames))
        return
    sheet = workbook[matching_sheet]
    english_column = 2
    estonian_column = find_column_by_header(
        sheet, ["Estonian", "Eesti", "Estonian (Estonia)"], 8
    )
    translation_map = build_translation_map(sheet, english_column, estonian_column)
    if not translation_map:
        workbook.close()
        test.fail("No English/Estonian translation pairs found. Check English column B and Estonian column.")
        return
    clickButton(
        waitForObject(
            names.numpad_pushButton_8_NumpadButton
        )
    )
    clickButton(
        waitForObject(
            names.numpad_pushButton_8_NumpadButton
        )
    )
    clickButton(
        waitForObject(
            names.numpad_pushButton_8_NumpadButton
        )
    )
    clickButton(
        waitForObject(
            names.numpad_pushButton_8_NumpadButton
        )
    )
    clickButton(
        waitForObject(
            names.mainFrame_okButton_QPushButton
        )
    )
    clickButton(
        waitForObject(
            names.homeScreen_settingsButton_QPushButton
        )
    )
    clickTab(
        waitForObject(
            names.settingsScreen_tbSystemInfo_TabWidget
        ),
        "INFO"
    )
    clickButton(
        waitForObject(
            names.gbSystemInformation_pbLanguageSettingsUpdate_QPushButton
        )
    )
    clickButton(waitForObject(names.languagesFrame_pbEstonia_QPushButton))
    clickButton(
        waitForObject(
            names.buttonsFrame_acceptButton_QPushButton
        )
    )
    snooze(3)
    tabWidget = waitForObject(
        names.settingsScreen_tbSystemInfo_TabWidget
    )
    tabWidget.setCurrentIndex(2)
    snooze(3)
    tabWidget.setCurrentIndex(2)
    snooze(3)
    sendEvent(
        "QMoveEvent",
        waitForObject(
            names.o_ScreenSwitcher
        ),
        0,
        -66,
        1285,
        233
    )
    clickButton(
        waitForObject(
            names.settingsScreen_addButton_QPushButton
        )
    )
    verify_excel_string(
        names.centralWidget_lbOperation_QLabel,
        translation_map,
        "ADD PRESET",
        "Add Preset Operation"
    )
    verify_excel_string(
        names.centralWidget_pbName_QPushButton,
        translation_map,
        "PRESET NAME",
        "Preset Name"
    )
    verify_excel_string(
        names.averagePowerWidget_nameLabel_QLabel,
        translation_map,
        "Average Power",
        "Average Power"
    )
    verify_excel_string(
        names.averagePowerWidget_unitsOfMeasureLabel_QLabel,
        translation_map,
        "W",
        "Average Power Unit"
    )
    verify_excel_string(
        names.pulseEnergyWidget_nameLabel_QLabel,
        translation_map,
        "Pulse Energy",
        "Pulse Energy"
    )
    mouseClick(
        waitForObject(
            names.centralWidget_frame_QFrame
        ),
        271,
        13,
        Qt.NoModifier,
        Qt.LeftButton
    )
    verify_excel_string(
        names.pulseEnergyWidget_unitsOfMeasureLabel_QLabel,
        translation_map,
        "J",
        "Pulse Energy Unit"
    )
    verify_excel_string(
        names.frequencyWidget_nameLabel_QLabel,
        translation_map,
        "Frequency",
        "Frequency"
    )
    verify_excel_string(
        names.frequencyWidget_unitsOfMeasureLabel_QLabel,
        translation_map,
        "Hz",
        "Frequency Unit"
    )
    verify_excel_string(
        names.modeButton_nameLabel_QLabel_3,
        translation_map,
        "Regular Pulse",
        "Regular Pulse"
    )
    verify_excel_string(
        names.peakPowerButton_nameLabel_QLabel_2,
        translation_map,
        "Peak Power",
        "Peak Power"
    )
    verify_excel_string(
        names.pulseModeWidget_saveButton_QPushButton,
        translation_map,
        "ADD",
        "Add Button"
    )
    mouseClick(
        waitForObject(
            names.pulseModeWidget_buttonsWidget_QWidget
        ),
        508,
        46,
        Qt.NoModifier,
        Qt.LeftButton
    )
    mouseClick(
        waitForObject(
            names.pulseModeWidget_buttonsWidget_QWidget
        ),
        411,
        1,
        Qt.NoModifier,
        Qt.LeftButton
    )
    clickButton(
        waitForObject(
            names.centralWidget_pbName_QPushButton
        )
    )
    type(
        waitForObject(
            names.keyboardPopup_wdEdit_WdLineEdit
        ),
        "preset one"
    )
    mouseClick(
        waitForObject(
            names.keyboardPopup_lbText_QLabel
        ),
        77,
        26,
        Qt.NoModifier,
        Qt.LeftButton
    )
    clickButton(
        waitForObject(
            names.pulseModeWidget_saveButton_QPushButton
        )
    )
    mouseClick(
        waitForObject(
            names.tabSettingsPresetsUserStonePresets_presetListView_PresetListView
        ),
        546,
        72,
        Qt.NoModifier,
        Qt.LeftButton
    )
    clickButton(
        waitForObject(
            names.settingsScreen_editButton_QPushButton
        )
    )
    mouseClick(
        waitForObject(
            names.tabSettingsPresetsUserStonePresets_presetListView_PresetListView
        ),
        494,
        66,
        Qt.NoModifier,
        Qt.LeftButton
    )
    verify_excel_string(
        names.centralWidget_lbOperation_QLabel,
        translation_map,
        "EDIT PRESET",
        "Edit Preset Operation"
    )
    snooze(6)
    actualEditName = str(
        waitForObjectExists(
            names.centralWidget_pbName_QPushButton
        ).text
    )
    test.compare(
        actualEditName,
        "preset one",
        "Edit Preset Name"
    )
    verify_excel_string(
        names.averagePowerWidget_nameLabel_QLabel,
        translation_map,
        "Average Power",
        "Edit - Average Power"
    )
    # ========================================================
    # EDIT - AVERAGE POWER UNIT
    # ========================================================
    verify_excel_string(
        names.averagePowerWidget_unitsOfMeasureLabel_QLabel,
        translation_map,
        "W",
        "Edit - Average Power Unit"
    )
    verify_excel_string(
        names.pulseEnergyWidget_nameLabel_QLabel,
        translation_map,
        "Pulse Energy",
        "Edit - Pulse Energy"
    )
    # ========================================================
    # EDIT - PULSE ENERGY UNIT
    # ========================================================
    verify_excel_string(
        names.pulseEnergyWidget_unitsOfMeasureLabel_QLabel,
        translation_map,
        "J",
        "Edit - Pulse Energy Unit"
    )
    verify_excel_string(
        names.frequencyWidget_nameLabel_QLabel,
        translation_map,
        "Frequency",
        "Edit - Frequency"
    )
    verify_excel_string(
        names.frequencyWidget_unitsOfMeasureLabel_QLabel,
        translation_map,
        "Hz",
        "Edit - Frequency Unit"
    )
    verify_excel_string(
        names.modeButton_nameLabel_QLabel_3,
        translation_map,
        "Regular Pulse",
        "Edit - Regular Pulse"
    )
    verify_excel_string(
        names.peakPowerButton_nameLabel_QLabel_2,
        translation_map,
        "Peak Power",
        "Edit - Peak Power"
    )
    verify_excel_string(
        names.pulseModeWidget_saveButton_QPushButton,
        translation_map,
        "UPDATE",
        "Update Button"
    )
    mouseClick(
        waitForObject(
            names.pulseModeWidget_buttonsWidget_QWidget
        ),
        531,
        22,
        Qt.NoModifier,
        Qt.LeftButton
    )
    mouseClick(
        waitForObject(
            names.pulseModeWidget_buttonsWidget_QWidget
        ),
        509,
        30,
        Qt.NoModifier,
        Qt.LeftButton
    )
    clickButton(
        waitForObject(
            names.pulseModeWidget_saveButton_QPushButton
        )
    )
    clickButton(
        waitForObject(
            names.settingsScreen_deletePresetButton_QPushButton
        )
    )
    mouseClick(
        waitForObject(
            names.tabSettingsPresetsUserStonePresets_presetListView_PresetListView
        ),
        553,
        80,
        Qt.NoModifier,
        Qt.LeftButton
    )
    verify_excel_string(
        names.centralWidget_lbDialogTitle_QLabel,
        translation_map,
        "DELETE PRESET",
        "Delete Dialog Title"
    )
    snooze(6)
    actualDeleteName = str(
        waitForObjectExists(
            names.centralWidget_questionLabel_QLabel
        ).text
    )
    test.compare(
        actualDeleteName,
        "preset one",
        "Delete Preset Name"
    )
    snooze(6)
    actualInfo = str(
        waitForObjectExists(
            names.infoFrame_extInfoLabel_QLabel
        ).text
    )
    expectedInfo = "6 J     3 Hz     18 W"
    test.compare(
        actualInfo,
        expectedInfo,
        "Preset Information"
    )
   
    
    verify_excel_string(
        names.buttonsFrame_rejectButton_QPushButton,
        translation_map,
        "CANCEL",
        "Cancel Button"
    )
  

    
    

    
    verify_excel_string(
        names.centralWidget_acceptButton_QPushButton,
        translation_map,
        "DELETE",
        "Delete Button"
    )
  
    clickButton(waitForObject(names.buttonsFrame_acceptButton_QPushButton_2))
    
    workbook.close()
