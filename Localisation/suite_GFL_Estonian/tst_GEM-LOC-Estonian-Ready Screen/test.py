# -*- coding: utf-8 -*-

import names
import re
import unicodedata
from openpyxl import load_workbook
import os

TESTCASE_DIR = os.path.dirname(os.path.abspath(__file__))
 
EXCEL_PATH = os.path.abspath(
    os.path.join( TESTCASE_DIR,
        "..",
        "shared",
        "testdata",
        "Gemini Split strings.xlsx"
    )
)

# EXCEL_PATH = (
# "/home/ezen/Test_Automation_Gemini/Localisation/"
# "suite_GFL_Estonian/"
# "tst_GEM-LOC-Estonian-PeakPower-SoftTissueCharacters/"
# "testdata/Gemini Split strings.xlsx"
# )

EXCEL_SHEET_NAME = "Treatment Modes"



ENGLISH_COLUMN = 2



Estonian_COLUMN = 8


def normalize_text(value):

    if value is None:
        return ""
    
    text = str(value)

    text = re.sub(
        r"<[^>]*>",
        " ",
        text
    )

    text = text.replace("\n", " ")
    text = text.replace("\r", " ")
    text = text.replace("\t", " ")
    text = text.replace("\xa0", " ")
    text = text.replace("\u2007", " ")
    text = text.replace("\u202f", " ")

    text = text.replace("™", "")
    text = text.replace("®", "")
    text = text.replace("©", "")

    try:
    
        text = unicodedata.normalize(
            "NFKC",
            text
        )
    
    except Exception:
        pass
    
    text = re.sub(
        r"\s+",
        " ",
        text
    ).strip()

    text = re.sub(
        r"\s+([.,;:!?])",
        r"\1",
        text
    )
    
    return text.casefold()



def comparison_key(value):


    normalized = normalize_text(value)
    
    return re.sub(
        r"\s+",
        "",
        normalized
    )
    


def load_excel_data():


    workbook = None
    
    try:
    
        workbook = load_workbook(
            EXCEL_PATH,
            read_only=True,
            data_only=True
        )

    
        if EXCEL_SHEET_NAME not in workbook.sheetnames:
    
            test.fail(
                "Excel sheet not found: "
                + EXCEL_SHEET_NAME
                + " | Available sheets: "
                + ", ".join(workbook.sheetnames)
            )
    
            workbook.close()
    
            return None, None
    
        sheet = workbook[
            EXCEL_SHEET_NAME
        ]

  
        Estonian_data = {}
    
        for row in sheet.iter_rows(
            min_row=2,
            min_col=ENGLISH_COLUMN,
            max_col=Estonian_COLUMN
        ):
    
            english_value = row[0].value
    
            Estonian_value = row[
                Estonian_COLUMN - ENGLISH_COLUMN
            ].value
    
    
            if english_value is None:
                continue
    
            english_key = comparison_key(
                english_value
            )
    
            if english_key == "":
                continue
    
    
            if Estonian_value is None:
                Estonian_value = ""
    
            Estonian_value = str(
                Estonian_value
            ).strip()
    
    
            if english_key not in Estonian_data:
    
                Estonian_data[
                    english_key
                ] = Estonian_value
    
            else:
    
                existing_value = Estonian_data[
                    english_key
                ]
    
                if (
                    comparison_key(existing_value) == ""
                    and
                    comparison_key(Estonian_value) != ""
                ):
    
                    Estonian_data[
                        english_key
                    ] = Estonian_value
    
    
        if not Estonian_data:
    
            test.fail(
                "No English/Estonian data found in Excel. "
                "Please check the sheet and column configuration."
            )
    
            workbook.close()
    
            return None, None
    
        return workbook, Estonian_data
    
    except Exception as e:
    
        test.fail(
            "Unable to load/read Estonian localization Excel: "
            + str(e)
        )
    
        if workbook is not None:
    
            try:
                workbook.close()
    
            except Exception:
                pass
    
        return None, None
    
    
    
def get_Estonian_text(
    Estonian_data,
    english_text
    ):
    
    
    search_text = comparison_key(
        english_text
    )
    
    
    if search_text not in Estonian_data:
    
        test.fail(
            "English text not found in Excel Column B: "
            + str(english_text)
        )
    
        return None
    
    Estonian_value = Estonian_data[
        search_text
    ]
    
    
    if comparison_key(
        Estonian_value
    ) == "":
    
        test.fail(
            "Estonian translation is empty for English text: "
            + str(english_text)
        )
    
        return None
    
    return Estonian_value
    
    
    
def verify_text(
    Estonian_data,
    object_name,
    english_text,
    description
    ):
    
      
    snooze(2)
    
    
    expected_text = get_Estonian_text(
        Estonian_data,
        english_text
    )
    
    if expected_text is None:
    
        test.fail(
            description
            + " skipped because Excel lookup failed"
        )
    
        return False
    
    
    try:
    
        obj = waitForObjectExists(
            getattr(
                names,
                object_name
            )
        )
    
    except Exception as e:
    
        test.fail(
            description
            + " | Object not found: "
            + object_name
            + " | "
            + str(e)
        )
    
        return False
    
    try:
    
        actual_text = obj.text
    
    except Exception as e:
    
        test.fail(
            description
            + " | Unable to read object text | "
            + str(e)
        )
    
        return False
    
    
    actual_normalized = comparison_key(
        actual_text
    )
    
    expected_normalized = comparison_key(
        expected_text
    )
    
    # --------------------------------------------------------
    # Compare
    # --------------------------------------------------------
    
    test.compare(
        actual_normalized,
        expected_normalized,
        description
        + " - Estonian text verification"
    )
    
    if actual_normalized == expected_normalized:
        return True
    
    test.fail(
        # description
        # + " | Estonian text mismatch"
        # + " | Actual: "
        # + str(actual_text)
        # + " | Expected: "
        # + str(expected_text)
    )
    
    return False
    
    
    
def find_text_recursively(
    parent_object,
    expected_normalized
    ):
    
    
    try:
    
        if not parent_object.visible:
            return None
    
    except Exception:
        pass
    
    
    try:
    
        current_text = comparison_key(
            parent_object.text
        )
    
        if current_text == expected_normalized:
            return parent_object
    
    except Exception:
        pass
    
    
    try:
    
        children = object.children(
            parent_object
        )
    
    except Exception:
    
        children = []
    
    
    for child in children:
    
        found = find_text_recursively(
            child,
            expected_normalized
        )
    
        if found is not None:
            return found
    
    return None
    
    
    
def verify_treatment_name(
    Estonian_data,
    english_text,
    description
    ):
    
    
    
    
    expected_text = get_Estonian_text(
        Estonian_data,
        english_text
    )
    
    if expected_text is None:
    
        test.fail(
            description
            + " skipped because Excel lookup failed"
        )
    
        return False
    
    expected_normalized = comparison_key(
        expected_text
    )
    
    
    try:
    
        screen_switcher = waitForObjectExists(
            names.o_ScreenSwitcher
        )
    
    except Exception as e:
    
        test.fail(
            description
            + " | ScreenSwitcher not found | "
            + str(e)
        )
    
        return False
    
    
    found_object = find_text_recursively(
        screen_switcher,
        expected_normalized
    )
    
    if found_object is None:
    
        test.fail(
            "{} | Estonian treatment text not found "
            "on screen: [{}]".format(
                description,
                expected_text
            )
        )
    
        return False
    
    
    try:
    
        actual_text = comparison_key(
            found_object.text
        )
    
    except Exception:
    
        actual_text = ""
    
    
    test.compare(
        actual_text,
        expected_normalized,
        description
        + " - Estonian treatment verification"
    )
    
    return actual_text == expected_normalized
    
    
    
def verify_left_treatment_name(
    Estonian_data,
    english_text,
    description
    ):
    
    
    return verify_treatment_name(
        Estonian_data,
        english_text,
        description
    )
    
    
    
def verify_right_treatment_name(
    Estonian_data,
    english_text,
    description
    ):
    
    
    return verify_treatment_name(
        Estonian_data,
        english_text,
        description
    )
    
    
    
def main():
    
    
    workbook = None
    Estonian_data = None
    
    
    result = load_excel_data()
    
    if result is None:
    
        test.fail(
            "load_excel_data() returned None"
        )
    
        return
    
    
    try:
    
        workbook, Estonian_data = result
    
    except Exception as e:
    
        test.fail(
            "Unable to unpack Excel data: "
            + str(e)
        )
    
        return
    
    if workbook is None or Estonian_data is None:
    
        test.fail(
            "Estonian Excel data could not be loaded"
        )
    
        return
    
    try:
    
    
        for i in range(4):
    
            clickButton(
                waitForObject(
                    names.numpad_pushButton_8_NumpadButton
                )
            )
    
            snooze(1)
    
        clickButton(
            waitForObject(
                names.mainFrame_okButton_QPushButton
            )
        )
    
    
    
        clickButton(
            waitForObject(
                names.homeScreen_settingsButton_QPushButton
            )
        )
    
    
    
        clickTab(
            waitForObject(
                names.settingsScreen_tbSystemInfo_TabWidget
            ),
            "INFO"
        )
    
    
    
        clickButton(
            waitForObject(
                names.gbSystemInformation_pbLanguageSettingsUpdate_QPushButton
            )
        )
    
    
    
        clickButton(
            waitForObject(
                names.languagesFrame_pbEstonia_QPushButton
            )
        )
    
    
    
        clickButton(
            waitForObject(
                names.buttonsFrame_acceptButton_QPushButton
            )
        )
    
    
    
        clickButton(
            waitForObject(
                names.buttonsBar_homeButton_QPushButton
            )
        )
    
    
    
        clickButton(
            waitForObject(
                names.homeScreen_pbSoftTissueQuickStart_QToolButton
            )
        )
    
    
    
        mouseClick(
            waitForObjectItem(
                names.tabLeftMainPresets_leftPresetsView_PresetListView,
                "_1"
            ),
            238,
            58,
            Qt.NoModifier,
            Qt.LeftButton
        )
    
    
     
        mouseClick(
            waitForObjectItem(
                names.tabRightMainPresets_rightPresetsView_PresetListView,
                "_1"
            ),
            148,
            32,
            Qt.NoModifier,
            Qt.LeftButton
        )
    
    
    
        mouseClick(
            waitForObject(
                names.quickStartScreen_wdBottom_QWidget
            ),
            1167,
            9,
            Qt.NoModifier,
            Qt.LeftButton
        )
    
    
    
        sendEvent(
            "QMoveEvent",
            waitForObject(
                names.o_ScreenSwitcher
            ),
            0,
            -51,
            1286,
            61
        )
    
    
    
        clickButton(
            waitForObject(
                names.quickStartScreen_pbContinue_QPushButton
            )
        )
    
    
    
        clickButton(
            waitForObject(
                names.mainFrame_okButton_QPushButton_2
            )
        )
    
    
    
        clickButton(
            waitForObject(
                names.stateSwitch_readyButton_QPushButton
            )
        )
    
    
    
        verify_left_treatment_name(
            Estonian_data,
            "PROSTATE ENUCLEATION-(HIGH ENERGY)",
            "Left Treatment Name"
        )
    
    
    
        verify_text(
            Estonian_data,
            "averagePowerWidgetEmission_nameLabel_QLabel",
            "Average Power",
            "Left Average Power"
        )
    
    
    
        verify_text(
            Estonian_data,
            "pulseEnergyWidgetEmission_nameLabel_QLabel",
            "Pulse Energy",
            "Left Pulse Energy"
        )
    
    
    
        verify_text(
            Estonian_data,
            "frequencyWidgetEmission_nameLabel_QLabel",
            "Frequency",
            "Left Frequency"
        )
    
    
    
        verify_text(
            Estonian_data,
            "modeButtonEmission_nameLabel_QLabel",
            "Pulse Mode",
            "Left Pulse Mode"
        )
    
    
    
        verify_text(
            Estonian_data,
            "peakPowerButtonEmission_nameLabel_QLabel",
            "Peak Power",
            "Left Peak Power"
        )
    
    
    
        verify_text(
            Estonian_data,
            "totalEnergyFrame_nameLabel_QLabel",
            "Total Energy",
            "Left Total Energy"
        )
    
    
    
        verify_text(
            Estonian_data,
            "totalTimeFrame_nameLabel_QLabel",
            "Total Lasing Time",
            "Left Total Lasing Time"
        )
    
    
    
        verify_text(
            Estonian_data,
            "averagePowerWidgetEmission_unitsOfMeasureLabel_QLabel",
            "W",
            "Left Power Unit"
        )
    
    
    
        verify_text(
            Estonian_data,
            "pulseEnergyWidgetEmission_unitsOfMeasureLabel_QLabel",
            "J",
            "Left Energy Unit"
        )
    
    
    
        verify_text(
            Estonian_data,
            "frequencyWidgetEmission_unitsOfMeasureLabel_QLabel",
            "Hz",
            "Left Frequency Unit"
        )
    
    
    
        verify_text(
            Estonian_data,
            "fiberFrame_nameLabel_QLabel",
            "Fiber",
            "Left Fiber"
        )
    
    
    
        verify_text(
            Estonian_data,
            "fiberUsesFrame_nameLabel_QLabel",
            "Fiber Uses Remaining",
            "Left Fiber Uses Remaining"
        )
    
    
    
        verify_text(
            Estonian_data,
            "bottomBar_aimingBeamButton_AimingBeamModeButton",
            "Aiming Beam",
            "Left Aiming Beam"
        )
    
    
    
        verify_text(
            Estonian_data,
            "stateSwitch_standbyButton_QPushButton",
            "Standby",
            "Left Standby"
        )
    
    
    
        verify_text(
            Estonian_data,
            "stateSwitch_readyButton_QPushButton",
            "Ready",
            "Left Ready"
        )
    
    
    
        verify_right_treatment_name(
            Estonian_data,
            "PROSTATE ENUCLEATION-(HIGH ENERGY)",
            "Right Treatment Name"
        )
    
    
    
        verify_text(
            Estonian_data,
            "averagePowerWidgetEmission_nameLabel_QLabel_2",
            "Average Power",
            "Right Average Power"
        )
    
    
    
        verify_text(
            Estonian_data,
            "pulseEnergyWidgetEmission_nameLabel_QLabel_2",
            "Pulse Energy",
            "Right Pulse Energy"
        )
    
    
    
        verify_text(
            Estonian_data,
            "frequencyWidgetEmission_nameLabel_QLabel_2",
            "Frequency",
            "Right Frequency"
        )
    
    
    
        verify_text(
            Estonian_data,
            "modeButtonEmission_nameLabel_QLabel_2",
            "Pulse Mode",
            "Right Pulse Mode"
        )
    
    
    
        verify_text(
            Estonian_data,
            "peakPowerButtonEmission_nameLabel_QLabel_2",
            "Peak Power",
            "Right Peak Power"
        )
    
  
    
        verify_text(
            Estonian_data,
            "averagePowerWidgetEmission_unitsOfMeasureLabel_QLabel_2",
            "W",
            "Right Power Unit"
        )
    
    
    
    
        verify_text(
            Estonian_data,
            "pulseEnergyWidgetEmission_unitsOfMeasureLabel_QLabel_2",
            "J",
            "Right Energy Unit"
        )
    
   
    
        verify_text(
            Estonian_data,
            "frequencyWidgetEmission_unitsOfMeasureLabel_QLabel_2",
            "Hz",
            "Right Frequency Unit"
        )
    
    
    
        clickButton(
            waitForObject(
                names.stateSwitch_readyButton_QPushButton
            )
        )
    
    
       
        
    
        clickButton(
            waitForObject(
                names.stateSwitch_readyButton_QPushButton
            )
        )
    
    except Exception as e:
    
        test.fail(
            "Estonian Ready Screen test failed: "
            + str(e)
        )
    
    finally:
    
      
        if workbook is not None:
    
            try:
                workbook.close()
    
            except Exception:
                pass
    
