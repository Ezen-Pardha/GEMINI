# -*- coding: utf-8 -*-

import names
import re
import unicodedata
from openpyxl import load_workbook
import os 

TESTCASE_DIR = os.path.dirname(os.path.abspath(__file__))
 
EXCEL_PATH = os.path.abspath(
    os.path.join( TESTCASE_DIR,
        "..",
        "shared",
        "testdata",
        "Gemini Split strings.xlsx"
    )
)


SHEET_NAME = "Soft tissue"
ENGLISH_COLUMN = 2
FINNISH_COLUMN = 9


def normalize(text):

    if text is None:
        return ""

    text = str(text)

    text = re.sub(
        r"<br\s*/?>",
        "",
        text,
        flags=re.IGNORECASE
    )

    text = re.sub(r"<[^>]+>", "", text)

    text = unicodedata.normalize("NFKC", text)

    text = (
        text.replace("\n", "")
        .replace("\r", "")
        .replace("\t", "")
        .replace("\xa0", "")
        .replace("\u2007", "")
        .replace("\u202f", "")
        .replace("™", "")
        .replace("®", "")
        .replace("©", "")
    )

    text = re.sub(r"\s+", "", text)

    return text.casefold()


def load_excel():

    try:
        workbook = load_workbook(
            EXCEL_PATH,
            data_only=True,
            read_only=True
        )

    except Exception as e:
        test.fail(
            "Unable to load Finnish localization Excel: "
            + str(e)
        )
        return None, None

    if SHEET_NAME not in workbook.sheetnames:
        test.fail(
            "Excel sheet not found: "
            + SHEET_NAME
        )
        workbook.close()
        return None, None

    return workbook, workbook[SHEET_NAME]


def get_finnish_values(sheet, english_text):

    expected_english = normalize(english_text)
    finnish_values = []

    for row in sheet.iter_rows(
        min_row=2,
        min_col=ENGLISH_COLUMN,
        max_col=FINNISH_COLUMN,
        values_only=True
    ):

        english_value = row[0]

        if english_value is None:
            continue

        if normalize(english_value) != expected_english:
            continue

        finnish_value = row[
            FINNISH_COLUMN - ENGLISH_COLUMN
        ]

        if finnish_value is None:
            continue

        finnish_value = str(finnish_value).strip()

        if finnish_value:
            finnish_values.append(finnish_value)

    return finnish_values


def get_actual_text(obj, description):

    actual_value = None

    for property_name in (
        "text",
        "plainText",
        "windowTitle",
        "accessibleName"
    ):

        try:
            actual_value = getattr(obj, property_name)

            if callable(actual_value):
                actual_value = actual_value()

            if actual_value is not None and str(actual_value).strip():
                break

        except Exception:
            actual_value = None

    if actual_value is None or not str(actual_value).strip():

        try:
            actual_value = obj.property("text")
        except Exception:
            actual_value = None

    actual_text = normalize(actual_value)

    if not actual_text:
        test.fail(
            description
            + " | Application text is empty or unavailable"
        )
        return None

    return actual_text


def verify_text(sheet, object_name, english_text, description):

    snooze(2)

    expected_values = get_finnish_values(
        sheet,
        english_text
    )

    if not expected_values:
        test.fail(
            description
            + " | English string not found in Excel Column B: "
            + str(english_text)
        )
        return False

    expected_normalized = [
        normalize(value)
        for value in expected_values
    ]

    try:
        obj = waitForObjectExists(object_name)

    except Exception as e:
        test.fail(
            description
            + " | Unable to find object: "
            + str(e)
        )
        return False

    actual_text = get_actual_text(
        obj,
        description
    )

    if actual_text is None:
        return False

    for expected_text in expected_normalized:

        if actual_text == expected_text:
            test.compare(
                actual_text,
                expected_text,
                description + " - Finnish text verification"
            )
            return True

    test.fail(
        description
        + " | Expected Finnish value from Excel Column I: "
        + " / ".join(expected_values)
        + " | Actual: "
        + actual_text
    )

    return False


def get_central_widget():

    try:
        central_widget = findObject(
            {
                "name": "centralWidget",
                "type": "QFrame",
                "window": ":o_ScreenSwitcher"
            }
        )

        for _ in range(20):

            try:
                if central_widget.visible:
                    return central_widget
            except Exception:
                pass

            snooze(0.5)

    except Exception:
        pass

    return None


def main():

    workbook, sheet = load_excel()

    if workbook is None or sheet is None:
        return

    try:

        clickButton(
            waitForObject(
                names.numpad_pushButton_8_NumpadButton
            )
        )

        doubleClick(
            waitForObject(
                names.numpad_pushButton_8_NumpadButton
            ),
            94,
            77,
            Qt.NoModifier,
            Qt.LeftButton
        )

        doubleClick(
            waitForObject(
                names.numpad_pushButton_8_NumpadButton
            ),
            94,
            77,
            Qt.NoModifier,
            Qt.LeftButton
        )

        clickButton(
            waitForObject(
                names.mainFrame_okButton_QPushButton
            )
        )

        clickButton(
            waitForObject(
                names.homeScreen_settingsButton_QPushButton
            )
        )

        clickTab(
            waitForObject(
                names.settingsScreen_tbSystemInfo_TabWidget
            ),
            "INFO"
        )

        clickButton(
            waitForObject(
                names.gbSystemInformation_pbLanguageSettingsUpdate_QPushButton
            )
        )

        clickButton(
            waitForObject(
                names.languagesFrame_pbFinland_QPushButton
            )
        )

        clickButton(
            waitForObject(
                names.buttonsFrame_acceptButton_QPushButton
            )
        )

        clickButton(
            waitForObject(
                names.buttonsBar_homeButton_QPushButton
            )
        )

        clickButton(
            waitForObject(
                names.homeScreen_pbSoftTissueQuickStart_QToolButton
            )
        )

        mouseClick(
            waitForObjectItem(
                names.tabLeftMainPresets_leftPresetsView_PresetListView,
                "_1"
            ),
            377,
            65,
            Qt.NoModifier,
            Qt.LeftButton
        )

        mouseClick(
            waitForObjectItem(
                names.tabRightMainPresets_rightPresetsView_PresetListView,
                "_1"
            ),
            150,
            37,
            Qt.NoModifier,
            Qt.LeftButton
        )

        sendEvent(
            "QMoveEvent",
            waitForObject(names.o_ScreenSwitcher),
            0,
            -50,
            1288,
            292
        )

        clickButton(
            waitForObject(
                names.quickStartScreen_pbContinue_QPushButton
            )
        )

        clickButton(
            waitForObject(
                names.mainFrame_okButton_QPushButton_2
            )
        )

        mouseClick(
            waitForObject(
                names.modeButton_nameLabel_QLabel
            ),
            115,
            15,
            Qt.NoModifier,
            Qt.LeftButton
        )

        verify_text(
            sheet,
            names.centralWidget_TitleLabel_QLabel_2,
            "EXIT QUICK START MODE?",
            "Quick Start Exit Title"
        )

        verify_text(
            sheet,
            names.centralWidget_MessageLabel_QLabel_2,
            "Please confirm you want to exit quick state mode.",
            "Quick Start Exit Message"
        )

        verify_text(
            sheet,
            names.centralWidget_CancelButton_QPushButton_3,
            "CANCEL",
            "Quick Start Cancel Button"
        )

        verify_text(
            sheet,
            names.centralWidget_DisableButton_QPushButton_2,
            "CONFIRM",
            "Quick Start Confirm Button"
        )

        clickButton(
            waitForObject(
                names.centralWidget_DisableButton_QPushButton_2
            )
        )

        clickButton(
            waitForObject(
                names.buttonsBar_homeButton_QPushButton_2
            )
        )

        clickButton(
            waitForObject(
                names.homeScreen_pbSoftTissueAssistant_QToolButton
            )
        )

        clickButton(
            waitForObject(
                names.gbSoftTissueProcedures_pbIncision_QPushButton
            )
        )

        clickButton(
            waitForObject(
                names.softTissueAssistantModeScreen_pbContinue_QPushButton
            )
        )

        mouseDrag(
            waitForObject(
                names.indexPicker_valuePicker_IntervalSlider
            ),
            195,
            15,
            211,
            29,
            1,
            Qt.LeftButton
        )

        verify_text(
            sheet,
            names.centralWidget_TitleLabel_QLabel_2,
            "EXIT GUIDED MODE?",
            "Guided Mode Exit Title"
        )

        verify_text(
            sheet,
            names.centralWidget_MessageLabel_QLabel_2,
            "Please confirm you want to exit guided mode.",
            "Guided Mode Exit Message"
        )

        verify_text(
            sheet,
            names.centralWidget_CancelButton_QPushButton_3,
            "CANCEL",
            "Guided Mode Cancel Button"
        )

        test.vp("Exit Guided Mode")

        verify_text(
            sheet,
            names.centralWidget_DisableButton_QPushButton_2,
            "CONFIRM",
            "Guided Mode Confirm Button"
        )

        clickButton(
            waitForObject(
                names.centralWidget_DisableButton_QPushButton_2
            )
        )

    finally:
        workbook.close()
