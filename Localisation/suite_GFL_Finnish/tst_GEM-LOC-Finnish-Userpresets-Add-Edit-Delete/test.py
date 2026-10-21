# -*- coding: utf-8 -*-

import names
import re
import unicodedata
from openpyxl import load_workbook
import os


TESTCASE_DIR = os.path.dirname(os.path.abspath(__file__))
 
EXCEL_PATH = os.path.abspath(
    os.path.join( TESTCASE_DIR,
        "..",
        "shared",
        "testdata",
        "Gemini Split strings.xlsx"
    )
)




def normalize_text(value):
    if value is None:
        return ""

    text = str(value)
    text = re.sub(r"<[^>]*>", " ", text)
    text = text.replace("\u00A0", " ")
    text = text.replace("\u2007", " ")
    text = text.replace("\u202F", " ")
    text = re.sub(r"[\r\n\t]+", " ", text)
    text = text.replace("™", "").replace("®", "").replace("©", "")

    try:
        text = unicodedata.normalize("NFKC", text)
    except Exception:
        pass

    text = re.sub(r"\s+", " ", text)
    text = re.sub(r"\s+([.,;:!?])", r"\1", text)

    return text.strip().casefold()


def comparison_key(value):
    return re.sub(r"\s+", "", normalize_text(value))


def get_excel_value(sheet, row, column):
    value = sheet.cell(row=row, column=column).value
    return "" if value is None else str(value)


def find_column_by_header(sheet, possible_headers, fallback_column):
    expected_headers = {
        normalize_text(header) for header in possible_headers
    }

    for row in sheet.iter_rows(
        min_row=1,
        max_row=min(sheet.max_row, 10)
    ):
        for cell in row:
            if cell.value is None:
                continue

            if normalize_text(cell.value) in expected_headers:
                return cell.column

    return fallback_column


def find_finnish_column(sheet):
    return find_column_by_header(
        sheet,
        [
            "Finnish",
            "Suomi",
            "Finnish (Finland)",
            "Suomi (Suomi)",
            "Finland"
        ],
        8
    )


def find_english_column(sheet):
    return find_column_by_header(
        sheet,
        [
            "English",
            "English (United States)",
            "English (United Kingdom)",
            "English (US)",
            "en-US",
            "en-GB"
        ],
        2
    )


def find_excel_row(sheet, english_column, english_text):
    """
    Find the Excel row dynamically using the English source text.

    No fixed Excel row numbers are required.
    Matching ignores case, HTML tags and whitespace differences.
    """

    if isinstance(english_text, str):
        candidates = [english_text]
    else:
        candidates = list(english_text)

    candidate_keys = [
        comparison_key(candidate)
        for candidate in candidates
        if candidate
    ]

    # First, try an exact normalized match.
    for row_number in range(1, sheet.max_row + 1):
        english_value = get_excel_value(
            sheet,
            row_number,
            english_column
        )

        if not english_value.strip():
            continue

        english_key = comparison_key(english_value)

        if english_key in candidate_keys:
            return row_number

    # If the English source contains additional text, try a
    # normalized containment match.
    for row_number in range(1, sheet.max_row + 1):
        english_value = get_excel_value(
            sheet,
            row_number,
            english_column
        )

        if not english_value.strip():
            continue

        english_key = comparison_key(english_value)

        for candidate_key in candidate_keys:
            if (
                candidate_key
                and (
                    candidate_key in english_key
                    or english_key in candidate_key
                )
            ):
                return row_number

    test.fail(
        "English source string not found in Excel: "
        + " / ".join(candidates)
    )

    return None


def verify_excel_string(
    object_name,
    sheet,
    english_column,
    finnish_column,
    english_text,
    description
):
    snooze(2)

    row = find_excel_row(
        sheet,
        english_column,
        english_text
    )

    if row is None:
        return False

    try:
        obj = waitForObjectExists(object_name)
        actual = get_object_text(obj)
    except Exception as error:
        test.fail(
            description
            + " | Object not found: "
            + str(error)
        )
        return False

    expected = get_excel_value(
        sheet,
        row,
        finnish_column
    )

    if not actual.strip():
        test.fail(description + " | UI text is empty")
        return False

    if not expected.strip():
        test.fail(
            description
            + " | Finnish translation is empty for English: "
            + str(english_text)
        )
        return False

    test.compare(
        comparison_key(actual),
        comparison_key(expected),
        description
    )

    return True


def get_object_text(obj):
    for attribute in (
        "text",
        "title",
        "windowTitle",
        "accessibleName"
    ):
        try:
            value = getattr(obj, attribute)

            if value is not None and str(value).strip():
                return str(value)

        except Exception:
            pass

    return ""


def main():

    workbook = load_workbook(
        EXCEL_PATH,
        read_only=True,
        data_only=True
    )

    try:
        sheet = workbook["Treatment Modes"]

        english_column = find_english_column(sheet)
        finnish_column = find_finnish_column(sheet)

        # Login
        for _ in range(4):
            clickButton(
                waitForObject(
                    names.numpad_pushButton_8_NumpadButton
                )
            )

        clickButton(
            waitForObject(
                names.mainFrame_okButton_QPushButton
            )
        )

        # Navigate to language settings
        clickButton(
            waitForObject(
                names.homeScreen_settingsButton_QPushButton
            )
        )

        clickTab(
            waitForObject(
                names.settingsScreen_tbSystemInfo_TabWidget
            ),
            "INFO"
        )

        clickButton(
            waitForObject(
                names.gbSystemInformation_pbLanguageSettingsUpdate_QPushButton
            )
        )

        clickButton(
            waitForObject(
                names.languagesFrame_pbFinland_QPushButton
            )
        )

        clickButton(
            waitForObject(
                names.buttonsFrame_acceptButton_QPushButton
            )
        )

        snooze(3)

        tab_widget = waitForObject(
            names.settingsScreen_tbSystemInfo_TabWidget
        )

        tab_widget.setCurrentIndex(2)
        snooze(3)

        tab_widget.setCurrentIndex(2)
        snooze(3)

        sendEvent(
            "QMoveEvent",
            waitForObject(names.o_ScreenSwitcher),
            0,
            -66,
            1285,
            233
        )

        # Add preset
        clickButton(
            waitForObject(
                names.settingsScreen_addButton_QPushButton
            )
        )

        verify_excel_string(
            names.centralWidget_lbOperation_QLabel,
            sheet,
            english_column,
            finnish_column,
            ["ADD PRESET", "ADD USER PRESET", "ADD"],
            "Add Preset Operation"
        )

        verify_excel_string(
            names.centralWidget_pbName_QPushButton,
            sheet,
            english_column,
            finnish_column,
            ["Preset Name", "Name"],
            "Preset Name"
        )

        verify_excel_string(
            names.averagePowerWidget_nameLabel_QLabel,
            sheet,
            english_column,
            finnish_column,
            "Average Power",
            "Average Power"
        )

        verify_excel_string(
            names.averagePowerWidget_unitsOfMeasureLabel_QLabel,
            sheet,
            english_column,
            finnish_column,
            ["W", "Watts"],
            "Average Power Unit"
        )

        verify_excel_string(
            names.pulseEnergyWidget_nameLabel_QLabel,
            sheet,
            english_column,
            finnish_column,
            "Pulse Energy",
            "Pulse Energy"
        )

        mouseClick(
            waitForObject(
                names.centralWidget_frame_QFrame
            ),
            271,
            13,
            Qt.NoModifier,
            Qt.LeftButton
        )

        verify_excel_string(
            names.pulseEnergyWidget_unitsOfMeasureLabel_QLabel,
            sheet,
            english_column,
            finnish_column,
            ["J", "Joules"],
            "Pulse Energy Unit"
        )

        verify_excel_string(
            names.frequencyWidget_nameLabel_QLabel,
            sheet,
            english_column,
            finnish_column,
            "Frequency",
            "Frequency"
        )

        verify_excel_string(
            names.frequencyWidget_unitsOfMeasureLabel_QLabel,
            sheet,
            english_column,
            finnish_column,
            ["Hz", "Hertz"],
            "Frequency Unit"
        )

        verify_excel_string(
            names.modeButton_nameLabel_QLabel_3,
            sheet,
            english_column,
            finnish_column,
            ["Regular Pulse", "Regular pulse"],
            "Regular Pulse"
        )

        verify_excel_string(
            names.peakPowerButton_nameLabel_QLabel_2,
            sheet,
            english_column,
            finnish_column,
            "Peak Power",
            "Peak Power"
        )

        verify_excel_string(
            names.pulseModeWidget_saveButton_QPushButton,
            sheet,
            english_column,
            finnish_column,
            ["ADD", "ADD PRESET", "SAVE"],
            "Add Button"
        )

        mouseClick(
            waitForObject(
                names.pulseModeWidget_buttonsWidget_QWidget
            ),
            508,
            46,
            Qt.NoModifier,
            Qt.LeftButton
        )

        mouseClick(
            waitForObject(
                names.pulseModeWidget_buttonsWidget_QWidget
            ),
            411,
            1,
            Qt.NoModifier,
            Qt.LeftButton
        )

        clickButton(
            waitForObject(
                names.centralWidget_pbName_QPushButton
            )
        )

        type(
            waitForObject(
                names.keyboardPopup_wdEdit_WdLineEdit
            ),
            "preset one"
        )

        mouseClick(
            waitForObject(
                names.keyboardPopup_lbText_QLabel
            ),
            77,
            26,
            Qt.NoModifier,
            Qt.LeftButton
        )

        clickButton(
            waitForObject(
                names.pulseModeWidget_saveButton_QPushButton
            )
        )

        mouseClick(
            waitForObject(
                names.tabSettingsPresetsUserStonePresets_presetListView_PresetListView
            ),
            546,
            72,
            Qt.NoModifier,
            Qt.LeftButton
        )

        # Edit preset
        clickButton(
            waitForObject(
                names.settingsScreen_editButton_QPushButton
            )
        )

        mouseClick(
            waitForObject(
                names.tabSettingsPresetsUserStonePresets_presetListView_PresetListView
            ),
            494,
            66,
            Qt.NoModifier,
            Qt.LeftButton
        )

        verify_excel_string(
            names.centralWidget_lbOperation_QLabel,
            sheet,
            english_column,
            finnish_column,
            ["EDIT PRESET", "EDIT USER PRESET", "EDIT"],
            "Edit Preset Operation"
        )

        snooze(3)

        actual_edit_name = str(
            waitForObjectExists(
                names.centralWidget_pbName_QPushButton
            ).text
        )

        test.compare(
            comparison_key(actual_edit_name),
            comparison_key("preset one"),
            "Edit Preset Name"
        )

        verify_excel_string(
            names.averagePowerWidget_nameLabel_QLabel,
            sheet,
            english_column,
            finnish_column,
            "Average Power",
            "Edit - Average Power"
        )

        verify_excel_string(
            names.averagePowerWidget_unitsOfMeasureLabel_QLabel,
            sheet,
            english_column,
            finnish_column,
            ["W", "Watts"],
            "Edit - Average Power Unit"
        )

        verify_excel_string(
            names.pulseEnergyWidget_nameLabel_QLabel,
            sheet,
            english_column,
            finnish_column,
            "Pulse Energy",
            "Edit - Pulse Energy"
        )

        verify_excel_string(
            names.pulseEnergyWidget_unitsOfMeasureLabel_QLabel,
            sheet,
            english_column,
            finnish_column,
            ["J", "Joules"],
            "Edit - Pulse Energy Unit"
        )

        verify_excel_string(
            names.frequencyWidget_nameLabel_QLabel,
            sheet,
            english_column,
            finnish_column,
            "Frequency",
            "Edit - Frequency"
        )

        verify_excel_string(
            names.frequencyWidget_unitsOfMeasureLabel_QLabel,
            sheet,
            english_column,
            finnish_column,
            ["Hz", "Hertz"],
            "Edit - Frequency Unit"
        )

        verify_excel_string(
            names.modeButton_nameLabel_QLabel_3,
            sheet,
            english_column,
            finnish_column,
            ["Regular Pulse", "Regular pulse"],
            "Edit - Regular Pulse"
        )

        verify_excel_string(
            names.peakPowerButton_nameLabel_QLabel_2,
            sheet,
            english_column,
            finnish_column,
            "Peak Power",
            "Edit - Peak Power"
        )

        verify_excel_string(
            names.pulseModeWidget_saveButton_QPushButton,
            sheet,
            english_column,
            finnish_column,
            ["UPDATE", "UPDATE PRESET", "SAVE"],
            "Update Button"
        )

        mouseClick(
            waitForObject(
                names.pulseModeWidget_buttonsWidget_QWidget
            ),
            531,
            22,
            Qt.NoModifier,
            Qt.LeftButton
        )

        mouseClick(
            waitForObject(
                names.pulseModeWidget_buttonsWidget_QWidget
            ),
            509,
            30,
            Qt.NoModifier,
            Qt.LeftButton
        )

        clickButton(
            waitForObject(
                names.pulseModeWidget_saveButton_QPushButton
            )
        )

        # Delete preset
        clickButton(
            waitForObject(
                names.settingsScreen_deletePresetButton_QPushButton
            )
        )

        mouseClick(
            waitForObject(
                names.tabSettingsPresetsUserStonePresets_presetListView_PresetListView
            ),
            553,
            80,
            Qt.NoModifier,
            Qt.LeftButton
        )

        verify_excel_string(
            names.centralWidget_lbDialogTitle_QLabel,
            sheet,
            english_column,
            finnish_column,
            ["DELETE PRESET", "DELETE USER PRESET", "DELETE"],
            "Delete Dialog Title"
        )

        snooze(3)

        actual_delete_name = str(
            waitForObjectExists(
                names.centralWidget_questionLabel_QLabel
            ).text
        )

        test.compare(
            comparison_key(actual_delete_name),
            comparison_key("preset one"),
            "Delete Preset Name"
        )

        snooze(3)

        actual_info = str(
            waitForObjectExists(
                names.infoFrame_extInfoLabel_QLabel
            ).text
        )

        expected_info = "6 J     3 Hz     18 W"

        test.compare(
            comparison_key(actual_info),
            comparison_key(expected_info),
            "Preset Information"
        )

        verify_excel_string(
            names.buttonsFrame_rejectButton_QPushButton,
            sheet,
            english_column,
            finnish_column,
            ["CANCEL", "CLOSE"],
            "Cancel Button"
        )

        verify_excel_string(
            names.centralWidget_acceptButton_QPushButton,
            sheet,
            english_column,
            finnish_column,
            ["DELETE", "DELETE PRESET"],
            "Delete Button"
        )

       

    finally:
        clickButton(
            waitForObject(
                names.centralWidget_acceptButton_QPushButton
            )
        )
        workbook.close()