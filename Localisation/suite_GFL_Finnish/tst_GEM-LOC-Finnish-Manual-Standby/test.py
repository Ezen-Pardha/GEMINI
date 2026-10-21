# -*- coding: utf-8 -*-

import names
import re
import unicodedata
from openpyxl import load_workbook
import os 

TESTCASE_DIR = os.path.dirname(os.path.abspath(__file__))
 
EXCEL_PATH = os.path.abspath(
    os.path.join( TESTCASE_DIR,
        "..",
        "shared",
        "testdata",
        "Gemini Split strings.xlsx"
    )
)



EXCEL_SHEET_NAME = "Treatment Modes"
ENGLISH_COLUMN = 2
FINNISH_COLUMN = 9


def normalize_text(value):

    if value is None:
        return ""

    text = str(value)

    text = re.sub(
        r"<br\s*/?>",
        "",
        text,
        flags=re.IGNORECASE
    )

    text = re.sub(r"<[^>]*>", "", text)

    text = unicodedata.normalize("NFKC", text)

    text = (
        text.replace("\n", "")
        .replace("\r", "")
        .replace("\t", "")
        .replace("\xa0", "")
        .replace("\u2007", "")
        .replace("\u202f", "")
        .replace("™", "")
        .replace("®", "")
        .replace("©", "")
    )

    text = re.sub(r"\s+", "", text)

    return text.casefold()


def load_excel_data():

    try:
        workbook = load_workbook(
            EXCEL_PATH,
            read_only=True,
            data_only=True
        )
    except Exception as e:
        test.fail(
            "Unable to open Excel file: "
            + EXCEL_PATH
            + " | Error: "
            + str(e)
        )
        return None, {}

    if EXCEL_SHEET_NAME not in workbook.sheetnames:
        test.fail(
            "Excel sheet not found: "
            + EXCEL_SHEET_NAME
        )
        workbook.close()
        return None, {}

    sheet = workbook[EXCEL_SHEET_NAME]
    finnish_data = {}

    for row in sheet.iter_rows(
        min_col=ENGLISH_COLUMN,
        max_col=FINNISH_COLUMN,
        values_only=True
    ):

        english_value = row[0]
        finnish_value = row[FINNISH_COLUMN - ENGLISH_COLUMN]

        if english_value is None:
            continue

        english_key = normalize_text(english_value)

        if not english_key:
            continue

        if finnish_value is None:
            continue

        finnish_value = str(finnish_value).strip()

        if not finnish_value:
            continue

        finnish_data[english_key] = finnish_value

    return workbook, finnish_data


def get_finnish_text(finnish_data, english_text):

    key = normalize_text(english_text)

    if key not in finnish_data:
        test.fail(
            "English text not found in Excel Column B: "
            + str(english_text)
        )
        return None

    return finnish_data[key]


def get_actual_text(obj, description):

    actual_value = None

    for property_name in (
        "text",
        "plainText",
        "windowTitle",
        "accessibleName"
    ):

        try:
            actual_value = getattr(obj, property_name)

            if callable(actual_value):
                actual_value = actual_value()

            if actual_value is not None and str(actual_value) != "":
                break

        except Exception:
            actual_value = None

    if actual_value is None or str(actual_value) == "":
        try:
            actual_value = obj.property("text")
        except Exception:
            actual_value = None

    actual_text = normalize_text(actual_value)

    if not actual_text:
        test.fail(
            "Application text is empty or unavailable for: "
            + description
        )
        return None

    return actual_text


def verify_text(
    finnish_data,
    object_name,
    english_text,
    description
):

    snooze(2)

    expected_text = get_finnish_text(
        finnish_data,
        english_text
    )

    if expected_text is None:
        return False

    expected_normalized = normalize_text(expected_text)

    if not expected_normalized:
        test.fail(
            "Expected Finnish translation is empty for: "
            + description
        )
        return False

    try:
        obj = waitForObjectExists(
            getattr(names, object_name)
        )
    except Exception as e:
        test.fail(
            "Unable to find application object: "
            + object_name
            + " | Error: "
            + str(e)
        )
        return False

    actual_text = get_actual_text(
        obj,
        description
    )

    if actual_text is None:
        return False

    result = actual_text == expected_normalized

    test.compare(
        actual_text,
        expected_normalized,
        description + " - Finnish text verification"
    )

    if not result:
        test.fail(
            description
            + " mismatch | Actual: "
            + actual_text
            + " | Expected: "
            + expected_normalized
        )

    return result


def verify_manual_tab(
    finnish_data,
    object_name,
    english_text,
    description
):

    snooze(2)

    expected_text = get_finnish_text(
        finnish_data,
        english_text
    )

    if expected_text is None:
        return False

    expected_normalized = normalize_text(expected_text)

    try:
        manual_tab_object = getattr(
            names,
            object_name
        ).copy()

        if "text" in manual_tab_object:
            del manual_tab_object["text"]

        obj = waitForObjectExists(manual_tab_object)

    except Exception as e:
        test.fail(
            "Unable to find Manual tab: "
            + description
            + " | Error: "
            + str(e)
        )
        return False

    actual_text = get_actual_text(
        obj,
        description
    )

    if actual_text is None:
        return False

    result = actual_text == expected_normalized

    test.compare(
        actual_text,
        expected_normalized,
        description + " - Finnish text verification"
    )

    if not result:
        test.fail(
            description
            + " mismatch | Actual: "
            + actual_text
            + " | Expected: "
            + expected_normalized
        )

    return result


def main():

    workbook, finnish_data = load_excel_data()

    if workbook is None:
        return

    try:

        doubleClick(
            waitForObject(
                names.numpad_pushButton_8_NumpadButton,
                100461
            ),
            96,
            56,
            Qt.NoModifier,
            Qt.LeftButton
        )

        snooze(1)

        doubleClick(
            waitForObject(
                names.numpad_pushButton_8_NumpadButton
            ),
            96,
            56,
            Qt.NoModifier,
            Qt.LeftButton
        )

        snooze(1)

        clickButton(
            waitForObject(
                names.mainFrame_okButton_QPushButton
            )
        )

        snooze(2)

        clickButton(
            waitForObject(
                names.homeScreen_settingsButton_QPushButton
            )
        )

        snooze(2)

        clickTab(
            waitForObject(
                names.settingsScreen_tbSystemInfo_TabWidget
            ),
            "INFO"
        )

        snooze(2)

        clickButton(
            waitForObject(
                names.gbSystemInformation_pbLanguageSettingsUpdate_QPushButton
            )
        )

        snooze(2)

        clickButton(
            waitForObject(
                names.languagesFrame_pbFinland_QPushButton
            )
        )

        snooze(1)

        clickButton(
            waitForObject(
                names.buttonsFrame_acceptButton_QPushButton
            )
        )

        snooze(3)

        clickButton(
            waitForObject(
                names.buttonsBar_homeButton_QPushButton
            )
        )

        snooze(2)

        clickButton(
            waitForObject(
                names.gbExpert_expertButton_QPushButton
            )
        )

        snooze(3)

        verify_manual_tab(
            finnish_data,
            "mANUEL_TabItem",
            "MANUAL",
            "Left Manual Tab"
        )

        verify_text(
            finnish_data,
            "averagePowerWidget_nameLabel_QLabel_2",
            "AVERAGE POWER",
            "Left Average Power"
        )

        verify_text(
            finnish_data,
            "averagePowerWidget_unitsOfMeasureLabel_QLabel_2",
            "W",
            "Left Power Unit"
        )

        verify_text(
            finnish_data,
            "pulseEnergyWidget_nameLabel_QLabel_2",
            "PULSE ENERGY",
            "Left Pulse Energy"
        )

        verify_text(
            finnish_data,
            "pulseEnergyWidget_unitsOfMeasureLabel_QLabel_2",
            "J",
            "Left Energy Unit"
        )

        verify_text(
            finnish_data,
            "frequencyWidget_nameLabel_QLabel_2",
            "FREQUENCY",
            "Left Frequency"
        )

        verify_text(
            finnish_data,
            "frequencyWidget_unitsOfMeasureLabel_QLabel_2",
            "Hz",
            "Left Frequency Unit"
        )

        verify_text(
            finnish_data,
            "modeButton_nameLabel_QLabel",
            "REGULAR PULSE",
            "Left Pulse Mode"
        )

        verify_text(
            finnish_data,
            "peakPowerButton_nameLabel_QLabel",
            "PEAK POWER",
            "Left Peak Power"
        )

        verify_text(
            finnish_data,
            "pulseModeWidget_saveButton_QPushButton_2",
            "SAVE",
            "Left Save Button"
        )

        verify_text(
            finnish_data,
            "totalEnergyFrame_nameLabel_QLabel",
            "TOTAL ENERGY",
            "Left Total Energy"
        )

        verify_text(
            finnish_data,
            "totalTimeFrame_nameLabel_QLabel",
            "TOTAL LASING TIME",
            "Left Total Lasing Time"
        )

        verify_manual_tab(
            finnish_data,
            "mANUEL_TabItem_2",
            "MANUAL",
            "Right Manual Tab"
        )

        verify_text(
            finnish_data,
            "averagePowerWidget_nameLabel_QLabel_3",
            "AVERAGE POWER",
            "Right Average Power"
        )

        verify_text(
            finnish_data,
            "averagePowerWidget_unitsOfMeasureLabel_QLabel_3",
            "W",
            "Right Power Unit"
        )

        verify_text(
            finnish_data,
            "pulseEnergyWidget_nameLabel_QLabel_3",
            "PULSE ENERGY",
            "Right Pulse Energy"
        )

        verify_text(
            finnish_data,
            "pulseEnergyWidget_unitsOfMeasureLabel_QLabel_3",
            "J",
            "Right Energy Unit"
        )

        verify_text(
            finnish_data,
            "frequencyWidget_nameLabel_QLabel_3",
            "FREQUENCY",
            "Right Frequency"
        )

        verify_text(
            finnish_data,
            "pulseModeWidget_saveButton_QPushButton_3",
            "SAVE",
            "Right Save Button"
        )

        verify_text(
            finnish_data,
            "peakPowerButton_nameLabel_QLabel_3",
            "PEAK POWER",
            "Right Peak Power"
        )

        verify_text(
            finnish_data,
            "modeButton_nameLabel_QLabel_2",
            "REGULAR PULSE",
            "Right Pulse Mode"
        )

        verify_text(
            finnish_data,
            "stateSwitch_readyButton_QPushButton",
            "READY",
            "Ready Button"
        )

        verify_text(
            finnish_data,
            "stateSwitch_standbyButton_QPushButton",
            "STANDBY",
            "Standby Button"
        )

        verify_text(
            finnish_data,
            "fiberFrame_nameLabel_QLabel",
            "FIBER",
            "Fiber"
        )

        verify_text(
            finnish_data,
            "fiberUsesFrame_nameLabel_QLabel",
            "FIBER USES REMAINING",
            "Fiber Uses Remaining"
        )

        verify_text(
            finnish_data,
            "bottomBar_aimingBeamButton_AimingBeamModeButton",
            "AIMING BEAM",
            "Aiming Beam"
        )

    finally:
        workbook.close()