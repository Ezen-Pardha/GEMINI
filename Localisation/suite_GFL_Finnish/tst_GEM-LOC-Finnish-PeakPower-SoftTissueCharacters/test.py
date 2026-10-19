# -*- coding: utf-8 -*-

import names
import openpyxl
import re
import unicodedata
import os 


TESTCASE_DIR = os.path.dirname(os.path.abspath(__file__))
 
EXCEL_PATH = os.path.abspath(
    os.path.join( TESTCASE_DIR,
        "..",
        "shared",
        "testdata",
        "Gemini Split strings.xlsx"
    )
)
def normalize(text):

    if text is None:
        return ""

    text = str(text)

    text = re.sub(
        r"<br\s*/?>",
        "",
        text,
        flags=re.IGNORECASE
    )

    text = re.sub(
        r"<[^>]+>",
        "",
        text
    )

    text = text.replace("\xa0", " ")
    text = text.replace("\u2007", " ")
    text = text.replace("\u202f", " ")

    text = unicodedata.normalize("NFKC", text)

    text = re.sub(r"\s+", "", text)

    return text.casefold()


def load_finnish_excel(EXCEL_PATH):

    try:
        workbook = openpyxl.load_workbook(
            EXCEL_PATH,
            data_only=True,
            read_only=True
        )
    except Exception as e:
        test.fail(
            "Unable to open Excel file: "
            + str(EXCEL_PATH)
            + " | Error: "
            + str(e)
        )
        return None, {}

    if "Treatment Modes" not in workbook.sheetnames:
        workbook.close()
        test.fail(
            "Worksheet 'Treatment Modes' does not exist in Excel file."
        )
        return None, {}

    sheet = workbook["Treatment Modes"]

    finnish_data = {}

    for row in sheet.iter_rows(
        min_col=2,
        max_col=9,
        values_only=True
    ):

        english_value = row[0]
        finnish_value = row[7]

        if english_value is None:
            continue

        english_key = normalize(english_value)

        if not english_key:
            continue

        if finnish_value is None:
            continue

        finnish_value = str(finnish_value).strip()

        if not finnish_value:
            continue

        finnish_data[english_key] = finnish_value

    return workbook, finnish_data


def get_finnish_text(finnish_data, english_text):

    key = normalize(english_text)

    if key not in finnish_data:
        test.fail(
            "English text not found in Excel: "
            + str(english_text)
        )
        return None

    return finnish_data[key]


def get_object_text(object_name):

    obj = waitForObjectExists(object_name)

    for property_name in ("text", "plainText", "title"):

        try:
            value = getattr(obj, property_name)

            if callable(value):
                value = value()

            if value is not None:
                return str(value)

        except Exception:
            pass

    return ""


def verify_text(finnish_data, object_name, english_text):

    try:
        actual_text = get_object_text(object_name)

    except Exception as e:
        test.fail(
            "Unable to read application object: "
            + str(object_name)
            + " | Error: "
            + str(e)
        )
        return False

    expected_text = get_finnish_text(
        finnish_data,
        english_text
    )

    if expected_text is None:
        return False

    actual_normalized = normalize(actual_text)
    expected_normalized = normalize(expected_text)

    result = actual_normalized == expected_normalized

    test.compare(
        actual_normalized,
        expected_normalized,
        "Finnish text verification: " + str(english_text)
    )

    if not result:
        test.fail(
            "Finnish text mismatch for: "
            + str(english_text)
            + " | Actual: "
            + actual_text
            + " | Expected: "
            + expected_text
        )

    return result


def main():

   

    workbook = None

    clickButton(
        waitForObject(
            names.numpad_pushButton_8_NumpadButton
        )
    )

    clickButton(
        waitForObject(
            names.numpad_pushButton_8_NumpadButton
        )
    )

    clickButton(
        waitForObject(
            names.numpad_pushButton_8_NumpadButton
        )
    )

    clickButton(
        waitForObject(
            names.numpad_pushButton_8_NumpadButton
        )
    )

    snooze(6)

    clickButton(
        waitForObject(
            names.mainFrame_okButton_QPushButton
        )
    )

    clickButton(
        waitForObject(
            names.homeScreen_settingsButton_QPushButton
        )
    )

    clickTab(
        waitForObject(
            names.settingsScreen_tbSystemInfo_TabWidget
        ),
        "INFO"
    )

    clickButton(
        waitForObject(
            names.gbSystemInformation_pbLanguageSettingsUpdate_QPushButton
        )
    )

    clickButton(
        waitForObject(
            names.languagesFrame_pbFinland_QPushButton
        )
    )

    clickButton(
        waitForObject(
            names.buttonsFrame_acceptButton_QPushButton
        )
    )

    clickButton(
        waitForObject(
            names.settingsScreen_pbBack_QPushButton
        )
    )

    clickButton(
        waitForObject(
            names.homeScreen_pbSoftTissueQuickStart_QToolButton
        )
    )

    mouseClick(
        waitForObjectItem(
            names.tabLeftMainPresets_leftPresetsView_PresetListView,
            "_1"
        ),
        305,
        35,
        Qt.NoModifier,
        Qt.LeftButton
    )

    mouseClick(
        waitForObjectItem(
            names.tabRightMainPresets_rightPresetsView_PresetListView,
            "_1"
        ),
        299,
        49,
        Qt.NoModifier,
        Qt.LeftButton
    )

    clickButton(
        waitForObject(
            names.quickStartScreen_pbContinue_QPushButton
        )
    )

    clickButton(
        waitForObject(
            names.mainFrame_okButton_QPushButton_2
        )
    )

    mouseClick(
        waitForObject(
            names.peakPowerButton_nameLabel_QLabel
        ),
        34,
        34,
        Qt.NoModifier,
        Qt.LeftButton
    )

    snooze(6)

    workbook, finnish_data = load_finnish_excel(EXCEL_PATH)

    if workbook is None:
        return

    try:

        snooze(6)

        verify_text(
            finnish_data,
            names.centralWidget_lbTitle_QLabel,
            "SELECT PEAK POWER"
        )

        verify_text(
            finnish_data,
            names.centralWidget_pw500_QPushButton,
            "High"
        )

        verify_text(
            finnish_data,
            names.centralWidget_pw250_QPushButton,
            "Medium"
        )

        verify_text(
            finnish_data,
            names.centralWidget_pw125_QPushButton,
            "Low"
        )

        verify_text(
            finnish_data,
            names.buttonsFrame_rejectButton_QPushButton,
            "CANCEL"
        )

        mouseClick(
            waitForObject(
                names.centralWidget_mainWidget_QWidget
            ),
            460,
            549,
            Qt.NoModifier,
            Qt.LeftButton
        )

        clickButton(
            waitForObject(
                names.buttonsFrame_rejectButton_QPushButton
            )
        )

        clickButton(
            waitForObject(
                names.buttonsBar_homeButton_QPushButton_2
            )
        )

        clickButton(
            waitForObject(
                names.homeScreen_pbSoftTissueAssistant_QToolButton
            )
        )

        clickButton(
            waitForObject(
                names.gbSoftTissueProcedures_pbIncision_QPushButton
            )
        )

        clickButton(
            waitForObject(
                names.softTissueAssistantModeScreen_pbContinue_QPushButton
            )
        )

        clickButton(
            waitForObject(
                names.treatmentCharacteristicsBar_pbEdit_QPushButton
            )
        )

        snooze(6)

        verify_text(
            finnish_data,
            names.centralWidget_titleLabel_QLabel,
            "SOFT TISSUE CHARACTERISTICS"
        )

        verify_text(
            finnish_data,
            names.centralWidget_titleLabel_QLabel_2,
            "PROCEDURE"
        )

        verify_text(
            finnish_data,
            names.centralWidget_valueLabel_QLabel,
            "Incision"
        )

        verify_text(
            finnish_data,
            names.centralWidget_rejectButton_QPushButton,
            "EDIT"
        )

        verify_text(
            finnish_data,
            names.centralWidget_acceptButton_QPushButton,
            "OK"
        )

    finally:
        workbook.close()