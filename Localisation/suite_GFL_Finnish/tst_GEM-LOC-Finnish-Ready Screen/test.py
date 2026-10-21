# -*- coding: utf-8 -*-

import names
import re
import unicodedata
from openpyxl import load_workbook
import os 

TESTCASE_DIR = os.path.dirname(os.path.abspath(__file__))
 
EXCEL_PATH = os.path.abspath(
    os.path.join( TESTCASE_DIR,
        "..",
        "shared",
        "testdata",
        "Gemini Split strings.xlsx"
    )
)

EXCEL_SHEET_NAME = "Treatment Modes"
ENGLISH_COLUMN = 2
FINNISH_COLUMN = 9


def normalize_text(value):
    if value is None:
        return ""

    text = str(value)
    text = re.sub(r"<[^>]*>", " ", text)
    text = text.replace("\n", " ")
    text = text.replace("\r", " ")
    text = text.replace("\t", " ")
    text = text.replace("\xa0", " ")
    text = text.replace("\u2007", " ")
    text = text.replace("\u202f", " ")
    text = text.replace("™", "")
    text = text.replace("®", "")
    text = text.replace("©", "")

    try:
        text = unicodedata.normalize("NFKC", text)
    except Exception:
        pass

    text = re.sub(r"\s+", " ", text).strip()
    text = re.sub(r"\s+([.,;:!?])", r"\1", text)

    return text.casefold()


def comparison_key(value):
    normalized = normalize_text(value)
    return re.sub(r"\s+", "", normalized)


def load_excel_data():
    workbook = None

    try:
        workbook = load_workbook(
            EXCEL_PATH,
            read_only=True,
            data_only=True
        )

        if EXCEL_SHEET_NAME not in workbook.sheetnames:
            test.fail(
                "Excel sheet not found: "
                + EXCEL_SHEET_NAME
                + " | Available sheets: "
                + ", ".join(workbook.sheetnames)
            )
            workbook.close()
            return None, None

        sheet = workbook[EXCEL_SHEET_NAME]
        finnish_data = {}

        for row in sheet.iter_rows(
            min_row=2,
            min_col=ENGLISH_COLUMN,
            max_col=FINNISH_COLUMN
        ):
            english_value = row[0].value
            finnish_value = row[
                FINNISH_COLUMN - ENGLISH_COLUMN
            ].value

            if english_value is None:
                continue

            english_key = comparison_key(english_value)

            if not english_key:
                continue

            if finnish_value is None:
                finnish_value = ""

            finnish_value = str(finnish_value).strip()

            if english_key not in finnish_data:
                finnish_data[english_key] = finnish_value
            elif (
                not comparison_key(finnish_data[english_key])
                and comparison_key(finnish_value)
            ):
                finnish_data[english_key] = finnish_value

        if not finnish_data:
            test.fail(
                "No English/Finnish data found in Excel. "
                "Please check the sheet and column configuration."
            )
            workbook.close()
            return None, None

        return workbook, finnish_data

    except Exception as e:
        test.fail(
            "Unable to load/read Finnish localization Excel: "
            + str(e)
        )

        if workbook is not None:
            try:
                workbook.close()
            except Exception:
                pass

        return None, None


def get_finnish_text(finnish_data, english_text):
    search_text = comparison_key(english_text)

    if search_text not in finnish_data:
        test.fail(
            "English text not found in Excel Column B: "
            + str(english_text)
        )
        return None

    finnish_value = finnish_data[search_text]

    if not comparison_key(finnish_value):
        test.fail(
            "Finnish translation is empty for English text: "
            + str(english_text)
        )
        return None

    return finnish_value


def get_object_text(obj):
    for attribute in (
        "text",
        "title",
        "windowTitle",
        "accessibleName"
    ):
        try:
            value = getattr(obj, attribute)

            if value is not None and str(value).strip():
                return str(value)

        except Exception:
            pass

    return ""


def verify_text(
    finnish_data,
    object_name,
    english_text,
    description
):
    snooze(2)

    expected_text = get_finnish_text(
        finnish_data,
        english_text
    )

    if expected_text is None:
        test.fail(
            description
            + " skipped because Excel lookup failed"
        )
        return False

    try:
        obj = waitForObjectExists(
            getattr(names, object_name)
        )
    except Exception as e:
        test.fail(
            description
            + " | Object not found: "
            + object_name
            + " | "
            + str(e)
        )
        return False

    actual_text = get_object_text(obj)

    if not actual_text:
        test.fail(description + " | UI text is empty")
        return False

    actual_normalized = comparison_key(actual_text)
    expected_normalized = comparison_key(expected_text)

    test.compare(
        actual_normalized,
        expected_normalized,
        description + " - Finnish text verification"
    )

    if actual_normalized == expected_normalized:
        return True

    test.fail(
        description
        + " | Finnish text mismatch"
        + " | Actual: "
        + str(actual_text)
        + " | Expected: "
        + str(expected_text)
    )

    return False


def find_text_recursively(parent_object, expected_normalized):
    try:
        if not parent_object.visible:
            return None
    except Exception:
        pass

    try:
        current_text = comparison_key(
            parent_object.text
        )

        if current_text == expected_normalized:
            return parent_object
    except Exception:
        pass

    try:
        children = object.children(parent_object)
    except Exception:
        children = []

    for child in children:
        found = find_text_recursively(
            child,
            expected_normalized
        )

        if found is not None:
            return found

    return None


def verify_treatment_name(
    finnish_data,
    english_text,
    description
):
    snooze(3)

    expected_text = get_finnish_text(
        finnish_data,
        english_text
    )

    if expected_text is None:
        test.fail(
            description
            + " skipped because Excel lookup failed"
        )
        return False

    expected_normalized = comparison_key(expected_text)

    try:
        screen_switcher = waitForObjectExists(
            names.o_ScreenSwitcher
        )
    except Exception as e:
        test.fail(
            description
            + " | ScreenSwitcher not found | "
            + str(e)
        )
        return False

    found_object = find_text_recursively(
        screen_switcher,
        expected_normalized
    )

    if found_object is None:
        test.fail(
            "{} | Finnish treatment text not found on screen: [{}]".format(
                description,
                expected_text
            )
        )
        return False

    actual_text = get_object_text(found_object)
    actual_normalized = comparison_key(actual_text)

    test.compare(
        actual_normalized,
        expected_normalized,
        description + " - Finnish treatment verification"
    )

    return actual_normalized == expected_normalized


def verify_left_treatment_name(
    finnish_data,
    english_text,
    description
):
    return verify_treatment_name(
        finnish_data,
        english_text,
        description
    )


def verify_right_treatment_name(
    finnish_data,
    english_text,
    description
):
    return verify_treatment_name(
        finnish_data,
        english_text,
        description
    )


def main():
    workbook = None
    finnish_data = None

    result = load_excel_data()

    if result is None:
        test.fail("load_excel_data() returned None")
        return

    try:
        workbook, finnish_data = result
    except Exception as e:
        test.fail(
            "Unable to unpack Excel data: "
            + str(e)
        )
        return

    if workbook is None or finnish_data is None:
        test.fail("Finnish Excel data could not be loaded")
        return

    try:
        for i in range(4):
            clickButton(
                waitForObject(
                    names.numpad_pushButton_8_NumpadButton
                )
            )
            snooze(1)

        clickButton(
            waitForObject(
                names.mainFrame_okButton_QPushButton
            )
        )

        clickButton(
            waitForObject(
                names.homeScreen_settingsButton_QPushButton
            )
        )

        clickTab(
            waitForObject(
                names.settingsScreen_tbSystemInfo_TabWidget
            ),
            "INFO"
        )

        clickButton(
            waitForObject(
                names.gbSystemInformation_pbLanguageSettingsUpdate_QPushButton
            )
        )

        clickButton(
            waitForObject(
                names.languagesFrame_pbFinland_QPushButton
            )
        )

        clickButton(
            waitForObject(
                names.buttonsFrame_acceptButton_QPushButton
            )
        )

        clickButton(
            waitForObject(
                names.buttonsBar_homeButton_QPushButton
            )
        )

        clickButton(
            waitForObject(
                names.homeScreen_pbSoftTissueQuickStart_QToolButton
            )
        )

        mouseClick(
            waitForObjectItem(
                names.tabLeftMainPresets_leftPresetsView_PresetListView,
                "_1"
            ),
            238,
            58,
            Qt.NoModifier,
            Qt.LeftButton
        )

        mouseClick(
            waitForObjectItem(
                names.tabRightMainPresets_rightPresetsView_PresetListView,
                "_1"
            ),
            148,
            32,
            Qt.NoModifier,
            Qt.LeftButton
        )

        mouseClick(
            waitForObject(
                names.quickStartScreen_wdBottom_QWidget
            ),
            1167,
            9,
            Qt.NoModifier,
            Qt.LeftButton
        )

        sendEvent(
            "QMoveEvent",
            waitForObject(names.o_ScreenSwitcher),
            0,
            -51,
            1286,
            61
        )

        clickButton(
            waitForObject(
                names.quickStartScreen_pbContinue_QPushButton
            )
        )

        clickButton(
            waitForObject(
                names.mainFrame_okButton_QPushButton_2
            )
        )

        clickButton(
            waitForObject(
                names.stateSwitch_readyButton_QPushButton
            )
        )

        verify_left_treatment_name(
            finnish_data,
            "PROSTATE ENUCLEATION-(HIGH ENERGY)",
            "Left Treatment Name"
        )

        verify_text(
            finnish_data,
            "averagePowerWidgetEmission_nameLabel_QLabel",
            "Average Power",
            "Left Average Power"
        )

        verify_text(
            finnish_data,
            "pulseEnergyWidgetEmission_nameLabel_QLabel",
            "Pulse Energy",
            "Left Pulse Energy"
        )

        verify_text(
            finnish_data,
            "frequencyWidgetEmission_nameLabel_QLabel",
            "Frequency",
            "Left Frequency"
        )

        verify_text(
            finnish_data,
            "modeButtonEmission_nameLabel_QLabel",
            "Pulse Mode",
            "Left Pulse Mode"
        )

        verify_text(
            finnish_data,
            "peakPowerButtonEmission_nameLabel_QLabel",
            "Peak Power",
            "Left Peak Power"
        )

        verify_text(
            finnish_data,
            "totalEnergyFrame_nameLabel_QLabel",
            "Total Energy",
            "Left Total Energy"
        )

        verify_text(
            finnish_data,
            "totalTimeFrame_nameLabel_QLabel",
            "Total Lasing Time",
            "Left Total Lasing Time"
        )

        verify_text(
            finnish_data,
            "averagePowerWidgetEmission_unitsOfMeasureLabel_QLabel",
            "W",
            "Left Power Unit"
        )

        verify_text(
            finnish_data,
            "pulseEnergyWidgetEmission_unitsOfMeasureLabel_QLabel",
            "J",
            "Left Energy Unit"
        )

        verify_text(
            finnish_data,
            "frequencyWidgetEmission_unitsOfMeasureLabel_QLabel",
            "Hz",
            "Left Frequency Unit"
        )

        verify_text(
            finnish_data,
            "fiberFrame_nameLabel_QLabel",
            "Fiber",
            "Left Fiber"
        )

        verify_text(
            finnish_data,
            "fiberUsesFrame_nameLabel_QLabel",
            "Fiber Uses Remaining",
            "Left Fiber Uses Remaining"
        )

        verify_text(
            finnish_data,
            "bottomBar_aimingBeamButton_AimingBeamModeButton",
            "Aiming Beam",
            "Left Aiming Beam"
        )

        verify_text(
            finnish_data,
            "stateSwitch_standbyButton_QPushButton",
            "Standby",
            "Left Standby"
        )

        verify_text(
            finnish_data,
            "stateSwitch_readyButton_QPushButton",
            "Ready",
            "Left Ready"
        )

        verify_right_treatment_name(
            finnish_data,
            "PROSTATE ENUCLEATION-(HIGH ENERGY)",
            "Right Treatment Name"
        )

        verify_text(
            finnish_data,
            "averagePowerWidgetEmission_nameLabel_QLabel_2",
            "Average Power",
            "Right Average Power"
        )

        verify_text(
            finnish_data,
            "pulseEnergyWidgetEmission_nameLabel_QLabel_2",
            "Pulse Energy",
            "Right Pulse Energy"
        )

        verify_text(
            finnish_data,
            "frequencyWidgetEmission_nameLabel_QLabel_2",
            "Frequency",
            "Right Frequency"
        )

        verify_text(
            finnish_data,
            "modeButtonEmission_nameLabel_QLabel_2",
            "Pulse Mode",
            "Right Pulse Mode"
        )

        verify_text(
            finnish_data,
            "peakPowerButtonEmission_nameLabel_QLabel_2",
            "Peak Power",
            "Right Peak Power"
        )

        verify_text(
            finnish_data,
            "averagePowerWidgetEmission_unitsOfMeasureLabel_QLabel_2",
            "W",
            "Right Power Unit"
        )

        verify_text(
            finnish_data,
            "pulseEnergyWidgetEmission_unitsOfMeasureLabel_QLabel_2",
            "J",
            "Right Energy Unit"
        )

        verify_text(
            finnish_data,
            "frequencyWidgetEmission_unitsOfMeasureLabel_QLabel_2",
            "Hz",
            "Right Frequency Unit"
        )

        clickButton(
            waitForObject(
                names.stateSwitch_readyButton_QPushButton
            )
        )

       

        clickButton(
            waitForObject(
                names.stateSwitch_readyButton_QPushButton
            )
        )
        test.vp("Readyscreen")
    except Exception as e:
        test.fail(
            "Finnish Ready Screen test failed: "
            + str(e)
        )

    finally:
        if workbook is not None:
            try:
                workbook.close()
            except Exception:
                pass