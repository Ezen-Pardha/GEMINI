
# -*- coding: utf-8 -*-

import names
import openpyxl
import re

EXCEL_PATH = (
    "/home/ntc/Test_Automation_Gemini/Localisation/"
    "suite_GFL_Czech/tst_Home Screen/"
    "testdata/Gemini Split strings.xlsx"
)

SHEET_NAME = "Stone-Guided"


def normalize(text):
    if text is None:
        return ""

    text = str(text)

    # Normalize HTML line breaks
    text = re.sub(
        r"<br\s*/?>",
        " ",
        text,
        flags=re.IGNORECASE
    )

    # Remove HTML tags
    text = re.sub(r"<[^>]+>", " ", text)

    # Normalize whitespace
    text = text.replace("\xa0", " ")
    text = " ".join(text.split())

    # Normalize dash characters
    text = text.replace("–", "-")
    text = text.replace("—", "-")
    text = text.replace("−", "-")

    # Normalize spaces around angle brackets
    text = re.sub(r"\s*<\s*", "<", text)
    text = re.sub(r"\s*>\s*", ">", text)

    return text.strip().casefold()


def ignore_text(text):
    key = normalize(text)

    if not key:
        return True

    # Ignore any string containing digits
    if re.search(r"\d", key):
        return True

    # Ignore symbol-only strings
    if re.fullmatch(r"[-_=.<>/%µμ:+]+", key):
        return True

    return False


def get_text_lines(obj):
    lines = []

    # Only collect visible objects
    try:
        if not obj.visible:
            return lines
    except Exception:
        pass

    # Read text from the current object
    try:
        value = obj.text

        if value is not None:
            text = str(value).strip()

            if text:
                text = re.sub(
                    r"<br\s*/?>",
                    "\n",
                    text,
                    flags=re.IGNORECASE
                )

                text = re.sub(r"<[^>]+>", " ", text)

                for line in text.splitlines():
                    line = " ".join(line.split()).strip()

                    if line:
                        lines.append(line)

    except Exception:
        pass

    # Read text from child objects
    try:
        for child in object.children(obj):
            lines.extend(get_text_lines(child))
    except Exception:
        pass

    return lines


def load_czech_excel(excel_path):
    workbook = openpyxl.load_workbook(
        excel_path,
        data_only=True,
        read_only=True
    )

    if SHEET_NAME not in workbook.sheetnames:
        available_sheets = ", ".join(workbook.sheetnames)
        workbook.close()

        test.fail(
            "FAIL: Worksheet not found: "
            + SHEET_NAME
            + " | Available worksheets: "
            + available_sheets
        )

        return None, {}

    sheet = workbook[SHEET_NAME]

    # Column B = English
    # Column E = Czech
    czech_data = {}

    for row in sheet.iter_rows(
        min_row=2,
        min_col=2,
        max_col=5,
        values_only=True
    ):
        english = row[0]
        czech = row[3]

        if english is None or czech is None:
            continue

        english = str(english).strip()
        czech = str(czech).strip()

        if not english or not czech:
            continue

        key = normalize(czech)

        # Do not use numeric strings as translation keys
        if ignore_text(czech):
            continue

        if key:
            czech_data[key] = {
                "english": english,
                "czech": czech
            }

    return workbook, czech_data


def verify_screen_strings(screen_object, czech_data):
    screen_lines = get_text_lines(screen_object)

    screen_strings = {}

    # Collect unique eligible visible strings
    for text in screen_lines:
        key = normalize(text)

        # Ignore strings containing digits and symbol-only strings
        if ignore_text(text):
            continue

        if key not in screen_strings:
            screen_strings[key] = text.strip()

    pass_count = 0
    fail_count = 0

    # Verify every collected screen string
    for key in sorted(screen_strings):
        screen_text = screen_strings[key]

        if key in czech_data:
            data = czech_data[key]

            test.passes(
                "PASS: "
                + data["english"]
                + " -> "
                + data["czech"]
            )

            pass_count += 1

        else:
            test.fail(
                "FAIL: Visible screen string is missing "
                "from Czech Excel: "
                + screen_text
            )

            fail_count += 1

    total = len(screen_strings)

    # Final result
    if total == 0:
        test.fail(
            "FAIL: No eligible visible strings found on screen"
        )

    elif fail_count > 0:
        test.fail(
            "CZECH LOCALIZATION VALIDATION FAILED: "
            + str(fail_count)
            + " unmatched string(s) out of "
            + str(total)
        )

    else:
        test.passes(
            "CZECH LOCALIZATION VALIDATION PASSED: "
            + str(pass_count)
            + " unique visible string(s) verified"
        )

    return pass_count, fail_count


def main():
    workbook = None

    try:
        # LOGIN

        sendEvent(
            "QMoveEvent",
            waitForObject(names.o_ScreenSwitcher),
            57, 167, 718, 194
        )

        for _ in range(4):
            clickButton(
                waitForObject(
                    names.numpad_pushButton_8_NumpadButton
                )
            )

        snooze(2)

        clickButton(
            waitForObject(
                names.mainFrame_okButton_QPushButton
            )
        )

        # CHANGE LANGUAGE TO CZECH

        clickButton(
            waitForObject(
                names.homeScreen_settingsButton_QPushButton
            )
        )

        clickTab(
            waitForObject(
                names.settingsScreen_tbSystemInfo_TabWidget
            ),
            "INFO"
        )

        clickButton(
            waitForObject(
                names.gbSystemInformation_pbLanguageSettingsUpdate_QPushButton
            )
        )

        clickButton(
            waitForObject(
                names.languagesFrame_pbCzech_QPushButton
            )
        )

        clickButton(
            waitForObject(
                names.buttonsFrame_acceptButton_QPushButton
            )
        )

        clickButton(
            waitForObject(
                names.settingsScreen_pbBack_QPushButton
            )
        )

        # OPEN STONE ASSISTANT

        clickButton(
            waitForObject(
                names.homeScreen_pbStoneAssistant_QToolButton
            )
        )

        # LOAD EXCEL

        workbook, czech_data = load_czech_excel(
            EXCEL_PATH
        )

        if workbook is None:
            return

        if not czech_data:
            test.fail(
                "FAIL: No eligible Czech translations found "
                "in worksheet: "
                + SHEET_NAME
            )
            return

        # STONE-GUIDED SELECTION

        clickButton(
            waitForObject(
                names.locationFrame_kidneyPcnlLocationButton_QPushButton
            )
        )

        mouseClick(
            waitForObject(
                names.locationFrame_label_3_QLabel
            ),
            189, 20,
            Qt.NoModifier,
            Qt.LeftButton
        )

        clickButton(
            waitForObject(
                names.flowRateFrame_lowFlowRateButton_QPushButton
            )
        )

        mouseClick(
            waitForObject(
                names.flowRateFrame_label_4_QLabel
            ),
            202, 28,
            Qt.NoModifier,
            Qt.LeftButton
        )

        clickButton(
            waitForObject(
                names.hardnessFrame_largeHardnessButton_QPushButton
            )
        )

        clickButton(
            waitForObject(
                names.hardnessFrame_mediumHardnessButton_QPushButton
            )
        )

        mouseClick(
            waitForObject(
                names.hardnessFrame_label_6_QLabel
            ),
            196, 24,
            Qt.NoModifier,
            Qt.LeftButton
        )

        snooze(1)

        # GET STONE-GUIDED SCREEN

        stone = waitForObject(
            names.stoneAssistantModeScreen_StoneAssistantModeScreen
        )

        # VERIFY SCREEN STRINGS
        # Numeric-containing strings are ignored.
        # All other collected strings must exist in Czech Excel.

        verify_screen_strings(
            stone,
            czech_data
        )

    except Exception as e:
        test.fail(
            "CZECH LOCALIZATION TEST EXECUTION ERROR: "
            + str(e)
        )

    finally:
        if workbook is not None:
            workbook.close()
