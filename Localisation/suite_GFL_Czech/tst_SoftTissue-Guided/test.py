
# -*- coding: utf-8 -*-

import names
import openpyxl
import re
import os

TESTCASE_DIR = os.path.dirname(os.path.abspath(__file__))

EXCEL_PATH = os.path.abspath(
    os.path.join(
        TESTCASE_DIR,
        "..",
        "shared",
        "testdata",
        "Gemini Split strings.xlsx"
    )
)

SHEET_NAME = "Soft Tissue-Guided"


def normalize(text):
    text = str(text or "")

    text = re.sub(r"<[^>]+>", "", text)
    text = text.replace("\u00a0", " ")
    text = " ".join(text.split())

    return text.casefold()


def get_all_texts(obj):
    texts = []

    try:
        text = str(obj.text).strip()

        if text:
            if "<html>" in text.lower():
                text = re.sub(r"<[^>]+>", "", text)
                text = " ".join(text.split())

            if text:
                texts.append(text)

    except Exception:
        pass

    try:
        children = object.children(obj)

        for child in children:
            texts.extend(get_all_texts(child))

    except Exception:
        pass

    return texts


def ignore_text(text):
    text = str(text).strip()

    if not text:
        return True

    if re.fullmatch(r"[-_=.]+", text):
        return True

    return False


def main():

    if not os.path.isfile(EXCEL_PATH):
        test.fail(
            "Excel file not found: {}".format(EXCEL_PATH)
        )
        return

    workbook = openpyxl.load_workbook(
        EXCEL_PATH,
        data_only=True,
        read_only=True
    )

    try:
        if SHEET_NAME not in workbook.sheetnames:
            test.fail(
                "Worksheet not found: {}".format(SHEET_NAME)
            )
            return

        sheet = workbook[SHEET_NAME]

        # Column B = English
        # Column E = Czech

        excel_strings = {}

        for row in sheet.iter_rows(
            min_row=2,
            values_only=True
        ):

            if len(row) < 5:
                continue

            english_value = row[1]
            czech_value = row[4]

            if czech_value is None:
                continue

            english = (
                str(english_value).strip()
                if english_value is not None
                else ""
            )

            czech = str(czech_value).strip()

            if ignore_text(czech):
                continue

            key = normalize(czech)

            if key not in excel_strings:
                excel_strings[key] = []

            excel_strings[key].append(english)

        # Login

        doubleClick(
            waitForObject(names.numpad_pushButton_8_NumpadButton),
            68, 52,
            Qt.NoModifier,
            Qt.LeftButton
        )

        doubleClick(
            waitForObject(names.numpad_pushButton_8_NumpadButton),
            116, 46,
            Qt.NoModifier,
            Qt.LeftButton
        )

        clickButton(
            waitForObject(names.mainFrame_okButton_QPushButton)
        )

        # Settings

        clickButton(
            waitForObject(names.homeScreen_settingsButton_QPushButton)
        )

        clickTab(
            waitForObject(names.settingsScreen_tbSystemInfo_TabWidget),
            "INFO"
        )

        clickButton(
            waitForObject(
                names.gbSystemInformation_pbLanguageSettingsUpdate_QPushButton
            )
        )

        clickButton(
            waitForObject(names.languagesFrame_pbCzech_QPushButton)
        )

        clickButton(
            waitForObject(names.buttonsFrame_acceptButton_QPushButton)
        )

        clickButton(
            waitForObject(names.settingsScreen_pbBack_QPushButton)
        )

        # Open Soft Tissue Assistant

        clickButton(
            waitForObject(
                names.homeScreen_pbSoftTissueAssistant_QToolButton
            )
        )

        mouseClick(
            waitForObject(
                names.softTissueAssistantModeScreen_gbSoftTissueProcedures_QFrame
            ),
            90, 150,
            Qt.NoModifier,
            Qt.LeftButton
        )

        # Collect screen text

        soft_tissue = waitForObject(
            names.softTissueAssistantModeScreen_SoftTissueAssistantModeScreen
        )

        actual_texts = get_all_texts(soft_tissue)

        # Remove duplicates

        unique_texts = []
        seen = set()

        for text in actual_texts:

            if ignore_text(text):
                continue

            key = normalize(text)

            if key not in seen:
                seen.add(key)
                unique_texts.append(text)

        # Validate translations

        total_pass = 0
        total_fail = 0

        for text in unique_texts:

            key = normalize(text)

            if key in excel_strings:

                english_values = excel_strings[key]

                english = next(
                    (value for value in english_values if value),
                    ""
                )

                if english:
                    test.passes(
                        "PASS: {} -> {}".format(english, text)
                    )
                else:
                    test.passes(
                        "PASS: Czech string found in Excel -> {}".format(
                            text
                        )
                    )

                total_pass += 1

            else:

                test.fail(
                    "FAIL: Screen string not found in Czech Excel: {}".format(
                        text
                    )
                )

                total_fail += 1

        # Final summary

        test.log(
            "Localization Summary | Total: {} | PASS: {} | FAIL: {}".format(
                len(unique_texts),
                total_pass,
                total_fail
            )
        )

    finally:
        workbook.close()
