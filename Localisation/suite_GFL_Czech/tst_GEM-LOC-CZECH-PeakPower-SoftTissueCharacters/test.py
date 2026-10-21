# -*- coding: utf-8 -*-

import names
import openpyxl
import re
import time
import os

TESTCASE_DIR = os.path.dirname(os.path.abspath(__file__))
EXCEL_PATH = os.path.abspath(os.path.join(
    TESTCASE_DIR, "..", "shared", "testdata", "Gemini Split strings.xlsx"
))

SHEET_NAME = "Treatment Modes"
ENGLISH_COLUMN = 2
CZECH_COLUMN = 5


def normalize(text):
    if text is None:
        return ""
    text = str(text)
    text = re.sub(r"<br\s*/?>", " ", text, flags=re.IGNORECASE)
    text = re.sub(r"<[^>]+>", " ", text)
    text = text.replace("\xa0", " ")
    return " ".join(text.split()).strip().casefold()


def load_Czech_excel(excel_path):
    workbook = None
    try:
        start_time = time.time()
        workbook = openpyxl.load_workbook(
            excel_path, data_only=True, read_only=True
        )
        load_time = time.time() - start_time

        if SHEET_NAME not in workbook.sheetnames:
            test.fail("Worksheet '" + SHEET_NAME + "' does not exist in: " + excel_path)
            workbook.close()
            return None, None, {}

        sheet = workbook[SHEET_NAME]
        Czech_data = {}

        # Range B:E returns four values: row[0] is B, row[3] is E.
        for row in sheet.iter_rows(
            min_row=1,
            min_col=ENGLISH_COLUMN,
            max_col=CZECH_COLUMN,
            values_only=True
        ):
            if len(row) < (CZECH_COLUMN - ENGLISH_COLUMN + 1):
                continue

            english_value = row[0]
            Czech_value = row[CZECH_COLUMN - ENGLISH_COLUMN]

            if english_value is None or Czech_value is None:
                continue

            english_key = normalize(english_value)
            Czech_text = str(Czech_value).strip()

            if not english_key or not Czech_text:
                continue

            if english_key not in Czech_data:
                Czech_data[english_key] = []
            if Czech_text not in Czech_data[english_key]:
                Czech_data[english_key].append(Czech_text)

        if not Czech_data:
            test.fail("No English/Czech translations loaded. Check worksheet and columns B/E.")
            workbook.close()
            return None, None, {}

        return workbook, sheet, Czech_data

    except Exception as e:
        if workbook is not None:
            try:
                workbook.close()
            except Exception:
                pass
        test.fail("Unable to load Czech localization Excel: " + str(e))
        return None, None, {}


def get_Czech_text(Czech_data, english_text):
    key = normalize(english_text)
    values = Czech_data.get(key, [])

    if isinstance(values, list):
        for value in values:
            if value:
                return value
        return ""
    return str(values) if values else ""


def verify_text(Czech_data, object_name, english_text):
    try:
        obj = waitForObjectExists(object_name, 10000)
        actual_value = getattr(obj, "text", "")
        actual_text = "" if actual_value is None else str(actual_value).strip()
    except Exception as e:
        test.fail(
            "Unable to read application object: " + str(object_name)
            + " | Error: " + str(e)
        )
        return False

    expected_text = get_Czech_text(Czech_data, english_text)
    if not expected_text:
        test.fail(
            "English text not found in Excel Column B, or Czech translation "
            "is empty in Column E: " + str(english_text)
        )
        return False

    actual_normalized = normalize(actual_text)
    expected_normalized = normalize(expected_text)

    if actual_normalized == expected_normalized:
        test.compare(
            actual_normalized,
            expected_normalized,
            "Czech text verification: " + str(english_text) + " -> " + expected_text
        )
        return True

    test.fail(
        "Czech text mismatch for: " + str(english_text)
        + " | Actual: " + actual_text
        + " | Expected: " + expected_text
    )
    return False


def main():
    workbook, sheet, Czech_data = load_Czech_excel(EXCEL_PATH)
    if workbook is None or sheet is None or not Czech_data:
        return

    try:
        clickButton(waitForObject(names.numpad_pushButton_8_NumpadButton))
        clickButton(waitForObject(names.numpad_pushButton_8_NumpadButton))
        clickButton(waitForObject(names.numpad_pushButton_8_NumpadButton))
        clickButton(waitForObject(names.numpad_pushButton_8_NumpadButton))
        snooze(6)

        clickButton(waitForObject(names.mainFrame_okButton_QPushButton))
        clickButton(waitForObject(names.homeScreen_settingsButton_QPushButton))
        clickTab(waitForObject(names.settingsScreen_tbSystemInfo_TabWidget), "INFO")
        clickButton(waitForObject(
            names.gbSystemInformation_pbLanguageSettingsUpdate_QPushButton
        ))
        clickButton(waitForObject(names.languagesFrame_pbCzech_QPushButton))
        clickButton(waitForObject(names.buttonsFrame_acceptButton_QPushButton))
        clickButton(waitForObject(names.settingsScreen_pbBack_QPushButton))

        clickButton(waitForObject(names.homeScreen_pbSoftTissueQuickStart_QToolButton))
        mouseClick(
            waitForObjectItem(
                names.tabLeftMainPresets_leftPresetsView_PresetListView, "_1"
            ),
            305, 35, Qt.NoModifier, Qt.LeftButton
        )
        mouseClick(
            waitForObjectItem(
                names.tabRightMainPresets_rightPresetsView_PresetListView, "_1"
            ),
            299, 49, Qt.NoModifier, Qt.LeftButton
        )
        clickButton(waitForObject(names.quickStartScreen_pbContinue_QPushButton))
        clickButton(waitForObject(names.mainFrame_okButton_QPushButton_2))

        mouseClick(
            waitForObject(names.peakPowerButton_nameLabel_QLabel),
            34, 34, Qt.NoModifier, Qt.LeftButton
        )
        snooze(2)

        verify_text(Czech_data, names.centralWidget_lbTitle_QLabel, "SELECT PEAK POWER")
        verify_text(Czech_data, names.centralWidget_pw500_QPushButton, "High")
        verify_text(Czech_data, names.centralWidget_pw250_QPushButton, "Medium")
        verify_text(Czech_data, names.centralWidget_pw125_QPushButton, "Low")
        verify_text(Czech_data, names.buttonsFrame_rejectButton_QPushButton, "CANCEL")

        mouseClick(
            waitForObject(names.centralWidget_mainWidget_QWidget),
            460, 549, Qt.NoModifier, Qt.LeftButton
        )
        clickButton(waitForObject(names.buttonsFrame_rejectButton_QPushButton))
        clickButton(waitForObject(names.buttonsBar_homeButton_QPushButton_2))
        clickButton(waitForObject(names.homeScreen_pbSoftTissueAssistant_QToolButton))
        clickButton(waitForObject(names.gbSoftTissueProcedures_pbIncision_QPushButton))
        clickButton(waitForObject(names.softTissueAssistantModeScreen_pbContinue_QPushButton))
        clickButton(waitForObject(names.treatmentCharacteristicsBar_pbEdit_QPushButton))
        snooze(2)

        verify_text(Czech_data, names.centralWidget_titleLabel_QLabel, "SOFT TISSUE CHARACTERISTICS")
        verify_text(Czech_data, names.centralWidget_titleLabel_QLabel_2, "PROCEDURE")
        verify_text(Czech_data, names.centralWidget_valueLabel_QLabel, "Incision")
        verify_text(Czech_data, names.centralWidget_rejectButton_QPushButton, "EDIT")

       
        verify_text(Czech_data, names.centralWidget_acceptButton_QPushButton, "OK")

    finally:
        workbook.close()
