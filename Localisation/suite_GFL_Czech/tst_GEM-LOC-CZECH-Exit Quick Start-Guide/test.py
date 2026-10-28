# -*- coding: utf-8 -*-

import names
import openpyxl
import re
import os


TESTCASE_DIR = os.path.dirname(os.path.abspath(__file__))

EXCEL_PATH = os.path.abspath(
    os.path.join(
        TESTCASE_DIR,
        "..",
        "shared",
        "testdata",
        "Gemini Split strings.xlsx"
    )
)

ENGLISH_COLUMN = 2
CZECH_COLUMN = 5
SHEET_NAME = "Soft tissue"


def normalize(text):
    if text is None:
        return ""

    text = str(text)
    text = re.sub(r"<[^>]+>", " ", text)
    text = " ".join(text.split())

    text = re.sub(r"\s*<\s*", " < ", text)
    text = re.sub(r"\s*>\s*", " > ", text)

    return text.strip().casefold()


def load_excel():
    try:
        workbook = openpyxl.load_workbook(
            EXCEL_PATH,
            data_only=True,
            read_only=True
        )

        if SHEET_NAME not in workbook.sheetnames:
            test.fail("Excel sheet not found: " + SHEET_NAME)
            workbook.close()
            return None, None

        return workbook, workbook[SHEET_NAME]

    except Exception as e:
        test.fail(
            "Unable to load Czech localization Excel: " + str(e)
        )
        return None, None


def build_excel_lookup(sheet):
    """
    Read the worksheet once and build a normal Python dictionary.
    Dictionary .get() is used later; ReadOnlyWorksheet.get() is never used.
    """
    excel_lookup = {}

    for row in sheet.iter_rows(min_row=2, values_only=True):
        if len(row) < max(ENGLISH_COLUMN, CZECH_COLUMN):
            continue

        english_value = row[ENGLISH_COLUMN - 1]
        czech_value = row[CZECH_COLUMN - 1]

        if english_value is None or czech_value is None:
            continue

        english_key = normalize(english_value)
        czech_translation = str(czech_value).strip()

        if not english_key or not czech_translation:
            continue

        excel_lookup.setdefault(english_key, [])

        if czech_translation not in excel_lookup[english_key]:
            excel_lookup[english_key].append(czech_translation)

    return excel_lookup


def get_CZECH_values(excel_lookup, english_text):
    english_key = normalize(english_text)
    return excel_lookup.get(english_key, [])


def verify_text(excel_lookup, object_name, english_text, description):
    try:
        obj = waitForObjectExists(object_name, 10000)
        actual_value = getattr(obj, "text", "")

        if actual_value is None:
            actual_value = ""

        actual = str(actual_value).strip()

    except Exception as e:
        test.fail(
            description + " | Unable to read UI object: " + str(e)
        )
        return

    expected_values = get_CZECH_values(excel_lookup, english_text)

    if not expected_values:
        test.fail(
            description
            + " | English string not found in Excel Column B: "
            + english_text
        )
        return

    actual_normalized = normalize(actual)

    for expected in expected_values:
        expected_normalized = normalize(expected)

        if actual_normalized == expected_normalized:
            test.compare(
                actual_normalized,
                expected_normalized,
                description
                + " | English: "
                + english_text
                + " | Czech: "
                + expected
            )
            return

    test.fail(
        description
        + " | Expected Czech value from Excel Column "
        + str(CZECH_COLUMN)
        + ": "
        + " / ".join(expected_values)
        + " | Actual UI text: "
        + actual
    )


def get_central_widget():
    try:
        central_widget = findObject(
            {
                "name": "centralWidget",
                "type": "QFrame",
                "window": ":o_ScreenSwitcher"
            }
        )

        for _ in range(20):
            try:
                if central_widget.visible:
                    return central_widget
            except Exception:
                pass

            snooze(0.5)

        test.log("centralWidget was found but is not visible")
        return None

    except Exception as e:
        test.log("centralWidget not found: " + str(e))
        return None


def main():
    workbook, sheet = load_excel()

    if workbook is None or sheet is None:
        return

    try:
        # Build the lookup once. All UI comparisons use this dictionary.
        excel_lookup = build_excel_lookup(sheet)

        if not excel_lookup:
            test.fail(
                "No English/Czech translations were loaded from Excel. "
                "Check the sheet and column settings."
            )
            return

        clickButton(
            waitForObject(
                names.numpad_pushButton_8_NumpadButton
            )
        )

        doubleClick(
            waitForObject(
                names.numpad_pushButton_8_NumpadButton
            ),
            94,
            77,
            Qt.NoModifier,
            Qt.LeftButton
        )

        doubleClick(
            waitForObject(
                names.numpad_pushButton_8_NumpadButton
            ),
            94,
            77,
            Qt.NoModifier,
            Qt.LeftButton
        )

        clickButton(
            waitForObject(
                names.mainFrame_okButton_QPushButton
            )
        )

        clickButton(
            waitForObject(
                names.homeScreen_settingsButton_QPushButton
            )
        )

        clickTab(
            waitForObject(
                names.settingsScreen_tbSystemInfo_TabWidget
            ),
            "INFO"
        )

        clickButton(
            waitForObject(
                names.gbSystemInformation_pbLanguageSettingsUpdate_QPushButton
            )
        )

        clickButton(
            waitForObject(
                names.languagesFrame_pbCzech_QPushButton
            )
        )

        clickButton(
            waitForObject(
                names.buttonsFrame_acceptButton_QPushButton
            )
        )

        clickButton(
            waitForObject(
                names.buttonsBar_homeButton_QPushButton
            )
        )

        clickButton(
            waitForObject(
                names.homeScreen_pbSoftTissueQuickStart_QToolButton
            )
        )

        mouseClick(
            waitForObjectItem(
                names.tabLeftMainPresets_leftPresetsView_PresetListView,
                "_1"
            ),
            377,
            65,
            Qt.NoModifier,
            Qt.LeftButton
        )

        mouseClick(
            waitForObjectItem(
                names.tabRightMainPresets_rightPresetsView_PresetListView,
                "_1"
            ),
            150,
            37,
            Qt.NoModifier,
            Qt.LeftButton
        )

        sendEvent(
            "QMoveEvent",
            waitForObject(
                names.o_ScreenSwitcher
            ),
            0,
            -50,
            1288,
            292
        )

        clickButton(
            waitForObject(
                names.quickStartScreen_pbContinue_QPushButton
            )
        )

        clickButton(
            waitForObject(
                names.mainFrame_okButton_QPushButton_2
            )
        )

        mouseClick(
            waitForObject(
                names.modeButton_nameLabel_QLabel
            ),
            115,
            15,
            Qt.NoModifier,
            Qt.LeftButton
        )

        verify_text(
            excel_lookup,
            names.centralWidget_TitleLabel_QLabel_2,
            "EXIT QUICK START MODE?",
            "Quick Start Exit Title"
        )

        verify_text(
            excel_lookup,
            names.centralWidget_MessageLabel_QLabel_2,
            "Please confirm you want to exit quick state mode.",
            "Quick Start Exit Message"
        )

        verify_text(
            excel_lookup,
            names.centralWidget_CancelButton_QPushButton_3,
            "CANCEL",
            "Quick Start Cancel Button"
        )

        verify_text(
            excel_lookup,
            names.centralWidget_DisableButton_QPushButton_2,
            "CONFIRM",
            "Quick Start Confirm Button"
        )

        clickButton(
            waitForObject(
                names.centralWidget_DisableButton_QPushButton_2
            )
        )

        clickButton(
            waitForObject(
                names.buttonsBar_homeButton_QPushButton_2
            )
        )

        clickButton(
            waitForObject(
                names.homeScreen_pbSoftTissueAssistant_QToolButton
            )
        )

        clickButton(
            waitForObject(
                names.gbSoftTissueProcedures_pbIncision_QPushButton
            )
        )

        clickButton(
            waitForObject(
                names.softTissueAssistantModeScreen_pbContinue_QPushButton
            )
        )

        mouseDrag(
            waitForObject(
                names.indexPicker_valuePicker_IntervalSlider
            ),
            195,
            15,
            211,
            29,
            1,
            Qt.LeftButton
        )

        verify_text(
            excel_lookup,
            names.centralWidget_TitleLabel_QLabel_2,
            "EXIT GUIDED MODE?",
            "Guided Mode Exit Title"
        )

        verify_text(
            excel_lookup,
            names.centralWidget_MessageLabel_QLabel_2,
            "Please confirm you want to exit guided mode.",
            "Guided Mode Exit Message"
        )

        verify_text(
            excel_lookup,
            names.centralWidget_CancelButton_QPushButton_3,
            "CANCEL",
            "Guided Mode Cancel Button"
        )

        verify_text(
            excel_lookup,
            names.centralWidget_DisableButton_QPushButton_2,
            "CONFIRM",
            "Guided Mode Confirm Button"
        )

        # test.vp("Exit Guided Mode")

        clickButton(
            waitForObject(
                names.centralWidget_DisableButton_QPushButton_2
            )
        )

    finally:
        workbook.close()
