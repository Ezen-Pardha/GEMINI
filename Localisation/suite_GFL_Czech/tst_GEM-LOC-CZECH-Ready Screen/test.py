# # -*- coding: utf-8 -*-
#
# import names
#
#
# def main():
#     clickButton(waitForObject(names.numpad_pushButton_8_NumpadButton))
#     clickButton(waitForObject(names.numpad_pushButton_8_NumpadButton))
#     clickButton(waitForObject(names.numpad_pushButton_8_NumpadButton))
#     clickButton(waitForObject(names.numpad_pushButton_8_NumpadButton))
#     clickButton(waitForObject(names.mainFrame_okButton_QPushButton))
#     clickButton(waitForObject(names.homeScreen_settingsButton_QPushButton))
#     clickTab(waitForObject(names.settingsScreen_tbSystemInfo_TabWidget), "INFO")
#     clickButton(waitForObject(names.gbSystemInformation_pbLanguageSettingsUpdate_QPushButton))
#     clickButton(waitForObject(names.languagesFrame_pbCzech_QPushButton))
#     clickButton(waitForObject(names.buttonsFrame_acceptButton_QPushButton))
#     clickButton(waitForObject(names.buttonsBar_homeButton_QPushButton))
#     clickButton(waitForObject(names.homeScreen_pbSoftTissueQuickStart_QToolButton))
#     mouseClick(waitForObjectItem(names.tabLeftMainPresets_leftPresetsView_PresetListView, "_1"), 238, 58, Qt.NoModifier, Qt.LeftButton)
#     mouseClick(waitForObjectItem(names.tabRightMainPresets_rightPresetsView_PresetListView, "_1"), 148, 32, Qt.NoModifier, Qt.LeftButton)
#     mouseClick(waitForObject(names.quickStartScreen_wdBottom_QWidget), 1167, 9, Qt.NoModifier, Qt.LeftButton)
#     sendEvent("QMoveEvent", waitForObject(names.o_ScreenSwitcher), 0, -51, 1286, 61)
#     clickButton(waitForObject(names.quickStartScreen_pbContinue_QPushButton))
#     clickButton(waitForObject(names.mainFrame_okButton_QPushButton_2))
#     clickButton(waitForObject(names.stateSwitch_readyButton_QPushButton))
#     test.compare(waitForObjectExists(names.eNUKLEACE_PROSTATY_VYSOK_ENERGIE_TabItem).text, "ENUKLEACE PROSTATY-(VYSOKÁ ENERGIE)")
#     test.compare(str(waitForObjectExists(names.averagePowerWidgetEmission_nameLabel_QLabel).text), "PRŮMĚRNÝ\nVÝKON")
#     test.compare(str(waitForObjectExists(names.frequencyWidgetEmission_nameLabel_QLabel).text), "FREKVENCE")
#     test.compare(str(waitForObjectExists(names.averagePowerWidgetEmission_unitsOfMeasureLabel_QLabel).text), "W")
#     test.compare(str(waitForObjectExists(names.pulseEnergyWidgetEmission_unitsOfMeasureLabel_QLabel).text), "J")
#     test.compare(str(waitForObjectExists(names.frequencyWidgetEmission_unitsOfMeasureLabel_QLabel).text), "Hz")
#     clickButton(waitForObject(names.stateSwitch_readyButton_QPushButton))
#     test.compare(str(waitForObjectExists(names.modeButtonEmission_nameLabel_QLabel).text), "BĚŽNÝ\nIMPULS")
#     test.compare(str(waitForObjectExists(names.peakPowerButtonEmission_nameLabel_QLabel).text), "ŠPIČKOVÝ\nVÝKON")
#     test.compare(str(waitForObjectExists(names.totalEnergyFrame_nameLabel_QLabel).text), "CELKOVÁ\nENERGIE")
#     test.compare(str(waitForObjectExists(names.totalTimeFrame_nameLabel_QLabel).text), "CELKOVÁ DOBA\nLASEROVÁNÍ")
#     test.compare(waitForObjectExists(names.eNUKLEACE_PROSTATY_VYSOK_ENERGIE_TabItem_2).text, "ENUKLEACE PROSTATY-(VYSOKÁ ENERGIE)")
#     test.compare(str(waitForObjectExists(names.averagePowerWidgetEmission_nameLabel_QLabel_2).text), "PRŮMĚRNÝ\nVÝKON")
#     clickButton(waitForObject(names.stateSwitch_readyButton_QPushButton))
#     test.compare(str(waitForObjectExists(names.pulseEnergyWidgetEmission_nameLabel_QLabel_2).text), "PULSNÍ\nENERGIE")
#     test.compare(str(waitForObjectExists(names.frequencyWidgetEmission_nameLabel_QLabel_2).text), "FREKVENCE")
#     test.compare(str(waitForObjectExists(names.averagePowerWidgetEmission_unitsOfMeasureLabel_QLabel_2).text), "W")
#     test.compare(str(waitForObjectExists(names.pulseEnergyWidgetEmission_unitsOfMeasureLabel_QLabel_2).text), "J")
#     test.compare(str(waitForObjectExists(names.frequencyWidgetEmission_unitsOfMeasureLabel_QLabel_2).text), "Hz")
#     test.compare(str(waitForObjectExists(names.peakPowerButtonEmission_nameLabel_QLabel_2).text), "ŠPIČKOVÝ\nVÝKON")
#     clickButton(waitForObject(names.stateSwitch_readyButton_QPushButton))
#     test.compare(str(waitForObjectExists(names.modeButtonEmission_nameLabel_QLabel_2).text), "BĚŽNÝ\nIMPULS")
#     test.compare(str(waitForObjectExists(names.fiberFrame_nameLabel_QLabel).text), "VLÁKNO")
#     test.compare(str(waitForObjectExists(names.fiberUsesFrame_nameLabel_QLabel).text), "ZBÝVAJÍCÍ\nPOUŽITÍ")
#     test.compare(str(waitForObjectExists(names.bottomBar_aimingBeamButton_AimingBeamModeButton).text), "Zaměřovací\npaprsek")
#     test.compare(str(waitForObjectExists(names.stateSwitch_standbyButton_QPushButton).text), "POHOTOVOSTNÍ REŽIM")
#     test.compare(str(waitForObjectExists(names.stateSwitch_readyButton_QPushButton).text), "PŘIPRAVENO")

# -*- coding: utf-8 -*-

import names
import openpyxl
from openpyxl import load_workbook




EXCEL_PATH = (
    "/home/ezen/Test_Automation_Gemini/Localisation/"
    "suite_GFL_Czech/"
    "tst_GEM-LOC-CZECH-Ready Screen/"
    "testdata/Ready Screen Czech.xlsx"
)




def normalize_text(value):

    if value is None:
        return ""

    return " ".join(str(value).split())





def get_expected_text(sheet, excel_row):

    value = sheet.cell(
        row=excel_row,
        column=2
    ).value

    return normalize_text(value)




def verify_text(sheet, object_name, excel_row, description):

 
    

    snooze(2)

    

  

    obj = waitForObjectExists(
        getattr(names, object_name)
    )



    actual_text = normalize_text(
        obj.text
    )

 


    expected_text = get_expected_text(
        sheet,
        excel_row
    )

  
   

    test.compare(
        actual_text,
        expected_text,
        description + " text verification"
    )



def main():


    workbook = load_workbook(
        EXCEL_PATH
    )

    sheet = workbook["Sheet1"]

    

    for i in range(4):

        clickButton(
            waitForObject(
                names.numpad_pushButton_8_NumpadButton
            )
        )

        snooze(2)


    clickButton(
        waitForObject(
            names.mainFrame_okButton_QPushButton
        )
    )

    snooze(2)


    clickButton(
        waitForObject(
            names.homeScreen_settingsButton_QPushButton
        )
    )

    snooze(2)




    

    clickTab(
        waitForObject(
            names.settingsScreen_tbSystemInfo_TabWidget
        ),
        "INFO"
    )

    test.ocrTextPresent("DEFAULT")

    # test.breakpoint()
    snooze(2)


    clickButton(
        waitForObject(
            names.gbSystemInformation_pbLanguageSettingsUpdate_QPushButton
        )
    )

    snooze(2)

    

    clickButton(
        waitForObject(
            names.languagesFrame_pbCzech_QPushButton
        )
    )

    snooze(2)

    clickButton(
        waitForObject(
            names.buttonsFrame_acceptButton_QPushButton
        )
    )

    snooze(2)

    clickButton(
        waitForObject(
            names.buttonsBar_homeButton_QPushButton
        )
    )

    snooze(2)


    clickButton(
        waitForObject(
            names.homeScreen_pbSoftTissueQuickStart_QToolButton
        )
    )

    snooze(2)


   

    mouseClick(
        waitForObjectItem(
            names.tabLeftMainPresets_leftPresetsView_PresetListView,
            "_1"
        ),
        238,
        58,
        Qt.NoModifier,
        Qt.LeftButton
    )

    snooze(2)


  

    mouseClick(
        waitForObjectItem(
            names.tabRightMainPresets_rightPresetsView_PresetListView,
            "_1"
        ),
        148,
        32,
        Qt.NoModifier,
        Qt.LeftButton
    )

    snooze(2)


   

    mouseClick(
        waitForObject(
            names.quickStartScreen_wdBottom_QWidget
        ),
        1167,
        9,
        Qt.NoModifier,
        Qt.LeftButton
    )

    snooze(2)


    
    

    sendEvent(
        "QMoveEvent",
        waitForObject(
            names.o_ScreenSwitcher
        ),
        0,
        -51,
        1286,
        61
    )

    snooze(2)


    

    clickButton(
        waitForObject(
            names.quickStartScreen_pbContinue_QPushButton
        )
    )

    snooze(2)


   
    clickButton(
        waitForObject(
            names.mainFrame_okButton_QPushButton_2
        )
    )

    snooze(2)


   

    clickButton(
        waitForObject(
            names.stateSwitch_readyButton_QPushButton
        )
    )

    snooze(2)


   
    verify_text(
        sheet,
        "eNUKLEACE_PROSTATY_VYSOK_ENERGIE_TabItem",
        2,
        "Left Treatment Name"
    )


    # --------------------------------------------------------
    # LEFT AVERAGE POWER
    # Excel Row 3
    # --------------------------------------------------------

    verify_text(
        sheet,
        "averagePowerWidgetEmission_nameLabel_QLabel",
        3,
        "Left Average Power"
    )


    # --------------------------------------------------------
    # LEFT PULSE ENERGY
    # Excel Row 4
    # --------------------------------------------------------

    verify_text(
        sheet,
        "pulseEnergyWidgetEmission_nameLabel_QLabel",
        4,
        "Left Pulse Energy"
    )


 

    verify_text(
        sheet,
        "frequencyWidgetEmission_nameLabel_QLabel",
        5,
        "Left Frequency"
    )


  


    verify_text(
        sheet,
        "modeButtonEmission_nameLabel_QLabel",
        6,
        "Left Pulse Mode"
    )


  

    verify_text(
        sheet,
        "peakPowerButtonEmission_nameLabel_QLabel",
        7,
        "Left Peak Power"
    )


 

    verify_text(
        sheet,
        "totalEnergyFrame_nameLabel_QLabel",
        8,
        "Left Total Energy"
    )




    verify_text(
        sheet,
        "totalTimeFrame_nameLabel_QLabel",
        9,
        "Left Total Time"
    )




    verify_text(
        sheet,
        "averagePowerWidgetEmission_unitsOfMeasureLabel_QLabel",
        10,
        "Left Power Unit"
    )


 

    verify_text(
        sheet,
        "pulseEnergyWidgetEmission_unitsOfMeasureLabel_QLabel",
        11,
        "Left Energy Unit"
    )


 

    verify_text(
        sheet,
        "frequencyWidgetEmission_unitsOfMeasureLabel_QLabel",
        12,
        "Left Frequency Unit"
    )



    verify_text(
        sheet,
        "fiberFrame_nameLabel_QLabel",
        13,
        "Left Fiber"
    )


 

    verify_text(
        sheet,
        "fiberUsesFrame_nameLabel_QLabel",
        14,
        "Left Remaining Uses"
    )


  

    verify_text(
        sheet,
        "bottomBar_aimingBeamButton_AimingBeamModeButton",
        15,
        "Left Aiming Beam"
    )


 

    verify_text(
        sheet,
        "stateSwitch_standbyButton_QPushButton",
        16,
        "Left Standby"
    )


 

    verify_text(
        sheet,
        "stateSwitch_readyButton_QPushButton",
        17,
        "Left Ready"
    )


   

    
 
    verify_text(
        sheet,
        "eNUKLEACE_PROSTATY_VYSOK_ENERGIE_TabItem_2",
        2,
        "Right Treatment Name"
    )


 

    verify_text(
        sheet,
        "averagePowerWidgetEmission_nameLabel_QLabel_2",
        3,
        "Right Average Power"
    )


  

    verify_text(
        sheet,
        "pulseEnergyWidgetEmission_nameLabel_QLabel_2",
        4,
        "Right Pulse Energy"
    )


    # --------------------------------------------------------
    # RIGHT FREQUENCY
    # Excel Row 5
    # --------------------------------------------------------

    verify_text(
        sheet,
        "frequencyWidgetEmission_nameLabel_QLabel_2",
        5,
        "Right Frequency"
    )




    verify_text(
        sheet,
        "modeButtonEmission_nameLabel_QLabel_2",
        6,
        "Right Pulse Mode"
    )


  

    verify_text(
        sheet,
        "peakPowerButtonEmission_nameLabel_QLabel_2",
        7,
        "Right Peak Power"
    )


    
    verify_text(
        sheet,
        "averagePowerWidgetEmission_unitsOfMeasureLabel_QLabel_2",
        10,
        "Right Power Unit"
    )


   
 

    verify_text(
        sheet,
        "pulseEnergyWidgetEmission_unitsOfMeasureLabel_QLabel_2",
        11,
        "Right Energy Unit"
    )



    verify_text(
        sheet,
        "frequencyWidgetEmission_unitsOfMeasureLabel_QLabel_2",
        12,
        "Right Frequency Unit"
    )


 


  
 

    clickButton(
        waitForObject(
            names.stateSwitch_readyButton_QPushButton
        )
    )

    snooze(2)


   
    clickButton(waitForObject(names.stateSwitch_readyButton_QPushButton))
    test.vp("Readyscreen")

    

    
    

