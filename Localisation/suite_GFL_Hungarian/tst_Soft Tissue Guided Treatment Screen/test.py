
# -*- coding: utf-8 -*-

import names
import openpyxl
import re
import os

TESTCASE_DIR = os.path.dirname(os.path.abspath(__file__))

EXCEL_PATH = os.path.abspath(
    os.path.join(
        TESTCASE_DIR,
        "..",
        "shared",
        "testdata",
        "Gemini Split strings.xlsx"
    )
)


# EXCEL_PATH = "/home/ntc/Test_Automation_Gemini/Localisation/suite_GFL_Czech/tst_Home Screen/testdata/Gemini Split strings.xlsx"

SHEET_NAME = "Treatment Modes"
LANGUAGE_COLUMN = 12
ENGLISH_COLUMN = 1


def normalize(text):
    if text is None:
        return ""

    text = str(text)
    text = re.sub(r"<br\s*/?>", " ", text, flags=re.IGNORECASE)
    text = re.sub(r"<[^>]+>", " ", text)
    text = text.replace("\xa0", " ")
    text = " ".join(text.split())
    text = re.sub(r"\s*<\s*", "<", text)
    text = re.sub(r"\s*>\s*", ">", text)

    return text.strip().casefold()


def ignore_text(text):
    key = normalize(text)

    if not key:
        return True

    # Ignore standalone numeric values and measurements.
    if re.fullmatch(
        r"\d+(?:[.,]\d+)?\s*(?:μm|µm|mm|cm|m|w|hz|j|s|ms|%)?",
        key,
        flags=re.IGNORECASE
    ):
        return True

    if re.fullmatch(r"[\d\s.,:+\-/<>=%µμ]+", key):
        return True

    if re.fullmatch(r"[-_=.<>/%µμ:+]+", key):
        return True

    return False


def get_object_texts(obj):
    texts = []

    try:
        if obj.visible is False:
            return texts
    except Exception:
        pass

    for attribute in ("text", "title"):
        try:
            value = getattr(obj, attribute)

            if value is None:
                continue

            value = str(value).strip()
            value = re.sub(
                r"<br\s*/?>", " ", value,
                flags=re.IGNORECASE
            )
            value = re.sub(r"<[^>]+>", " ", value)
            value = " ".join(value.split())

            if value and not ignore_text(value):
                texts.append(value)

        except Exception:
            pass

    return texts


def get_screen_texts(obj, visited=None):
    if visited is None:
        visited = set()

    texts = []

    try:
        obj_id = id(obj)

        if obj_id in visited:
            return texts

        visited.add(obj_id)
    except Exception:
        pass

    texts.extend(get_object_texts(obj))

    try:
        children = object.children(obj)
    except Exception:
        children = []

    for child in children:
        texts.extend(get_screen_texts(child, visited))

    return texts


def load_excel_strings():
    workbook = openpyxl.load_workbook(
        EXCEL_PATH,
        data_only=True,
        read_only=True
    )

    try:
        if SHEET_NAME not in workbook.sheetnames:
            raise Exception(
                "Excel sheet not found: " + SHEET_NAME
            )

        sheet = workbook[SHEET_NAME]
        excel_strings = {}

        for row in sheet.iter_rows(min_row=2, values_only=True):
            if len(row) <= LANGUAGE_COLUMN:
                continue

            english = row[ENGLISH_COLUMN]
            localized = row[LANGUAGE_COLUMN]

            if english is None or localized is None:
                continue

            english = str(english).strip()
            localized = str(localized).strip()

            if not english or not localized:
                continue

            key = normalize(localized)

            if key:
                excel_strings[key] = {
                    "english": english,
                    "localized": localized
                }

        return excel_strings

    finally:
        workbook.close()


def verify_visible_strings(
    *objects,
    screen_name,
    excel_strings
):
    runtime_texts = []
    visited = set()

    # Collect text from every supplied object and its descendants.
    for obj in objects:
        if obj is not None:
            runtime_texts.extend(
                get_screen_texts(obj, visited)
            )

    unique_texts = {}

    for text in runtime_texts:
        key = normalize(text)

        if key and key not in unique_texts:
            unique_texts[key] = text

    test.log(
        "Collected screen text entries: "
        + str(len(runtime_texts))
    )

    test.log(
        "Unique screen strings: "
        + str(len(unique_texts))
    )

    if not unique_texts:
        test.fail(
            "No readable screen text collected from "
            + screen_name
        )
        return

    pass_count = 0
    fail_count = 0

    for key in sorted(unique_texts):
        screen_text = unique_texts[key]

        if key in excel_strings:
            data = excel_strings[key]

            test.passes(
                "PASS: "
                + data["english"]
                + " -> "
                + data["localized"]
            )

            pass_count += 1

        else:
            test.fail(
                "FAIL: Screen string not found in Hungary Excel: "
                + screen_text
            )

            fail_count += 1

    test.log("------------------------------------------")
    test.log(
        screen_name
        + " Localization Result: PASS = "
        + str(pass_count)
        + ", FAIL = "
        + str(fail_count)
    )
    test.log("------------------------------------------")

    if fail_count == 0 and pass_count > 0:
        test.passes(
            "Hungary "
            + screen_name.upper()
            + " LOCALIZATION VALIDATION PASSED"
        )
    else:
        test.fail(
            "Hungary "
            + screen_name.upper()
            + " LOCALIZATION VALIDATION FAILED"
        )


def main():

    # LOGIN
    sendEvent(
        "QMoveEvent",
        waitForObject(names.o_ScreenSwitcher),
        34, 185, 582, 201
    )

    clickButton(waitForObject(names.numpad_pushButton_8_NumpadButton))
    clickButton(waitForObject(names.numpad_pushButton_8_NumpadButton))
    clickButton(waitForObject(names.numpad_pushButton_8_NumpadButton))
    clickButton(waitForObject(names.numpad_pushButton_8_NumpadButton))
    clickButton(waitForObject(names.mainFrame_okButton_QPushButton))

    # CHANGE LANGUAGE TO Hungary
    clickButton(waitForObject(names.homeScreen_settingsButton_QPushButton))
    clickTab(
        waitForObject(names.settingsScreen_tbSystemInfo_TabWidget),
        "INFO"
    )
    clickButton(
        waitForObject(
            names.gbSystemInformation_pbLanguageSettingsUpdate_QPushButton
        )
    )
    clickButton(waitForObject(names.languagesFrame_pbHungary_QPushButton))
    clickButton(waitForObject(names.buttonsFrame_acceptButton_QPushButton))
    clickButton(waitForObject(names.settingsScreen_pbBack_QPushButton))

    # OPEN SOFT TISSUE ASSISTANT
    clickButton(
        waitForObject(names.homeScreen_pbSoftTissueAssistant_QToolButton)
    )
    clickButton(
        waitForObject(names.gbSoftTissueProcedures_pbIncision_QPushButton)
    )
    clickButton(
        waitForObject(
            names.softTissueAssistantModeScreen_pbContinue_QPushButton
        )
    )

    # TREATMENT SCREEN - RECORDED OBJECTS
    mouseClick(
        waitForObject(names.leftPedalWidget_wdPresetTabs_TabWidget),
        297, 32, Qt.NoModifier, Qt.LeftButton
    )

    mouseClick(
        waitForObject(names.averagePowerWidget_unitsOfMeasureLabel_QLabel_2),
        164, 27, Qt.NoModifier, Qt.LeftButton
    )

    # mouseClick(
    #     waitForObject(names.averagePowerWidget_nameLabel_QLabel_2),
    #     30, 10, Qt.NoModifier, Qt.LeftButton
    # )

    mouseClick(
        waitForObject(names.pulseModeWidget_buttonsWidget_QWidget),
        566, 41, Qt.NoModifier, Qt.LeftButton
    )

    mouseClick(
        waitForObject(names.pulseModeWidget_buttonsWidget_QWidget_2),
        96, 36, Qt.NoModifier, Qt.LeftButton
    )

    mouseClick(
        waitForObject(names.rightPedalWidget_wdPresetTabs_TabWidget),
        303, 18, Qt.NoModifier, Qt.LeftButton
    )

    mouseClick(
        waitForObject(names.topBar_buttonsBar_ButtonsBar),
        875, 33, Qt.NoModifier, Qt.LeftButton
    )

    mouseClick(
        waitForObject(names.bottomBar_lbPedalStatusIcon_QLabel),
        53, 46, Qt.NoModifier, Qt.LeftButton
    )

    snooze(1)

    # LOAD Hungary EXCEL CATALOG
    excel_strings = load_excel_strings()

    test.log(
        "Hungary Excel catalog loaded: "
        + str(len(excel_strings))
        + " translations"
    )

    # VALIDATE TEXT FROM MULTIPLE UI OBJECTS
    verify_visible_strings(
        waitForObject(names.o_ScreenSwitcher),
        waitForObject(names.treatmentScreen_bottomBar_QFrame),
        waitForObject(names.leftPedalWidget_wdPresetTabs_TabWidget),
        waitForObject(names.rightPedalWidget_wdPresetTabs_TabWidget),
        waitForObject(names.topBar_buttonsBar_ButtonsBar),
        waitForObject(names.bottomBar_lbPedalStatusIcon_QLabel),
        waitForObject(names.averagePowerWidget_unitsOfMeasureLabel_QLabel_2),
        # waitForObject(names.averagePowerWidget_nameLabel_QLabel_2),
        waitForObject(names.pulseModeWidget_buttonsWidget_QWidget),
        waitForObject(names.pulseModeWidget_buttonsWidget_QWidget_2),
        screen_name="Soft Tissue Treatment Screen",
        excel_strings=excel_strings
    )
